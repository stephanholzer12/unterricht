import json, os, io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE as T
from PIL import Image
SRC="/home/claude/out/26-27-2AB-K1.2-Bindungen-TH-Kadetten.pptx"; OUT="/home/claude/handy/"
p=Presentation(SRC); E=914400
def flat(shs):
    for s in shs:
        if s.shape_type==T.GROUP: yield from flat(s.shapes)
        else: yield s
def runs(par):
    out=[]
    for r in par.runs:
        if not r.text: continue
        b=r.font.bold if r.font.bold is not None else False
        out.append({"t":r.text,"b":bool(b)})
    return out
def paras(s):
    ps=[]
    for par in s.text_frame.paragraphs:
        rr=runs(par)
        if rr and "".join(x["t"] for x in rr).strip(): ps.append(rr)
    return ps
def has_fill(s):
    try: return s.fill.type is not None and s.fill.type==1
    except Exception: return False
nimg=0
def save_img(s, slide_i):
    global nimg
    nimg+=1; ext=s.image.ext
    im=Image.open(io.BytesIO(s.image.blob))
    if ext=="gif":
        try:
            im.seek(im.n_frames-1)
        except Exception: pass
    im=im.convert("RGBA")
    bg=Image.new("RGBA",im.size,(255,255,255,255)); bg.alpha_composite(im)
    f=f"img/s{slide_i:02d}_{nimg}.png"; bg.convert("RGB").save(OUT+f)
    d={"type":"img","f":f,"w":s.width/E,"h":s.height/E}
    if ext=="gif":
        g=f"img/s{slide_i:02d}_{nimg}.gif"; open(OUT+g,"wb").write(s.image.blob); d["gif"]=g
    return d
slides=[]
for i,sl in enumerate(p.slides,1):
    lay=sl.slide_layout.name
    shapes=list(flat(sl.shapes))
    title=""; items=[]; containers=[]
    for s in shapes:
        n=s.name
        if n.startswith("Titel") and s.has_text_frame: title=s.text_frame.text.strip(); continue
        if n.startswith(("Slide Number","Buchseite")) :
            if n.startswith("Buchseite") and s.has_text_frame: items.append({"type":"pill","t":s.text_frame.text.replace("\n"," "),"x":s.left/E,"y":s.top/E,"w":s.width/E,"h":s.height/E})
            continue
        x,y,w,h=s.left/E,s.top/E,s.width/E,s.height/E
        if s.shape_type==T.PICTURE:
            if w<0.7 and h<0.7: continue   # Icons
            if "GIF Hintergrund" in n: continue
            d=save_img(s,i); d.update(x=x,y=y,name=n); items.append(d); continue
        if s.has_table:
            rows=[[ "\n".join(c.text_frame.text.split("\n")) for c in r.cells] for r in s.table.rows]
            items.append({"type":"table","rows":rows,"x":x,"y":y,"w":w,"h":h,"cw":[c.width/E for c in s.table.columns]}); continue
        if n.startswith("Antwortfeld"):
            t=s.text_frame.text.split("\n") if s.has_text_frame else ["Deine Antwort"]
            items.append({"type":"answer","label":t[0],"hint":t[1] if len(t)>1 else "","x":x,"y":y,"w":w,"h":h}); continue
        if s.has_text_frame and paras(s):
            items.append({"type":"text","p":paras(s),"x":x,"y":y,"w":w,"h":h,"name":n,"filled":has_fill(s)}); continue
        if not s.has_text_frame or not s.text_frame.text.strip():
            if w>1.2 and h>0.5 and not n.startswith(("Linie","Kreis","Pfeil")): containers.append({"x":x,"y":y,"w":w,"h":h,"name":n})
    # Zuordnung zu Containern (Karten)
    def inside(it,c):
        cx,cy=it["x"]+it.get("w",0)/2,it["y"]+it.get("h",0)/2
        return c["x"]-0.05<=cx<=c["x"]+c["w"]+0.05 and c["y"]-0.05<=cy<=c["y"]+c["h"]+0.05
    # kleinste passende Container zuerst
    containers.sort(key=lambda c:c["w"]*c["h"])
    blocks=[]; used=set()
    for c in containers:
        mem=[k for k,it in enumerate(items) if k not in used and it["type"] in("text","img") and inside(it,c)]
        if not mem: continue
        for k in mem: used.add(k)
        blocks.append({"type":"card","x":c["x"],"y":c["y"],"items":sorted([items[k] for k in mem],key=lambda a:(round(a["y"],1),a["x"])),"name":c["name"]})
    for k,it in enumerate(items):
        if k not in used: blocks.append(it)
    blocks.sort(key=lambda b:(round(b["y"]*2)/2,b["x"]))
    slides.append({"i":i,"layout":lay,"title":title,"blocks":blocks})
json.dump(slides,open(OUT+"slides.json","w"),ensure_ascii=False,indent=1)
print(len(slides),"slides", nimg,"imgs")
