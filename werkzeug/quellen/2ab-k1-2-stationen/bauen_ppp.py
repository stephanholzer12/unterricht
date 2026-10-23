"""Folien für die Stationen K1.2 (2AB, Stunde 3–4): VS-Lehrer und VS-Kadetten mit dem BHAK-Folienmaster.
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/2ab-k1-2-stationen/bauen_ppp.py <BHAK_Folienmaster_Vorlage.pptx> <Zielordner>
Die Folien können auch einzeln in die TH-PPP K1.2 kopiert werden (nach »Stunde 2«).
Lehrer: Lösungen per Klick, ausgeblendete Lehrer-Infos. Kadetten: gleiche Folien, leere Antwortfelder, keine Lehrer-Infos."""
import copy, io, os, sys, uuid
import qrcode
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import daten as D

MASTER, OUT = sys.argv[1], sys.argv[2]
PRE = "26-27-2AB-K1.2b-Station"
VER = "v3"
OLIVE, INK, CARD, GREEN, ORANGE, TEAL, RUST, MUTED, WHITE, PALE = (RGBColor.from_string(x) for x in
    ["3B4A2F", "1F2A1C", "EEF3E4", "7FA83A", "D9772B", "2E7D8C", "8C3B2E", "5E6652", "FFFFFF", "F7F9F2"])
RED = RGBColor.from_string("C0001E")
V = {v["id"]: v for v in D.versuche()}
IMG = D.UNIT + "vs-img/"


# ---------------------------------------------------------------- Helfer
def qr_png(url):
    q = qrcode.QRCode(border=1, box_size=10); q.add_data(url); q.make(fit=True)
    b = io.BytesIO(); q.make_image(fill_color="#1F2A1C", back_color="white").save(b, "PNG"); b.seek(0); return b


class Deck:
    def __init__(self, lehrer):
        self.lehrer = lehrer
        self.p = Presentation(MASTER)
        lst = self.p.slides._sldIdLst
        self.pill_proto = None
        for sh in self.p.slides[2].shapes:
            if sh.name == "Buchseite": self.pill_proto = copy.deepcopy(sh._element)
        for sid in list(lst):  # Beispiel-Folien der Vorlage entfernen
            self.p.part.drop_rel(sid.get(qn("r:id"))); lst.remove(sid)
        self.L = {l.name: l for l in self.p.slide_layouts}
        self.sections = []  # (Name, [sldId])
        for part in self.p.part.package.iter_parts():  # Fußzeile »Fach · Jahrgang · Thema« einmal anpassen
            if "slideLayout" in str(part.partname) or "slideMaster" in str(part.partname):
                for t in part._element.iter(qn("a:t")):
                    if t.text and "Fach · Jahrgang · Thema" in t.text: t.text = "NAWI · 2AB · K1.2 Bindungsarten"

    def section(self, name):
        self.sections.append((name, []))

    def slide(self, layout, title, hidden=False):
        s = self.p.slides.add_slide(self.L[layout])
        for ph in s.placeholders:
            if ph.placeholder_format.type == 1: ph.text_frame.text = title  # TITLE
        if hidden: s._element.set("show", "0")
        self.sections[-1][1].append(self.p.slides._sldIdLst[-1].get("id"))
        s._anim = []
        return s

    def save(self, path):
        for s in self.p.slides:
            if getattr(s, "_anim", None): add_clicks(s, s._anim)
        # Abschnitte
        prs = self.p.part._element
        ext_lst = prs.find(qn("p:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(prs, qn("p:extLst"))
        P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        for old in ext_lst.findall(qn("p:ext")):
            if old.get("uri") == "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}": ext_lst.remove(old)
        ext = etree.SubElement(ext_lst, qn("p:ext")); ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        sl = etree.SubElement(ext, "{%s}sectionLst" % P14, nsmap={"p14": P14})
        for name, ids in self.sections:
            se = etree.SubElement(sl, "{%s}section" % P14); se.set("name", name); se.set("id", "{%s}" % str(uuid.uuid4()).upper())
            il = etree.SubElement(se, "{%s}sldIdLst" % P14)
            for i in ids: etree.SubElement(il, "{%s}sldId" % P14).set("id", i)
        self.p.save(path)


def add_clicks(slide, groups):
    """groups: Liste von Klick-Schritten, jeder Schritt = Liste von Shapes, die gemeinsam einblenden (Verblassen)."""
    nid = [2]
    def cid():
        nid[0] += 1; return str(nid[0])
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    clicks = ""
    for grp in groups:
        inner = ""
        for k, sh in enumerate(grp):
            spid = sh.shape_id
            inner += (f'<p:par><p:cTn id="{cid()}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{"clickEffect" if k == 0 else "withEffect"}">'
                      f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                      f'<p:set><p:cBhvr><p:cTn id="{cid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                      f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                      f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                      f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid()}" dur="400"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                      f'</p:childTnLst></p:cTn></p:par>')
        clicks += (f'<p:par><p:cTn id="{cid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{cid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'
                   f'</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{"".join(f'<p:bldP spid="{sh.shape_id}" grpId="0" animBg="1"/>' for grp in groups for sh in grp)}</p:bldLst></p:timing>')
    el = etree.fromstring(xml)
    sld = slide._element
    ext = sld.find(qn("p:extLst"))
    if ext is not None: ext.addprevious(el)
    else: sld.append(el)


def box(s, x, y, w, h, fill=None, line=None, radius=True, name=None, lw=1.25):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: shp.adjustments[0] = 0.08
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if name: shp.name = name
    tf = shp.text_frame; tf.word_wrap = True
    for side in ("margin_left", "margin_right"): setattr(tf, side, Inches(0.1))
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    return shp


def text(target, parts, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4, font=None):
    """parts: Liste von Absätzen; ein Absatz = str oder Liste von (Text, {opts})."""
    tf = target.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    first = True
    for para in parts:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(after)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.color.rgb = o.get("color", color)
            if o.get("italic"): f.italic = True
            if font or o.get("font"): f.name = o.get("font", font)
    return target


def tb(s, x, y, w, h, parts, **kw):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.margin_left = t.text_frame.margin_right = Inches(0.05)
    return text(t, parts, **kw)


def pill(s, label, x=3.55, y=5.18, w=1.9, kind="theory"):
    el = copy.deepcopy(d.pill_proto)
    ids = [int(x) for x in s.shapes._spTree.xpath("//p:cNvPr/@id")]
    el.xpath(".//p:cNvPr")[0].set("id", str(max(ids + [1]) + 1))
    s.shapes._spTree.append(el)
    shp = s.shapes[-1]
    shp.left, shp.top, shp.width = Inches(x), Inches(y), Inches(w)
    col = {"theory": RGBColor.from_string("6E7F4E"), "book": TEAL, "versuch": RUST}[kind]
    shp.fill.solid(); shp.fill.fore_color.rgb = col
    r = shp.text_frame.paragraphs[0].runs
    r[0].text = label
    for extra in r[1:]: extra.text = ""
    return shp


def tag(s, label, color=OLIVE, x=7.9, y=0.95, w=1.6):
    t = box(s, x, y, w, 0.34, fill=color)
    text(t, [label], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return t


def grid(s, x, y, colw, rows, rowh, head=True, size=16, fills=None, hidden_cols=(), sizes=None):
    """Tabelle aus einzelnen Rechtecken (Zellen lassen sich einzeln per Klick einblenden). Gibt Zellen-Matrix zurück."""
    cells = []
    yy = y
    for ri, row in enumerate(rows):
        xx = x; rc = []
        h = rowh[ri] if isinstance(rowh, (list, tuple)) else rowh
        for ci, val in enumerate(row):
            hd = head and ri == 0
            fill = OLIVE if hd else (fills[ri] if fills else (CARD if ri % 2 else WHITE))
            c = box(s, xx, yy, colw[ci], h, fill=fill, line=RGBColor.from_string("B9C4A6"), radius=False, lw=0.75)
            sz = (sizes[ci] if sizes else size)
            parts = val if isinstance(val, list) else [val]
            text(c, parts, size=sz, color=WHITE if hd else INK, bold=hd or (ci == 0), anchor=MSO_ANCHOR.MIDDLE, after=0)
            rc.append(c); xx += colw[ci]
        cells.append(rc); yy += h
    return cells


def answer(s, x, y, w, h, title, body, lehrer, hint="Deine Antwort – im Versuchsprotokoll eintragen"):
    """Lehrer: orange Lösungskasten (per Klick). Kadetten: leeres Antwortfeld."""
    if lehrer:
        b = box(s, x, y, w, h, fill=WHITE, line=ORANGE, lw=2)
        text(b, [[(title, {"bold": True, "color": ORANGE})], body], size=18)
        s._anim.append([b])
    else:
        b = box(s, x, y, w, h, fill=PALE, line=RGBColor.from_string("B9C4A6"), lw=1.5)
        text(b, [[(hint, {"italic": True, "color": MUTED})]], size=16)
    return b


def ppt_qr(s, url, x, y, size, step, label, color=OLIVE, arrow=False):
    s.shapes.add_picture(qr_png(url), Inches(x), Inches(y), Inches(size), Inches(size))
    c = box(s, x + size / 2 - 0.2, y - 0.25, 0.4, 0.4, fill=color)
    c.auto_shape_type  # noqa
    text(c, [step], size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    tb(s, x - 0.25, y + size + 0.02, size + 0.5, 0.35, [label], size=16, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)
    tb(s, x - 0.3, y + size + 0.36, size + 0.6, 0.5, [url], size=8, color=MUTED, align=PP_ALIGN.CENTER, after=0)


# ---------------------------------------------------------------- Inhalte
STATIONEN = [
 ("A", "haut", "Wasser mit Haut", ["Wasser auf die Münze tropfen – Tropfen zählen", "dasselbe mit Seifenwasser", "Büroklammer auf Filterpapier aufs Wasser legen", "1 Tropfen Waschmittel dazu"],
  ("Buch S. 26", "theory"), "Spülen! Seifenreste verderben den Versuch."),
 ("B", "wasserstrahl", "Der scheue Wasserstrahl", ["dünnen Wasserstrahl aufdrehen", "Ballon ungerieben hinhalten", "Ballon 20-mal reiben", "langsam dem Strahl nähern"],
  ("Versuch im Buch S. 27", "versuch"), "Ballon trocken halten – nass ist er entladen."),
 ("C", "mischen", "Wer mischt sich mit wem?", ["1 Wasser + Glycerin · 2 Wasser + Öl", "4 Öl + Spatelspitze Salz", "verschließen, 10 s schütteln, 1 min warten", "Glas 2 + 3 Tropfen Waschmittel = Nr. 3"],
  ("Buch S. 22–23, 27", "theory"), "Öl in den Sammelbehälter – nicht in den Ausguss!"),
 ("D", "kristall", "Salzkörner und Metall unter Druck", ["Salz, Zucker, Alaun unter der Lupe (K1.2-V2)", "Alaun im Mörser drücken (K1.2-V3)", "Zinnkorn platt drücken", "Kupferstreifen biegen"],
  ("Buch S. 21, 25", "theory"), "Schutzbrille! Drücken, nicht schlagen."),
]


KURZFRAGE = {"C · Wer mischt sich mit wem?": "Warum mischen sich Wasser und Glycerin – Öl und Wasser aber nicht?",
             "D · Salz und Metall": "Warum zerbricht Alaun, während sich Zinn und Kupfer nur verformen?"}


def build(lehrer):
    global d
    d = Deck(lehrer)
    L = lehrer
    # ================= Start
    d.section("Start")
    s = d.slide("Titelfolie", "Bindungsarten im Versuch")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2:
            text(ph, ["NAWI · 2AB · K1.2 Die Bindungsarten · Stunde 3–4 · Stationen A–D · Buch S. 19–28",
                      [(f"Stand {D.STAND} · {VER} · " + ("Lehrer" if L else "Kadetten"), {"size": 12})]], size=18, color=RGBColor.from_string("C9D6B0"))

    s = d.slide("Nur Titel", "Dein Lernpfad für Stunde 3–4")
    for i, (url, step, lab, col) in enumerate([(D.URL["versuche"], "V", "Versuche · Protokoll", RUST), (D.URL["vertiefung"], "1", "Vertiefung 11–17", OLIVE),
                                               (D.URL["training"], "2", "Test-Training", OLIVE), (D.URL["lernseite"], "?", "Theorie-Seite", TEAL)]):
        ppt_qr(s, url, 0.75 + i * 2.25, 1.45, 1.75, step, lab, col)
    tb(s, 0.5, 4.75, 9, 0.4, ["Ein Protokoll-Handy pro Gruppe scannt »Versuche«. Die Vertiefung macht jeder am eigenen Handy."], size=16, color=MUTED, after=0)

    s = d.slide("Nur Titel", "Dein Fahrplan")
    tb(s, 0.5, 1.0, 9, 0.4, [[("Dafür sind 4 Stunden eingeplant. ", {"bold": True, "color": OLIVE}), ("Stunde 1–2 hast du schon geschafft.", {})]], size=18, after=0)
    rows = [["Wann", "Das machen wir", "Aufgaben", "Abgabe"]]
    kurz = {"Stunde 1": "Theorie-Seite Teil 1–3", "Stunde 2": "Theorie-Seite Teil 4–5 + Überblick"}
    for z in D.FAHRPLAN["zeilen"]:
        ab = z["abgabe"].replace("Arbeitsblatt-PDF in LMS", "PDF in LMS").replace("Versuchs-PDF (Gruppe) + Vertiefungs-PDF (jeder) in LMS", "2 PDFs in LMS")
        rows.append([z["wann"], kurz.get(z["wann"], z["lesen"].replace("Versuchsseite: ", "")), z["aufgaben"], ab])
    GREY = RGBColor.from_string("DCDFD6")
    grid(s, 0.5, 1.42, [1.45, 4.45, 1.85, 1.25], rows, [0.38, 0.42, 0.42, 0.66, 0.92, 0.42], size=16, fills=[OLIVE, GREY, GREY, WHITE, CARD, WHITE])
    tb(s, 0.5, 4.7, 7.5, 0.3, ["Grau = schon erledigt"], size=12, color=MUTED, after=0)

    def abgabe():
        s = d.slide("Nur Titel", "Fertig? So gibst du ab")
        steps = [("1", "Versuchsseite: Namen der Gruppe eintragen → »PDF erstellen« → ein PDF pro Gruppe in LMS"),
                 ("2", "Vertiefung: Name und Klasse → »PDF erstellen« → jeder sein PDF in LMS"),
                 ("3", "Startet kein Download? »Teilen / Speichern« oder »PDF öffnen« tippen – oder die Seite im normalen Browser öffnen")]
        for i, (n, t) in enumerate(steps):
            c = box(s, 0.5, 1.2 + i * 1.12, 0.6, 0.6, fill=OLIVE); text(c, [n], size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
            tb(s, 1.25, 1.12 + i * 1.12, 5.6, 1.0, [t], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
        s.shapes.add_picture(qr_png(D.URL["versuche"]), Inches(7.25), Inches(1.15), Inches(1.0), Inches(1.0))
        tb(s, 6.95, 2.15, 1.6, 0.3, ["Versuche"], size=14, bold=True, color=RUST, align=PP_ALIGN.CENTER, after=0)
        s.shapes.add_picture(qr_png(D.URL["vertiefung"]), Inches(7.25), Inches(2.75), Inches(1.0), Inches(1.0))
        tb(s, 6.95, 3.75, 1.6, 0.3, ["Vertiefung"], size=14, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)
        tb(s, 0.5, 4.6, 8.5, 0.45, ["Abgabe bis zur nächsten NAWI-Stunde."], size=18, bold=True, color=RUST, after=0)
    abgabe()

    if L:
        s = d.slide("Nur Titel", "Lehrer-Info · Lehrplanbezug", hidden=True)
        tag(s, "AUSGEBLENDET", RUST, x=7.6, w=1.9)
        tb(s, 0.5, 1.4, 9, 3.6, [
            [("Jahresplan E3–E5 ", {"bold": True, "color": OLIVE}), ("(je 2 Std., zusammen 6) · Buch S. 19–28 · hier: 2 Std. Theorie + 2 Std. Versuche", {})],
            [("K02 ", {"bold": True}), ("Entstehung chemischer Bindungen erklären → Vertiefung 11, 14, 15", {})],
            [("K03 ", {"bold": True}), ("Bindung ↔ Stoffeigenschaften → Stationen A–D, Demo K1.2-V4, Vertiefung 13, 16", {})],
            [("K07 ", {"bold": True}), ("Experimente durchführen und dokumentieren → Protokoll am Handy", {})],
            [("SDG 3 · Schulprofil: ", {"bold": True}), ("Elektrolyte nach dem Marsch, Tarnschminke, Gewehrlauf im Winter", {})],
            [("Abweichung: ", {"bold": True, "color": RUST}), ("4 statt 6 Stunden; Versuche als Stationen statt verteilt", {})]], size=17, after=6)

        s = d.slide("Nur Titel", "Lehrer-Info · Stundenvorschlag", hidden=True)
        tb(s, 0.5, 0.98, 9, 0.35, ["AUSGEBLENDET · Kern in 4 Stunden · Stunde 1–2 erledigt · Rest optional"], size=13, bold=True, color=RUST, after=0)
        rows = [["Std.", "Theorie", "Kern-Versuch", "Aufgaben", "Methode"],
                ["1–2", "Teil 1–5 Theorie-Seite", "–", "1–10 · Quiz", "Handy + Erklären"],
                ["3", "Sicherheit, Stationsplan", "Gr. 1–4: A–D", "Gr. 5–8: 11–16", "ohne Handy: Stationen"],
                ["4", "Demo, Besprechung, Teststoff", "K1.2-V4 Demo · Gr. 5–8: A–D", "Gr. 1–4: 11–16", "ohne Handy: Demo + Plenum"],
                ["Zusatz", "Für Schnelle", "V1, V5, V6, V10, V11 (heute nicht)", "17 · Test-Training", "Handy"]]
        grid(s, 0.5, 1.4, [0.8, 2.35, 2.35, 1.6, 1.9], rows, [0.4, 0.6, 0.6, 0.75, 0.75], size=14)
        tb(s, 0.5, 4.6, 9, 0.5, ["KERN: Stationen A–D + Demo · OPTIONAL: Aufgabe 17, Zusatz-Versuche · je Station 7 min, Wechsel auf Zeichen"], size=13, color=MUTED, after=0)

        s = d.slide("Nur Titel", "Lehrer-Info · Die 5 größten Stolpersteine", hidden=True)
        rows = [["Kadetten sagen …", "Richtig ist …"]] + [[a, b] for a, b in D.STOLPERSTEINE]
        grid(s, 0.5, 1.08, [3.4, 5.6], rows, [0.34] + [0.7] * 5, size=14)

        for nr, (tit, items) in enumerate([
            ("Herrichten 1/2", [
                ("A · Wasser mit Haut (je Station eine Wanne)", "Petrischale (+1 Ersatz) · Tropfpipette · Becherglas 100 mL · 2 Büroklammern (Büro!) · Filterpapier · Waschmittel in Tropfflasche · Münzen bringen die Kadetten"),
                ("B · Wasserstrahl", "2 Luftballone aufgeblasen · Reibetuch #41 – Station ans Waschbecken"),
                ("C · Mischen", "1 Gestell + 32 Reagenzgläser · Parafilm-Streifen · Pipette · Spatel · blaues Wasser (400 mL) · Glycerin (50-mL-Becherglas) · Olivenöl · Salz · Waschmittel · Sammelbehälter Öl")]),
            ("Herrichten 2/2", [
                ("D · Salz und Metall", "Mörser + Pistill · Alu-Unterlegplatte · 2 Lupen (+ 1 Mikroskop) · 6 Objektträger · Spatel · Salz · Zucker · Alaun (grob) · Zinn gekörnt · Kupferstreifen · Schutzbrillen"),
                ("Demo K1.2-V4", "Multimeter (Batterie prüfen, 200 kΩ) · 2 Bechergläser 100 mL · dest. Wasser · Kupferblech · Salz · Zucker"),
                ("Vor der Stunde", "Wannen bestücken (ca. 15 min) · QR-Codes am Beamer · Teststoff in LMS stellen · Gruppen 1–8 einteilen")])]):
            s = d.slide("Nur Titel", "Lehrer-Info · " + tit, hidden=True)
            for i, (h, t) in enumerate(items):
                c = box(s, 0.5, 1.1 + i * 1.3, 9.0, 1.18, fill=CARD)
                text(c, [[(h, {"bold": True, "color": OLIVE, "size": 17})], t], size=15, after=2)
            tb(s, 0.5, 5.0, 8, 0.3, ["Mit Fundort (Schrank, Lehrmittel-Nr.) und Bestand: VS-Fachwissen S. 1"], size=11, color=MUTED, after=0)

    # ================= Stunde 3
    d.section("Stunde 3")
    s = d.slide("Nur Titel", "Sicherheit an den Stationen")
    items = [("Schutzbrille", "an Station C und D – auch beim Fotografieren"), ("Nichts kosten", "auch nicht Salz oder Zucker – danach Hände waschen"),
             ("Handy weg vom Tisch", "erst arbeiten, dann tippen – nie beim Schütteln oder Drücken"), ("Öl", "in den Sammelbehälter – nicht in den Ausguss"),
             ("Aufräumen", "Station so verlassen, wie ihr sie vorgefunden habt")]
    for i, (h, t) in enumerate(items):
        c = box(s, 0.5, 1.12 + i * 0.78, 0.5, 0.5, fill=ORANGE if i < 4 else GREEN)
        text(c, ["!" if i < 4 else "✓"], size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        tb(s, 1.15, 1.08 + i * 0.78, 8.3, 0.6, [[(h + ": ", {"bold": True, "color": OLIVE}), (t, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    tb(s, 0.5, 5.0, 8, 0.3, ["Heute kein Brenner, kein Spiritus, keine Steckdose."], size=12, color=MUTED, after=0)

    s = d.slide("Nur Titel", "So läuft es heute")
    for i, st in enumerate("ABCD"):
        x = 0.7 + i * 1.55
        c = box(s, x, 1.25, 1.15, 1.15, fill=[TEAL, TEAL, OLIVE, OLIVE][i] if False else OLIVE)
        text(c, [[("Station", {"size": 13})], [(st, {"size": 32, "bold": True})]], color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        if i < 3: tb(s, x + 1.15, 1.55, 0.4, 0.5, ["→"], size=24, bold=True, color=ORANGE, align=PP_ALIGN.CENTER, after=0)
    tb(s, 0.7, 2.45, 5.8, 0.4, ["7 Minuten pro Station · Wechsel auf das Zeichen"], size=18, bold=True, color=ORANGE, after=0)
    c = box(s, 6.85, 1.15, 2.65, 1.75, fill=CARD)
    text(c, [[("Vertiefung", {"bold": True, "color": OLIVE})], "Aufgaben 11–17 jeder am eigenen Handy"], size=18)
    tb(s, 0.5, 3.0, 9, 2.0, [
        [("Stunde 3: ", {"bold": True, "color": OLIVE}), ("Gruppe 1–4 an den Stationen · Gruppe 5–8 Vertiefung", {})],
        [("Stunde 4: ", {"bold": True, "color": OLIVE}), ("umgekehrt", {})],
        [("Start: ", {"bold": True, "color": OLIVE}), ("Gruppe 1 + 5 bei A · 2 + 6 bei B · 3 + 7 bei C · 4 + 8 bei D", {})],
        [("Ein Protokoll-Handy pro Gruppe", {"bold": True})]], size=20, after=4)

    for st, vid, tit, steps, (pl, kind), warn in STATIONEN:
        v = V[vid]
        nr = v["nr"] if st != "D" else "K1.2-V2 + V3"
        s = d.slide("Nur Titel", f"Station {st} · {nr}")
        tag(s, f"KERN · Station {st}", OLIVE, x=7.55, w=1.95)
        tb(s, 0.5, 0.98, 7.0, 0.45, [tit], size=22, bold=True, color=OLIVE, after=0)
        s.shapes.add_picture(IMG + f"d_{v['skizze']}.png", Inches(0.5), Inches(1.5), Inches(4.4), Inches(4.4 * 629 / 900))
        for i, t in enumerate(steps):
            c = box(s, 5.15, 1.5 + i * 0.68, 0.45, 0.45, fill=TEAL)
            text(c, [str(i + 1)], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
            tb(s, 5.7, 1.43 + i * 0.68, 3.85, 0.62, [t], size=18, anchor=MSO_ANCHOR.MIDDLE, after=0)
        w = box(s, 5.15, 4.28, 4.35, 0.62, fill=RGBColor.from_string("FBE9DC"))
        text(w, [[("! ", {"bold": True, "color": ORANGE}), (warn, {})]], size=15, anchor=MSO_ANCHOR.MIDDLE, after=0)
        pill(s, pl, x=0.5, w=2.4, kind=kind)
        tb(s, 3.0, 5.17, 3.6, 0.3, ["Alle Schritte und Fragen: Versuchsseite"], size=11, color=MUTED, after=0)

    s = d.slide("Nur Titel", "Vertiefung · Aufgaben im Überblick")
    rows = [["Nr.", "Aufgabe", "Was ist zu tun?", "Quelle"],
            ["11", "Bindungs-Detektiv", "Bindungsart antippen", "S. 20–25"],
            ["12", "Vorhersage Demo", "leitet / leitet nicht tippen", "S. 21, 25"],
            ["13", "Mythos oder Fakt?", "entscheiden + 1 Satz", "S. 21–27"],
            ["14", "Wasser-Molekül", "Lücken füllen", "S. 23, 26"],
            ["15", "So entsteht Kochsalz", "Schritte ordnen", "S. 20–21"],
            ["16", "Elektrolytgetränk", "2 kurze Antworten", "S. 21"],
            ["17", "Science oder Fiction? (OPTIONAL)", "Film-Check", "S. 26"]]
    grid(s, 0.5, 1.1, [0.7, 3.5, 3.1, 1.7], rows, 0.48, size=16)

    # ================= Stunde 4
    d.section("Stunde 4")
    s = d.slide("Nur Titel", "Demo K1.2-V4 · Leitet es den Strom?")
    tag(s, "KERN · Demo", OLIVE, x=7.55, w=1.95)
    tb(s, 0.5, 0.98, 7.0, 0.4, ["Erst tippen (Hand heben), dann messen!"], size=18, bold=True, color=ORANGE, after=0)
    rows = [["Probe", "Dein Tipp", "Ergebnis", "Warum?"]] + [[a, "", b if L else "", c if L else ""] for a, b, c in D.DEMO]
    cells = grid(s, 0.5, 1.45, [2.2, 1.2, 1.85, 3.75], rows, [0.4] + [0.52] * 6, size=16)
    if L:
        for r in cells[1:]: s._anim.append([r[2], r[3]])
    pill(s, "Buch S. 21, 25", x=0.5, w=2.2)

    for half in (D.BESPRECHUNG[:2], D.BESPRECHUNG[2:]):
        s = d.slide("Nur Titel", "Was habt ihr herausgefunden?")
        for i, (st, q, a) in enumerate(half):
            x = 0.5 + i * 4.6
            tb(s, x, 1.0, 4.4, 0.4, [st], size=16, bold=True, color=TEAL, after=0)
            tb(s, x, 1.35, 4.4, 1.1, [KURZFRAGE.get(st, q)], size=20, bold=True, color=INK, after=0)
            short = a if len(a) < 230 else a[:a.rfind(".", 0, 230) + 1]
            answer(s, x, 2.55, 4.4, 2.45, "Lösung", short, L)
        pill(s, "Buch S. 21–27", x=0.5, w=2.2)

    s = d.slide("Nur Titel", "1. Test · Donnerstag, 5. November")
    tb(s, 0.5, 1.0, 6.2, 4.0, [
        [("Stoff: ", {"bold": True, "color": OLIVE}), ("Schulbuch S. 10–28 und alle Unterlagen", {})],
        "• Atombau und Periodensystem (Wiederholung)",
        "• Atommodelle von Dalton bis Bohr, Orbitalmodell",
        "• Ionen-, Atom- und Metallbindung",
        "• Wasserstoffbrücken, Van-der-Waals-Kräfte",
        [("Teststoff-Blatt: ", {"bold": True, "color": OLIVE}), ("in LMS", {})]], size=20, after=6)
    s.shapes.add_picture(qr_png(D.URL["training"]), Inches(7.05), Inches(1.15), Inches(2.1), Inches(2.1))
    tb(s, 6.8, 3.3, 2.6, 0.4, ["Test-Training"], size=18, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)
    tb(s, 6.6, 3.7, 3.0, 0.6, [D.URL["training"]], size=8, color=MUTED, align=PP_ALIGN.CENTER, after=0)

    # ================= Abgabe
    d.section("Abgabe")
    abgabe()

    name = f"{PRE}-VS-{'Lehrer' if L else 'Kadetten'}-{VER}.pptx"
    d.save(os.path.join(OUT, name))
    print("ok", name, len(d.p.slides._sldIdLst), "Folien")


build(True)
build(False)
