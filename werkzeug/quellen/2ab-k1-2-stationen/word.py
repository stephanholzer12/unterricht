"""Word-Dateien für die Stationen K1.2 (2AB, Stunde 3–4): VS-AB, VS-Loesung, VS-Fachwissen.
Aufruf im Repository-Ordner: python3 werkzeug/quellen/2ab-k1-2-stationen/word.py <Zielordner>
Layout wie Muster-VS-AB (Kopf mit Wappen, Titelzeile mit olivgrüner Linie), Icons aus vs-img/, Daten aus versuche.html + daten.py."""
import os, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import daten as D

OUT = sys.argv[1]
PRE = "26-27-2AB-K1.2b-Station"
VER = "v1"
OLIVE, TEAL, RED, MUTED, RUST, INK = RGBColor(0x3B, 0x4A, 0x2F), RGBColor(0x2E, 0x7D, 0x8C), RGBColor(0xC0, 0x00, 0x00), RGBColor(0x5E, 0x66, 0x52), RGBColor(0x8C, 0x3B, 0x2E), RGBColor(0x1F, 0x2A, 0x1C)
LOGO = os.path.join(D.ROOT, "werkzeug/logo.png")
IMG = D.UNIT + "vs-img/"
V = {v["id"]: v for v in D.versuche()}
ORDER = D.KERN_IDS


# ---------------------------------------------------------------- Helfer
def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcol); tcPr.append(s)


def borders(table, color="B9C4A6", size=6):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:color"), color); b.append(el)
    tblPr.append(b)


def noborders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblPr.append(b)


def pbottom(p, color="3B4A2F", sz=12):
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "1"); b.set(qn("w:color"), color); bd.append(b); pPr.append(bd)


def para(doc_or_cell, text="", size=10.5, bold=False, color=INK, italic=False, after=2, before=0, align=None, keep=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.05
    if keep: pf.keep_with_next = True
    if align: p.alignment = align
    if text: run(p, text, size, bold, color, italic)
    return p


def run(p, text, size=10.5, bold=False, color=INK, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color; r.font.name = "Calibri"
    return r


def label(p, lab, text, size=10.5, color=INK, labcol=OLIVE):
    run(p, lab, size, True, labcol); run(p, text, size, False, color)


def widths(table, cms):
    for row in table.rows:
        for c, w in zip(row.cells, cms): c.width = Cm(w)


def cellp(cell, text="", size=10, bold=False, color=INK, align=None):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    if align: p.alignment = align
    if text: run(p, text, size, bold, color)
    return p


def lines(doc, n):
    for _ in range(n):
        p = para(doc, "", after=0); p.paragraph_format.line_spacing = 1.0; pbottom(p, "B9C4A6", 4)
        p.paragraph_format.space_before = Pt(7)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_doc(kopf, titel, fuss):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.4); sec.footer_distance = Cm(0.6)
    t = doc.add_table(rows=1, cols=2); noborders(t); widths(t, [1.6, 15.8])
    t.cell(0, 0).paragraphs[0].add_run().add_picture(LOGO, width=Cm(1.2))
    c = t.cell(0, 1); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cellp(c, kopf, 8.5, True, TEAL); p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); run(p2, titel, 16, True, OLIVE)
    p = para(doc, "", after=4); pbottom(p, "3B4A2F", 14)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(fp, fuss + " · Seite ", 8, False, MUTED)
    r = fp.add_run(); r.font.size = Pt(8); r.font.color.rgb = MUTED
    for tag, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tag); r._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
    return doc


def box(doc, title, rows_fn, fill="EEF3E4", line="B9C4A6"):
    t = doc.add_table(rows=1, cols=1); borders(t, line, 8); widths(t, [17.4]); c = t.cell(0, 0); shade(c, fill)
    cellp(c, title, 10, True, OLIVE); rows_fn(c)
    para(doc, "", after=2)
    return t


def table(doc, head, rows, cms, size=9.5, headfill="3B4A2F", zebra=True):
    t = doc.add_table(rows=1, cols=len(head)); borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; shade(c, headfill); cellp(c, h, size, True, RGBColor(0xFF, 0xFF, 0xFF))
    for k, r in enumerate(rows):
        cells = t.add_row().cells
        for i, x in enumerate(r):
            if isinstance(x, tuple): cellp(cells[i], x[0], size, x[1] if len(x) > 1 else False, x[2] if len(x) > 2 else INK)
            else: cellp(cells[i], str(x), size)
            if zebra and k % 2 == 1: shade(cells[i], "F6F9F0")
    widths(t, cms)
    return t


def fahrplan_box(doc):
    def fill(c):
        p = cellp(c, "", 9.5); p = c.add_paragraph(); run(p, "Dafür sind 4 Stunden eingeplant. ", 9.5, True, OLIVE); run(p, D.FAHRPLAN["text"], 9.5)
        it = c.add_table(rows=1, cols=4); borders(it, "D3DCC2", 4)
        for i, h in enumerate(["Wann", "Das machen wir", "Aufgaben", "Abgabe"]): cellp(it.rows[0].cells[i], h, 8.5, True, OLIVE)
        for z in D.FAHRPLAN["zeilen"]:
            cs = it.add_row().cells
            cellp(cs[0], z["wann"], 8.5, True); cellp(cs[1], z["lesen"] + (" · " + z["wie"] if z.get("wie") else ""), 8.5); cellp(cs[2], z["aufgaben"], 8.5); cellp(cs[3], z["abgabe"], 8.5)
        widths(it, [2.2, 9.2, 3.0, 2.6])
    box(doc, "Dein Fahrplan", fill)


def versuch_titel(doc, v):
    p = para(doc, "", after=3, keep=True); p.paragraph_format.page_break_before = True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    run(p, v["nr"] + "  ", 13, True, OLIVE); run(p, v["titel"].upper(), 13, True, OLIVE)
    tag = ("KERN · Station " + v["station"]) if v.get("station") else ("KERN · Lehrer-Demo" if v.get("demo") else "ZUSATZ")
    run(p, "\t" + tag, 9, True, TEAL); pbottom(p, "B9C4A6", 6)


def mat_table(doc, v):
    para(doc, "Geräte & Chemikalien", 10.5, True, OLIVE, after=1, keep=True)
    m = v["mat"]; n = (len(m) + 2) // 3
    t = doc.add_table(rows=n, cols=3); borders(t, "B9C4A6", 6)
    for i, (txt, ic) in enumerate(m):
        c = t.cell(i // 3, i % 3); p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        f = IMG + f"i_{ic}.png"
        if os.path.exists(f): p.add_run().add_picture(f, height=Cm(0.6))
        run(p, "  " + txt, 9)
    widths(t, [5.8, 5.8, 5.8])


def beob_table(doc, v, sol=None):
    b = v["beob"]; para(doc, b["titel"], 10.5, True, OLIVE, after=1, before=3, keep=True)
    ng = max(len(z) for z in b["zeilen"])
    rows = []
    for i, z in enumerate(b["zeilen"]):
        r = list(z) + [""] * (len(b["kopf"]) - len(z))
        if sol: r[-1] = (sol[i], False, RED)
        rows.append(r)
    nc = len(b["kopf"]); first = 4.2 if nc > 2 else 5.0
    cms = [first] + [(17.4 - first) / (nc - 1)] * (nc - 1) if nc == 2 or ng == 1 else [1.4, 4.4] + [(17.4 - 5.8) / (nc - 2)] * (nc - 2)
    t = table(doc, b["kopf"], rows, cms, 9.5, "6E7F4E", zebra=False)
    if not sol:
        for row in t.rows[1:]:
            row.height = Cm(0.55)


def versuch_seite(doc, v, loesung=False):
    L = D.LOESUNG[v["id"]]
    versuch_titel(doc, v)
    para(doc, v["intro"], 10.5, False, INK, italic=True, after=2)
    if v.get("sicher"):
        p = para(doc, "", after=2); label(p, "Sicherheit: ", v["sicher"], 9.5, INK, RUST)
    p = para(doc, "", after=3)
    label(p, "Versuch im Schulbuch: " if v["buch"].get("versuch") else "Theorie im Schulbuch: ", v["buch"]["kap"], 9.5, INK, TEAL)
    mat_table(doc, v)
    if not loesung:  # Lösungs-Word ohne Entsorgung (steht im VS-AB und im VS-Fachwissen), damit jeder Versuch auf eine Seite passt
        p = para(doc, "", after=2, before=3); label(p, "Entsorgung und Reinigung: ", v["entsorgung"], 9)
    para(doc, "Durchführung", 10.5, True, OLIVE, after=1, keep=True)
    t = doc.add_table(rows=1, cols=2); noborders(t); widths(t, [10.6, 6.8])
    c = t.cell(0, 0); first = True
    for i, s in enumerate(v["schritte"], 1):
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.5)
        run(p, f"{i}  ", 9.5, True, OLIVE); run(p, s, 9.5)
    pc = t.cell(0, 1).paragraphs[0]; pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.add_run().add_picture(IMG + f"d_{v['skizze']}.png", width=Cm(4.8 if loesung else 6.4))
    beob_table(doc, v, L["beob"] if loesung else None)
    for key, tit in [("erk", "Erkenntnis"), ("interp", "Interpretation"), ("schluss", "Schlussfolgerung")]:
        para(doc, tit, 10.5, True, OLIVE, after=0, before=3, keep=True)
        para(doc, v[key], 9.5, False, MUTED, after=1, keep=True)
        if loesung: para(doc, L[key], 9.5, False, RED, after=1)
        else: lines(doc, 2)
    if loesung:
        inf = L["info"]
        def fill(cc):
            f, vb, tp, fe, q = inf
            for lab, x in [("Form & Dauer: ", f + " · Quelle: " + q.replace("Idee: ", "")), ("Vorbereitung: ", vb),
                           ("Tipps: ", tp + " Typische Fehler: " + fe.replace("Fehler: ", ""))]:
                p = cc.add_paragraph(); p.paragraph_format.space_after = Pt(0); label(p, lab, x, 8.5)
        box(doc, "LEHRER-INFO " + v["nr"], fill, "FBF1E6", "E0B48A")


def lehrplan_box(doc):
    def fill(c):
        for lab, x in [("Einheit: ", "Jahresplan E3–E5 (Ionen-, Atom-, Metallbindung, zwischenmolekulare Kräfte) · 6 Std. laut Jahresplan, hier Stunde 3–4 der Einheit K1.2"),
                       ("K02 ", "die Entstehung von chemischen Bindungen erklären – Stationen A–D, Besprechung, Vertiefung 11, 14, 15"),
                       ("K03 ", "den Zusammenhang von chemischer Bindung und Stoffeigenschaften erkennen – alle Stationen, Demo K1.2-V4, Vertiefung 12, 13, 16"),
                       ("K07 ", "einfache Experimente durchführen und dokumentieren – Versuchsprotokoll am Handy (PDF in LMS)"),
                       ("N01 · SDG 3 ", "Elektrolyte im Körper (Vertiefung 16), Umweltschutz Tenside (Station A)"),
                       ("Abweichung: ", "Versuche K1.2-V2 und K1.2-V8 vereinfacht (ohne Trocknen, ohne Brennspiritus und Paraffinöl), K1.2-V4 als Lehrer-Demo, weil nur 1 Multimeter vorhanden ist.")]:
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1); label(p, lab, x, 9)
    box(doc, "Lehrplanbezug", fill)


def stationsplan(doc, size=9.5):
    para(doc, "So läuft es: Stationenbetrieb in Stunde 3–4", 11, True, OLIVE, after=2, before=2)
    rows = [["Station A", "K1.2-V9 Wasser mit Haut", "Waschbecken-nah, Petrischale", "Gruppe 1 / 5"],
            ["Station B", "K1.2-V7 Der scheue Wasserstrahl", "am Waschbecken", "Gruppe 2 / 6"],
            ["Station C", "K1.2-V8 Wer mischt sich mit wem?", "Schutzbrille!", "Gruppe 3 / 7"],
            ["Station D", "K1.2-V2 Würfel im Salzstreuer + K1.2-V3 Spröde oder verformbar?", "Mörser, Lupe/Mikroskop", "Gruppe 4 / 8"],
            ["Lehrer-Demo", "K1.2-V4 Piepst es? Der Multimeter-Check", "Stunde 4, Plenum", "alle"]]
    table(doc, ["Station", "Versuch", "Wo / Hinweis", "Start"], rows, [2.4, 8.2, 4.0, 2.8], size)
    p = para(doc, "", after=2, before=3)
    run(p, "8 Gruppen à 3 (eine zu viert). ", size, True); run(p, "Gruppe 1–4: Stunde 3 Stationen, Stunde 4 Vertiefung – Gruppe 5–8 umgekehrt. Pro Station 7 Minuten, dann weiter A → B → C → D → A. Wer wartet oder früher fertig ist: Vertiefung (Aufgaben 11–17) am eigenen Handy. Protokoll am Handy: ", size)
    run(p, D.URL["versuche"], size - 1, False, TEAL)


# ---------------------------------------------------------------- VS-AB
fuss = f"K1.2 Bindungsarten · Stationen · Stand {D.STAND} · {VER}"
doc = new_doc("VERSUCHE · NAWI · 2. JAHRGANG · K1.2", "Die Bindungsarten – Stationen", fuss)
fahrplan_box(doc)
stationsplan(doc)
p = para(doc, "", after=2, before=4)
run(p, "So wertest du jeden Versuch aus: ", 10, True, OLIVE); run(p, "Erkenntnis – Was habt ihr herausgefunden? · Interpretation – Warum ist das so? Erklärt mit den Teilchen. · Schlussfolgerung – Was bedeutet das für Alltag, Sport oder Bundesheer?", 10)
p = para(doc, "", after=2); run(p, "Sicherheit im NAWI-Saal: ", 10, True, RUST); run(p, "Schutzbrille an Station C und D · nichts kosten · Verschüttetes sofort aufwischen · Öl in den Sammelbehälter, nicht in den Ausguss · beim Arbeiten Handy weglegen, Fotos nur aus sicherem Abstand.", 10)
for vid in ORDER:
    versuch_seite(doc, V[vid])
f_ab = f"{PRE}-VS-AB-{VER}.docx"; doc.save(os.path.join(OUT, f_ab))

# ---------------------------------------------------------------- VS-Loesung
doc = new_doc("LÖSUNGEN VERSUCHE · NAWI · 2. JAHRGANG · K1.2", "Die Bindungsarten – Stationen: Lösungen & Lehrer-Infos", fuss + " · Lösungen")
lehrplan_box(doc)
fahrplan_box(doc)
stationsplan(doc)
para(doc, "Die Lehrer-Infos stehen auch als ausgeblendete Folien in der Stationen-PPP (Lehrer). Die Zusatz-Versuche K1.2-V1, V5, V6, V10, V11 bleiben auf der Versuchsseite – heute nicht dran (Brenner, Hitze, Spiritus).", 9, False, MUTED, italic=True, before=3)
for vid in ORDER:
    versuch_seite(doc, V[vid], loesung=True)
f_lo = f"{PRE}-VS-Loesung-{VER}.docx"; doc.save(os.path.join(OUT, f_lo))

# ---------------------------------------------------------------- VS-Fachwissen
doc = new_doc("NUR FÜR DICH · NAWI · 2. JAHRGANG · K1.2", "Stationen Bindungsarten – Fachwissen & Herrichten", fuss + " · Fachwissen")
para(doc, "Herrichte-Liste (1 Satz pro Station, 8 Gruppen wandern durch; Mengen gelten für beide Stunden)", 11, True, OLIVE, after=2)
rows = []
for r in D.HERRICHTEN:
    ic, name, inv, a, b, c, d, demo, bestand, hint = r
    rows.append([("", ), name, inv, str(a), str(b), str(c), str(d), str(demo), bestand, (hint, False, RED if "fehlt" in hint else INK)])
t = table(doc, ["", "Gerät / Chemikalie", "Inventar / Ort", "A", "B", "C", "D", "Demo", "Bestand", "Hinweis"], rows,
          [0.9, 3.6, 2.9, 0.75, 0.75, 1.05, 0.75, 1.1, 1.6, 4.0], 7.5)
for i, r in enumerate(D.HERRICHTEN):
    f = IMG + f"i_{r[0]}.png"
    if os.path.exists(f):
        p = t.rows[i + 1].cells[0].paragraphs[0]; p.add_run().add_picture(f, height=Cm(0.42))
for row in t.rows: row.height = Cm(0.45)
para(doc, "Fehlmenge: Büroklammern (nicht im Inventar) – 2 Stück aus dem Büro. Münzen bringen die Kadetten mit. Alles andere ist laut Inventar vorhanden.", 9, True, RED, before=3)
para(doc, "Checkliste »Vor der Stunde« (ca. 15 Minuten, alles wiederverwendbar)", 11, True, OLIVE, after=1, before=4)
for x in ["Vier Aufbewahrungswannen beschriften (A–D) und das Material laut Liste einräumen – nach der Stunde Deckel drauf, für das nächste Jahr aufheben.",
          "Station A und B möglichst nahe am Waschbecken; Station C mit Sammelbehälter »organische Reste« und Sammelbecher für Reagenzgläser.",
          "Parafilm-Streifen schneiden, Waschmittel in eine Tropfflasche, blaues Wasser (400-mL-Becherglas) ansetzen.",
          "Multimeter testen (Batterie, Bereich 200 kΩ), drei Bechergläser für die Demo mit dest. Wasser, Salz- und Zuckerlösung.",
          "QR-Codes Versuchsseite und Vertiefung am Beamer (Folie »Dein Lernpfad«); Teststoff-Datei bereitlegen (LMS)."]:
    p = para(doc, "", after=1); run(p, "☐  ", 10, True, OLIVE); run(p, x, 9.5)

page_break(doc)
para(doc, "Stundenkarten", 13, True, OLIVE, after=2)
for tit, steps in [("Stunde 3 · Stationen Teil 1", ["Start (5 min): Sicherheit, Gruppen 1–8 einteilen, QR-Code Versuchsseite (ein Protokoll-Handy pro Gruppe) und Vertiefung.",
                                                     "Gruppe 1–4 an den Stationen A–D, 4 Runden à 7 Minuten, Wechsel auf dein Zeichen (A → B → C → D).",
                                                     "Gruppe 5–8 arbeiten einzeln an der Vertiefung (Aufgaben 11–16, Prüfen-Knopf).",
                                                     "Letzte 5 Minuten: Stationen aufräumen und für Stunde 4 neu bestücken (Reagenzgläser, Filterpapier)."]),
                   ("Stunde 4 · Demo, Stationen Teil 2, Besprechung, Teststoff", ["Demo K1.2-V4 Multimeter gleich zu Beginn (ca. 6 min, ohne Handy): erst Vorhersage per Handzeichen, dann messen, Ergebnisse ins Protokoll.",
                                                     "Wechsel: Gruppe 5–8 an die Stationen, Gruppe 1–4 Vertiefung – wieder 4 Runden à 7 Minuten.",
                                                     "Besprechung (ca. 8 min, ohne Handy): je Station eine Frage (Folien »Was habt ihr herausgefunden?«).",
                                                     "Teststoff 05.11. austeilen bzw. in LMS (2 min) + QR-Code Test-Training; Abgabe der PDFs bis zur nächsten NAWI-Stunde."])]:
    def fill(c, steps=steps):
        for i, s in enumerate(steps, 1):
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1); run(p, f"{i}  ", 9.5, True, OLIVE); run(p, s, 9.5)
        p = c.add_paragraph(); label(p, "Phase ohne Handy: ", "Versuche an den Stationen (Handy nur zum Protokollieren); in Stunde 4 Demo und Besprechung im Plenum.", 9)
        p = c.add_paragraph(); label(p, "Links: ", D.URL["versuche"] + " · " + D.URL["vertiefung"], 8.5)
    box(doc, tit, fill)

para(doc, "Die Stationen auf einen Blick", 11, True, OLIVE, after=2, before=2)
table(doc, ["Station", "Kernaussage", "Material (Wanne)"], [
    ["A · K1.2-V9", "Wasserstoffbrücken → Oberflächenspannung; Tenside schwächen sie", "Petrischale, Pipette, Büroklammer, Waschmittel"],
    ["B · K1.2-V7", "Wasser ist ein Dipol – geladener Ballon zieht den Strahl an", "Ballon, Reibetuch, Waschbecken"],
    ["C · K1.2-V8", "Ähnliches löst sich in Ähnlichem; Ionen nur in polarem Wasser", "Reagenzgläser, Glycerin, Olivenöl, Salz, Waschmittel"],
    ["D · K1.2-V2/V3", "Ionengitter: Würfelform und spröde; Metall: verformbar (Elektronengas)", "Lupe/Mikroskop, Salze, Mörser, Zinn, Kupfer"],
    ["Demo · K1.2-V4", "Strom braucht frei bewegliche Ladungen (Elektronen oder Ionen)", "Multimeter, 3 Bechergläser"]], [3.0, 8.6, 5.8], 9)

page_break(doc)
para(doc, "Was hinter den Versuchen steckt", 13, True, OLIVE, after=2)
HINTER = {
 "haut": "Wasser hat mit rund 72 mN/m eine der höchsten Oberflächenspannungen aller Flüssigkeiten. Ursache: Jedes Molekül kann bis zu vier Wasserstoffbrücken bilden; an der Oberfläche fehlen die Nachbarn nach oben, die Moleküle werden nach innen gezogen. Tenside reichern sich an der Oberfläche an und senken die Spannung auf etwa ein Drittel – die Klammer sinkt, der »Wasserberg« auf der Münze bricht früher.",
 "wasserstrahl": "Nur für dich: Streng genommen wirkt neben der Ausrichtung der Dipole auch Influenz (Ladungsverschiebung im leitfähigen Leitungswasser). Für die 2. Klasse reicht die Dipol-Erklärung des Schulbuchs (S. 27). Gegenprobe: Ein dünner Öl- oder Benzinstrahl wird kaum abgelenkt.",
 "mischen": "Glycerin (Propan-1,2,3-triol) hat drei OH-Gruppen und bildet viele Wasserstoffbrücken – es mischt sich in jedem Verhältnis mit Wasser. Olivenöl (Triglyceride, lange Kohlenwasserstoffketten) ist unpolar. Ionen werden in Wasser hydratisiert: Die Dipole umhüllen Na⁺ mit der O-Seite, Cl⁻ mit der H-Seite; im Öl fehlt diese Anziehung. Tenside bilden Mizellen um Öltröpfchen (Emulsion).",
 "kristall": "Im NaCl-Gitter ist jedes Na⁺ von 6 Cl⁻ umgeben und umgekehrt (kubisch-flächenzentriert). Die Würfelform der Körner zeigt diese Ordnung direkt. Kaliumalaun kristallisiert oktaedrisch (Achtflächner), Zucker (Saccharose) ist ein Molekülkristall mit länglicher (monokliner) Form.",
 "sproede": "Spaltbarkeit von Salzen: Eine Verschiebung um eine halbe Gitterkonstante bringt gleich geladene Ionen nebeneinander, die Abstoßung spaltet den Kristall. Metalle verformen sich plastisch über Gleitebenen und Versetzungen; das Elektronengas ist richtungslos, die Bindung bleibt erhalten.",
 "multimeter": "Das Multimeter misst im Widerstandsbereich mit wenigen Volt aus der Batterie – völlig ungefährlich. Destilliertes Wasser zeigt meist »OL« oder einige MΩ (Eigendissoziation, gelöstes CO₂). Leitungswasser leitet schon etwas (gelöste Mineralsalze) – schöner Nebeneffekt für die Frage »Warum ist Wasser in der Steckdose gefährlich?«.",
}
for vid in ORDER:
    v = V[vid]; L = D.LOESUNG[vid]
    para(doc, v["nr"] + " " + v["titel"] + (" · Station " + v["station"] if v.get("station") else " · Lehrer-Demo"), 10.5, True, OLIVE, after=1, before=3, keep=True)
    para(doc, HINTER[vid], 9.5, after=1)
    p = para(doc, "", after=1); label(p, "Erwartet: ", L["erk"], 9.5)
    p = para(doc, "", after=1); label(p, "Ablauf & Wegräumen: ", L["info"][1] + " " + L["info"][2], 9)
    p = para(doc, "", after=1); label(p, "Typische Fehler: ", L["info"][3], 9)

para(doc, "Die 5 größten Stolpersteine", 11, True, OLIVE, after=2, before=5, keep=True)
table(doc, ["Fehlvorstellung", "Richtig ist (so zeigst du es)"], [[a, b] for a, b in D.STOLPERSTEINE], [6.0, 11.4], 9)
para(doc, "Besprechung in Stunde 4 (ca. 8 Minuten)", 11, True, OLIVE, after=2, before=5, keep=True)
table(doc, ["Station", "Frage an die Klasse", "Erwartete Antwort"], [[a, b, c] for a, b, c in D.BESPRECHUNG], [3.0, 5.6, 8.8], 9)
para(doc, "Quellen und geprüfte Bücher", 11, True, OLIVE, after=2, before=5, keep=True)
for x in ["Jahresplan v33, Blatt »Versuche«: V2.5/V2.7 (Leitfähigkeit) → Demo K1.2-V4; V2.6 (Kristalle) → K1.2-V2 vereinfacht; V2.8 (Oberflächenspannung: Münze, Büroklammer, Spülmittel) → K1.2-V9.",
          "Trauner Chemie 2. Klasse S. 19–28 (Theorie); »Versuch – Der scheue Wasserstrahl« S. 27 → K1.2-V7.",
          "Naturwissenschaften HAK 1 (Trauner): Beobachtung »Oberflächenspannung« mit der Münze → Schritt 1 an Station A.",
          "Chemie 2 (Westermann) E 14.01 Tenside mit Milch und Lebensmittelfarbe – nicht übernommen (Milch nicht im Inventar), schöne Variante für später.",
          "Chemie 1 (Westermann) S. 66 (Löslichkeit polar/unpolar) → Interpretation Station C; Expedition Chemie 4 Praxisteil V2.2 (Leitfähigkeit von Lösungen) → Demo.",
          "Geprüft ohne Übernahme: El-Mo-Experimente, Chemie muss stimmen 7–8, Treffpunkt Chemie Arbeitsheft (keine einfacheren Varianten mit Inventar-Material)."]:
    p = para(doc, "", after=1); run(p, "·  ", 9.5, True, OLIVE); run(p, x, 9)
f_fw = f"{PRE}-VS-Fachwissen-{VER}.docx"; doc.save(os.path.join(OUT, f_fw))
print(f_ab, f_lo, f_fw)
