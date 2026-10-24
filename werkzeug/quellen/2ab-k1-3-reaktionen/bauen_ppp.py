"""Präsentationen NAWI 2AB · K1.3 Stoffumwandlung mit dem BHAK-Folienmaster: TH-Lehrer, TH-Kadetten, VS-Lehrer, VS-Kadetten.
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/2ab-k1-3-reaktionen/bauen_ppp.py <BHAK_Folienmaster_Vorlage.pptx> <Zielordner>
(.potx vorher in .pptx umwandeln: Content-Type »template.main« → »presentation.main«.)
Theorie-Folien kommen aus theorie_daten.py – dieselbe Quelle wie die Theorie-Seite (GitHub und PPP gleich).
Lehrer: Lösungen per Klick, ausgeblendete Lehrer-Infos. Kadetten: gleiche Folien, leere Antwortfelder, keine Lehrer-Infos."""
import copy, io, os, re, sys, uuid
import qrcode
from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, HERE)
import inhalt as I
import theorie_daten as TD
import versuche_daten as VD
import herrichten as H

MASTER, OUT = sys.argv[1], sys.argv[2]
UNIT = os.path.join(ROOT, I.ORDNER) + "/"
VER = "v5"
OLIVE, INK, CARD, GREEN, ORANGE, TEAL, RUST, MUTED, WHITE, PALE = (RGBColor.from_string(x) for x in
    ["3B4A2F", "1F2A1C", "EEF3E4", "7FA83A", "D9772B", "2E7D8C", "8C3B2E", "5E6652", "FFFFFF", "F7F9F2"])
RED = RGBColor.from_string("A33A2B")
GREY = RGBColor.from_string("DCDFD6")
FONT_R = ImageFont.truetype("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf", 100)
FONT_B = ImageFont.truetype("/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf", 100)
URL = {k: I.URL + k + ".html" for k in ("theorie", "aufgaben", "versuche", "quiz")}
X0, X1, Y0, Y1 = 0.5, 9.5, 1.12, 5.08


# ================================================================ Messen
def plain(t):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", t)


def lines_needed(t, size, w_in, bold=False):
    """Zeilen, die ein Text bei Schriftgröße size (pt) in einer Box der Breite w_in (Zoll) braucht (Calibri ≈ Carlito)."""
    font = FONT_B if bold else FONT_R
    maxw = (w_in - 0.22) * 72 / size * 100 * 0.97
    n = 0
    for para in plain(t).split("\n"):
        words, cur = para.split(" "), ""
        n += 1
        for wd in words:
            test = (cur + " " + wd).strip()
            if font.getlength(test) > maxw and cur:
                n += 1; cur = wd
            else:
                cur = test
    return n


def th(t, size, w_in, bold=False, ls=1.17, pad=0.1):
    return lines_needed(t, size, w_in, bold) * size * ls / 72 + pad


# ================================================================ Deck
def qr_png(url):
    q = qrcode.QRCode(border=1, box_size=10); q.add_data(url); q.make(fit=True)
    b = io.BytesIO(); q.make_image(fill_color="#1F2A1C", back_color="white").save(b, "PNG"); b.seek(0); return b


class Deck:
    def __init__(self, lehrer, fuss):
        self.lehrer = lehrer
        self.p = Presentation(MASTER)
        lst = self.p.slides._sldIdLst
        self.pill_proto = None
        for sh in self.p.slides[2].shapes:
            if sh.name == "Buchseite": self.pill_proto = copy.deepcopy(sh._element)
        for sid in list(lst):
            self.p.part.drop_rel(sid.get(qn("r:id"))); lst.remove(sid)
        self.L = {l.name: l for l in self.p.slide_layouts}
        self.sections = []
        for part in self.p.part.package.iter_parts():
            if "slideLayout" in str(part.partname) or "slideMaster" in str(part.partname):
                for t in part._element.iter(qn("a:t")):
                    if t.text and "Fach · Jahrgang · Thema" in t.text: t.text = fuss

    def section(self, name):
        self.sections.append((name, []))

    def slide(self, layout, title, hidden=False):
        s = self.p.slides.add_slide(self.L[layout])
        for ph in s.placeholders:
            if ph.placeholder_format.type == 1: ph.text_frame.text = title
        if hidden: s._element.set("show", "0")
        self.sections[-1][1].append(self.p.slides._sldIdLst[-1].get("id"))
        s._anim = []
        return s

    def save(self, path):
        for s in self.p.slides:
            if getattr(s, "_anim", None): add_clicks(s, s._anim)
        prs = self.p.part._element
        ext_lst = prs.find(qn("p:extLst"))
        if ext_lst is None: ext_lst = etree.SubElement(prs, qn("p:extLst"))
        P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        for old in ext_lst.findall(qn("p:ext")):
            if old.get("uri") == "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}": ext_lst.remove(old)
        ext = etree.SubElement(ext_lst, qn("p:ext")); ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        sl = etree.SubElement(ext, "{%s}sectionLst" % P14, nsmap={"p14": P14})
        for name, ids in self.sections:
            se = etree.SubElement(sl, "{%s}section" % P14); se.set("name", name); se.set("id", "{%s}" % str(uuid.uuid4()).upper())
            il = etree.SubElement(se, "{%s}sldIdLst" % P14)
            for i in ids: etree.SubElement(il, "{%s}sldId" % P14).set("id", i)
        self.p.save(path)


def add_clicks(slide, groups):
    nid = [2]
    def cid():
        nid[0] += 1; return str(nid[0])
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    clicks = ""
    for grp in groups:
        inner = ""
        for k, sh in enumerate(grp):
            spid = sh.shape_id
            inner += (f'<p:par><p:cTn id="{cid()}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{"clickEffect" if k == 0 else "withEffect"}">'
                      f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                      f'<p:set><p:cBhvr><p:cTn id="{cid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                      f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                      f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                      f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid()}" dur="400"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                      f'</p:childTnLst></p:cTn></p:par>')
        clicks += (f'<p:par><p:cTn id="{cid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{cid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'
                   f'</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{"".join(f"<p:bldP spid=\"{sh.shape_id}\" grpId=\"0\" animBg=\"1\"/>" for grp in groups for sh in grp)}</p:bldLst></p:timing>')
    el = etree.fromstring(xml)
    sld = slide._element
    ext = sld.find(qn("p:extLst"))
    if ext is not None: ext.addprevious(el)
    else: sld.append(el)


def box(s, x, y, w, h, fill=None, line=None, radius=True, lw=1.25):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: shp.adjustments[0] = 0.08
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    return shp


def runs_md(t):
    out, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*", t):
        if m.start() > pos: out.append((t[pos:m.start()], {}))
        out.append((m.group(1), {"bold": True})); pos = m.end()
    if pos < len(t): out.append((t[pos:], {}))
    return out or [("", {})]


def text(target, parts, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4):
    tf = target.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    first = True
    for para in parts:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(after)
        runs = para if isinstance(para, list) else runs_md(para)
        for t, o in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.color.rgb = o.get("color", color)
            if o.get("italic"): f.italic = True
    return target


def tb(s, x, y, w, h, parts, **kw):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.margin_left = t.text_frame.margin_right = Inches(0.05)
    return text(t, parts, **kw)


def pill(s, label, x=3.55, y=5.18, w=2.2, kind="theory"):
    el = copy.deepcopy(d.pill_proto)
    ids = [int(v) for v in s.shapes._spTree.xpath("//p:cNvPr/@id")]
    el.xpath(".//p:cNvPr")[0].set("id", str(max(ids + [1]) + 1))
    s.shapes._spTree.append(el)
    shp = s.shapes[-1]
    shp.left, shp.top, shp.width = Inches(x), Inches(y), Inches(w)
    col = {"theory": RGBColor.from_string("6E7F4E"), "book": TEAL, "versuch": RED}[kind]
    shp.fill.solid(); shp.fill.fore_color.rgb = col
    r = shp.text_frame.paragraphs[0].runs
    r[0].text = label
    for extra in r[1:]: extra.text = ""
    return shp


def kind_of(src):
    return "versuch" if src.startswith("Versuch im Buch") else ("book" if src.startswith("Buch-Aufgabe") else "theory")


def tag(s, label, color=OLIVE, x=7.6, y=0.95, w=1.9):
    t = box(s, x, y, w, 0.32, fill=color)
    text(t, [label], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return t


def grid(s, x, y, colw, rows, rowh, head=True, size=16, fills=None):
    cells, yy = [], y
    for ri, row in enumerate(rows):
        xx, rc = x, []
        h = rowh[ri] if isinstance(rowh, (list, tuple)) else rowh
        for ci, val in enumerate(row):
            hd = head and ri == 0
            fill = OLIVE if hd else (fills[ri] if fills else (CARD if ri % 2 else WHITE))
            c = box(s, xx, yy, colw[ci], h, fill=fill, line=RGBColor.from_string("B9C4A6"), radius=False, lw=0.75)
            text(c, [val] if isinstance(val, str) else val, size=size, color=WHITE if hd else INK, bold=hd, anchor=MSO_ANCHOR.MIDDLE, after=0)
            rc.append(c); xx += colw[ci]
        cells.append(rc); yy += h
    return cells


def auto_grid(s, x, y, colw, rows, size=16, head=True, minh=0.36, fills=None):
    """Tabelle mit gemessenen Zeilenhöhen. Gibt (Zellen, Gesamthöhe) zurück."""
    hs = [max(minh, max(th(str(c), size, w, bold=(head and ri == 0), pad=0.1) for c, w in zip(r, colw))) for ri, r in enumerate(rows)]
    return grid(s, x, y, colw, rows, hs, head=head, size=size, fills=fills), sum(hs)


def table_h(colw, rows, size=16, head=True, minh=0.36):
    return sum(max(minh, max(th(str(c), size, w, bold=(head and ri == 0), pad=0.1) for c, w in zip(r, colw))) for ri, r in enumerate(rows))


def answer(s, x, y, w, h, title, body, lehrer, hint="Deine Antwort – am Arbeitsblatt eintragen", size=20):
    if lehrer:
        b = box(s, x, y, w, h, fill=WHITE, line=ORANGE, lw=2)
        text(b, ([[(title, {"bold": True, "color": ORANGE, "size": 16})]] if title else []) + (body if isinstance(body, list) else [body]), size=size, after=2)
        s._anim.append([b])
    else:
        b = box(s, x, y, w, h, fill=PALE, line=RGBColor.from_string("B9C4A6"), lw=1.5)
        text(b, [[(hint, {"italic": True, "color": MUTED, "size": 16})]], size=16)
    return b


def ppt_qr(s, url, x, y, size, step, label, color=OLIVE, linksize=8):
    s.shapes.add_picture(qr_png(url), Inches(x), Inches(y), Inches(size), Inches(size))
    c = box(s, x + size / 2 - 0.2, y - 0.25, 0.4, 0.4, fill=color)
    text(c, [step], size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    tb(s, x - 0.3, y + size + 0.02, size + 0.6, 0.32, [label], size=16, bold=True, color=color, align=PP_ALIGN.CENTER, after=0)
    tb(s, x - 0.45, y + size + 0.34, size + 0.9, 0.5, [url], size=linksize, color=MUTED, align=PP_ALIGN.CENTER, after=0)


def picture(s, path, x, y, w=None, h=None):
    from PIL import Image
    im = Image.open(path); iw, ih = im.size
    if w and h:
        k = min(w / iw, h / ih); w2, h2 = iw * k, ih * k
        return s.shapes.add_picture(path, Inches(x + (w - w2) / 2), Inches(y), Inches(w2), Inches(h2)), h2
    if w: return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(w * ih / iw)), w * ih / iw
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(h * iw / ih), Inches(h)), h


# ================================================================ Theorie-Blöcke → Folien
def blk_h(b, w):
    k = b[0]
    if k == "ul": return sum(th("• " + x, 20, w, pad=0.06) for x in b[1]) + 0.06
    if k == "ol": return sum(th("1. " + x, 20, w, pad=0.06) for x in b[1]) + 0.06
    if k == "merke": return th(b[1], 20, w - 0.1, bold=True, pad=0.22)
    if k == "hint": return th(b[1], 20, w, pad=0.1)
    if k == "table":
        cw = col_widths(b[1], w); return table_h(cw, b[1])
    if k == "tiles":
        n = len(b[1]); per = 3 if n in (3, 6) else (n if n <= 3 else 2); tw = (w - 0.15 * (per - 1)) / per
        rows = [b[1][i:i + per] for i in range(0, n, per)]
        return sum(max(th(h_, 20, tw, True, pad=0.05) + sum(th(x, 20, tw, pad=0.04) for x in (bb if isinstance(bb, list) else [bb])) + 0.15 for h_, bb in r) + 0.12 for r in rows)
    if k == "video": return 0.95
    if k == "versuch": return th(b[2], 20, w - 0.3, True, pad=0.05) + th(b[3], 20, w - 0.3, pad=0.05) + 0.55
    if k == "reveal": return sum(th(q, 20, w, True, pad=0.05) + max(0.75, th(a, 20, w - 0.2, pad=0.25)) + 0.1 for q, a in b[1])
    if k == "karten": return 0
    raise ValueError(k)


def col_widths(rows, w):
    n = len(rows[0]); lens = [max(len(plain(str(r[i]))) for r in rows) for i in range(n)]
    tot = sum(max(l, 8) for l in lens); cw = [w * max(l, 8) / tot for l in lens]
    return cw


def draw_blk(s, b, x, y, w, L):
    k = b[0]
    if k in ("ul", "ol"):
        h = blk_h(b, w); t = tb(s, x, y, w, h, [("• " if k == "ul" else f"{i + 1}. ") + it for i, it in enumerate(b[1])], size=20, after=4); return h
    if k == "merke":
        h = blk_h(b, w); c = box(s, x, y, w, h - 0.05, fill=OLIVE); text(c, [b[1]], size=20, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0); return h
    if k == "hint":
        h = blk_h(b, w); tb(s, x, y, w, h, [[(t_, dict(o, color=MUTED)) for t_, o in runs_md(b[1])]], size=20); return h
    if k == "table":
        cw = col_widths(b[1], w); cells, h = auto_grid(s, x, y, cw, b[1]); return h + 0.08
    if k == "tiles":
        n = len(b[1]); per = 3 if n in (3, 6) else (n if n <= 3 else 2); tw = (w - 0.15 * (per - 1)) / per; yy = y
        for i in range(0, n, per):
            r = b[1][i:i + per]
            rh = max(th(h_, 20, tw, True, pad=0.05) + sum(th(x_, 20, tw, pad=0.04) for x_ in (bb if isinstance(bb, list) else [bb])) + 0.15 for h_, bb in r)
            for j, (h_, bb) in enumerate(r):
                c = box(s, x + j * (tw + 0.15), yy, tw, rh, fill=CARD)
                text(c, [[(h_, {"bold": True, "color": OLIVE})]] + list(bb if isinstance(bb, list) else [bb]), size=20, after=2)
            yy += rh + 0.12
        return yy - y
    if k == "video":
        t_, q, du, u = I.VIDEOS[b[1]]
        c = box(s, x, y, w, 0.88, fill=RGBColor.from_string("FFF4EA"), line=ORANGE, lw=1.5)
        text(c, [[("Video: " + t_, {"bold": True})], [(q + ("" if du == "kurz" else " · " + du), {"size": 16, "color": MUTED})], [(u, {"size": 11, "color": MUTED})]], size=20, after=0)
        s.shapes.add_picture(qr_png(u), Inches(x + w - 0.84), Inches(y + 0.04), Inches(0.8), Inches(0.8))
        return 0.95
    if k == "versuch":
        h = blk_h(b, w)
        c = box(s, x, y, w, h - 0.05, fill=RGBColor.from_string("FBF1EE"), line=RED, lw=2)
        text(c, [[(b[1] + "  ", {"bold": True, "color": RED}), (b[2], {"bold": True})], b[3], [("Protokoll: Versuchsseite (QR auf »Dein Lernpfad«)", {"size": 14, "color": MUTED})]], size=20, after=2)
        return h
    if k == "reveal":
        yy = y
        for q, a in b[1]:
            hq = th(q, 20, w, True, pad=0.05); tb(s, x, yy, w, hq, [q], size=20, bold=True); yy += hq
            ha = max(0.75, th(a, 20, w - 0.2, pad=0.25))
            answer(s, x, yy, w, ha, "Lösung", a, L, hint="Erst selbst überlegen – Lösung auf der Theorie-Seite: »Lösung anzeigen«"); yy += ha + 0.1
        return yy - y
    raise ValueError(k)


def zerlegen(blocks):
    out = []
    for b in blocks:
        if b[0] == "ul" and len(b[1]) > 1: out += [("ul", [x]) for x in b[1]]
        elif b[0] == "tiles":
            n = len(b[1]); per = 3 if n in (3, 6) else 2
            out += [("tiles", b[1][i:i + per]) if per == len(b[1][i:i + per]) or n <= per else ("tiles", b[1][i:i + per]) for i in range(0, n, per)]
        elif b[0] == "reveal" and len(b[1]) > 1: out += [("reveal", [x]) for x in b[1]]
        else: out.append(b)
    return out


def theorie_folien(item, L):
    """Theorie-Folie(n) aus einem Eintrag von theorie_daten (Bild links, Text rechts; zu viel Inhalt → Fortsetzungsfolie)."""
    blocks = zerlegen(item["blocks"])
    img = next((b for b in blocks if b[0] == "img"), None)
    rest = [b for b in blocks if b[0] not in ("img", "karten")]
    karten = next((b for b in blocks if b[0] == "karten"), None)
    khint = None
    if karten and all(b[0] == "hint" for b in rest): khint, rest = rest[0][1], []
    pages, cur, used = [], [], 0
    if img:
        # Bild auf eigener Hälfte; Text rechts. Ist der Text zu lang, Bild-Folie + Text-Folie(n).
        cap_w, txt_w = 4.4, 4.35
        hs = [blk_h(b, txt_w) for b in rest]
        if sum(hs) + 0.1 * len(hs) <= Y1 - Y0:
            pages.append(("bild", img, rest))
        else:
            pages.append(("bild", img, []))
            img = None
    if rest and (not pages or pages[0][2] == []):
        avail = Y1 - Y0
        for b in rest:
            h = blk_h(b, X1 - X0) + 0.1
            if cur and used + h > avail + 0.15:
                pages.append(("text", None, cur)); cur, used = [], 0
            cur.append(b); used += h
        if cur: pages.append(("text", None, cur))
    out = []
    for n, (art, im, bl) in enumerate(pages):
        title = item["title"] + (f" ({n + 1})" if len(pages) > 1 and n else "")
        s = d.slide("Nur Titel", title)
        if art == "bild":
            path = UNIT + im[1].replace(".gif", ".gif")
            if bl:
                pic, hh = picture(s, path, X0, Y0, 4.4, 3.15)
                tb(s, X0, Y0 + hh + 0.05, 4.4, min(0.7, 4.98 - Y0 - hh), [im[3]], size=14, color=MUTED, align=PP_ALIGN.CENTER, after=0)
                y = Y0
                for b in bl: y += draw_blk(s, b, 5.15, y, 4.35, L) + 0.1
            else:
                cap_h = th(im[3], 18, 9.0, pad=0.05)
                pic, hh = picture(s, path, X0 + 0.6, Y0, 7.8, min(3.5, Y1 - Y0 - cap_h - 0.05))
                tb(s, X0, Y0 + hh + 0.05, 9.0, cap_h, [im[3]], size=18, color=MUTED, align=PP_ALIGN.CENTER, after=0)
        else:
            y = Y0
            for b in bl: y += draw_blk(s, b, X0, y, X1 - X0, L) + 0.1
        if item.get("pill"): pill(s, item["pill"], kind="versuch" if item["pill"].startswith("Versuch im Buch") else "theory")
        out.append(s)
    if karten:
        for ki, part in enumerate((karten[1][:4], karten[1][4:])):
            s = d.slide("Nur Titel", item["title"] + (" (2)" if ki else ""))
            rows = [["Frage", "Lösung (für den Fragenden)"]] + [[q, a if L else ""] for q, a in part]
            y0 = Y0
            if khint and ki == 0:
                hh = th(khint, 18, 9); tb(s, X0, Y0, 9, hh, [khint], size=18, color=MUTED); y0 += hh
            cells, _ = auto_grid(s, X0, y0, [4.6, 4.4], rows, size=16)
            if L:
                for r in cells[1:]: s._anim.append([r[1]])
            pill(s, item["pill"])
    return out


# ================================================================ Aufgaben
def fields_rows(t):
    """Für Auswahl-Aufgaben: [(Feld-Label, richtige Antwort)]"""
    rows = []
    for f in t.get("fields", []):
        c = (t.get("chk") or {}).get(f[0], {})
        rows.append((f[1], c.get("sel")))
    return rows


def sol_lines(t):
    out = []
    for head, items in t["sol"]:
        for x in items: out.append(x if head in ("Lösung", "Beispiel", "Bewertung") else head + " " + x)
    return out


def teile_der_frage(q):
    """Frage in Einleitung und Teilfragen a), b) … zerlegen (für die Folien kurz und gegliedert)."""
    m = re.search(r"(?:(?<=\s)|^)a\)\s", q)
    if not m: return q.strip(), {}
    stem, rest = q[:m.start()].strip(), q[m.start():]
    parts = {}
    for mm in re.finditer(r"(?:(?<=\s)|^)([a-h])\)\s(.*?)(?=(?:\s(?:Zusatz:\s)?[a-h]\)\s)|$)", rest, re.S):
        parts[mm.group(1)] = re.sub(r"\s*(Zusatz|Zusatz \(freiwillig\)):?\s*$", "", mm.group(2).strip())
    zus = re.search(r"Zusatz:\s*([a-h])\)", rest)
    if zus and zus.group(1) in parts: parts[zus.group(1)] = "(freiwillig) " + parts[zus.group(1)]
    return stem, parts


def loesung_je_feld(t):
    """Lösung je Feld: Auswahl → richtige Option, sonst die Lösungszeile mit demselben Buchstaben. Rest = übrige Lösungszeilen."""
    sol = sol_lines(t); used = set(); out = {}
    chk = t.get("chk") or {}
    for f in t.get("fields", []):
        k, lab = f[0], f[1]
        c = chk.get(k, {})
        if k not in chk and "tipp" in k:
            out[k] = ("Vorhersage – zählt nicht", "tipp"); continue
        mm = re.match(r"\s*\(?([a-h])\)", lab)
        line = None
        if mm:
            for n, x in enumerate(sol):
                if n not in used and re.match(r"\s*" + mm.group(1) + r"\)", x):
                    line = re.sub(r"^\s*[a-h]\)(?:\s*[–-]\s*[a-h]\))?\s*", "", x); used.add(n); break
        if "sel" in c and c["sel"] is not None: out[k] = (c["sel"], "sel")
        elif line: out[k] = (line, "text")
        elif "num" in c: out[k] = (str(c["num"]).replace(".", ",") + (" " + c["unit"] if c.get("unit") else ""), "num")
        else: out[k] = (None, "text")
    rest = [x for n, x in enumerate(sol) if n not in used]
    return out, rest


def quelle_badge(s, src, x, y):
    kind = kind_of(src)
    col = {"theory": RGBColor.from_string("6E7F4E"), "book": TEAL, "versuch": RED}[kind]
    a, _, b = src.rpartition(" S. ")
    c = box(s, x, y, 1.25, 0.72, fill=col)
    text(c, [[(a or src, {"size": 12, "bold": True, "color": WHITE})], [("S. " + b if a else "", {"size": 16, "bold": True, "color": WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)


def aufgabe(t, L):
    """Aufgaben-Folie im Stil der 1AB-Messen-PPP: Layout »Aufgabe« (hellgrün), Auftragskarte mit Nummer und Quelle,
    darunter jede Teilfrage als eigene Zeile mit Antwortkasten (Lehrer: Lösung per Klick, Kadetten: leer)."""
    tit = f"Aufgabe {t['nr']}: " + re.sub(r"\s*\((Demo|Versuch|Zusatz)[^)]*\)\s*$", "", t["title"])
    stem, parts = teile_der_frage(t["q"])
    sols, rest = loesung_je_feld(t)
    fields = [f for f in t.get("fields", [])]
    chk = t.get("chk") or {}
    seite = [0]

    def neue_folie():
        seite[0] += 1
        sl = d.slide("Aufgabe", tit + (f" ({seite[0]})" if seite[0] > 1 else ""))
        for ph in sl.placeholders:
            if ph.placeholder_format.type == 1:
                n_ = len(ph.text_frame.text)
                if n_ > 34:
                    for r_ in ph.text_frame.paragraphs[0].runs: r_.font.size = Pt(24 if n_ <= 42 else 21)
        lab = ("ZUSATZ" if t["tag"] == "Zusatz" or not t["h"] else t["tag"].upper()) + f" · ca. {t['zeit']} min"
        mz = re.search(r"\(((?:Demo|Versuch) [^)]*)\)", t["title"])
        if mz: lab += " · " + mz.group(1)
        if L and t.get("besprechen"): lab += " · BESPRECHEN"
        lab = lab.replace("ca. ", "")
        pill(sl, lab, x=2.75 if len(lab) > 30 else 3.35, w=4.6 if len(lab) > 30 else 3.1, kind="theory" if t["h"] and t["tag"] != "Zusatz" else "book")
        return sl

    s = neue_folie()
    # gemeinsame Auswahl (alle Felder mit denselben Optionen) in die Karte
    opts = [tuple(f[2]) for f in fields if len(f) > 2 and isinstance(f[2], (list, tuple))]
    gemeinsam = opts[0] if opts and len({frozenset(o) for o in opts}) == 1 and len(opts) == len(fields) else None
    zuletzt = [None]
    kt = stem if (stem and (parts or fields)) else t["q"]
    if t["type"] == "sketch": kt += " (mit der Hand, Foto ins Arbeitsblatt)" if "Foto" not in kt else ""
    lang = th(kt, 20, 6.45) > 2.0
    size = 20 if (lang or th(kt, 20, 6.45) <= 1.45) else 18
    zus = []
    if gemeinsam: zus = [[("Wähle: ", {"bold": True, "color": TEAL, "size": 18}), (" · ".join(gemeinsam), {"size": 18})]]
    ch = max(0.9, th(kt, size, 6.45, pad=0.16) + (th("Wähle: " + " · ".join(gemeinsam), 18, 6.45, pad=0.02) if gemeinsam else 0))
    if lang:
        size = 20
        while size > 16 and th(kt, size, 6.45, pad=0.16) > Y1 - Y0 - 0.05: size -= 1
        ch = Y1 - Y0 + 0.02
    has_img = bool(t.get("img"))
    card = box(s, X0, Y0 - 0.02, 9.0, ch, fill=WHITE, line=RGBColor.from_string("C9D3B6"), lw=1)
    ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(X0 + 0.14), Inches(Y0 - 0.02 + ch / 2 - 0.3), Inches(0.6), Inches(0.6))
    ov.fill.solid(); ov.fill.fore_color.rgb = GREEN; ov.line.fill.background(); ov.shadow.inherit = False
    ov.text_frame.margin_left = ov.text_frame.margin_right = ov.text_frame.margin_top = ov.text_frame.margin_bottom = 0; ov.text_frame.word_wrap = False
    text(ov, [str(t["nr"])], size=20 if t["nr"] < 10 else 17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    tb(s, X0 + 0.88, Y0 + 0.03, 6.45, ch - 0.1, [kt] + zus, size=size, anchor=MSO_ANCHOR.MIDDLE, after=2)
    quelle_badge(s, t["src"], X0 + 9.0 - 1.37, Y0 - 0.02 + ch / 2 - 0.36)
    y = Y0 + ch + 0.14
    xr, wr = X0, 9.0
    if lang: y = Y1 + 1  # langer Auftrag: Karte füllt die Folie, Antworten auf der nächsten
    if has_img:
        pic, hh = picture(s, UNIT + t["img"], X0, y, 4.1, Y1 - y)
        xr, wr = X0 + 4.3, 4.7

    def zeile(label, ans, art, opt_txt):
        nonlocal s, y, xr, wr
        stacked = has_img or art == "text"
        lw_ = wr if stacked else wr * 0.58
        aw = wr if stacked else wr - lw_ - 0.12
        lh = th(label, 18, lw_ - 0.1, pad=0.06) + (th(opt_txt, 16, lw_ - 0.1, pad=0.02) if opt_txt else 0)
        atxt = ans if (L and ans) else ""
        ah = max(0.42, th(atxt or "x", 18, aw - 0.2, bold=True, pad=0.14)) if art != "text" else max(0.5, th(atxt or "x", 18, aw - 0.2, pad=0.16))
        rh = (lh + ah + 0.04) if stacked else max(lh, ah)
        if y + rh > Y1 + 0.02:
            s = neue_folie(); y = Y0; xr, wr = X0, 9.0
            if has_img: pass
        parts_ = [[(label, {"size": 18, "bold": True, "color": INK})]]
        if opt_txt: parts_.append([(opt_txt, {"size": 16, "color": MUTED})])
        tb(s, xr, y, lw_, lh, parts_, size=18, after=1)
        ax, ay = (xr, y + lh + 0.02) if stacked else (xr + lw_ + 0.12, y)
        frame = box(s, ax, ay, aw, ah, fill=WHITE, line=TEAL if not L else ORANGE, lw=1.5)
        frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH  # gestrichelt wie in der 1AB-PPP
        if L and ans:
            col = MUTED if art == "tipp" else RED
            a_ = tb(s, ax + 0.05, ay + 0.02, aw - 0.1, ah - 0.04, [ans], size=18, bold=art != "text", color=col, anchor=MSO_ANCHOR.MIDDLE, after=0)
            if art != "tipp": s._anim.append([a_])
        y += rh + 0.1

    for f in fields:
        k = f[0]
        mm = re.match(r"\s*\(?([a-h])\)", f[1])
        lab = f[1]
        if mm and mm.group(1) in parts and len(parts[mm.group(1)]) <= 170:
            lab = mm.group(1) + ") " + parts[mm.group(1)]
        ans, art = sols.get(k, (None, "text"))
        if art == "text" and "kw" in chk.get(k, {}) and not chk[k].get("man") and ans and len(ans) < 60: art = "num"
        o_txt = None
        if len(f) > 2 and isinstance(f[2], (list, tuple)) and not gemeinsam and frozenset(f[2]) != zuletzt[0]:
            o_txt = "Wähle: " + " · ".join(f[2]); zuletzt[0] = frozenset(f[2])
        zeile(lab, ans, art, o_txt)
    if not fields:
        zeile("Eure Antwort", " ".join(rest) if rest else "", "text", None); rest = []
    if L and rest:
        txt = rest
        h = sum(th(x, 18, 8.7, pad=0.04) for x in txt) + 0.45
        if y + h > Y1 + 0.02:
            if all(len(f) > 2 for f in fields) and h > 1.6: return
            s = neue_folie(); y = Y0
        b = box(s, X0, y, 9.0, min(h, Y1 - y + 0.02), fill=WHITE, line=ORANGE, lw=2)
        text(b, [[("Lösung" + (" · BESPRECHEN" if t.get("besprechen") else ""), {"bold": True, "color": ORANGE, "size": 16})]] + txt, size=18, after=2)
        s._anim.append([b])


def ueberblick(nrs, titel):
    s = d.slide("Nur Titel", "Aufgaben im Überblick · " + titel)
    rows = [["Nr.", "Aufgabe", "Was ist zu tun?", "Quelle", "Zeit"]]
    was = {"multi": "auswählen bzw. kurz schreiben", "sketch": "mit der Hand zeichnen/schreiben, Foto", "text": "schreiben"}
    for t in I.T:
        if t["nr"] in nrs:
            rows.append([str(t["nr"]), t["title"], was[t["type"]], t["src"], f"{t['zeit']} min"])
    cells, h = auto_grid(s, X0, Y0, [0.55, 3.75, 2.4, 1.6, 0.7], rows, size=16, minh=0.34)
    tb(s, X0, Y0 + h + 0.05, 9, 0.35, [f"Zusammen ca. {sum(t['zeit'] for t in I.T if t['nr'] in nrs)} min · Arbeitsblatt: QR auf »Dein Lernpfad«"], size=16, color=MUTED)


def weisst_du_noch(st):
    s = d.slide("Nur Titel", "Weißt du noch?")
    y = Y0
    for q, a in I.WEISST_DU_NOCH[st]:
        tb(s, X0, y, 9, 0.4, [q], size=20, bold=True); y += 0.42
        answer(s, X0, y, 9, 0.62, "", a, d.lehrer, hint="Mini-Whiteboard: aufschreiben und hochhalten"); y += 0.72
    tb(s, X0, y, 9, 0.35, ["ohne Handy: Antworten auf das Mini-Whiteboard schreiben und auf Zeichen hochhalten"], size=16, color=MUTED)


def stopp(st, L):
    s = d.slide("Nur Titel", "Stopp – kannst du das?")
    draw_blk(s, ("reveal", I.STOPPFRAGEN[st]), X0, Y0, 9, L)


def jetzt(txt):
    s = d.slide("Nur Titel", txt)
    c = box(s, 1.5, 2.0, 7.0, 1.2, fill=GREEN); text(c, [txt.replace("Jetzt: ", "Jetzt: ") + " →"], size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    s.shapes.add_picture(qr_png(URL["aufgaben"]), Inches(4.35), Inches(3.35), Inches(1.3), Inches(1.3))
    tb(s, 1.5, 4.68, 7, 0.3, [URL["aufgaben"]], size=10, color=MUTED, align=PP_ALIGN.CENTER)


# ================================================================ Versuche
VS = {v["nr"]: v for v in VD.VS}


def vtag(v):
    return ("KERN" if v["kern"] else "ZUSATZ") + (" · Station " + v["station"] if v.get("station") else (" · Demo" if v.get("demo") else ""))


def versuch_folien(nr, L, kurz=False):
    v = VS[nr]
    pk = "versuch" if v["buch"]["versuch"] else "theory"
    pl = ("Versuch im Buch " if v["buch"]["versuch"] else "Buch ") + v["buch"]["s"]
    # 1 Aufbau, Material, Sicherheit
    s = d.slide("Nur Titel", f"{nr} · {v['titel']}")
    tag(s, vtag(v), OLIVE if v["kern"] else MUTED, x=6.9, w=2.6)
    picture(s, UNIT + "vs-img/d_" + v["skizze"] + ".png", X0, Y0 + 0.3, 4.6, 3.6)
    rows = [["Geräte & Chemikalien"]] + [[m[0]] for m in v["mat"]]
    cells, h = auto_grid(s, 5.3, Y0 + 0.3, [4.2], rows, size=16, minh=0.3)
    hs = th("Sicherheit: " + v["sicher"], 16, 4.2, pad=0.15)
    sich_spaeter = Y0 + 0.35 + h + hs > Y1
    if not sich_spaeter:
        c = box(s, 5.3, Y0 + 0.35 + h, 4.2, hs, fill=RGBColor.from_string("F6E3DF"), line=RUST)
        text(c, [[("Sicherheit: ", {"bold": True, "color": RUST}), (v["sicher"], {})]], size=16, after=0)
    pill(s, pl, kind=pk, w=2.4)
    # 2 Durchführung
    s = d.slide("Nur Titel", f"{nr} · Durchführung")
    tag(s, f"ca. {v['zeit']} min · {v['form'].split(' (')[0][:24]}", OLIVE if v["kern"] else MUTED, x=5.9, w=3.6)
    y = Y0 + 0.3
    if sich_spaeter:
        hs = th("Sicherheit: " + v["sicher"], 18, 9, pad=0.12)
        c = box(s, X0, y, 9, hs, fill=RGBColor.from_string("F6E3DF"), line=RUST)
        text(c, [[("Sicherheit: ", {"bold": True, "color": RUST}), (v["sicher"], {})]], size=18, after=0); y += hs + 0.08
    for i, st in enumerate(v["schritte"]):
        h = th(st, 20, 8.35, pad=0.06)
        if y + h > Y1:
            pill(s, pl, kind=pk, w=2.4); s = d.slide("Nur Titel", f"{nr} · Durchführung (2)"); y = Y0 + 0.1
        c = box(s, X0, y + 0.04, 0.42, 0.42, fill=TEAL); text(c, [str(i + 1)], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        tb(s, X0 + 0.55, y, 8.45, h, [st], size=20); y += h + 0.04
    if y + 0.45 < Y1:
        tb(s, X0, min(y + 0.05, Y1 - 0.45), 9, 0.45, [[("Entsorgung: ", {"bold": True, "color": RUST}), (v["entsorgung"], {})]], size=14, color=MUTED)
    pill(s, pl, kind=pk, w=2.4)
    # 3 Messwerte & Erkenntnis
    s = d.slide("Nur Titel", f"{nr} · {v['beob']['titel']} & Erkenntnis")
    kopf = v["beob"]["kopf"]; ng = max(len(z) for z in v["beob"]["zeilen"])
    erw = dict((r[0].split(" · ")[0], r[1]) for r in v["erwartet"])
    rows = [kopf]
    for i, z in enumerate(v["beob"]["zeilen"]):
        e = v["erwartet"][i][1] if i < len(v["erwartet"]) else ""
        rows.append(z + [e if L else ""] + [""] * (len(kopf) - ng - 1))
    cw = [3.2, 5.8] if len(kopf) == 2 else [3.0, 3.6, 2.4]
    cells, h = auto_grid(s, X0, Y0 + 0.1, cw, rows, size=16, minh=0.42)
    if L:
        for r in cells[1:]: s._anim.append(r[ng:])
    y = Y0 + 0.2 + h
    hq = th("Erkenntnis: " + v["erk"], 20, 9, pad=0.05)
    if Y1 - y - hq < max(0.75, th(v["erk_l"], 20, 8.8, pad=0.4)):
        pill(s, pl, kind=pk, w=2.4); s = d.slide("Nur Titel", f"{nr} · Erkenntnis"); y = Y0 + 0.05
    tb(s, X0, y, 9, hq, [[("Erkenntnis: ", {"bold": True, "color": OLIVE}), (v["erk"], {})]], size=20); y += hq
    answer(s, X0, y, 9, Y1 - y, "Erkenntnis", v["erk_l"], L, hint="Deine Antwort – im Versuchsprotokoll eintragen", size=20)
    pill(s, pl, kind=pk, w=2.4)
    # 4 Interpretation & Schlussfolgerung
    teile = [("Interpretation", v["interp"], v["interp_l"]), ("Schlussfolgerung", v["schluss"], v["schluss_l"])]
    hs = [th(lab + ": " + q, 20, 9, pad=0.05) + max(0.9, th(a, 20, 8.8, pad=0.35)) + 0.1 for lab, q, a in teile]
    gruppen = [teile] if sum(hs) <= Y1 - Y0 else [[teile[0]], [teile[1]]]
    for g in gruppen:
        s = d.slide("Nur Titel", f"{nr} · " + " & ".join(x[0] for x in g))
        y = Y0 + 0.05
        for lab, q, a in g:
            hq = th(lab + ": " + q, 20, 9, pad=0.05)
            tb(s, X0, y, 9, hq, [[(lab + ": ", {"bold": True, "color": OLIVE}), (q, {})]], size=20); y += hq
            ha = max(0.9, th(a, 20, 8.8, pad=0.35)) if len(g) == 2 else Y1 - y
            answer(s, X0, y, 9, ha, lab, a, L, hint="Deine Antwort – im Versuchsprotokoll eintragen", size=20); y += ha + 0.1
        pill(s, pl, kind=pk, w=2.4)
    if L:
        s = d.slide("Nur Titel", f"Lehrer-Info · {nr}", hidden=True)
        tag(s, "AUSGEBLENDET", RUST, x=7.6, w=1.9)
        parts = [[("Form & Dauer: ", {"bold": True, "color": OLIVE}), (f"{v['form']} · ca. {v['zeit']} min · Jahresplan: {v['jp']}", {})]]
        parts += [[("Tipp: ", {"bold": True, "color": OLIVE}), (x, {})] for x in v["lehrer"]]
        parts += [[("Typischer Fehler: ", {"bold": True, "color": RUST}), (x, {})] for x in v["fehler"]]
        parts += [[("Quelle: ", {"bold": True, "color": MUTED}), (v["quelle"], {})]]
        size = 16
        while sum(th("".join(t for t, _ in p), size, 9, pad=0.04) for p in parts) > Y1 - Y0 and size > 12: size -= 1
        tb(s, X0, Y0 + 0.2, 9, Y1 - Y0 - 0.2, parts, size=size, after=4)


def sicherheit():
    s = d.slide("Nur Titel", "Sicherheit bei den Versuchen")
    y = Y0 + 0.05
    for i, (h_, t_) in enumerate(VD.SICHERHEIT):
        c = box(s, X0, y + 0.06, 0.5, 0.5, fill=ORANGE)
        text(c, ["!"], size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        hh = th(h_ + ": " + t_, 20, 8.3, pad=0.08)
        tb(s, X0 + 0.65, y, 8.35, hh, [[(h_ + ": ", {"bold": True, "color": OLIVE}), (t_, {})]], size=20); y += max(0.66, hh + 0.05)
    if y < 4.75: tb(s, X0, y, 9, 0.3, ["Brenner nur bei den Lehrer-Demos K1.3-V3 und V9. Haare zusammenbinden, wenn eine Flamme brennt."], size=14, color=MUTED)


def herrichten_folien():
    rows, slots = H.tabelle()
    head = ["Gerät / Chemikalie", "Fundort", "gesamt", "Bestand", "fehlt"]
    body = [[r["name"], r["ort"], str(r["gesamt"]), "–" if r["bestand"] is None else str(r["bestand"]), r["fehlt"]] for r in rows]
    cw = [2.9, 3.6, 0.8, 0.85, 0.85]
    chunks, cur = [], []
    for r in body:
        if cur and table_h(cw, [head] + cur + [r], size=12, minh=0.28) > 3.45: chunks.append(cur); cur = []
        cur.append(r)
    chunks.append(cur)
    for k, ch in enumerate(chunks):
        s = d.slide("Nur Titel", "Lehrer-Info · Herrichten" + (f" ({k + 1}/{len(chunks)})" if len(chunks) > 1 else ""), hidden=True)
        tag(s, "AUSGEBLENDET", RUST, x=7.6, w=1.9)
        cells, h = auto_grid(s, X0, Y0 + 0.15, cw, [head] + ch, size=12, minh=0.28)
        for r, row in zip(cells[1:], ch):
            if row[4] not in ("", "–"):
                for p_ in r[4].text_frame.paragraphs:
                    for rn in p_.runs: rn.font.color.rgb = RUST; rn.font.bold = True
        tb(s, X0, min(Y0 + 0.2 + h, 4.8), 9, 0.3, ["6 Gruppen zu viert · gesamt = höchster gleichzeitiger Bedarf (Stationen 1 und 2 laufen gleichzeitig) · Details pro Versuch: VS-Fachwissen S. 1"], size=11, color=MUTED)


# ================================================================ Start-Folien
def lernpfad():
    s = d.slide("Nur Titel", "Dein Lernpfad")
    for i, (k, step, lab, col) in enumerate([("theorie", "1", "Theorie", OLIVE), ("aufgaben", "2", "Arbeitsblatt", OLIVE), ("quiz", "3", "Quiz", OLIVE), ("versuche", "V", "Versuche", RED)]):
        ppt_qr(s, URL[k], 0.75 + i * 2.25, 1.45, 1.6, step, lab, col)
    tb(s, X0, 4.62, 9, 0.4, ["Theorie lesen → Aufgaben im Arbeitsblatt → Quiz. Versuche: ein Protokoll-Handy pro Gruppe."], size=16, color=MUTED, after=0)


def fahrplan_folien():
    import json
    u = [x for x in json.load(open(os.path.join(ROOT, "werkzeug/einheiten.json"))) if x["ordner"] == I.ORDNER][0]
    z = u["fahrplan"]["zeilen"]
    n = sum(1 for r in z if r["wann"].startswith("Stunde"))
    cw = [1.15, 3.25, 1.6, 1.7, 1.3]
    mk = lambda r: [r["wann"] + ("\n" + r["datum"] if r.get("datum") else ""), r["lesen"], r.get("aufgaben", "–") + ("\n" + r["zeit"] if r.get("zeit") else ""),
                    (r.get("wie") or "–").replace("ohne Handy: ", ""), r.get("abgabe", "–")]
    schnell = [r for r in z if not r["wann"].startswith("Stunde")]
    z = [r for r in z if r["wann"].startswith("Stunde")]
    gruppen, cur = [], []
    for r in z:
        if cur and table_h(cw, [["Wann", "Das machen wir", "Aufgaben", "Wie (ohne Handy)", "Abgabe"]] + [mk(x) for x in cur + [r]]) > 3.55:
            gruppen.append(cur); cur = []
        cur.append(r)
    gruppen.append(cur)
    for part, rr in enumerate(gruppen, 1):
        s = d.slide("Nur Titel", f"Dein Fahrplan ({part}/{len(gruppen)})")
        tb(s, X0, 1.0, 9, 0.4, [[(f"Dafür sind {n} Stunden eingeplant. ", {"bold": True, "color": OLIVE}), ("Nach Stunde 5 gibst du beide PDFs in LMS ab.", {})]], size=16, after=0)
        rows = [["Wann", "Das machen wir", "Aufgaben", "Wie (ohne Handy)", "Abgabe"]] + [mk(r) for r in rr]
        cells, h = auto_grid(s, X0, 1.45, cw, rows, size=16, minh=0.34)
        if part == len(gruppen) and schnell:
            tb(s, X0, 1.5 + h, 9, 0.4, [[("Für Schnelle: ", {"bold": True, "color": OLIVE}), (f"Aufgaben {schnell[0]['aufgaben']} – freiwillig (OPTIONAL)", {})]], size=18)


def abgabe_folie():
    s = d.slide("Nur Titel", "Fertig? So gibst du ab")
    steps = [("1", "Arbeitsblatt öffnen (QR) → Name und Klasse eintragen"), ("2", "»PDF erstellen« tippen → PDF in LMS hochladen"),
             ("3", "Versuchsseite: Namen der Gruppe + dein Nachname → eigene PDF in LMS"),
             ("4", "Startet kein Download? »Teilen / Speichern« oder »PDF öffnen« tippen bzw. die Seite im Browser öffnen. Geht LMS nicht, schick das PDF per E-Mail an Stephan.")]
    y = 1.12
    for n, t in steps:
        c = box(s, X0, y + 0.04, 0.5, 0.5, fill=OLIVE); text(c, [n], size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        h = th(t, 18 if n == "4" else 20, 5.8, pad=0.06); tb(s, X0 + 0.65, y, 5.8, h, [t], size=18 if n == "4" else 20); y += max(0.65, h + 0.08)
    s.shapes.add_picture(qr_png(URL["aufgaben"]), Inches(7.35), Inches(1.15), Inches(1.15), Inches(1.15))
    tb(s, 6.9, 2.3, 2.1, 0.3, ["Arbeitsblatt"], size=14, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)
    s.shapes.add_picture(qr_png(URL["versuche"]), Inches(7.35), Inches(2.8), Inches(1.15), Inches(1.15))
    tb(s, 6.9, 3.95, 2.1, 0.3, ["Versuche"], size=14, bold=True, color=RED, align=PP_ALIGN.CENTER, after=0)
    tb(s, X0, 4.62, 9, 0.4, ["Abgabe: nach Stunde 5 (Di 17.11.)"], size=18, bold=True, color=RUST, after=0)


def lehrer_infos():
    s = d.slide("Nur Titel", "Lehrer-Info · Lehrplanbezug", hidden=True)
    tag(s, "AUSGEBLENDET", RUST, x=7.6, w=1.9)
    tb(s, X0, 1.25, 9, 3.8, [
        [("Jahresplan E6 + E7 ", {"bold": True, "color": OLIVE}), ("(je 3 Std. = 6) · Buch S. 29–37, 98 · hier 5 Std. in Buch-Reihenfolge (Stunde 2 = Teststunde)", {})],
        [("K04 ", {"bold": True}), ("Masse und Energie bei Reaktionen → A1, A9–A12, V1, V4–V6, V8 · K05 Gleichungen: A3, A4, V4 (Teilgleichungen selbst ergänzen)", {})],
        [("K05 ", {"bold": True}), ("Redox als Elektronenabgabe/-aufnahme, Gleichungen → A2–A8, A15, V1–V4, V9", {})],
        [("N01, N02 · P3 · SDG 7 ", {"bold": True}), ("Fachsprache/Diagramme · Wirtschaft (Katalysator, Korrosion) · Umwelt (Autokat, Streusalz) · Energie", {})],
        [("Schulprofil: ", {"bold": True}), ("Handwärmer im Biwak, Kältekompresse im Sport, Rostschutz an Gerät und Fahrzeug", {})],
        [("Abweichung: ", {"bold": True, "color": RUST}), ("5 statt 6 Std.; Oxidationszahlen nur Zusatz; Enthalpie nur als Vorzeichen von ΔH (steht im Buch S. 33)", {})]], size=17, after=6)
    s = d.slide("Nur Titel", "Lehrer-Info · Stundenvorschlag", hidden=True)
    tb(s, X0, 0.98, 9, 0.3, ["AUSGEBLENDET · Kern in 5 Stunden (Jahresplan: 6) · Rest optional"], size=12, bold=True, color=RUST, after=0)
    rows = [["Std.", "Theorie", "Kern-Versuch", "Aufgaben", "Zeit", "Methode"],
            ["1 · 03.11.", "Einstieg, Teil 1 Redox", "V1 Demo 12 min, V2 Ansatz parallel", "1–3 · Zusatz 4", "13 min + 8 Reserve", "Demo mit Vorhersage"],
            ["2 · 05.11.", "Test 20 min · Teil 2 Rost", "V3 Demo 8 min", "5–7 · Zusatz 8", "10 min + 7 Reserve", "Test, Demo mit Vorhersage"],
            ["3 · 10.11.", "V2 auswerten, Teil 3 Energie", "V6 Demo 8 min · Gr. 1–3: V4 + V5", "Gr. 4–6: 9–11 · Zusatz 12", "Stationen 3 × 11 min", "Demo, Stationenbetrieb"],
            ["4 · 12.11.", "Teil 4 Katalysator", "V7 Demo 10 min · Gr. 4–6: V4 + V5", "13–14 · Gr. 1–3: 9–11", "Stationen 3 × 11 min", "Demo mit Vorhersage"],
            ["5 · 17.11.", "Partnercheck, Teil 5 Tempo", "V8 Gruppen 18 min · Protokolle 7 min", "18, 16–17 · Zusatz 15", "9 min + 5 Reserve", "Partnercheck, Gruppenversuch"],
            ["Zusatz", "Für Schnelle", "V9 (Lehrer, Abzug)", "19–23 · Kahoot", "–", "–"],
            ["Besprechung", "Stunde 1, 2, 5 je 5–7 min", "", "BESPRECHEN: 3, 11, 13 · keine HÜ", "", "Plenum"]]
    auto_grid(s, X0, 1.32, [1.0, 2.0, 1.75, 1.75, 0.8, 1.7], rows, size=13, minh=0.3)
    s = d.slide("Nur Titel", "Lehrer-Info · Die 5 größten Stolpersteine", hidden=True)
    rows = [["Stolperstein", "Richtig ist …"]] + [[a, plain(b)] for a, b in I.STOLPERSTEINE]
    auto_grid(s, X0, 1.1, [2.4, 6.6], rows, size=14)
    herrichten_folien()


def titelfolie(titel, unter, L):
    s = d.slide("Titelfolie", titel)
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2:
            text(ph, [unter, [(f"Stand {I.STAND} · {VER} · " + ("Lehrer" if L else "Kadetten"), {"size": 12})]], size=18, color=RGBColor.from_string("C9D6B0"))


# ================================================================ TH-Deck
def build_th(L):
    global d
    d = Deck(L, "NAWI · 2AB · K1.3 Stoffumwandlung")
    d.section("Start")
    titelfolie("Stoffumwandlung", "NAWI · 2AB · K1.3 · Redox, Energie, Katalysatoren · Buch S. 29–37, 98", L)
    lernpfad(); fahrplan_folien(); abgabe_folie()
    if L: lehrer_infos()
    d.section("Einstieg")
    for it in TD.EINSTIEG: theorie_folien(it, L)
    s = d.slide("Nur Titel", "Das kannst du nach dieser Einheit")
    tb(s, X0, Y0, 9, Y1 - Y0, [f"{i + 1}. " + z for i, z in enumerate(I.LERNZIELE)], size=20, after=8)
    parts = {p[0]: p[3] for p in TD.PARTS}
    item = lambda pid, iid: [x for x in parts[pid] if isinstance(x, dict) and x["id"] == iid][0]

    d.section("Stunde 1")
    theorie_folien(item("teil1", "lavoisier"), L); versuch_folien("K1.3-V1", L)
    for iid in ("redox", "elektronen", "mittel"): theorie_folien(item("teil1", iid), L)
    versuch_folien("K1.3-V2", L)
    stopp(1, L); ueberblick([1, 2, 3, 4], "Stunde 1")
    for t in I.T[0:4]: aufgabe(t, L)

    d.section("Stunde 2")
    weisst_du_noch(2)
    s = d.slide("Abschnitt", "1. Test · 20 Minuten")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2: ph.text_frame.text = "Stunde 2 · Do 05.11."
    theorie_folien(item("teil2", "alltag"), L); versuch_folien("K1.3-V3", L)
    for iid in ("rost", "schutz", "gleichung"): theorie_folien(item("teil2", iid), L)
    stopp(2, L); ueberblick([5, 6, 7, 8], "Stunde 2")
    for t in I.T[4:8]: aufgabe(t, L)

    d.section("Stunde 3")
    weisst_du_noch(3)
    s = d.slide("Nur Titel", "Rost-Ansätze K1.3-V2 auswerten")
    picture(s, UNIT + "vs-img/d_v2.png", X0, Y0, 4.6, 3.4)
    rows = [["Glas", "Ergebnis"]] + [[r[0], r[1] if L else ""] for r in VS["K1.3-V2"]["erwartet"]]
    cells, _ = auto_grid(s, 5.3, Y0, [1.9, 2.3], rows, size=14)
    if L:
        for r in cells[1:]: s._anim.append([r[1]])
    tb(s, 5.3, 4.2, 4.2, 0.8, ["Vergleicht mit eurer Vorhersage und tragt das Ergebnis im Versuchsprotokoll ein."], size=16, color=MUTED)
    pill(s, "Buch S. 98")
    theorie_folien(item("teil3", "exoendo"), L); theorie_folien(item("teil3", "stationen"), L)
    sicherheit()
    for nr in ("K1.3-V6", "K1.3-V4", "K1.3-V5"): versuch_folien(nr, L)
    for iid in ("diagramm", "aktivierung", "erhalt"): theorie_folien(item("teil3", iid), L)
    stopp(3, L); ueberblick([9, 10, 11, 12], "Stunde 3–4 (Gruppen ohne Station)")
    for t in I.T[8:12]: aufgabe(t, L)

    d.section("Stunde 4")
    weisst_du_noch(4)
    theorie_folien(item("teil4", "kat"), L); versuch_folien("K1.3-V7", L)
    for iid in ("enzyme", "autokat", "industrie"): theorie_folien(item("teil4", iid), L)
    stopp(4, L); ueberblick([13, 14], "Stunde 4")
    for t in I.T[12:14]: aufgabe(t, L)

    d.section("Stunde 5")
    weisst_du_noch(5)
    for iid in ("tempo", "rgt", "staub"): theorie_folien(item("teil5", iid), L)
    versuch_folien("K1.3-V8", L)
    theorie_folien(item("teil5", "partner"), L)
    stopp(5, L); ueberblick([15, 16, 17, 18], "Stunde 5")
    for t in I.T[14:18]: aufgabe(t, L)
    s = d.slide("Nur Titel", "Quiz: Reaktions-Check")
    s.shapes.add_picture(qr_png(URL["quiz"]), Inches(3.6), Inches(1.25), Inches(2.8), Inches(2.8))
    tb(s, X0, 4.15, 9, 0.4, [URL["quiz"]], size=12, color=MUTED, align=PP_ALIGN.CENTER)
    tb(s, X0, 4.5, 9, 0.5, ["20 Fragen · auch als Kahoot/Blooket am Beamer (Team-Modus ohne Handy)"], size=18, align=PP_ALIGN.CENTER)

    d.section("Zum Schluss")
    rows = I.FACHBEGRIFFE
    for k, part in enumerate((rows[1:8], rows[8:])):
        s = d.slide("Nur Titel", "Fachbegriffe" + (" (2)" if k else ""))
        auto_grid(s, X0, Y0, [2.6, 6.4], [rows[0]] + part, size=16)
    s = d.slide("Nur Titel", "Das Wichtigste in 5 Sätzen")
    tb(s, X0, Y0, 9, Y1 - Y0, [f"{i + 1}. " + z for i, z in enumerate(I.WICHTIGSTE)], size=19, after=6)
    s = d.slide("Nur Titel", "Kann ich das?")
    rows = [["Ich kann …", "sicher", "fast", "noch nicht"]] + [[plain(z), "☐", "☐", "☐"] for z in I.LERNZIELE]
    auto_grid(s, X0, Y0, [6.3, 0.9, 0.9, 0.9], rows, size=16)

    d.section("Für Schnelle und Neugierige")
    for it in TD.SCHNELLE:
        if isinstance(it, dict):
            sl = theorie_folien(it, L)
            for x in sl: tag(x, "OPTIONAL", MUTED, x=6.6, y=4.78, w=1.7)
    ueberblick([19, 20, 21, 22, 23], "Für Schnelle (OPTIONAL)")
    for t in I.T[18:]: aufgabe(t, L)

    d.section("Abgabe")
    abgabe_folie()
    name = f"{I.PRAEFIX}-TH-{'Lehrer' if L else 'Kadetten'}-{VER}.pptx"
    d.save(os.path.join(OUT, name)); print("ok", name, len(d.p.slides._sldIdLst), "Folien")


# ================================================================ VS-Deck
def build_vs(L):
    global d
    d = Deck(L, "NAWI · 2AB · K1.3 Versuche")
    d.section("Start")
    titelfolie("Versuche zur Stoffumwandlung", "NAWI · 2AB · K1.3 · K1.3-V1 bis V9 · Buch S. 29–37, 98", L)
    sicherheit()
    s = d.slide("Nur Titel", "Versuche im Überblick")
    rows = [["ID", "Versuch", "Wann", "Form", "Einsatz"]]
    for v in VD.VS:
        rows.append([v["nr"], v["titel"], f"Stunde {v['stunde']}" if v["stunde"] else "Zusatz", v["form"].split(" (")[0], "KERN" if v["kern"] else "ZUSATZ"])
    auto_grid(s, X0, Y0 - 0.05, [1.05, 3.4, 0.95, 1.85, 0.85], rows, size=12, minh=0.3)
    s.shapes.add_picture(qr_png(URL["versuche"]), Inches(8.45), Inches(1.15), Inches(1.0), Inches(1.0))
    if L: herrichten_folien()
    for v in VD.VS:
        d.section(v["nr"]); versuch_folien(v["nr"], L)
    d.section("Abgabe"); abgabe_folie()
    name = f"{I.PRAEFIX}-VS-{'Lehrer' if L else 'Kadetten'}-{VER}.pptx"
    d.save(os.path.join(OUT, name)); print("ok", name, len(d.p.slides._sldIdLst), "Folien")


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for L in (True, False): build_th(L); build_vs(L)
