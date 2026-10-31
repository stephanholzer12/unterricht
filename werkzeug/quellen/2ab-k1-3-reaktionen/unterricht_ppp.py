"""Unterrichts-PPP K1.3 Stoffumwandlung (TH-Unterricht) – kurze, geführte PPP für den Frontalunterricht im Saal (Regelwerk bhak-ppp-gestaltung Abschnitt 10, 13, 14).
Zeigt, was die Kadetten in welcher Reihenfolge tun (Lesen mit Abschnitt, Demo/Versuch, Aufgaben), erklärt nur dort nochmals, wo es Schwierigkeiten gibt, und bespricht die heiklen Teilfragen.
Minuten aus plan.py. Aufruf im Repository-Ordner: python3 werkzeug/quellen/2ab-k1-3-reaktionen/unterricht_ppp.py <BHAK_Folienmaster_Vorlage.pptx> <Zielordner>"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import bauen_ppp as B
import inhalt as I, plan as P, theorie_daten as TD, versuche_daten as VD, fachwissen as F
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

ZIEL = sys.argv[2]
OLIVE, INK, TEAL, RUST, MUTED, WHITE, GREEN, ORANGE, CARD, RED = B.OLIVE, B.INK, B.TEAL, B.RUST, B.MUTED, B.WHITE, B.GREEN, B.ORANGE, B.CARD, B.RED
X0, W = 0.5, 9.0
YC = 4.52                         # Inhalt endet hier; darunter der Phasenhinweis (4.6–5.0), dann die Master-Fußzeile
UNIT = B.UNIT
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

FARBE = {"LESEN": TEAL, "KURZ KLÄREN": OLIVE, "AUFGABE": GREEN, "VERSUCH": RED, "BESPRECHEN": ORANGE, "BAUEN": GREEN,
         "START": MUTED, "TEST": MUTED, "STATION": RED, "ABSCHLUSS": RUST, "ABGABE": RUST, "WEG": MUTED, "EINSTIEG": ORANGE, "WIEDERHOLEN": TEAL, "ZU ZWEIT": GREEN}
TITEL = {it["id"]: it["title"] for p in TD.PARTS for it in p[3] if isinstance(it, dict)}
T = {t["nr"]: t for t in I.T}
VS = {v['nr']: v for v in VD.VS}


# ================================================================ Bausteine
def phasen(std):
    return [x for x in P.PLAN if x["std"] == std][0]["phasen"]


def minuten(std, idx):
    return f"{sum(phasen(std)[i][1] for i in idx)} min"


def hinweis(s, art, mins, danach):
    """Phasenhinweis unten links: was JETZT läuft und was danach kommt."""
    col = FARBE.get(art, OLIVE)
    lab = ("JETZT " + art if art not in ("START", "ABGABE", "TEST", "WEG") else art) + (" · " + mins if mins else "")
    w1 = max(2.2, min(4.6, B.FONT_B.getlength(lab) / 100 * 14 / 72 + 0.45))
    c = B.box(s, X0, 4.62, w1, 0.38, fill=col)
    B.text(c, [lab], size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
    if danach:
        B.tb(s, X0 + w1 + 0.1, 4.62, W - w1 - 0.1, 0.38, [[("Danach: ", {"bold": True, "color": col}), (danach, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE, after=0)


def karte(s, x, y, w, h, kopf, zeilen, akzent=GREEN, size=20):
    c = B.box(s, x, y, w, h, fill=CARD)
    band = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(0.08))
    band.fill.solid(); band.fill.fore_color.rgb = akzent; band.line.fill.background(); band.shadow.inherit = False
    c.text_frame.margin_top = Inches(0.14)
    B.text(c, [[(kopf, {"bold": True, "color": OLIVE})]] + list(zeilen), size=size, after=2)
    return c


def merke(s, y, txt, size=20, h=None):
    h = h or B.th(txt, size, W - 0.1, bold=True, pad=0.24)
    c = B.box(s, X0, y, W, h, fill=OLIVE)
    B.text(c, [txt], size=size, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return c, h


def frage_klick(s, y, frage, antwort, w=W, x=X0, size=20):
    """Frage fett, darunter Antwortkasten; die Antwort erscheint per Klick (Lehrer). Gibt neue y zurück."""
    hq = B.th(frage, size, w, True, pad=0.06)
    B.tb(s, x, y, w, hq, [frage], size=size, bold=True, after=0)
    ha = B.th(antwort, size, w - 0.2, pad=0.16)
    fr = B.box(s, x, y + hq, w, ha, fill=RGBColor.from_string("FFFFFF"), line=ORANGE, lw=1.5)
    a = B.tb(s, x + 0.08, y + hq + 0.02, w - 0.16, ha - 0.04, [antwort], size=size, color=RED, anchor=MSO_ANCHOR.MIDDLE, after=0)
    s._anim.append([a])
    return y + hq + ha + 0.1


def qr(s, key, x, y, size, label):
    s.shapes.add_picture(B.qr_png(B.URL[key]), Inches(x), Inches(y), Inches(size), Inches(size))
    B.tb(s, x - 0.3, y + size, size + 0.6, 0.3, [label], size=14, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)


def schritte(s, x, y, w, items, size=20, farbe=TEAL):
    for i, it in enumerate(items):
        h = B.th(it, size, w - 0.6, pad=0.06)
        c = B.box(s, x, y + 0.04, 0.42, 0.42, fill=farbe)
        B.text(c, [str(i + 1)], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, x + 0.55, y, w - 0.55, h, [it], size=size, after=0); y += max(0.5, h + 0.06)
    return y


PROBLEME = []


def titel_von(s):
    return next((ph.text_frame.text for ph in s.placeholders if ph.placeholder_format.type == 1), "?")


def check(s):
    """Kontrolle: nichts unter die Inhaltsgrenze (außer Phasenhinweis, Fußzeile, Pill)."""
    for sh in s.shapes:
        if sh.is_placeholder or sh.name == "Buchseite": continue
        unten = (sh.top + sh.height) / 914400
        if unten > YC + 0.02 and sh.top / 914400 < 4.6:
            PROBLEME.append(f"»{titel_von(s)}«: {sh.name} {sh.text_frame.text[:40] if sh.has_text_frame else ''!r} bis {unten:.2f} in")



# ================================================================ Folien
SPEC = {}
FOLIE = {}
T = {t["nr"]: t for t in I.T}
D = {p["std"]: p["datum"] for p in P.PLAN}
IMG = lambda n: UNIT + "img/" + n
SK = lambda n: UNIT + "vs-img/d_" + n + ".png"
HEUTE_ZIEL = {
    1: "Verstehen, was bei Oxidation und Reduktion passiert: Elektronen wandern – und beim Verbrennen wird ein Stoff nicht leichter.",
    2: "Erklären, wie Eisen rostet und wie man es schützt – und warum eine Redoxreaktion keinen Sauerstoff braucht.",
    3: "Unterscheiden, ob eine Reaktion Wärme abgibt oder aufnimmt – und ein Energiediagramm lesen.",
    4: "Erklären, was ein Katalysator tut – und was er nicht tut. Danach: Arbeitsblatt und Protokolle abgeben.",
}
STEP = {}   # std -> [(Beschriftung, Phasen-Indizes)] für die »Heute«-Folie


def heute_folie(std):
    s = d.slide("Nur Titel", f"Stunde {std} · Heute")
    B.tb(s, X0, 0.98, W, 0.32, [[(f"{D[std]} · " + next(p["titel"] for p in P.PLAN if p["std"] == std), {"bold": True, "color": TEAL})]], size=16, after=0)
    ziel = HEUTE_ZIEL[std]
    c = B.box(s, X0, 1.36, 7.0, B.th("Ziel: " + ziel, 18, 6.9, pad=0.12), fill=OLIVE)
    B.text(c, [[("Ziel: ", {"bold": True}), (ziel, {})]], size=18, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, after=0)
    y = 1.36 + c.height / 914400 + 0.08
    items = [f"{lab} · {minuten(std, idx)}" for lab, idx in STEP[std]]
    zwei = len(items) > 5
    cw = 3.0 if zwei else 6.0; sz = 16 if zwei else 17
    half = (len(items) + 1) // 2
    yy = y; y0 = y
    for i, it in enumerate(items):
        if zwei and i == half: yy = y0
        xx = X0 + (3.5 if zwei and i >= half else 0)
        hh = B.th(it, sz, cw, pad=0.02)
        cc = B.box(s, xx, yy + 0.02, 0.34, 0.34, fill=GREEN)
        B.text(cc, [str(i + 1)], size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, xx + 0.42, yy, cw, hh, [it], size=sz, after=0); yy += max(0.42, hh + 0.06)
    qr(s, "aufgaben", 7.75, 1.36, 1.05, "Arbeitsblatt")
    qr(s, "versuche" if std in (3, 4) else "theorie", 7.75, 2.85, 1.05, "Versuche" if std in (3, 4) else "Theorieteil")
    return s


def lesen_folie(std, buch, abschnitte, leitfrage):
    titel = "Lesen: " + (TITEL[abschnitte[0]] if len(abschnitte) == 1 else f"Theorieteil Teil {std}")
    if len(titel) > 40: titel = f"Lesen: Theorieteil Teil {std}"
    s = d.slide("Nur Titel", titel)
    y = 1.05
    bh = 0
    if buch:
        bh = max(0.62, B.th(buch, 20, W - 0.95, pad=0.2))
        B.box(s, X0, y, W, bh, fill=CARD, line=OLIVE, lw=1.5)
        ic = B.box(s, X0 + 0.1, y + bh / 2 - 0.27, 0.62, 0.54, fill=OLIVE)
        B.text(ic, [[("BUCH", {"size": 12, "bold": True, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, X0 + 0.82, y, W - 0.9, bh, [buch], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0); y += bh + 0.12
    B.tb(s, X0, y, 6.9, 0.36, [[("Im Theorieteil genau diese Abschnitte:", {"bold": True, "color": TEAL})]], size=18, after=0); y += 0.4
    for k, a in enumerate(abschnitte):
        xx = X0 + (k % 2) * 3.5; yy = y + (k // 2) * 0.62
        B.box(s, xx, yy, 3.4, 0.56, fill=RGBColor.from_string("E8F1F3"), line=TEAL, lw=1.25)
        B.tb(s, xx + 0.08, yy, 3.25, 0.56, [[("»" + TITEL[a] + "«", {"bold": True})]], size=16, anchor=MSO_ANCHOR.MIDDLE, after=0)
    y += 0.62 * ((len(abschnitte) + 1) // 2)
    qr(s, "theorie", 7.75, 1.05 + (bh + 0.12 if buch else 0), 1.2, "Theorieteil")
    y += 0.08
    hq = B.th("Leitfrage: " + leitfrage, 20, W - 0.1, pad=0.2)
    B.box(s, X0, y, W, hq, fill=RGBColor.from_string("FFF4EA"), line=ORANGE, lw=1.5)
    B.tb(s, X0 + 0.1, y, W - 0.2, hq, [[("Leitfrage: ", {"bold": True, "color": ORANGE}), (leitfrage, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


def auftrag_folie(titel, nrs, tun, erwartet, ohne_handy=False):
    s = d.slide("Aufgabe", titel)
    rows = [["Nr.", "Aufgabe", "Zeit", "So arbeitest du"]]
    for n in nrs: rows.append([str(n), B.plain(T[n]["title"]), f"{T[n]['zeit']} min", tun.get(n, "allein am Handy")])
    cells, h = B.auto_grid(s, X0, 1.05, [0.6, 4.2, 0.9, 3.3], rows, size=18, minh=0.42)
    y = 1.05 + h + 0.12
    if ohne_handy:
        c = B.box(s, X0, y, W, 0.45, fill=RGBColor.from_string("F6E3DF"), line=RUST)
        B.text(c, [[("Ohne Handy: ", {"bold": True, "color": RUST}), (ohne_handy, {})]], size=18, anchor=MSO_ANCHOR.MIDDLE, after=0); y += 0.55
    if erwartet:
        he = B.th("Am Ende steht da: " + erwartet, 20, W - 0.2, pad=0.18)
        B.box(s, X0, y, W, he, fill=RGBColor.from_string("FFFFFF"), line=GREEN, lw=1.5)
        B.tb(s, X0 + 0.1, y, W - 0.2, he, [[("Am Ende steht da: ", {"bold": True, "color": GREEN}), (erwartet, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


def versuch_folie(nr, titel, tipp, ergebnis, extra_sicher=None, bild=None):
    """Demo-/Versuchsfolie: Skizze, Kernanweisung, Sicherheit in Rot; Ergebnis per Klick."""
    v = VS[nr]; s = d.slide("Nur Titel", titel)
    B.picture(s, SK(v["skizze"]), X0, 1.0, 3.6, 3.0)
    y = 1.05
    y = frage_klick(s, y, tipp, ergebnis, w=5.2, x=4.3, size=19)
    sic = extra_sicher or v["sicher"]
    hs = B.th("Sicherheit: " + sic, 15, 5.2, pad=0.14)
    B.box(s, 4.3, min(max(y, 2.7), YC - hs), 5.2, hs, fill=RGBColor.from_string("F6E3DF"), line=RUST)
    B.tb(s, 4.38, min(max(y, 2.7), YC - hs), 5.04, hs, [[("Sicherheit: ", {"bold": True, "color": RUST}), (sic, {})]], size=15, anchor=MSO_ANCHOR.MIDDLE, after=0)
    B.tb(s, X0, 4.05, 3.6, 0.4, [v["form"].split(" (")[0] + (" · " + v["platz_kurz"] if v.get("platz_kurz") else "")], size=14, color=MUTED, after=0)
    B.pill(s, "Versuch " + nr, w=2.2, kind="versuch")
    return s


def beobachtung_folie(nr, titel, nachher=None):
    """Erwartete Beobachtung als Tabelle, Zeilen erscheinen per Klick."""
    v = VS[nr]; s = d.slide("Nur Titel", titel)
    rows = [["Was?", "Das sehen wir"]] + [[B.plain(a), B.plain(b)] for a, b in v["erwartet"]]
    size = 18 if len(rows) <= 5 else 16
    cells, h = B.auto_grid(s, X0, 1.05, [2.4, 6.6], rows, size=size, minh=0.44)
    for r in cells[1:]: s._anim.append([r[1]])
    if nachher:
        y = 1.05 + h + 0.12
        merke(s, y, nachher, size=19)
    B.pill(s, "Versuch " + nr, w=2.2, kind="versuch")
    return s


# ---------------------------------------------------------------- Stunde 1
def s1_einst():
    s = d.slide("Nur Titel", "Einstieg: Warum wird es warm?")
    B.picture(s, IMG("handwaermer.png"), X0, 1.0, 4.4, 3.4)
    karte(s, 5.1, 1.05, 4.4, 1.75, "Handwärmer im Winterbiwak", ["Tütchen auf – nach Minuten über 50 °C. Drinnen: **Eisenpulver**, Salz, Wasser."], ORANGE, size=19)
    y = frage_klick(s, 2.95, "Wird ein Stoff beim Verbrennen leichter?", "Das prüfen wir gleich mit der Waage – erst tippen!", w=4.4, x=5.1, size=19)
    B.pill(s, "Buch S. 29", w=2.2, kind="book")
    return s


def s1_v1_warum():
    s = d.slide("Nur Titel", "Warum wurde die Wolle schwerer?")
    karte(s, X0, 1.05, 4.4, 2.2, "Eisenwolle", ["Eisen + **Sauerstoff aus der Luft** → Eisenoxid", "Der Sauerstoff wird eingebaut – darum **mehr Masse**."], GREEN, size=19)
    karte(s, X0 + 4.6, 1.05, 4.4, 2.2, "Holz im Feuer", ["Es entweichen **Kohlenstoffdioxid und Wasserdampf** – unsichtbar.", "Die Asche wiegt weniger, die **Gesamtmasse bleibt gleich**."], TEAL, size=19)
    c, h = merke(s, 3.45, "Lavoisier (um 1775): Bei einer Reaktion geht keine Masse verloren – und genau daher kommt der Name »Oxidation«.", size=19)
    s._anim.append([c])
    return s


def s1_redox():
    s = d.slide("Nur Titel", "Kurz klären: Oxidation heute")
    rows = [["", "früher", "heute"], ["Oxidation", "Verbindung mit Sauerstoff", "Abgabe von Elektronen"], ["Reduktion", "Entzug von Sauerstoff", "Aufnahme von Elektronen"]]
    cells, h = B.auto_grid(s, X0, 1.0, [2.2, 3.4, 3.4], rows, size=18, minh=0.42)
    y = 1.0 + h + 0.1
    c, h2 = merke(s, y, "Merkhilfe: Bei der REduktion kommen Elektronen REin. Bei der Oxidation gehen sie weg.", size=18)
    y += h2 + 0.15
    y = frage_klick(s, y, "Rostet Eisen im Wasser – ist das noch Oxidation?", "Ja – Eisen gibt Elektronen ab. Sauerstoff ist nur ein möglicher Partner.", size=18)
    return s


def s1_mg():
    s = d.slide("Nur Titel", "Kurz klären: Magnesium und Sauerstoff")
    B.picture(s, IMG("redox_mg.gif"), X0, 1.0, 4.5, 3.4)
    rows = [["Teilgleichung", "Was passiert?"], ["Mg → Mg²⁺ + 2 e⁻", "Mg gibt 2 e⁻ ab (Oxidation)"], ["O + 2 e⁻ → O²⁻", "O nimmt 2 e⁻ auf (Reduktion)"]]
    cells, h = B.auto_grid(s, 5.2, 1.05, [2.0, 2.3], [[B.plain(c) for c in r] for r in rows], size=16, minh=0.5)
    y = frage_klick(s, 1.05 + h + 0.1, "Wer ist das Reduktionsmittel?", "Magnesium – es gibt Elektronen ab und reduziert so den Sauerstoff.", w=4.3, x=5.2, size=18)
    B.pill(s, "Buch S. 30–31", w=2.2, kind="book")
    return s


def s1_bespr():
    s = d.slide("Aufgabe", "Besprechen: Aufgaben 1–3")
    y = 1.05
    y = frage_klick(s, y, "A1 c) Woher kommt die zusätzliche Masse?", "Aus dem Sauerstoff der Luft – er wird im Eisenoxid eingebaut.")
    y = frage_klick(s, y, "A2 e) Warum »Redox«?", "Was ein Stoff abgibt, nimmt ein anderer auf – beides läuft gleichzeitig.")
    y = frage_klick(s, y, "A3 e/f) Oxidationsmittel – Reduktionsmittel?", "Sauerstoff nimmt auf = Oxidationsmittel · Magnesium gibt ab = Reduktionsmittel")
    return s


def s1_v2():
    s = d.slide("Nur Titel", "Ausblick: Rost-Ansatz – Tipp ab")
    B.picture(s, SK("v2"), X0, 1.0, 4.4, 3.4)
    karte(s, 5.1, 1.05, 4.4, 1.9, "Vier Gläser, vier Bedingungen", ["Wo rostet die Eisenwolle in einer Woche am stärksten – wo gar nicht?"], ORANGE, size=19)
    karte(s, 5.1, 3.1, 4.4, 1.3, "Heute", ["Tipp aufschreiben. Ausgewertet am **Di 10.11.**"], GREEN, size=19)
    B.pill(s, "Versuch K1.3-V2", w=2.2, kind="versuch")
    return s


# ---------------------------------------------------------------- Stunde 2
def s2_klaeren_rost():
    s = d.slide("Nur Titel", "Kurz klären: Wie Eisen rostet")
    B.picture(s, IMG("rost_bedingungen.png"), X0, 1.0, 4.6, 3.4)
    y = 1.05
    y = frage_klick(s, y, "Was braucht Eisen zum Rosten?", "Wasser und Sauerstoff – Salz beschleunigt es.", w=4.2, x=5.3, size=19)
    y = frage_klick(s, y, "Was passiert mit dem Eisen?", "Fe → Fe²⁺ + 2 e⁻: Es gibt Elektronen ab (Oxidation).", w=4.2, x=5.3, size=19)
    B.pill(s, "Buch S. 30, 98", w=2.2, kind="book")
    return s


def s2_klaeren_cu():
    s = d.slide("Nur Titel", "Kurz klären: Eisen in Kupfersulfat")
    rows = [["", "Teilgleichung"], ["Oxidation", "Fe → Fe²⁺ + 2 e⁻"], ["Reduktion", "Cu²⁺ + 2 e⁻ → Cu"], ["Redoxreaktion", "Fe + Cu²⁺ → Fe²⁺ + Cu"]]
    cells, h = B.auto_grid(s, X0, 1.05, [2.6, 6.4], [[B.plain(c) for c in r] for r in rows], size=20, minh=0.5)
    y = frage_klick(s, 1.05 + h + 0.15, "Wo ist hier der Sauerstoff – und ist es trotzdem eine Redoxreaktion?", "Es gibt keinen. Elektronen wandern von Fe zu Cu²⁺ – das reicht. (Station 1, Stunde 3)", size=19)
    B.pill(s, "Buch S. 30–31", w=2.2, kind="book")
    return s


def s2_bespr():
    s = d.slide("Aufgabe", "Besprechen: Aufgaben 5–7")
    y = 1.05
    y = frage_klick(s, y, "A5 d) Warum innen kupferrot, außen schwarz?", "Außen kommt Sauerstoff an das heiße Kupfer → Kupferoxid. Innen kaum Luft.")
    y = frage_klick(s, y, "A6 Eis schmelzen – Redoxreaktion?", "Nein – es entsteht kein neuer Stoff, kein Elektronenübergang. (Zucker lösen auch nicht.)")
    y = frage_klick(s, y, "A7 Wie schützt Verzinken das Eisen?", "Zink ist unedler und wird zuerst oxidiert – es »opfert« sich.")
    return s


def s2_ende():
    s = d.slide("Nur Titel", "Nächste Stunde: Stationen")
    karte(s, X0, 1.05, 4.4, 1.9, "Gruppen 1–3 (Di 10.11.)", ["Drei Runden an **Station 1 (Mobiltisch 1)** und **Station 2 (Schülertisch)**", "11 Minuten je Runde"], RED, size=18)
    karte(s, X0 + 4.6, 1.05, 4.4, 1.9, "Gruppen 4–6 (Di 10.11.)", ["Theorieteil Teil 3 lesen", "Aufgaben 9–11 am Handy"], TEAL, size=19)
    merke(s, 3.15, "Am Do 12.11. tauschen die Gruppen: 4–6 an die Stationen, 1–3 machen Aufgaben.", size=19)
    return s


# ---------------------------------------------------------------- Stunde 3
def s3_exo():
    s = d.slide("Nur Titel", "Kurz klären: exotherm – endotherm")
    karte(s, X0, 1.05, 4.4, 2.1, "exotherm – Energie geht hinaus", ["Die Reaktion **gibt Wärme ab**: es wird warm oder hell.", "Handwärmer, Lagerfeuer"], RED, size=19)
    karte(s, X0 + 4.6, 1.05, 4.4, 2.1, "endotherm – Energie geht hinein", ["Der Vorgang **nimmt Wärme auf**: es wird kalt.", "Fotosynthese, Kältemischung"], TEAL, size=19)
    y = frage_klick(s, 3.3, "Eis schmelzen und Schweiß verdunsten – exo oder endo?", "Beides endotherm: Sie nehmen Wärme auf – darum kühlt Schweiß.", size=19)
    B.pill(s, "Buch S. 32", w=2.2, kind="book")
    return s


def s3_diagramm():
    s = d.slide("Nur Titel", "Kurz klären: Das Energiediagramm")
    B.picture(s, IMG("diagramm_exo_endo.png"), X0, 1.0, 5.0, 3.4)
    y = 1.05
    y = frage_klick(s, y, "Was zeigt der Gipfel?", "Die Aktivierungsenergie – den Berg zum Start.", w=3.9, x=5.6, size=19)
    y = frage_klick(s, y, "Woran sieht man exotherm?", "Die Endstoffe liegen tiefer als die Ausgangsstoffe.", w=3.9, x=5.6, size=19)
    B.pill(s, "Buch S. 33–34", w=2.2, kind="book")
    return s


def runden(sp):
    """Rotationstabelle für drei Gruppen: (Station 1, Station 2, Aufgaben) je Runde."""
    a, b, c = sp
    return [["Runde", "Station 1 · V4 Mobiltisch 1", "Station 2 · V5 Schülertisch", "Aufgaben + Lesen am Platz"], ["1", f"Gruppe {a}", f"Gruppe {b}", f"Gruppe {c}"],
            ["2", f"Gruppe {c}", f"Gruppe {a}", f"Gruppe {b}"], ["3", f"Gruppe {b}", f"Gruppe {c}", f"Gruppe {a}"]]


def s3_stationen(sp, lesen_teil, aufg):
    s = d.slide("Nur Titel", "Stationen: so wechselt ihr")
    cells, h = B.auto_grid(s, X0, 1.05, [1.2, 2.7, 2.9, 2.2], runden(sp), size=16, minh=0.4)
    y = 1.05 + h + 0.15
    karte(s, X0, y, 4.4, 4.45 - y, "An der Station", ["messen, beobachten, **aufräumen**", "Erkenntnis am Handy eintragen"], RED, size=18)
    karte(s, X0 + 4.6, y, 4.4, 4.45 - y, "Freie Gruppe", [lesen_teil, aufg], TEAL, size=18)
    return s


def s3_sicher():
    s = d.slide("Nur Titel", "Sicherheit an den Stationen")
    karte(s, X0, 1.05, 4.4, 3.35, "Station 1 · Kupfersulfat (Mobiltisch 1)", ["**Schutzbrille und Handschuhe**", "reizt Augen und Haut, giftig für Wasserlebewesen", "Reste NUR in den Behälter »Schwermetall-Lösungen« – nie in den Ausguss"], RED, size=19)
    karte(s, X0 + 4.6, 1.05, 4.4, 3.35, "Station 2 · Zitronensäure-Pulver", ["**Schutzbrille**", "Pulver nicht einatmen, reizt die Augen – nichts kosten", "nach dem Versuch Hände waschen, Becher ausspülen"], ORANGE, size=19)
    return s


def s3_v4():
    v = VS["K1.3-V4"]; s = d.slide("Nur Titel", "Station 1: Eisenwolle in Kupfersulfat")
    B.picture(s, SK("v4"), X0, 1.0, 3.6, 3.4)
    y = schritte(s, 4.3, 1.05, 5.2, ["Temperatur der Lösung messen, aufschreiben", "Eisenwolle untertauchen, Thermometer bleibt drin", "2 min beobachten: Farbe der Wolle und der Lösung, dann Temperatur"], size=19, farbe=RED)
    B.tb(s, 4.3, min(y + 0.05, 3.9), 5.2, 0.4, ["Platz: Mobiltisch 1 · Ausspülen: Wasserleitung 2"], size=15, bold=True, color=OLIVE, after=0)
    B.pill(s, "Versuch K1.3-V4", w=2.2, kind="versuch")
    return s


def s3_v5():
    s = d.slide("Nur Titel", "Station 2: Zitronensäure und Natron (Pulver)")
    B.picture(s, SK("v5"), X0, 1.0, 3.6, 3.4)
    y = schritte(s, 4.3, 1.05, 5.2, ["25 ml Wasser: Temperatur messen", "Zitronensäure (5,0 g) und Natron (6,0 g) zugleich zugeben, rühren", "Sprudeln beobachten, Temperatur ablesen – wärmer oder kälter?"], size=19, farbe=ORANGE)
    B.tb(s, 4.3, min(y + 0.05, 3.9), 5.2, 0.4, ["Platz: Schülertisch (IKT-Tisch, nicht am Fenster)"], size=15, bold=True, color=OLIVE, after=0)
    B.pill(s, "Versuch K1.3-V5", w=2.2, kind="versuch")
    return s


def s3_ergebnis():
    s = d.slide("Nur Titel", "Sichern: warm oder kalt?")
    rows = [["Station", "Was ist passiert?", "Temperatur", "Art"], ["1 · Eisen + CuSO₄", "Kupfer scheidet sich ab", "steigt (einige °C)", "exotherm"], ["2 · Zitronensäure + Natron", "Gas sprudelt", "sinkt", "endotherm"]]
    cells, h = B.auto_grid(s, X0, 1.05, [2.8, 2.5, 1.9, 1.8], [[B.plain(c) for c in r] for r in rows], size=18, minh=0.5)
    for r in cells[1:]: s._anim.append([r[2], r[3]])
    y = frage_klick(s, 1.05 + h + 0.15, "Warum wird es bei Station 2 kalt, obwohl es sprudelt?", "Die Reaktion nimmt Wärme aus der Mischung und der Umgebung auf – endotherm.", size=19)
    return s


def s3_plaetze():
    s = d.slide("Nur Titel", "Wo steht was?")
    for i, (kopf, zl, col) in enumerate([("Lehrertisch", ["Demos (V6 heute, V7 am Do)", "Waage, Feuerzeug, Material"], OLIVE),
                                         ("Mobiltisch 1 · Station 1", ["V4 Eisenwolle in Kupfersulfat", "**Schutzbrille und Handschuhe**", "Sammelbehälter für Reste stehen hier"], RED),
                                         ("Schülertisch · Station 2", ["V5 Zitronensäure + Natron", "IKT-Tisch, nicht am Fenster"], ORANGE),
                                         ("Wasserleitung 1 und 2 · Ausspülen", ["Becher und Thermometer ausspülen", "Kupferlösung NIE in den Ausguss"], TEAL)]):
        karte(s, X0 + (i % 2) * 4.6, 1.05 + (i // 2) * 1.65, 4.4, 1.5, kopf, zl, col, size=17)
    B.tb(s, X0, 4.2, W, 0.3, ["Fenster: Vorhänge beachten – Flamme nur am Lehrertisch. Mobiltisch 2–4 bleiben frei (Reserve)."], size=14, color=MUTED, after=0)
    return s


# ---------------------------------------------------------------- Stunde 4
def s4_weisst():
    s = d.slide("Nur Titel", "Weißt du noch?")
    y = 1.05
    for q, a in I.WEISST_DU_NOCH[4]: y = frage_klick(s, y, q, a, size=19)
    return s


def s4_v7w():
    s = d.slide("Nur Titel", "Kurz klären: Was tat das Lithiumchlorid?")
    B.picture(s, IMG("diagramm_kat.png"), X0, 1.0, 4.4, 3.4)
    y = 1.05
    y = frage_klick(s, y, "Was tat das Lithiumchlorid?", "Katalysator: senkt die Aktivierungsenergie.", w=4.4, x=5.1, size=16)
    y = frage_klick(s, y, "Wozu der Würfel mit Kochsalz?", "Kontrolle: Nur das Lithiumchlorid wirkt, nicht jedes Pulver.", w=4.4, x=5.1, size=16)
    y = frage_klick(s, y, "Und im Körper?", "Enzyme: Zucker verbrennt bei 37 °C ohne Flamme.", w=4.4, x=5.1, size=16)
    B.pill(s, "Buch S. 34", w=2.2, kind="book")
    return s


def s4_tausch():
    s = d.slide("Nur Titel", "Heute tauschen die Gruppen")
    cells, h = B.auto_grid(s, X0, 1.05, [1.2, 2.7, 2.9, 2.2], runden((4, 5, 6)), size=16, minh=0.4)
    y = 1.05 + h + 0.15
    karte(s, X0, y, 4.4, 4.45 - y, "Gruppen 4–6", ["V4 am Mobiltisch 1, V5 am Schülertisch – genau wie am Di 10.11."], RED, size=18)
    karte(s, X0 + 4.6, y, 4.4, 4.45 - y, "Gruppen 1–3 (und Freie)", ["Teil 4 lesen, Aufgabe 14, Deutung V4/V5 im Protokoll, Aufgaben 9–11 fertig"], TEAL, size=18)
    return s


def s4_auto():
    s = d.slide("Nur Titel", "Kurz klären: Autokatalysator")
    B.picture(s, IMG("autokat.png"), X0, 1.0, 4.6, 3.4)
    y = 1.05
    y = frage_klick(s, y, "Ist der Katalysator ein Filter?", "Nein – er wandelt um: CO → CO₂, NOₓ → N₂.", w=4.2, x=5.3, size=19)
    y = frage_klick(s, y, "Wann arbeitet er schlecht?", "Beim Kaltstart – er muss erst heiß werden (ab etwa 300 °C).", w=4.2, x=5.3, size=19)
    B.pill(s, "Buch S. 34–35", w=2.2, kind="book")
    return s


def s4_bespr():
    s = d.slide("Aufgabe", "Besprechen: Aufgaben 13–14")
    y = 1.05
    y = frage_klick(s, y, "A13 c) Welche Rolle spielt das Lithiumchlorid?", "Katalysator: Es senkt die Aktivierungsenergie und wird nicht verbraucht.")
    y = frage_klick(s, y, "A13 d) Warum gab es den Würfel mit Kochsalz?", "Kontrolle: Nicht irgendein Pulver wirkt, sondern das Lithiumchlorid.")
    y = frage_klick(s, y, "A14 Stimmt: »Ein Katalysator wird verbraucht«?", "Stimmt nicht – er wird nicht verbraucht und ändert nur Weg und Tempo.")
    return s


# ---------------------------------------------------------------- Abschluss (Stunde 4)
def s5_abgabe():
    s = d.slide("Nur Titel", "Abgabe: PDF in LMS")
    schritte(s, X0, 1.1, 6.9, ["Arbeitsblatt: »PDF erstellen« → in LMS hochladen", "Versuchsseite: Protokoll fertig, dann PDF erstellen und hochladen", "Quiz zum Üben – nicht abgabepflichtig", "Tempo (Teil 5) und Aufgaben 16–18: freiwilliger Zusatz"], size=20, farbe=RUST)
    qr(s, "aufgaben", 7.75, 1.0, 0.85, "Arbeitsblatt")
    qr(s, "versuche", 7.75, 2.15, 0.85, "Versuche")
    qr(s, "quiz", 7.75, 3.3, 0.85, "Quiz")
    return s


def s5_wichtig(von=0, bis=3, titel="Das Wichtigste in 5 Sätzen (1/2)"):
    s = d.slide("Nur Titel", titel)
    items = [B.plain(x) for x in I.WICHTIGSTE][von:bis]
    y = 1.05
    for i, it in enumerate(items):
        hh = B.th(it, 18, 8.3, pad=0.04)
        c = B.box(s, X0, y + 0.03, 0.4, 0.4, fill=OLIVE)
        B.text(c, [str(von + i + 1)], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, X0 + 0.55, y, 8.4, hh, [it], size=18, after=0); y += max(0.5, hh + 0.1)
    return s


def s5_wichtig2(): return s5_wichtig(3, 5, "Das Wichtigste in 5 Sätzen (2/2)")


def zusatz_folie():
    s = d.slide("Nur Titel", "OPTIONAL: Für Schnelle", hidden=True)
    rows = [["Nr.", "Zusatzaufgabe", "Zeit"]] + [[str(n), B.plain(T[n]["title"]), f"{T[n]['zeit']} min"] for n in (4, 8, 12, 15, 16, 17, 18, 19, 20, 21, 22, 23)]
    cells, h = B.auto_grid(s, X0, 1.0, [0.7, 6.9, 1.4], rows, size=16, minh=0.36)
    B.tb(s, X0, 1.0 + h + 0.1, W, 0.6, ["Nur, wenn die Pflicht fertig ist. Aufgabe 23 (Video, Podcast oder Plakat) ist freiwillig – Abgabe als Datei in LMS."], size=16, after=0)
    return s


# ---------------------------------------------------------------- Ablauf je Stunde: (Schlüssel, Phasen-Indizes, Art, Schritt für »Heute«, Wozu?, Funktion)
SPEC = {
 1: [("s1_heute", [0], "START", None, "Ziel und Ablauf zeigen", lambda: heute_folie(1)),
     ("s1_einst", [1], "EINSTIEG", None, "Neugier wecken, Mythos per Handzeichen – noch nicht auflösen", s1_einst),
     ("s1_lesen", [2], "LESEN", "Lesen", "Redox-Begriffe selbst lesen – Stephan erklärt nicht vor",
      lambda: lesen_folie(1, "**Buch S. 30–31** (Redoxreaktionen)", ["lavoisier", "redox", "elektronen", "mittel"], "Was gibt ein Magnesium-Atom ab – und was nimmt Sauerstoff auf?")),
     ("s1_v1", [3], "VERSUCH", "Demo V1: Eisenwolle auf der Waage", "Tipp abgeben, dann Demo; Ergebnis per Klick",
      lambda: versuch_folie("K1.3-V1", "Demo V1: Eisenwolle auf der Waage", "Leichter, schwerer oder gleich? – Tipp per Handzeichen (Aufgabe 1 a)", "schwerer (z. B. 1,20 g → etwa 1,5 g)")),
     ("s1_v1w", [3], "VERSUCH", None, "Warum schwerer? Lavoisier, Masse bleibt erhalten", s1_v1_warum),
     ("s1_redox", [4], "KURZ KLÄREN", "Kurz klären", "Fehlvorstellung »Oxidation = Sauerstoff« ausräumen", s1_redox),
     ("s1_mg", [4], "KURZ KLÄREN", None, "Elektronenübergang an Magnesium sichtbar machen; Mittel macht es beim anderen", s1_mg),
     ("s1_a13", [5], "AUFGABE", "Aufgaben 1–3", "Begriffe anwenden; Besprechung gleich danach",
      lambda: auftrag_folie("Jetzt: Aufgaben 1–3", [1, 2, 3], {}, "Auswahlfelder mit »Prüfen«; A1 c) und A3 e) erklären in eigenen Worten.")),
     ("s1_bespr", [5], "BESPRECHEN", None, "Die drei häufigsten Fehler klären", s1_bespr),
     ("s1_v2", [6], "ABSCHLUSS", None, "Rost-Ansatz V2 (Lehrer + 2 Kadetten parallel angesetzt): Tipp aufschreiben lassen; Reserve", s1_v2)],
 2: [("s2_heute", [0, 1], "TEST", "Test K1.2", "Test K1.2 (20 min), danach Ziel von Stunde 2", lambda: heute_folie(2)),
     ("s2_lesen", [2], "LESEN", "Lesen", "Rost und Rostschutz selbst lesen",
      lambda: lesen_folie(2, "**Buch S. 30, 98** (Rost und Rostschutz)", ["alltag", "rost", "schutz", "gleichung"], "Was braucht Eisen zum Rosten – und warum schützt Zink?")),
     ("s2_v3", [3], "VERSUCH", "Demo V3: Kupferbrief", "Tipp abgeben, dann Demo; Ergebnis per Klick",
      lambda: versuch_folie("K1.3-V3", "Demo V3: Der Kupferbrief", "Wie sieht es innen aus? – Tipp per Handzeichen (Aufgabe 5 a)", "außen schwarz (Kupferoxid) · innen kupferrot glänzend")),
     ("s2_rost", [3], "KURZ KLÄREN", "Kurz klären", "Rosten als Redox: Fe gibt Elektronen ab", s2_klaeren_rost),
     ("s2_cu", [3], "KURZ KLÄREN", None, "Redox ohne Sauerstoffgas (Vorbereitung auf Station 1)", s2_klaeren_cu),
     ("s2_a57", [4], "AUFGABE", "Aufgaben 5–7", "Alltag erkennen und Rostschutz zuordnen",
      lambda: auftrag_folie("Jetzt: Aufgaben 5–7", [5, 6, 7], {}, "5: Beobachtung und Erklärung · 6: Redox-Detektiv · 7: Rostschutz zuordnen.")),
     ("s2_bespr", [4], "BESPRECHEN", None, "Die heiklen Fälle klären (Eis, Zucker, Verzinken)", s2_bespr),
     ("s2_ende", [5], "ABSCHLUSS", None, "Gruppen 1–6 für die Stationen einteilen; Reserve", s2_ende)],
 3: [("s3_heute", [0], "START", None, "Ziel und Ablauf zeigen", lambda: heute_folie(3)),
     ("s3_v2", [1], "VERSUCH", "Auswertung V2", "Rost-Ansätze auswerten, mit den Tipps vergleichen",
      lambda: beobachtung_folie("K1.3-V2", "Auswertung V2: Wo rostet es?", "Fazit: Wasser und Sauerstoff nötig, Salz beschleunigt – genau wie in der Theorie.")),
     ("s3_v6", [2], "DEMO", "Demo V6: Kältemischung", "Tipp, Messung, dabei exotherm/endotherm klären",
      lambda: versuch_folie("K1.3-V6", "Demo V6: Kältemischung", "Wird es mit Salz kälter, wärmer oder gleich? – Tipp per Handzeichen", "deutlich kälter: oft −10 °C oder tiefer; außen beschlägt der Becher")),
     ("s3_exo", [2], "KURZ KLÄREN", None, "exotherm und endotherm trennen", s3_exo),
     ("s3_diag", [2], "KURZ KLÄREN", None, "Energiediagramm lesen: Gipfel, Höhenunterschied", s3_diagramm),
     ("s3_stat", [3], "STATION", "Stationen erklären", "Rotation zeigen: 11 min je Runde",
      lambda: s3_stationen((1, 2, 3), "Theorieteil **Teil 3** lesen (Buch S. 32–34)", "Aufgaben 9–11")),
     ("s3_plaetze", [3], "STATION", None, "Plätze zeigen: Mobiltisch 1, Schülertisch, Wasserleitung 1 und 2", s3_plaetze),
     ("s3_sich", [3], "STATION", None, "Sicherheit an beiden Stationen laut vorlesen", s3_sicher),
     ("s3_v4", [4], "STATION", "Stationen V4 + V5", "Ablauf Station 1 zum Hinsehen während der Runden", s3_v4),
     ("s3_v5", [4], "STATION", None, "Ablauf Station 2", s3_v5),
     ("s3_a911", [4], "AUFGABE", None, "Aufgaben der freien Gruppe",
      lambda: auftrag_folie("Freie Gruppe: Aufgaben 9–11", [9, 10, 11], {11: "Skizze mit der Hand, dann Foto"}, "9: exo oder endo · 10: Diagramme lesen · 11: Lagerfeuer-Skizze + Erklärung.")),
     ("s3_erg", [5], "ABSCHLUSS", None, "Ergebnisse sichern: warm oder kalt; Reserve", s3_ergebnis)],
 4: [("s4_heute", [0], "START", None, "Ziel und Ablauf zeigen", lambda: heute_folie(4)),
     ("s4_weisst", [0], "WIEDERHOLEN", None, "Drei Fragen zur letzten Stunde (Lösung per Klick)", s4_weisst),
     ("s4_v7", [1], "DEMO", "Demo V7: Zuckerwürfel mit Lithiumchlorid", "Tipp, Demo (nur Lehrkraft zündet); Ergebnis per Klick",
      lambda: versuch_folie("K1.3-V7", "Demo V7: Was lässt den Zucker brennen?", "Welcher Würfel brennt mit Flamme? – Tipp (Aufgabe 13 a)", "Nur Würfel 2 (mit Lithiumchlorid) · Würfel 1 pur und Würfel 3 mit Kochsalz schmelzen und werden braun",
                                extra_sicher="Nur die Lehrkraft zündet. Würfel nur auf der Unterlegplatte, Abstand halten, Schutzbrille. Heißes Karamell nicht berühren. Löschen mit Becherglas, kein Wasser.")),
     ("s4_v7w", [1], "KURZ KLÄREN", "Kurz klären", "Lithiumchlorid = Katalysator; Kochsalz als Kontrolle; Enzyme im Körper", s4_v7w),
     ("s4_tausch", [2], "STATION", "Stationen: Gruppen tauschen", "Rotation zeigen; Gruppen 4–6 jetzt an die Stationen", s4_tausch),
     ("s4_a1314", [2], "AUFGABE", None, "Aufgaben der freien Gruppe (und aller nach dem Tausch)",
      lambda: auftrag_folie("Freie Gruppe: Aufgabe 14", [14], {14: "allein am Handy"}, "6 Aussagen: stimmt oder stimmt nicht – mit »Prüfen«.")),
     ("s4_auto", [3], "KURZ KLÄREN", "Autokatalysator", "Kein Filter, sondern Umwandlung; Kaltstart", s4_auto),
     ("s4_bespr", [3], "BESPRECHEN", "Besprechen", "Heikle Punkte A13/A14 klären", s4_bespr),
     ("s4_abg", [4], "ABGABE", "PDF abgeben", "Abgabe erklären; Tempo (Teil 5) ist freiwilliger Zusatz", s5_abgabe),
     ("s4_wicht", [5], "ABSCHLUSS", None, "Das Wichtigste in 5 Sätzen (Sätze 1–3); nur wenn Reserve bleibt", s5_wichtig),
     ("s4_wicht2", [5], "ABSCHLUSS", None, "Sätze 4–5", s5_wichtig2)],
}
for std, sp in SPEC.items():
    alle = sorted({i for x in sp for i in x[1]})
    assert alle == list(range(len(phasen(std)))), (std, alle)
    STEP[std] = [(x[3], x[1]) for x in sp if x[3]]

DANACH_STD = {1: "Stunde 2 am Do 05.11.: Test, dann Rost", 2: "Stunde 3 am Di 10.11.: Stationen", 3: "Stunde 4 am Do 12.11.: Katalysatoren", 4: "Ende der Einheit – ab Di 17.11. Säuren und Basen"}
KURZ = {"s1_einst": "Einstieg", "s1_mg": "Magnesium und Sauerstoff", "s2_rost": "Wie Eisen rostet", "s3_v4": "Station 1", "s3_plaetze": "Wo steht was?", "s4_v7w": "Was tat das Lithiumchlorid?", "s4_abg": "PDF abgeben", "s4_wicht2": "Das Wichtigste (2)", "s4_wicht": "Das Wichtigste (1)", "s4_tausch": "Gruppen tauschen", "s1_v1w": "Warum schwerer?", "s1_bespr": "Besprechen A1–A3", "s1_v2": "Ausblick Rost-Ansatz", "s2_bespr": "Besprechen A5–A7", "s2_cu": "Eisen in Kupfersulfat",
        "s2_ende": "Ausblick Stationen", "s3_exo": "exotherm – endotherm", "s3_diag": "Energiediagramm", "s3_sich": "Sicherheit", "s3_v5": "Station 2", "s3_a911": "Aufgaben 9–11",
        "s3_erg": "Ergebnisse sichern", "s4_weisst": "Weißt du noch?", "s4_a1314": "Aufgabe 14", "s4_bespr": "Besprechen A13/A14", "s1_heute": "Heute", "s2_heute": "Heute", "s3_heute": "Heute", "s4_heute": "Heute"}


def nachher(spec): return spec[3] or KURZ.get(spec[0], spec[0])


# ================================================================ Deck bauen
def bauen():
    global d
    d = B.Deck(True, "NAWI · 2AB · K1.3 Stoffumwandlung · Unterricht"); B.d = d
    d.section("Start")
    s = d.slide("Titelfolie", "Stoffumwandlung")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2:
            B.text(ph, ["NAWI · 2AB · K1.3 · Unterrichts-PPP: Redox, Energie, Katalysatoren · Buch S. 29–37",
                        [(f"Stand {I.STAND} · {VER} · TH-Unterricht (Ablauf) – zum Nachlesen: TH-Lehrer", {"size": 12})]], size=18, color=RGBColor.from_string("C9D6B0"))
    titel = s
    bedien = d.slide("Nur Titel", "Lehrer-Info · So führt die PPP", hidden=True)
    regie = {st["std"]: d.slide("Nur Titel", f"Regie · Stunde {st['std']} · {st['datum']}", hidden=True) for st in P.PLAN}
    material = d.slide("Nur Titel", "Lehrer-Info · Herrichten", hidden=True)
    stolper = d.slide("Nur Titel", "Lehrer-Info · Stolpersteine", hidden=True)
    quellen = d.slide("Nur Titel", "Lehrer-Info · Quellen", hidden=True)
    for std in (1, 2, 3, 4):
        d.section(f"Stunde {std}")
        sec = d.slide("Abschnitt", f"Stunde {std} · " + next(p["titel"] for p in P.PLAN if p["std"] == std).split(":")[0])
        for ph in sec.placeholders:
            if ph.placeholder_format.type == 2: ph.text_frame.text = D[std] + " · 2A und 2B"
        FOLIE[f"s{std}_sec"] = sec
        sp = SPEC[std]
        for j, (key, idx, art, schritt, wozu, fn) in enumerate(sp):
            s = fn(); FOLIE[key] = s
            nn = nachher(sp[j + 1]) if j + 1 < len(sp) else DANACH_STD[std]
            hinweis(s, art, minuten(std, idx), nn)
            check(s)
    d.section("Optional")
    zusatz = zusatz_folie()

    nr_von = {s.slide_id: i + 1 for i, s in enumerate(d.p.slides)}
    NR = lambda key: nr_von[FOLIE[key].slide_id]
    folgen = {std: [FOLIE[f"s{std}_sec"]] + [FOLIE[x[0]] for x in SPEC[std]] for std in SPEC}
    for st in P.PLAN:
        s = regie[st["std"]]; std = st["std"]
        rows = [["Minute", "Was läuft jetzt?", "Folie", "Wozu?"]]; t0 = 0
        zu = {}
        for key, idx, art, schritt, wozu, fn in SPEC[std]:
            for i in idx: zu.setdefault(i, []).append((key, wozu))
        for i, (name, mi, oh, a) in enumerate(st["phasen"]):
            keys = [x for x, _ in zu[i]]
            fol = ", ".join(dict.fromkeys(str(NR(x)) for x in keys))
            wz = zu[i][0][1] if a != "reserve" else "Puffer – nichts Neues beginnen"
            rows.append([f"{t0}–{t0 + mi}", name, fol, wz]); t0 += mi
        assert t0 == 50
        B.tb(s, X0, 0.98, W, 0.3, [[(f"2A und 2B · {st['datum']} · {st['titel']} · Benutzerdefinierte Präsentation »2AB {st['datum']} Stunde {std}«", {"bold": True, "color": TEAL})]], size=11, after=0)
        kn = P.KNAPP.get(std)
        hk = B.th("Wenn es knapp wird: " + kn, 10, W, pad=0.04) if kn else 0
        size = 11
        while size > 8 and B.table_h([0.8, 4.2, 0.7, 3.3], rows, size=size, minh=0.2) + hk > 3.8: size -= 1
        cells, h = B.auto_grid(s, X0, 1.28, [0.8, 4.2, 0.7, 3.3], rows, size=size, minh=0.2)
        if kn: B.tb(s, X0, 1.3 + h, W, hk, [[("Wenn es knapp wird: ", {"bold": True, "color": RUST}), (kn, {})]], size=10, after=0)
    B.tb(bedien, X0, 0.98, W, 4.1, [
        [("Eine Datei, beide Klassen (gleiche Termine). ", {"bold": True, "color": OLIVE}), ("F5 zeigt die sichtbaren Folien Stunde 1 bis 4. Pro Stunde: Bildschirmpräsentation → Benutzerdefinierte Bildschirmpräsentation → z. B. »2AB Do 05.11. Stunde 2«. Ohne Menü: Foliennummer tippen und Enter (Nummern stehen auf den Regiefolien).", {})],
        [("Jede Stunde läuft gleich: ", {"bold": True, "color": OLIVE}), ("»Heute« → »Lesen: …« (Buch und Abschnitte im Theorieteil, mit Leitfrage) → »Kurz klären« (nur Schwierigkeiten und Fehlvorstellungen, Antwort per Klick) → Versuch bzw. Demo (Tipp abgeben, dann Ergebnis per Klick) → Aufgaben → Besprechen (nur die heiklen Teilfragen).", {})],
        [("Unten links ", {"bold": True, "color": OLIVE}), ("steht immer, was JETZT läuft (mit Minuten aus dem Zeitplan) und was DANACH kommt.", {})],
        [("Stunde 3 und 4: ", {"bold": True, "color": OLIVE}), ("Stationen mit Rotation (11 min je Runde). Während die Stationen laufen, lässt du die Folie »Stationen: so wechselt ihr« stehen; die Ablauf-Folien V4/V5 sind zum Hinsehen gedacht.", {})],
        [("Nachschlagen statt zeigen: ", {"bold": True, "color": OLIVE}), ("Die vollständige Theorie (TH-Lehrer, 212 Folien), Theorieteil und private Lösungsmappe bleiben bestehen – hier steht nur das Wichtigste. Zusatzaufgaben: ausgeblendete letzte Folie.", {})],
        [("Zeitplan: ", {"bold": True, "color": OLIVE}), ("Die Zeiten stammen aus plan.py und stimmen mit der Spalte »Zeit« im Stundenvorschlag überein; Stunde 2 beginnt mit dem Test.", {})]],
        size=13, after=5)
    chk = [x for x in F.CHECKLISTE]
    B.tb(material, X0, 0.98, W, 4.1, ["· " + B.plain(x) for x in chk], size=11, after=3)
    B.tb(stolper, X0, 0.98, W, 4.1, [[(a + ": ", {"bold": True, "color": RED}), (B.plain(b), {})] for a, b in I.STOLPERSTEINE], size=13, after=5)
    B.tb(quellen, X0, 0.98, W, 4.1, ["· " + x for x in [
        "Inhalt und Zeiten: dieselbe Datenquelle wie Theorieteil, Arbeitsblatt, Versuchsseite und TH-Lehrer (werkzeug/quellen/2ab-k1-3-reaktionen: inhalt.py, theorie_daten.py, versuche_daten.py, plan.py).",
        "Lesestopps: Trauner NAWI 2, S. 29–37 (und S. 98); Abschnittstitel exakt wie im Theorieteil.",
        "Bilder: eigene Zeichnungen und Animationen (img, vs-img) – keine Buchabbildungen, darum auch als Datei weitergebbar.",
        "Gebaut und gerendert von Claude; Abspieltest in PowerPoint (Klick-Lösungen, GIF) steht aus."]], size=13, after=5)

    pres = d.p.part._element
    csl = etree.Element("{%s}custShowLst" % NS_P)
    anker = pres.find("{%s}defaultTextStyle" % NS_P)
    if anker is not None: anker.addprevious(csl)
    else: pres.append(csl)
    rid = {s.slide_id: el.get("{%s}id" % NS_R) for s, el in zip(d.p.slides, d.p.slides._sldIdLst)}
    nid = [0]

    def show(name, folien):
        cs = etree.SubElement(csl, "{%s}custShow" % NS_P, name=name, id=str(nid[0])); nid[0] += 1
        l = etree.SubElement(cs, "{%s}sldLst" % NS_P)
        for sl in folien: etree.SubElement(l, "{%s}sld" % NS_P).set("{%s}id" % NS_R, rid[sl.slide_id])

    show("2AB gesamt", [titel] + sum([folgen[std] for std in sorted(folgen)], []))
    for st in P.PLAN:
        show(f"2AB {st['datum']} Stunde {st['std']}", ([titel] if st["std"] == 1 else []) + folgen[st["std"]])
    show("Lehrer-Regiepläne", [bedien] + list(regie.values()))
    show("Optional: Für Schnelle", [zusatz])

    os.makedirs(ZIEL, exist_ok=True)
    name = f"{I.PRAEFIX}-TH-Unterricht-{VER}.pptx"
    d.save(os.path.join(ZIEL, name))
    sicht = [s for s in d.p.slides if s._element.get("show") != "0"]
    for x in PROBLEME: print("ÜBERLAUF", x)
    print("ok", name, len(d.p.slides._sldIdLst), "Folien,", len(sicht), "sichtbar ·", nid[0], "Präsentationen")
    for st in P.PLAN: print(" ", st["std"], st["datum"], "Folien", [nr_von[x.slide_id] for x in folgen[st["std"]]])


VER = "v2"
if __name__ == "__main__":
    bauen()
