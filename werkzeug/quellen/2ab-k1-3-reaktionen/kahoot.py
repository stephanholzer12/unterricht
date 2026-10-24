"""Klassenspiel K1.3: Quizfragen als Kahoot-/Blooket-Tabelle (Blätter Anleitung, Kahoot, Blooket).
Aufruf: python3 werkzeug/quellen/2ab-k1-3-reaktionen/kahoot.py <Zielordner>"""
import os, random, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import inhalt as I

OUT = sys.argv[1]; VER = "v1"
wb = Workbook()
a = wb.active; a.title = "Anleitung"
for r in [["Klassenspiel K1.3 Stoffumwandlung – " + str(len(I.QUIZ)) + " Fragen (dieselben wie im Handy-Quiz)"],
          ["Kahoot: kahoot.com → Erstellen → Quiz → »Tabelle importieren« → Blatt »Kahoot« als eigene .xlsx speichern und hochladen."],
          ["Blooket: blooket.com → Create → »Spreadsheet Import« → Blatt »Blooket« als CSV speichern und hochladen."],
          ["Gespielt wird am Lehrer-PC am Beamer. Kadetten nur mit Spitznamen, kein Konto. Ohne Handy: Team-Modus (ein Gerät pro Gruppe) oder ABCD-Karten."],
          ["Längen geprüft: Frage ≤ 120 Zeichen, Antwort ≤ 75 Zeichen, Zeit 20 s (Rechnung 30 s); richtige Antwort steht gemischt an Platz 1–4."],
          [f"Stand {I.STAND} · {VER}"]]:
    a.append(r)
a["A1"].font = Font(bold=True, size=13, color="3B4A2F"); a.column_dimensions["A"].width = 130
k = wb.create_sheet("Kahoot"); bl = wb.create_sheet("Blooket")
head = ["Nr.", "Frage (max. 120 Zeichen)", "Antwort 1", "Antwort 2", "Antwort 3", "Antwort 4", "Zeit (s)", "Richtig (1–4)", "Thema", "Buch"]
k.append(head)
bl.append(["Question #", "Question Text", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Time Limit (sec)", "Correct Answer(s)"])
rng = random.Random(13)
pos_used = []
for n, (thema, q, ans, e, p) in enumerate(I.QUIZ, 1):
    assert len(q) <= 120, (n, len(q), q)
    for x in ans: assert len(x) <= 75, (n, x)
    order = ans[:]; rng.shuffle(order)
    if pos_used[-2:] == [order.index(ans[0])] * 2:  # nicht dreimal hintereinander gleiche Position
        order = order[1:] + order[:1]
    pos = order.index(ans[0]) + 1; pos_used.append(pos - 1)
    zeit = 30 if any(c.isdigit() for c in q) and ("Wie lange" in q or "°C" in q) else 20
    k.append([n, q] + order + [zeit, pos, thema, p]); bl.append([n, q] + order + [zeit, pos])
for ws in (k, bl):
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="3B4A2F")
    for col, w in zip("ABCDEFGHIJ", [5, 70, 32, 32, 32, 32, 9, 12, 14, 11]): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
name = f"{I.PRAEFIX}-Kahoot-{VER}.xlsx"
wb.save(os.path.join(OUT, name))
from collections import Counter
print(name, "Positionen:", Counter(p + 1 for p in pos_used))
