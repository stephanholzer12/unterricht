"""Word-Dateien NAWI 4AB · K2.1–2.3 Schwingungen und Wellen: TH-AB, TH-Loesung, TH-Fachwissen, VS-AB, VS-Loesung, VS-Fachwissen.
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/4ab-k2-1-wellen/word.py <Zielordner>
Layout wie Muster-VS-AB (Kopf mit Wappen, Titelzeile mit olivgrüner Linie). Daten aus inhalt.py, versuche_daten.py, fachwissen.py."""
import json, os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, HERE)
import inhalt as I
import versuche_daten as VD
import herrichten as H
import fachwissen as F

OUT = sys.argv[1]
VER = I.V
PRE = I.PRAEFIX
OLIVE, TEAL, RED, MUTED, RUST, INK = RGBColor(0x3B, 0x4A, 0x2F), RGBColor(0x2E, 0x7D, 0x8C), RGBColor(0xC0, 0x00, 0x00), RGBColor(0x5E, 0x66, 0x52), RGBColor(0x8C, 0x3B, 0x2E), RGBColor(0x1F, 0x2A, 0x1C)
THEORY, BOOKRED = RGBColor(0x6E, 0x7F, 0x4E), RGBColor(0xA3, 0x3A, 0x2B)
LOGO = os.path.join(ROOT, "werkzeug/logo.png")
UNIT = os.path.join(ROOT, I.ORDNER) + "/"
IMG = UNIT + "vs-img/"
EINH = [u for u in json.load(open(os.path.join(ROOT, "werkzeug/einheiten.json"))) if u["ordner"] == I.ORDNER][0]


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcol); tcPr.append(s)


def borders(table, color="B9C4A6", size=6):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:color"), color); b.append(el)
    tblPr.append(b)


def noborders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblPr.append(b)


def pbottom(p, color="3B4A2F", sz=12):
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "1"); b.set(qn("w:color"), color); bd.append(b); pPr.append(bd)


def run(p, text, size=10.5, bold=False, color=INK, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color; r.font.name = "Calibri"
    return r


def md_runs(p, text, size=10.5, color=INK, bold=False):
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > pos: run(p, text[pos:m.start()], size, bold, color)
        run(p, m.group(1), size, True, color); pos = m.end()
    if pos < len(text): run(p, text[pos:], size, bold, color)


def para(doc, text="", size=10.5, bold=False, color=INK, italic=False, after=2, before=0, align=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.05
    if keep: pf.keep_with_next = True
    if align: p.alignment = align
    if text: (md_runs(p, text, size, color, bold) if not italic else run(p, text, size, bold, color, italic))
    return p


def label(p, lab, text, size=10.5, color=INK, labcol=OLIVE):
    run(p, lab, size, True, labcol); md_runs(p, text, size, color)


def widths(table, cms):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None: lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), cms): gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for c, w in zip(row.cells, cms): c.width = Cm(w)


def cellp(cell, text="", size=10, bold=False, color=INK, align=None):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    if align: p.alignment = align
    if text: md_runs(p, text, size, color, bold)
    return p


def lines(doc, n):
    for _ in range(n):
        p = para(doc, "", after=0); p.paragraph_format.line_spacing = 1.0; pbottom(p, "B9C4A6", 4)
        p.paragraph_format.space_before = Pt(6)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_doc(kopf, titel, fuss, quer=False):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.4); sec.footer_distance = Cm(0.6)
    t = doc.add_table(rows=1, cols=2); noborders(t); widths(t, [1.6, 15.8])
    t.cell(0, 0).paragraphs[0].add_run().add_picture(LOGO, width=Cm(1.2))
    c = t.cell(0, 1); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cellp(c, kopf, 8.5, True, TEAL); p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); run(p2, titel, 16, True, OLIVE)
    p = para(doc, "", after=4); pbottom(p, "3B4A2F", 14)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(fp, fuss + " · Seite ", 8, False, MUTED)
    r = fp.add_run(); r.font.size = Pt(8); r.font.color.rgb = MUTED
    for tg, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tg:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tg); r._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
    return doc


def box(doc, title, rows_fn, fill="EEF3E4", line="B9C4A6"):
    t = doc.add_table(rows=1, cols=1); borders(t, line, 8); widths(t, [17.4]); c = t.cell(0, 0); shade(c, fill)
    cellp(c, title, 10, True, OLIVE); rows_fn(c)
    para(doc, "", after=2)
    return t


def table(doc, head, rows, cms, size=9.5, headfill="3B4A2F", zebra=True):
    t = doc.add_table(rows=1, cols=len(head)); borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; shade(c, headfill); cellp(c, h, size, True, RGBColor(0xFF, 0xFF, 0xFF))
    for k, r in enumerate(rows):
        cells = t.add_row().cells
        for i, x in enumerate(r):
            if isinstance(x, tuple): cellp(cells[i], x[0], size, x[1] if len(x) > 1 else False, x[2] if len(x) > 2 else INK)
            else: cellp(cells[i], str(x), size)
            if zebra and k % 2 == 1: shade(cells[i], "F6F9F0")
    widths(t, cms)
    return t


# ================================================================ gemeinsame Kästen
def fahrplan_box(doc):
    z = EINH["fahrplan"]["zeilen"]
    je = {}
    for r in z:
        if r.get("klasse") and r["wann"].startswith("Stunde"): je[r["klasse"]] = je.get(r["klasse"], 0) + 1
    def fill(c):
        p = c.add_paragraph(); run(p, "Dafür sind eingeplant: " + " · ".join(f"{k}: {n} Stunden" for k, n in je.items()) + ". ", 9.5, True, OLIVE); run(p, EINH["fahrplan"]["text"], 9.5)
        for k in list(je) + [None]:
            zz = [r for r in z if r.get("klasse") == k]
            if not zz: continue
            p = c.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1); run(p, k or "Beide Klassen", 9.5, True, OLIVE)
            it = c.add_table(rows=1, cols=5); borders(it, "D3DCC2", 4)
            for i, h in enumerate(["Wann", "Das machen wir", "Aufgaben · Zeit", "Wie", "Abgabe"]): cellp(it.rows[0].cells[i], h, 8, True, OLIVE)
            for r in zz:
                cs = it.add_row().cells
                cellp(cs[0], r["wann"] + (" · " + r["datum"] if r.get("datum") else ""), 8, True); cellp(cs[1], r["lesen"], 8)
                cellp(cs[2], r.get("aufgaben", "–") + (" · " + r["zeit"] if r.get("zeit") else ""), 8); cellp(cs[3], r.get("wie") or "–", 8); cellp(cs[4], r.get("abgabe", "–"), 8)
            widths(it, [2.2, 4.6, 4.2, 4.2, 1.8])
    box(doc, "Dein Fahrplan", fill)


def lehrplan_box(doc):
    def fill(c):
        for lab, x in F.LEHRPLAN:
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1); label(p, lab, x, 9)
    box(doc, "Lehrplanbezug", fill)


def stundenvorschlag(doc):
    para(doc, "Stundenvorschlag je Klasse · Jahresplan: 6 Stunden · 4AK 5 Termine (2 × nur 30 min), 4BK 4 Termine + optional 18.11.", 11, True, OLIVE, after=2, before=4, keep=True)
    for k, rows in F.STUNDENVORSCHLAG.items():
        para(doc, k, 10, True, TEAL, after=1, keep=True)
        table(doc, ["Std.", "Theorie", "Kern-Versuch", "Aufgaben", "Zeit (min)", "Methode (ohne Handy)"], rows, [2.0, 3.4, 3.6, 2.4, 2.2, 3.8], 8.5)
    para(doc, "KERN = Pflicht für alle (Aufgaben 1–13, Versuche K2.1-V1 bis K2.2-V2) · OPTIONAL = Zusatz 14–17 (4AK macht 16–17 in Stunde 4) · keine Hausübungen. "
              "Jede Stunde = 50 min, Summe je Stunde geprüft (plan.py). Besprechung: Aufgaben mit Tag BESPRECHEN (in der Lösungsmappe).", 8.5, False, MUTED, after=3)


def stolpersteine(doc):
    para(doc, "Die 5 größten Stolpersteine", 11, True, OLIVE, after=2, before=4, keep=True)
    table(doc, ["Stolperstein", "Richtig ist …"], [[a, b] for a, b in I.STOLPERSTEINE], [4.4, 13.0], 9)


def src_color(src):
    return BOOKRED if src.startswith("Versuch im Buch") else (TEAL if src.startswith("Buch-Aufgabe") else THEORY)


# ================================================================ TH-AB und TH-Loesung
def optionen(t, f):
    """Antwortmöglichkeiten wie am Handy gemischt (feste Reihenfolge nur, wenn das 4. Element False ist) – richtige Antwort nie immer zuerst."""
    import random
    if len(f) > 3 and f[3] is False: return list(f[2])
    o = list(f[2]); random.Random(t["id"] + f[0]).shuffle(o)
    if o[0] == f[2][0] and len(o) > 1: o = o[1:] + o[:1]
    return o


def aufgabe(doc, t, loesung):
    tb = doc.add_table(rows=1, cols=2); noborders(tb); widths(tb, [12.6, 4.8])
    p = cellp(tb.cell(0, 0), "", 11); run(p, f"Aufgabe {t['nr']}  ", 11, True, OLIVE); run(p, t["title"], 11, True, INK)
    p = cellp(tb.cell(0, 1), "", 8.5, align=WD_ALIGN_PARAGRAPH.RIGHT); run(p, t["src"], 8.5, True, src_color(t["src"]))
    p2 = tb.cell(0, 1).add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p2.paragraph_format.space_after = Pt(0)
    run(p2, f"{t['tag']} · ca. {t['zeit']} min" + (" (im Versuch)" if t.get("invs") else "") + (" · BESPRECHEN" if loesung and t.get("besprechen") else ""), 8.5, False, MUTED)
    para(doc, t["q"], 10, after=2)
    if t.get("hint"): para(doc, t["hint"], 8.5, False, MUTED, italic=True, after=2)
    if t.get("img"):
        p = para(doc, "", after=2, align=WD_ALIGN_PARAGRAPH.CENTER); p.add_run().add_picture(UNIT + t["img"], width=Cm(11))
    if loesung:
        for head, items in t["sol"]:
            for x in items:
                p = para(doc, "", after=1); run(p, ("" if head in ("Lösung", "Beispiel", "Bewertung") else head + " "), 9.5, True, RED); run(p, x, 9.5, False, RED)
        para(doc, "", after=3)
        return
    sel = [f for f in t.get("fields", []) if len(f) > 2]
    if t["type"] == "multi" and sel and len(sel) == len(t["fields"]) and len({tuple(f[2]) for f in sel}) == 1:
        opts = sel[0][2]
        rows = [[f[1], "   ".join("☐ " + o for o in opts)] for f in sel]
        table(doc, ["", "Kreuze an"], rows, [9.4, 8.0], 9, "6E7F4E", zebra=False)
    else:
        for f in t.get("fields", []):
            if len(f) > 2:
                p = para(doc, "", after=1); run(p, f[1] + ":  ", 9.5, True); run(p, "   ".join("☐ " + o for o in optionen(t, f)), 9)
            else:
                p = para(doc, "", after=0); run(p, f[1], 9.5, True)
                lines(doc, 2 if len(f[1]) < 60 else 3)
        if t["type"] == "sketch":
            para(doc, "Skizze bzw. Handschrift ins Heft – Foto im Arbeitsblatt am Handy einfügen.", 9, False, MUTED, italic=True)
    para(doc, "", after=3)


def th_docs():
    fuss = f"K2.1–2.3 Schwingungen und Wellen · Stand {I.STAND} · {VER}"
    for loesung in (False, True):
        doc = new_doc(("LÖSUNGEN · " if loesung else "") + "ARBEITSBLATT · NAWI · 4. JAHRGANG (4AK, 4BK) · K2.1–2.3",
                      "Schwingungen und Wellen" + (" · Lösungen" if loesung else ""), fuss + (" · Lösungen" if loesung else ""))
        if loesung:
            lehrplan_box(doc)
        if not loesung:
            p = para(doc, "", after=2); run(p, "Name: ", 10, True); run(p, "__________________________   ", 10); run(p, "Klasse: ", 10, True); run(p, "______", 10)
            p = para(doc, "", after=4); run(p, "Am Handy: ", 9, True, OLIVE); run(p, I.URL + "aufgaben.html", 9, False, TEAL)
        fahrplan_box(doc)
        if loesung:
            stundenvorschlag(doc); stolpersteine(doc)
        last = None
        for t in I.T:
            if t["h"] != last:
                para(doc, I.GRP[t["h"]], 11.5, True, OLIVE, after=3, before=6, keep=True).paragraph_format.keep_with_next = True
                last = t["h"]
            aufgabe(doc, t, loesung)
        if loesung:
            para(doc, "Partnercheck-Karten (Aufgabe 5)", 11, True, OLIVE, after=2, before=4, keep=True)
            table(doc, ["Frage", "Lösung"], [[q, a] for q, a in I.PARTNERKARTEN], [8.7, 8.7], 9)
        name = f"{PRE}-TH-{'Loesung' if loesung else 'AB'}-{VER}.docx"
        doc.save(os.path.join(OUT, name)); print(name)


# ================================================================ VS-AB und VS-Loesung
def versuch_titel(doc, v):
    p = para(doc, "", after=3, keep=True); p.paragraph_format.page_break_before = True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    run(p, v["nr"] + "  ", 13, True, OLIVE); run(p, v["titel"].upper(), 12, True, OLIVE)
    tag = ("KERN · Station " + v["station"]) if v.get("station") else ("KERN · Demo" if v.get("demo") else ("KERN" if v["kern"] else "ZUSATZ"))
    run(p, "\t" + tag, 9, True, TEAL); pbottom(p, "B9C4A6", 6)


def mat_table(doc, v):
    para(doc, "Geräte & Chemikalien", 10.5, True, OLIVE, after=1, keep=True)
    m = v["mat"]; n = (len(m) + 2) // 3
    t = doc.add_table(rows=n, cols=3); borders(t, "B9C4A6", 6)
    for i, mm in enumerate(m):
        c = t.cell(i // 3, i % 3); p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        f = IMG + f"i_{mm[1]}.png"
        if os.path.exists(f): p.add_run().add_picture(f, height=Cm(0.55))
        run(p, "  " + mm[0], 8.5)
    widths(t, [5.8, 5.8, 5.8])


def beob_table(doc, v, sol=False):
    b = v["beob"]; para(doc, b["titel"], 10.5, True, OLIVE, after=1, before=3, keep=True)
    ng = max(len(z) for z in b["zeilen"]); nc = len(b["kopf"])
    rows = []
    for i, z in enumerate(b["zeilen"]):
        r = list(z) + [""] * (nc - len(z))
        if sol and v.get("erwartet_tab"): r[ng:] = [(x, False, RED) for x in v["erwartet_tab"][i]]
        elif sol and i < len(v["erwartet"]): r[ng] = (v["erwartet"][i][1], False, RED)
        rows.append(r)
    cms = {2: [5.0, 12.4], 3: [4.6, 6.4, 6.4]}.get(nc) or [2.0] + [15.4 / (nc - 1)] * (nc - 1)
    t = table(doc, b["kopf"], rows, cms, 9, "6E7F4E", zebra=False)
    if not sol:
        for row in t.rows[1:]: row.height = Cm(0.6)


def versuch_seite(doc, v, loesung=False):
    versuch_titel(doc, v)
    para(doc, v["intro"], 10, False, INK, italic=True, after=2)
    p = para(doc, "", after=2); label(p, "Sicherheit: ", v["sicher"], 9, INK, RUST)
    p = para(doc, "", after=2)
    label(p, "Versuch im Schulbuch: " if v["buch"]["versuch"] else "Theorie im Schulbuch: ", v["buch"]["kap"], 9, INK, BOOKRED if v["buch"]["versuch"] else TEAL)
    p = para(doc, "", after=2); label(p, "Zeit: ", f"ca. {v['zeit']} min · {v['form']}", 9, INK, MUTED)
    mat_table(doc, v)
    p = para(doc, "", after=1, before=2); label(p, "Entsorgung und Reinigung: ", v["entsorgung"], 8.5)
    para(doc, "Durchführung", 10.5, True, OLIVE, after=1, keep=True)
    t = doc.add_table(rows=1, cols=2); noborders(t); widths(t, [10.8, 6.6])
    c = t.cell(0, 0); first = True
    for i, s in enumerate(v["schritte"], 1):
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.5)
        run(p, f"{i}  ", 9, True, OLIVE); run(p, s, 9)
    pc = t.cell(0, 1).paragraphs[0]; pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.add_run().add_picture(IMG + f"d_{v['skizze']}.png", width=Cm(4.6 if loesung else 6.2))
    beob_table(doc, v, loesung)
    for key, tit in [("erk", "Erkenntnis"), ("interp", "Interpretation"), ("schluss", "Schlussfolgerung")]:
        para(doc, tit, 10.5, True, OLIVE, after=0, before=2, keep=True)
        para(doc, v[key], 9, False, MUTED, after=1, keep=True)
        if loesung: para(doc, v[key + "_l"], 9, False, RED, after=1)
        else: lines(doc, 2)
    if loesung:
        def fill(cc):
            for lab, x in [("Form & Dauer: ", f"{v['form']} · ca. {v['zeit']} min · Jahresplan: {v['jp']}"),
                           ("Vorbereitung und Tipps: ", " ".join(v["lehrer"])), ("Typische Fehler: ", " ".join(v["fehler"])), ("Quelle: ", v["quelle"])]:
                p = cc.add_paragraph(); p.paragraph_format.space_after = Pt(0); label(p, lab, x, 8)
        box(doc, "LEHRER-INFO " + v["nr"], fill, "FBF1E6", "E0B48A")


def ablauf(doc):
    para(doc, "So laufen die Versuche", 11, True, OLIVE, after=2, before=2, keep=True)
    rows = [[v["nr"], v["titel"], v["wann"], v["form"], "KERN" if v["kern"] else "ZUSATZ"] for v in VD.VS]
    table(doc, ["ID", "Versuch", "Wann", "Form", "Einsatz"], rows, [1.9, 5.2, 4.6, 3.9, 1.8], 8.5)
    p = para(doc, "", after=2, before=3)
    run(p, "8 Gruppen (ca. 26 Kadetten: Gruppe 1–6 zu dritt, 7–8 zu viert). ", 9.5, True); run(p, "Alle Versuche laufen an EINEM Aufbau vorne: Die Lehrkraft baut auf bzw. zeigt, alle Gruppen beobachten und messen gleichzeitig (beim Pendel zwei Zeitnehmer je Gruppe, die Klasse zählt laut mit). Messwerte und Erkenntnis gleich eintragen, Interpretation und Schlussfolgerung in der Aufgabenzeit derselben Stunde – keine Hausübung. Protokoll am Handy: ", 9.5)
    run(p, I.URL + "versuche.html", 9, False, TEAL)
    p = para(doc, "", after=2); run(p, "So wertest du jeden Versuch aus: ", 9.5, True, OLIVE); run(p, "Erkenntnis – Was habt ihr herausgefunden? · Interpretation – Warum ist das so? Erklärt mit dem Modell (Pendel, Welle, Teilchen). · Schlussfolgerung – Was bedeutet das für Alltag, Sport oder Bundesheer?", 9.5)
    p = para(doc, "", after=2); run(p, "Sicherheit: ", 9.5, True, RUST); run(p, " · ".join(a + " – " + b for a, b in VD.SICHERHEIT), 9)


def vs_docs():
    fuss = f"K2.1–2.3 Schwingungen und Wellen · Versuche · Stand {I.STAND} · {VER}"
    doc = new_doc("VERSUCHE · NAWI · 4. JAHRGANG (4AK, 4BK) · K2.1–2.3", "Versuche zu Schwingungen und Wellen", fuss)
    fahrplan_box(doc); ablauf(doc)
    for v in VD.VS: versuch_seite(doc, v)
    f1 = f"{PRE}-VS-AB-{VER}.docx"; doc.save(os.path.join(OUT, f1))
    doc = new_doc("LÖSUNGEN VERSUCHE · NAWI · 4. JAHRGANG · K2.1–2.3", "Versuche zu Schwingungen und Wellen · Lösungen & Lehrer-Infos", fuss + " · Lösungen")
    lehrplan_box(doc); ablauf(doc)
    for v in VD.VS: versuch_seite(doc, v, loesung=True)
    f2 = f"{PRE}-VS-Loesung-{VER}.docx"; doc.save(os.path.join(OUT, f2))
    print(f1, f2)


# ================================================================ VS-Fachwissen
def vs_fachwissen():
    fuss = f"K2.1–2.3 Schwingungen und Wellen · Versuche · Stand {I.STAND} · {VER} · Fachwissen"
    doc = new_doc("NUR FÜR DICH · NAWI · 4. JAHRGANG · K2.1–2.3", "Versuche – Herrichten & Fachwissen", fuss)
    rows, slots = H.tabelle()
    para(doc, "Herrichte-Liste · 8 Gruppen (ca. 26 Kadetten) · ein gemeinsamer Aufbau – nur Handys je Gruppe", 10.5, True, OLIVE, after=2)
    nrs = [v["nr"] for v in VD.VS]
    head = ["Gerät / Material", "Fundort"] + [n.replace("K", "") for n in nrs] + ["ges.", "Best.", "fehlt"]
    body = []
    for r in rows:
        body.append([r["name"], (r["ort"], False, TEAL)] + [str(r["per"].get(n, "")) for n in nrs] + [str(r["gesamt"]), "–" if r["bestand"] is None else str(r["bestand"]),
                    (r["fehlt"], True, RED if r["fehlt"] not in ("–", "") else INK)])
    t = table(doc, head, body, [4.6, 4.2] + [1.3] * len(nrs) + [0.9, 0.9, 1.6], 7.5)
    for row in t.rows: row.height = Cm(0.4)
    para(doc, "ges. = höchster gleichzeitiger Bedarf (die Versuche laufen nie gleichzeitig). Bestand »–« = kein Inventargerät (Handys der Kadetten, Lehrer-PC). Mengen:", 8, False, MUTED, italic=True, before=2)
    for k, x in H.MENGEN.items():
        p = para(doc, "", after=0); run(p, "·  " + k + ": ", 8.5, True, OLIVE); run(p, x, 8.5)
    para(doc, "Checkliste »Vor der Stunde«", 10.5, True, OLIVE, after=1, before=4)
    for x in F.CHECKLISTE:
        p = para(doc, "", after=1); run(p, "☐  ", 10, True, OLIVE); run(p, x, 9)
    page_break(doc)
    para(doc, "Was hinter den Versuchen steckt", 13, True, OLIVE, after=2)
    for v in VD.VS:
        para(doc, f"{v['nr']} {v['titel']} · {v['form']}", 10.5, True, OLIVE, after=1, before=4, keep=True)
        for x in v["fach"]: para(doc, x, 9.5, after=1)
        p = para(doc, "", after=1); label(p, "Erwartet: ", " · ".join(a + ": " + b for a, b in v["erwartet"]), 9)
        p = para(doc, "", after=1); label(p, "Vorbereitung und Ablauf: ", " ".join(v["lehrer"]), 9)
        p = para(doc, "", after=1); label(p, "Nachbereitung und Wegräumen: ", v["entsorgung"], 9)
        p = para(doc, "", after=1); label(p, "Vorversuch (offen, nur Stephan): ", v["vorversuch"], 9, INK, RUST)
        p = para(doc, "", after=1); label(p, "Typische Fehler: ", " ".join(v["fehler"]), 9, INK, RUST)
        p = para(doc, "", after=1); label(p, "Quelle: ", v["quelle"] + " · Jahresplan: " + v["jp"], 8.5, MUTED, MUTED)
    f = f"{PRE}-VS-Fachwissen-{VER}.docx"; doc.save(os.path.join(OUT, f)); print(f)


# ================================================================ TH-Fachwissen
def th_fachwissen():
    fuss = f"K2.1–2.3 Schwingungen und Wellen · Stand {I.STAND} · {VER} · Fachwissen"
    doc = new_doc("NUR FÜR DICH · NAWI · 4. JAHRGANG · K2.1–2.3", "Schwingungen und Wellen – Fachwissen für die Lehrkraft", fuss)
    para(doc, "Stundenkarten", 13, True, OLIVE, after=2)
    for tit, steps, mat, ohne in F.STUNDENKARTEN:
        def fill(c, steps=steps, mat=mat, ohne=ohne):
            for i, s in enumerate(steps, 1):
                p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0.5); run(p, f"{i}  ", 9, True, OLIVE); run(p, s, 9)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); label(p, "Material: ", mat, 8.5)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); label(p, "Ohne Handy: ", ohne, 8.5)
        box(doc, tit, fill)
    p = para(doc, "", after=1); label(p, "Links: ", " · ".join(I.URL + k + ".html" for k in ("theorie", "aufgaben", "versuche", "quiz")), 8, TEAL)
    page_break(doc)
    para(doc, "Die Einheit auf einen Blick", 13, True, OLIVE, after=2)
    table(doc, ["Teil · Stunde", "Kernaussage", "Material"], F.AUF_EINEN_BLICK, [3.2, 8.8, 5.4], 9)
    lehrplan_box(doc); stundenvorschlag(doc)
    for kap in F.KAPITEL:
        para(doc, kap["titel"], 12, True, OLIVE, after=2, before=6, keep=True)
        for x in kap["text"]: para(doc, x, 10, after=2)
        if kap.get("bild"):
            pp = para(doc, "", after=2, align=WD_ALIGN_PARAGRAPH.CENTER); pp.add_run().add_picture(UNIT + kap["bild"], width=Cm(12))
        if kap.get("nur"):
            def fill(c, x=kap["nur"]):
                for y in x:
                    p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1); md_runs(p, y, 9.5)
            box(doc, "Nur für dich", fill, "ECECEC", "BDBDBD")
        if kap.get("stolper"):
            def fill(c, x=kap["stolper"]):
                p = c.add_paragraph(); md_runs(p, x, 9.5)
            box(doc, "Stolperstein", fill, "FBF1E6", "E0B48A")
    para(doc, "Rechenbeispiele", 12, True, OLIVE, after=2, before=6, keep=True)
    for x in F.RECHNEN: para(doc, x, 10, after=2)
    stolpersteine(doc)
    para(doc, "Häufige Fehlvorstellungen", 12, True, OLIVE, after=2, before=6, keep=True)
    table(doc, ["Fehlvorstellung", "Richtig ist", "So zeigst du es"], F.FEHLVORSTELLUNGEN, [4.8, 6.6, 6.0], 9)
    para(doc, "Fragen, die Kadetten stellen – mit Antworten", 12, True, OLIVE, after=2, before=6, keep=True)
    for q, a in F.FRAGEN:
        p = para(doc, "", after=0, keep=True); run(p, q, 10, True)
        para(doc, a, 9.5, after=3)
    para(doc, "Videos – worauf achten?", 12, True, OLIVE, after=2, before=6, keep=True)
    for k, (t_, q, d, u, *_) in I.VIDEOS.items():
        p = para(doc, "", after=0); run(p, t_ + " · " + q + ("" if d == "kurz" else " · " + d) + "  ", 9.5, True); run(p, u, 8.5, False, TEAL)
        para(doc, F.VIDEO_NOTIZ[k], 9, after=2)
    para(doc, "Ausblick", 12, True, OLIVE, after=2, before=6, keep=True)
    para(doc, F.AUSBLICK, 10, after=2)
    para(doc, "Quellen – geprüfte Bücher und was übernommen wurde", 12, True, OLIVE, after=2, before=6, keep=True)
    for x in F.QUELLEN:
        p = para(doc, "", after=1); run(p, "·  ", 9.5, True, OLIVE); md_runs(p, x, 9)
    f = f"{PRE}-TH-Fachwissen-{VER}.docx"; doc.save(os.path.join(OUT, f)); print(f)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    th_docs(); vs_docs(); vs_fachwissen(); th_fachwissen()
