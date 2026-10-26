"""Unterrichts-PPP K2.1–2.3 (kurz, zum Erklären im Saal) – v3 aus ChatGPTs Fassung v2.
Grundlage: 26-27-4AB-K2.1-2.3-Wellen-TH-Unterricht-v2.pptx (ChatGPT, 25.10.2026, liegt als vorlage_unterricht_v2.pptx daneben).
Übernimmt deren Folien unverändert im Kern und ergänzt nach Claudes Prüfung (Rückmeldung v1):
 · »Buch S.« statt »Theorie S.« (Regelwerk), Aufgaben-Folien im grünen Layout »Aufgabe«
 · Start: ausgeblendete Lehrer-Infos direkt nach dem Titel, Folie »Dein Lernpfad« mit QR-Codes
 · je Teil eine Folie »Heute« (Ziel, Aufgaben mit Zeit, Arbeitsblatt-QR) + ausgeblendete Lehrer-Info
   (Kernaussage · Leitfrage · Kadetten tun · Sicherung, Zeitstopp und Kürzung) – aus plan.py und inhalt.py
 · Ablauf je Klasse aus plan.py (mit 3 min Klassenbuch), benutzerdefinierte Präsentationen je Klasse und Stunde
 · keine Notizen: Quellen auf einer ausgeblendeten Folie, große Zusatzbeschriftungen bei kleinen Bildtexten
Aufruf: python3 unterricht_ppp.py <vorlage_unterricht_v2.pptx> <Zielordner>"""
import sys, os, copy, re
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
VORLAGE, ZIEL = sys.argv[1], sys.argv[2]
sys.argv = [sys.argv[0], "-", "-"]          # bauen_ppp liest MASTER/OUT aus argv – hier nur die Zeichenhilfen
import bauen_ppp as B
import inhalt as I, plan as P
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

VER = I.V if I.V >= "v3" else "v3"
p = Presentation(VORLAGE)
S0 = list(p.slides)                           # Originalfolien 1..34 (Index 0..33)
O = lambda n: S0[n - 1]                       # Originalfolie n
LAY = {l.name: l for l in p.slide_layouts}
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
OLIVE, INK, TEAL, RUST, MUTED, WHITE, GREEN = B.OLIVE, B.INK, B.TEAL, B.RUST, B.MUTED, B.WHITE, B.GREEN

# ------------------------------------------------------------ 1 Texte der Vorlage nachziehen
for sl in S0:
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if "Theorie S." in r.text: r.text = r.text.replace("Theorie S.", "Buch S.")
                if r.text.startswith("Theorie S"): r.text = r.text.replace("Theorie", "Buch", 1)


def setze(sl, name, neu):
    sh = [x for x in sl.shapes if x.name == name][0]
    tf = sh.text_frame; runs = [r for pa in tf.paragraphs for r in pa.runs]
    runs[0].text = neu
    for r in runs[1:]: r.text = ""
    for pa in list(tf.paragraphs)[1:]:
        if not "".join(r.text for r in pa.runs): pa._p.getparent().remove(pa._p)


setze(O(13), "work", "6 min (4BK: 5 min, 2 Fragen je Kadett) · Erst antworten, dann Lösung zeigen · Rollen wechseln")
setze(O(28), "pdf", "Beide PDFs in LMS: 4AK 18.11., 4BK 13.11. Versuchs-PDF kommt vom Protokollführer (WhatsApp/Teams).")

# Aufgaben-Folien im Layout »Aufgabe« (hellgrün), Regelwerk Abschnitt 3 der PPP-Gestaltung
for n in (6, 13, 17, 21, 25, 26, 27):
    rel = [r for r in O(n).part.rels.values() if r.reltype.endswith("/slideLayout")][0]
    rel._target = LAY["Aufgabe"].part

# Notizen einsammeln (Quellen) und leeren – das Regelwerk sieht keine Notizen vor
QUELLEN = []
for sl in S0:
    if sl.has_notes_slide:
        for z in sl.notes_slide.notes_text_frame.text.splitlines():
            z = z.strip().replace("Theorie S.", "Buch S.")
            if z and z not in QUELLEN: QUELLEN.append(z)
        sl.notes_slide.notes_text_frame.text = ""

# große Zusatzbeschriftungen, wo die Bildtexte für den langen Saal zu klein sind
B.tb(O(4), 0.3, 4.16, 5.6, 0.34, [[("Außen: Umkehrpunkte · Mitte: Gleichgewichtslage", {"bold": True, "color": TEAL})]], size=16, after=0)
for sh in O(10).shapes:            # Legende der Dämpfungsgrafik groß nachsetzen
    if sh.shape_type == 13:
        B.tb(O(10), sh.left / 914400, (sh.top + sh.height) / 914400 - 0.05, sh.width / 914400, 0.38,
             [[("blau: schwach gedämpft · rot: stark gedämpft", {"bold": True, "color": TEAL})]], size=16, after=0); break


# ------------------------------------------------------------ 2 neue Folien
FUSS = [copy.deepcopy(sh._element) for sh in O(2).shapes if sh.name in ("footer-white", "footer-topic", "Buchseiten", "page-number")]


def neue(titel, pill="", hidden=False, layout="Nur Titel"):
    s = p.slides.add_slide(LAY[layout])
    ph = [x for x in s.placeholders if x.placeholder_format.type == 1]   # Titel-Platzhalter (idx 104 im BHAK-Master)
    if ph: ph[0].text_frame.text = titel
    tree = s.shapes._spTree
    for el in FUSS:
        e = copy.deepcopy(el)
        ids = [int(v) for v in tree.xpath(".//p:cNvPr/@id")] or [1]
        e.xpath(".//p:cNvPr")[0].set("id", str(max(ids) + 1))
        tree.append(e)
    for sh in s.shapes:
        if sh.name == "Buchseiten":
            runs = [r for pa in sh.text_frame.paragraphs for r in pa.runs]; runs[0].text = pill or "Unterricht"
            for r in runs[1:]: r.text = ""
    if hidden: s._element.set("show", "0")
    return s


def lernpfad():
    s = neue("Dein Lernpfad", "Handy-Seiten")
    for k, (key, lab, col) in enumerate([("theorie", "Theorie", OLIVE), ("aufgaben", "Arbeitsblatt", OLIVE), ("quiz", "Quiz", OLIVE), ("versuche", "Versuche", B.RED)]):
        B.ppt_qr(s, B.URL[key], 0.9 + k * 2.25, 1.45, 1.45, "V" if key == "versuche" else str(k + 1), lab, color=col, linksize=8)
    B.tb(s, 0.5, 4.3, 9, 0.45, ["Theorie = zum Nachlesen · Arbeitsblatt = Aufgaben und PDF · Versuche = Protokoll, ein Handy je Gruppe"], size=16, color=MUTED, align=PP_ALIGN.CENTER)
    return s


TEIL_INFO = {
 1: ("4AK + 4BK: Mi 04.11.", "Eine Schwingung beschreiben (A, T, f) und messen, wie die Länge die Schwingungsdauer verändert.", [1, 2, 3], ""),
 2: ("4AK: Mo 09.11. · 4BK: Fr 06.11.", "Erklären, warum Schwingungen abklingen und wann sie sich aufschaukeln (Resonanz).", [4, 5], ""),
 3: ("4AK: Mi 11.11. · 4BK: Fr 06.11.", "Quer- und Längswellen unterscheiden und mit v = λ · f rechnen.", [6, 7], ""),
 4: ("4AK: Mo 16.11. · 4BK: Fr 13.11.", "Auslenkungen mit Vorzeichen überlagern, stehende Wellen und das Prinzip von Huygens.", [8, 9],
     "4AK danach Zusatz 16–17 · 4BK danach Aufgabe 12–13 und Abgabe beider PDFs"),
 5: ("4AK: Mi 18.11. · 4BK: Mi 11.11.", "Elektromagnetische Wellen ordnen, eine Funk-Wellenlänge berechnen, Nutzen und Schutz begründen.", [10, 11, 12, 13],
     "4BK heute nur 10–11 · 12–13 am Fr 13.11. nach Teil 4 · 4AK: Abgabe beider PDFs am Ende"),
}
WIE = {2: "im Versuch", 4: "in der Demo", 5: "ohne Handy", 6: "im Versuch", 8: "Simulation", 9: "Handschrift + Foto", 12: "Handschrift + Foto", 13: "Handschrift + Foto", 10: "zuerst mündlich"}


def heute(t):
    wann, ziel, nrs, extra = TEIL_INFO[t]
    s = neue(f"Teil {t} · Heute", wann)
    B.tb(s, 0.5, 1.0, 9, 0.35, [[(wann, {"bold": True, "color": TEAL})]], size=16, after=0)
    c = B.box(s, 0.5, 1.42, 6.75, 0.8, fill=OLIVE); B.text(c, [[("Ziel: ", {"bold": True}), (ziel, {})]], size=18, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, after=0)
    rows = [["Nr.", "Aufgabe", "Zeit", "Wie"]]
    for n in nrs:
        tt = [x for x in I.T if x["nr"] == n][0]
        rows.append([str(n), tt["title"].split(" (")[0].replace(" für den Einsatz im Freien", ""), f"{tt['zeit']} min", WIE.get(n, "am Handy")])
    cells, h = B.auto_grid(s, 0.5, 2.38, [0.6, 3.7, 0.95, 1.5], rows, size=16, minh=0.36)
    s.shapes.add_picture(B.qr_png(B.URL["aufgaben"]), Inches(7.75), Inches(1.4), Inches(1.45), Inches(1.45))
    B.tb(s, 7.45, 2.88, 2.05, 0.3, ["Arbeitsblatt"], size=14, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)
    if extra: B.tb(s, 7.4, 3.3, 2.2, 1.7, [[("Danach:", {"bold": True, "color": RUST})]] + ["· " + x for x in extra.split(" · ")], size=13, after=2)
    return s


def li_tabelle(s, rows, y=1.0, cw=(0.55, 2.35, 2.2, 2.05, 1.85), size=10):
    return B.auto_grid(s, 0.4, y, list(cw), [["Folie", "Kernaussage (kurz erklären)", "Leitfrage", "Kadetten tun", "Sicherung"]] + rows, size=size, minh=0.26)


# Erklärabschnitte je Teil: (Originalfolie(n), Kernaussage, Leitfrage, Tätigkeit, Sicherung)
ABSCHNITTE = {
 1: [((2,), "Vermutung sammeln, noch nicht auflösen.", "Wer schwingt schneller – leicht oder schwer?", "Handzeichen, 1 Satz Begründung", "Vermutung stehen lassen; Auflösung nach dem Versuch (Modell)"),
     ((4, 5), "Volle Schwingung = gleicher Ort UND gleiche Richtung; A = größte Auslenkung, T = Zeit für einen Zyklus, f = 1/T.", "Links → rechts: Wie viel einer Schwingung ist das?", "im Diagramm A und T zeigen", "Merkzeile f = 1/T"),
     ((6,), "–", "Welche Achse liest du ab?", "A1 am Handy (5 min)", "Lösung per Klick; Fehler »6 cm« ansprechen"),
     ((7,), "Im Modell hängt T nur von l (und g) ab, nicht von der Masse.", "Was verändert T – Masse oder Länge?", "zuhören, Formel nur ansehen", "»l bis zum Schwerpunkt«"),
     ((8, 9), "Nur eine Größe ändern: die Länge.", "Ist das Pendel bei 0,90 m schneller oder langsamer?", "messen (2 Zeitnehmer, 1 Protokoll), A2 im Versuch, dann A3", "Werte vergleichen (Modell 1,27 s / 1,90 s); Schaukelfrage mit dem Modell auflösen")],
 2: [((10,), "Dämpfung wandelt Schwingungsenergie in Wärme; die Feder speichert nur kurz.", "Warum schaukelt ein Auto mit kaputtem Dämpfer nach?", "kurz zu zweit antworten", "Merksatz Feder ≠ Dämpfer"),
     ((11,), "Anregung im eigenen Takt (nahe der Eigenfrequenz) → großer Ausschlag.", "Bei welchem Takt wird der Ausschlag am größten? (vorher tippen)", "Vorhersage, beobachten, A4a/b eintragen", "»Resonanz braucht den passenden Takt, Dämpfung begrenzt«"),
     ((12,), "Resonanz ≠ automatisch Einsturz; Broughton: Bolzen; Tacoma: Flattern.", "Stürzt jede Brücke im Gleichschritt ein?", "zuhören, 1 Rückfrage", "»Ob etwas bricht: Dämpfung und Festigkeit«"),
     ((13,), "–", "–", "A5 Partnercheck ohne Handy, dann eintragen", "Lösung je Frage per Klick")],
 3: [((14,), "Die Störung wandert, die Teilchen schwingen am Ort; quer ⟂, längs ∥ zur Ausbreitung.", "Was wandert: das Material oder die Störung?", "zuhören", "Merkzeile Folie 14"),
     ((15,), "–", "Was macht die Büroklammer? Was tut ein einzelner Kadett?", "Schnur + Kadetten-Welle quer/längs, A6", "A6b/c gemeinsam kontrollieren"),
     ((16,), "Waagrechte Achse zuerst lesen: Ort → λ, Zeit → T.", "Wo liest du eine Länge ab, wo eine Zeit?", "beide Diagramme zuordnen", "v = λ · f = λ / T"),
     ((17,), "Musterweg: Formel → Werte mit Einheit → Ergebnis mit Einheit.", "Welche Größe ist gesucht?", "A7 rechnen", "Lösung per Klick, Einheit kontrollieren")],
 4: [((18,), "Auslenkungen mit Vorzeichen addieren (lineares Modell); danach laufen die Pulse weiter.", "Was passiert mit +2 cm und −2 cm? (vorher tippen)", "Vorhersage, Animation, A8 (Simulation)", "+2 cm − 1 cm = +1 cm"),
     ((19,), "Nur Begriffe: Knoten ruhen, Bäuche schwingen; benachbarte Knoten λ/2.", "Wo schwingt nichts?", "zuhören (2 min)", "Vertiefung: Akustik"),
     ((20, 21), "Jeder Punkt der Front sendet eine Elementarwelle; die Einhüllende ist die neue Front.", "Wie entsteht die neue Front?", "A9 zeichnen, Foto, 1 Satz", "Gute Skizze: gleich große Halbkreise, Hülle vorne")],
 5: [((22,), "Schall braucht ein Medium, Licht und Funk nicht.", "Hört man eine Explosion im Weltraum?", "Science oder Fiction auflösen", "Merkzeile Folie 22"),
     ((23,), "Ordnung nach Frequenz/Wellenlänge; Anker: Funk, Wärmebild, Licht, UV.", "Wo liegt Wärmebild, wo Sonnenbrand?", "A10 zuerst mündlich zu zweit", "Reihenfolge gemeinsam"),
     ((24,), "Risiko = Art, Stärke, Dauer; nicht-ionisierend ≠ harmlos.", "Ist UV harmlos, weil es nicht ionisiert?", "zuhören", "Merkzeile Folie 24"),
     ((25,), "Musterweg: MHz → Hz, dann λ = c : f.", "Wie viele Hz sind 50 MHz?", "A11 rechnen und zuordnen", "Lösung per Klick"),
     ((26, 27), "–", "Woran erkennt man eine gute Korrektur bzw. Karte?", "A12 (Handschrift), A13 (Sicherheitskarte)", "2–3 Beispiele kurz zeigen (nur mit Zustimmung)")],
}
KUERZEN = {
 1: "Wenn knapp: A3 in der Besprechung gemeinsam rechnen. Während die Lehrkraft auf 0,90 m umstellt (ca. 2 min), rechnen alle Gruppen T für 0,40 m.",
 2: "4AK (nach Klassenbuch und Test ca. 27 min): Brückenfolie nur 2 min. " + P.KNAPP[("4BK", 2)].replace("Zeitstopp:", "4BK Zeitstopp:"),
 3: P.KNAPP[("4AK", 3)].replace("Zeitstopp:", "4AK Zeitstopp:"),
 4: "Stehende Welle (Folie 19) nur 2 min oder weglassen – kommt in der Akustik wieder. 4BK: A12 und A13 bleiben Pflicht; wenn knapp, A12 zuerst gemeinsam mündlich vorbesprechen, dann kurz schreiben.",
 5: P.KNAPP[("4BK", 3)].replace("Zeitstopp:", "4BK Zeitstopp:") + " 4AK: Abgabezeit (4 min) nicht kürzen.",
}

# ------------------------------------------------------------ 3 Reihenfolge festlegen (physische Ordnung = 4AK)
NEU = {}
NEU["li_bed"] = neue("Lehrer-Info · So nutzt du diese PPP", "Unterricht", hidden=True)
NEU["li_ak"] = neue("Lehrer-Info · Ablauf 4AK", "plan.py · 50 min je Stunde", hidden=True)
NEU["li_bk"] = neue("Lehrer-Info · Ablauf 4BK", "plan.py · 50 min je Stunde", hidden=True)
NEU["lernpfad"] = lernpfad()
for t in range(1, 6):
    NEU[f"heute{t}"] = heute(t)
    NEU[f"li{t}"] = neue(f"Lehrer-Info · Teil {t}", "nur Lehrer", hidden=True)
NEU["quellen"] = neue("Lehrer-Info · Quellen und Herkunft", "nur Lehrer", hidden=True)

ORDNUNG = ([O(1), NEU["li_bed"], NEU["li_ak"], NEU["li_bk"], O(32), O(33), O(34), NEU["lernpfad"],
            NEU["heute1"], NEU["li1"]] + [O(n) for n in range(2, 10)] +
           [NEU["heute2"], NEU["li2"]] + [O(n) for n in range(10, 14)] +
           [NEU["heute3"], NEU["li3"]] + [O(n) for n in range(14, 18)] +
           [NEU["heute4"], NEU["li4"]] + [O(n) for n in range(18, 22)] +
           [NEU["heute5"], NEU["li5"]] + [O(n) for n in range(22, 29)] + [NEU["quellen"]])
WEG = [O(29), O(30), O(31)]                   # alte Lehrer-Infos (Bedienung, Ablauf) – ersetzt durch neue aus plan.py

lst = p.slides._sldIdLst
el_von = {sl.part: el for sl, el in zip(list(p.slides), list(lst))}   # Schlüssel: Folien-Part (slide_id braucht die Liste)
for el in list(lst): lst.remove(el)
for sl in ORDNUNG: lst.append(el_von[sl.part])
for sl in WEG:                                # aus der Präsentation lösen
    p.part.drop_rel(el_von[sl.part].get("{%s}id" % NS["r"]))
NR = {sl.part: i + 1 for i, sl in enumerate(ORDNUNG)}
F = lambda n: NR[O(n).part]                   # neue Foliennummer der Originalfolie n

for sl in ORDNUNG:                            # Seitenzahlen = Foliennummer in PowerPoint
    for sh in sl.shapes:
        if sh.name == "page-number" and sh.has_text_frame:
            runs = [r for pa in sh.text_frame.paragraphs for r in pa.runs]
            if runs: runs[0].text = str(NR[sl.part]); [setattr(r, "text", "") for r in runs[1:]]

# ------------------------------------------------------------ 4 Folgen je Klasse und Stunde (benutzerdefinierte Präsentationen)
T = {1: [NEU["heute1"]] + [O(n) for n in range(2, 10)], 2: [NEU["heute2"]] + [O(n) for n in range(10, 14)],
     3: [NEU["heute3"]] + [O(n) for n in range(14, 18)], 4: [NEU["heute4"]] + [O(n) for n in range(18, 22)],
     5: [NEU["heute5"]] + [O(n) for n in range(22, 26)]}
TRANSFER = [O(26), O(27)]; ABSCHLUSS = [O(28)]
FOLGE = {
 ("4AK", 1): [O(1), NEU["lernpfad"]] + T[1], ("4AK", 2): T[2], ("4AK", 3): T[3], ("4AK", 4): T[4], ("4AK", 5): T[5] + TRANSFER + ABSCHLUSS,
 ("4BK", 1): [O(1), NEU["lernpfad"]] + T[1], ("4BK", 2): T[2] + T[3], ("4BK", 3): T[5], ("4BK", 4): T[4] + TRANSFER + ABSCHLUSS,
}
DATUM = {k: {st["std"]: st["datum"] for st in P.PLAN[k]} for k in P.PLAN}
pres = p.part._element
alt = pres.find("p:custShowLst", NS)
csl = etree.Element("{%s}custShowLst" % NS["p"])
if alt is not None: alt.addprevious(csl); pres.remove(alt)   # an derselben Stelle wie in der Vorlage (Schema-Reihenfolge)
else: pres.find("p:defaultTextStyle", NS).addprevious(csl)
rid = {sl.part: el_von[sl.part].get("{%s}id" % NS["r"]) for sl in ORDNUNG}


def show(name, folien, nid):
    cs = etree.SubElement(csl, "{%s}custShow" % NS["p"], name=name, id=str(nid))
    l = etree.SubElement(cs, "{%s}sldLst" % NS["p"])
    for sl in folien:
        etree.SubElement(l, "{%s}sld" % NS["p"]).set("{%s}id" % NS["r"], rid[sl.part])


nid = 0
for k in ("4AK", "4BK"):
    show(f"{k} gesamt", sum([FOLGE[(k, s)] for s in sorted(DATUM[k])], []), nid); nid += 1
for k in ("4AK", "4BK"):
    for s in sorted(DATUM[k]):
        show(f"{k} Stunde {s} · {DATUM[k][s]}", FOLGE[(k, s)], nid); nid += 1


def folien_text(liste):
    nrs = [NR[x.part] for x in liste]; out, a = [], None
    for i, n in enumerate(nrs):
        if a is None: a = n
        if i + 1 == len(nrs) or nrs[i + 1] != n + 1: out.append(str(a) if a == n else f"{a}–{n}"); a = None
    return ", ".join(out)


# ------------------------------------------------------------ 5 Lehrer-Infos füllen (jetzt mit fertigen Foliennummern)
s = NEU["li_bed"]
B.tb(s, 0.4, 0.98, 9.2, 4.1, [
    [("Eine Datei für beide Klassen. ", {"bold": True, "color": OLIVE}), ("F5 zeigt alle sichtbaren Folien in der Reihenfolge der 4AK (ausgeblendete Lehrer-Infos werden übersprungen).", {})],
    [("4BK oder eine einzelne Stunde: ", {"bold": True, "color": OLIVE}), ("Bildschirmpräsentation → Benutzerdefinierte Bildschirmpräsentation → z. B. »4BK Stunde 2 · Fr 06.11.«. Dort steht jede Stunde beider Klassen einzeln.", {})],
    [("Ohne Menü: ", {"bold": True, "color": OLIVE}), ("Im Vortrag die Foliennummer tippen und Enter – z. B. 4BK am Mi 11.11. direkt zu Folie " + str(NR[NEU['heute5'].part]) + " (Teil 5 · Heute).", {})],
    [("Aufbau jedes Teils: ", {"bold": True, "color": OLIVE}), ("»Heute«-Folie (Ziel, Aufgaben, Zeit, QR) → kurze Erklärfolien → Arbeitsauftrag. Die ausgeblendete Lehrer-Info danach nennt je Folie Kernaussage, Leitfrage, Tätigkeit und Sicherung sowie Zeitstopp und Kürzung.", {})],
    [("Lösungen ", {"bold": True, "color": OLIVE}), ("erscheinen per Klick erst nach der Arbeitszeit. Die ausführliche TH-PPP und die Theorie-Seite bleiben Nachschlagewerk (Fehlstunde, Wiederholung, Tipp-Links) – im Unterricht nicht zusätzlich durchlesen lassen.", {})],
    [("Frontal ", {"bold": True, "color": OLIVE}), ("je Abschnitt ca. 5–8 min erklären (Richtwert), dann sofort eine Tätigkeit. Selbst erarbeiten: Aufgaben, Versuche, Partnercheck, Zusatz für Schnelle.", {})]], size=14, after=6)

for k, key in (("4AK", "li_ak"), ("4BK", "li_bk")):
    s = NEU[key]; rows = []
    for st in P.PLAN[k]:
        ph = " · ".join(f"{n.split(' (')[0].replace('Beginn: ', '')} {m}" for n, m, o, a in st["phasen"])
        rows.append([f"{st['std']} · {st['datum']}", folien_text(FOLGE[(k, st['std'])]), ph])
    B.auto_grid(s, 0.4, 0.98, [1.15, 1.25, 6.8], [["Stunde", "Folien", "Phasen in Minuten (Summe je 50)"]] + rows, size=10, minh=0.26)
    B.tb(s, 0.4, 4.62, 9.2, 0.4, ["Jede Stunde beginnt mit ca. 3 min Klassenbuch. Pflicht: Aufgaben 1–13 und 4 Versuche · Zusatz 14–17 nur für Schnelle · keine Hausübung."
                                  + (" Zusatztermin Mi 18.11. ist optional – der Pflichtkern hängt nicht davon ab." if k == "4BK" else "")], size=11, color=MUTED)

for t in range(1, 6):
    s = NEU[f"li{t}"]; rows = []
    for nn, kern, frage, tun, sich in ABSCHNITTE[t]:
        rows.append([", ".join(str(F(x)) for x in nn), kern, frage, tun, sich])
    cells, h = li_tabelle(s, rows)
    B.tb(s, 0.4, min(1.06 + h, 4.45), 9.2, 0.6, [[("Zeitstopp und Kürzung: ", {"bold": True, "color": RUST}), (KUERZEN[t], {})]], size=10, after=0)

s = NEU["quellen"]
qs = ["Grundlage dieser Kurz-PPP: ChatGPT, Unterrichts-PPP v2 (25.10.2026); Prüfung und Überarbeitung v3: Claude.",
      "Zeichnungen und GIFs: eigene Darstellungen aus dem Claude-Paket (skizzen.py); Diagramme Folie " + ", ".join(str(F(n)) for n in (5, 6, 16)) + ": von ChatGPT als bearbeitbare Diagramme gesetzt.",
      "Keine Bilder aus dem Schulbuch eingefügt (die Buchfotos lagen Claude nicht vor – siehe Rückmeldung)."] + QUELLEN
B.tb(s, 0.4, 0.98, 9.2, 4.1, ["· " + q for q in qs[:14]], size=11, after=3)

os.makedirs(ZIEL, exist_ok=True)
name = f"{I.PRAEFIX}-TH-Unterricht-{VER}.pptx"
p.save(os.path.join(ZIEL, name))
sicht = [sl for sl in ORDNUNG if sl._element.get("show") != "0"]
print("ok", name, len(ORDNUNG), "Folien,", len(sicht), "sichtbar ·", nid, "Präsentationen")
for k in ("4AK", "4BK"):
    for st in sorted(DATUM[k]): print(" ", k, st, DATUM[k][st], "Folien", folien_text(FOLGE[(k, st)]))
