"""Private Lehrer-Seite (Lösungsmappe) K1.4 Erkenntnisgewinn – als Artifact, nur für Stephan. NICHT ins öffentliche Repository legen.
Aus derselben Datenquelle wie alle anderen Ausgaben (inhalt, versuche_daten, plan, fachwissen, herrichten).
Aufruf: python3 loesungsmappe.py <TH-Unterricht.pptx> <Ordner mit Folienbildern s-NN.jpg> <Ziel.html>
Die Folienbilder sind die sichtbaren Folien der TH-Unterricht (LibreOffice-Render), damit die Seite am Beamer-PC als Ersatz taugt."""
import base64, glob, html, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
import inhalt as I, versuche_daten as VD, plan as P, fachwissen as F, herrichten as H
from pptx import Presentation

PPTX, FOLDIR, ZIEL = sys.argv[1:4]
UNIT = os.path.join(ROOT, I.ORDNER) + "/"
e = lambda t: html.escape(str(t), quote=False)


def md(t):
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", e(t))


def b64(path, mime):
    return f"data:{mime};base64," + base64.b64encode(open(path, "rb").read()).decode()


def ul(items):
    return "<ul>" + "".join(f"<li>{md(x)}</li>" for x in items) + "</ul>"


def tab(head, rows, cls=""):
    h = "".join(f"<th>{e(x)}</th>" for x in head)
    b = "".join("<tr>" + "".join(f"<td>{md(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f"<div class='scroll'><table class='{cls}'><tr>{h}</tr>{b}</table></div>"


# ---------------------------------------------------------------- Überblick
offen = [
    "**Hallenprobe im Eingangsbereich** (vor Do 05.11. bzw. Mi 04.11.): Faltung in 10 min machbar? Fliegen A und B ohne Deckenkontakt? Wie weit ungefähr? Hält die Klammer? Ist der Bereich zur Stunde frei?",
    "**Klassengröße:** Geplant sind 8 Gruppen zu 3–4 Kadetten. Bei weniger Kadetten weniger Gruppen – dann sinkt auch die Messzeit (8 × 6 Würfe × 30 s ≈ 24 min ist eine Annahme, nicht gemessen).",
    "**Bandmaß #9 (20 m):** Im Inventar ist keine Zählung eingetragen – vorher nachsehen.",
    "**Papier:** 80-g-Kopierpapier ist nicht im Inventar (Kopierraum).",
    "**Apolin-Videos:** laut dir regelmäßig in dieser Einheit verwendet; Links aus der alten PPP übernommen – einmal am Schul-PC anspielen.",
    "**Einheit einblenden**, wenn es losgeht (1BK Mi 04.11.): GitHub → Actions → »Einheit ein-/ausblenden« → 1ab-nawi/k1-4-erkenntnis.",
    "**1BK-Flieger** liegen von Mi 04.11. bis Mo 09.11. in der Schachtel – vor dem Test auf Knicke prüfen.",
    "**Test am Fr 06.11.** (20 min) ist in beiden Klassen eingeplant. Teil 3 baut nicht auf dem Versuch auf (1BK hat ihn dann noch nicht gemacht).",
]
uber = "<h2>Was du vor Ort noch prüfen musst</h2>" + ul(offen)
uber += "<h2>Ablauf der Klassen</h2><p>" + md(P.FAHRPLAN_TEXT) + "</p>"
for k, rows in F.STUNDENVORSCHLAG.items():
    uber += f"<h3>{k}</h3>" + tab(["Std.", "Theorie", "Versuch", "Aufgaben", "Zeit", "Methode"], rows)
uber += "<h2>Stundenkarten (Minuten je Phase, Summe je 50)</h2>"
for k in I.KLASSEN:
    for st in P.PLAN[k]:
        t0, li = 0, []
        for name, mi, oh, a in st["phasen"]:
            li.append(f"<tr><td class='num'>{t0}–{t0 + mi}</td><td>{e(name)}{' <span class=chip>ohne Handy</span>' if oh and a not in ('org', 'reserve', 'test', 'bespr', 'erkl') else ''}</td></tr>"); t0 += mi
        kn = P.KNAPP.get((k, st["std"]))
        uber += (f"<details class='card'><summary><span class='nr'>{k}</span><span class='ttl'>Stunde {st['std']} · {e(st['datum'])} · {e(I.TEILNAME[st['teile'][0]])}</span></summary>"
                 f"<div class='body'><div class='scroll'><table>{''.join(li)}</table></div>" + (f"<p class='warnline'><b>Wenn es knapp wird:</b> {e(kn)}</p>" if kn else "") + "</div></details>")
uber += "<h2>Die 5 größten Stolpersteine</h2>" + "".join(f"<div class='fehler'><b class='k'>{e(a)}</b>{md(b)}</div>" for a, b in I.STOLPERSTEINE)
uber += "<h2>So bewertet die Seite</h2>" + ul([
    "**Automatisch geprüft:** Auswahlfelder (A1–A3, A5, A6 a–c, A9, A10 a–f, A11–A13 Auswahl) und Zahlen (A8 a–c, ±2 %, Einheit wird umgerechnet).",
    "**Offene Texte bewertest du:** A4 (Hypothese, Vorhersage, Gegenbefund), A6 d, A7, A9 e, A10 g, A12 e, A13 c, A14. Die Seite zeigt nur eine Schreibhilfe und gibt dafür keine Punkte.",
    "**PDF:** Kasten »Für LMS: x von y Punkten · Notenvorschlag« (88/75/62/50 %); Aufgaben mit angesehener Lösung zählen 0 Punkte.",
    "**Versuchsseite:** ein Protokoll je Gruppe, jedes Mitglied gibt seine eigene PDF ab (»Ich gebe ab als …«). Pflicht sind nur die drei Vergleichspaare; Mittelwert und Gegenprobe nicht.",
])
uber += "<h2>Links der Kadetten-Seiten</h2><ul>" + "".join(f"<li><span class='mono'>{I.URL}{k}.html</span></li>" for k in ("theorie", "aufgaben", "versuche", "quiz")) + "</ul>"

# ---------------------------------------------------------------- Herrichten (abhakbar)
herr = "<p>Zum Abhaken – die Haken bleiben nur in diesem Browser gespeichert. Mengen: Stück je Gruppe × 8 Gruppen + Reserve.</p>"
listen = [(f"{nr} · {tit}", rows) for nr, tit, rows in H.je_versuch()]
if len(VD.VS) > 1: listen.append(("Gesamt für alle Versuche", H.gesamt()))
else: listen[0] = (listen[0][0] + " (= Gesamtliste, nur ein Versuch)", listen[0][1])
n = 0
for tit, rows in listen:
    herr += f"<h2>{e(tit)}</h2>"
    for ab, lab in H.ABSCHNITTE:
        rr = [r for r in rows if r["abschnitt"] == ab]
        if not rr: continue
        herr += f"<h3>{e(lab)}</h3><ul class='check'>"
        for r in rr:
            n += 1
            fehlt = f" <span class='chip rot'>fehlt {e(r['fehlt'])}</span>" if r["fehlt"] not in ("–", "") else ""
            herr += (f"<li><label><input type='checkbox' id='h{n}' data-k='h{n}'> <b>{e(r['gesamt'])} ×</b> {e(r['name'])}{fehlt}"
                     f"<br><span class='klein'>{e(r['rechnung'])} · {e(r['ort'])} · Bestand: {e(r['bestand'])}</span></label></li>")
        herr += "</ul>"
herr += "<h2>Vor der Stunde</h2><ul class='check'>"
for x in F.CHECKLISTE:
    n += 1; herr += f"<li><label><input type='checkbox' id='h{n}' data-k='h{n}'> {md(x)}</label></li>"
herr += "</ul>"

# ---------------------------------------------------------------- Aufgaben
auf = []
for t in I.T:
    man = [k for k, v in (t.get("chk") or {}).items() if v.get("man")]
    tags = [f"<span class='tag{' zus' if t['h'] == 0 else ''}'>{e(t['tag'])} · {t['zeit']} min</span>"]
    if t.get("besprechen"): tags.append("<span class='tag besp'>BESPRECHEN</span>")
    if man: tags.append("<span class='tag bew'>du bewertest</span>")
    felder = "".join(f"<li>{e(f[1])}" + (f" <span class='klein'>({' · '.join(e(o) for o in f[2])})</span>" if len(f) > 2 and isinstance(f[2], (list, tuple)) else "") + "</li>" for f in t.get("fields", []))
    sol = "".join(f"<p>{'' if h in ('Lösung', 'Beispiel', 'Bewertung') else '<b>' + e(h) + '</b> '}{md(x)}</p>" for h, xs in t["sol"] for x in xs)
    bild = f"<div class='bild'><img src='{b64(UNIT + t['img'], 'image/png')}' alt=''></div>" if t.get("img") else ""
    auf.append(f"<details class='card' data-h='{t['h']}' data-s='{html.escape((t["title"] + " " + t["q"]).lower(), quote=True)}'><summary><span class='nr'>{t['nr']}</span><span class='ttl'>{e(t['title'])}</span>{''.join(tags)}</summary>"
               f"<div class='body'><p class='q'>{md(t['q'])}</p>{bild}<p class='klein'>{e(t['src'])}</p><ul class='felder'>{felder}</ul>"
               f"<div class='box sol'><b class='k'>Lösung</b>{sol}</div>"
               + (f"<div class='box bew'><b class='k'>Bewertungshilfe</b>Offene Felder ({', '.join(man)}): Die Seite prüft nur Stichwörter als Schreibhilfe – fachlich entscheidest du. Musterantwort oben.</div>" if man else "")
               + (f"<p class='klein'>Hinweis für Kadetten: {md(t['hint'])}</p>" if t.get("hint") else "") + "</div></details>")
filt = "".join(f"<button type='button' data-f='{h}' aria-pressed='false'>{e(I.GRP[h].split(' · ')[0] + (' · ' + I.GRP[h].split(' · ')[1] if h else ''))}</button>" for h in (1, 2, 3, 4, 0))
aufgaben = (f"<input class='search' type='search' id='suche' placeholder='Aufgaben durchsuchen' aria-label='Aufgaben durchsuchen'><div class='filter' id='filter'>{filt}</div>"
            + "".join(auf) + "<p class='leer' id='leer' hidden>Nichts gefunden.</p>"
            + "<h2>Partnerkarten (Aufgabe 7)</h2>" + tab(["Frage", "Lösung"], I.PARTNERKARTEN))

# ---------------------------------------------------------------- Versuch
v = VD.VS[0]
vs = (f"<div class='vsgrid'><div><h2>{e(v['nr'])} · {e(v['titel'])}</h2><p>{md(v['intro'])}</p><p class='klein'>{e(v['form'])} · ca. {v['zeit']} min · {e(v['wann'])}</p>"
      f"<div class='fehler'><b class='k'>Sicherheit</b>{md(v['sicher'])}</div></div><img src='{b64(UNIT + 'vs-img/d_' + v['skizze'] + '.png', 'image/png')}' alt='Messbahn im Eingangsbereich'></div>"
      "<h3>Durchführung</h3><ol>" + "".join(f"<li>{md(s)}</li>" for s in v["schritte"]) + "</ol>"
      + "<h3>Beispielwerte (erfunden – das Ergebnis ist offen)</h3>"
      + tab(v["beob"]["kopf"], [z + v["erwartet_tab"][i] for i, z in enumerate(v["beob"]["zeilen"][:5])])
      + "".join(f"<div class='box fach'><b class='k'>Erwartet: {e(a)}</b>{md(b)}</div>" for a, b in v["erwartet"])
      + "".join(f"<div class='box sol'><b class='k'>{lab}</b><p class='q'>{md(v[k])}</p><p>{md(v[k + '_l'])}</p></div>" for k, lab in (("erk", "Erkenntnis"), ("interp", "Interpretation"), ("schluss", "Schlussfolgerung")))
      + "<h3>Tipps für dich</h3>" + ul(v["lehrer"]) + "<h3>Typische Pannen</h3>" + "".join(f"<div class='fehler'>{md(x)}</div>" for x in v["fehler"])
      + "<h3>Fachlich genau</h3>" + "".join(f"<div class='box fach'>{md(x)}</div>" for x in v["fach"])
      + f"<div class='box bew'><b class='k'>Vorversuch (offen – nur du)</b>{md(v['vorversuch'])}</div><p class='klein'>Quelle: {e(v['quelle'])} · {e(v['jp'])}</p>")

# ---------------------------------------------------------------- Fachwissen
fw = "<h2>Lehrplanbezug</h2><ul>" + "".join(f"<li><b>{e(a)}</b> {md(b)}</li>" for a, b in F.LEHRPLAN) + "</ul>"
for kap in F.KAPITEL:
    fw += f"<h2>{e(kap['titel'])}</h2>" + "".join(f"<p>{md(x)}</p>" for x in kap["text"])
    if kap.get("bild"): fw += f"<div class='bild'><img src='{b64(UNIT + kap['bild'], 'image/png')}' alt=''></div>"
    if kap.get("nur"): fw += "<div class='box fach'><b class='k'>Nur für dich</b>" + "".join(f"<p>{md(y)}</p>" for y in kap["nur"]) + "</div>"
    if kap.get("stolper"): fw += f"<div class='fehler'><b class='k'>Stolperstein</b>{md(kap['stolper'])}</div>"
fw += "<h2>Rechenbeispiele</h2>" + "".join(f"<p>{md(x)}</p>" for x in F.RECHNEN)
fw += "<h2>Häufige Fehlvorstellungen</h2>" + tab(["Fehlvorstellung", "Richtig ist", "So zeigst du es"], F.FEHLVORSTELLUNGEN)
fw += "<h2>Fragen, die Kadetten stellen</h2>" + "".join(f"<p><b>{e(q)}</b><br>{md(a)}</p>" for q, a in F.FRAGEN)
fw += "<h2>Videos</h2>" + "".join(f"<p><b>{e(t_)}</b> · {e(q)}<br><span class='mono'>{e(u)}</span><br>{md(F.VIDEO_NOTIZ[k])}</p>" for k, (t_, q, d, u, *_) in I.VIDEOS.items())
fw += "<h2>Ausblick</h2><p>" + md(F.AUSBLICK) + "</p><h2>Quellen</h2>" + ul(F.QUELLEN)

# ---------------------------------------------------------------- Folien (TH-Unterricht als Bilder)
prs = Presentation(PPTX)
titel = [next((ph.text_frame.text for ph in s.placeholders if ph.placeholder_format.type == 1), "") for s in prs.slides if s._element.get("show") != "0"]
bilder = sorted(glob.glob(os.path.join(FOLDIR, "s-*.jpg")))
assert len(bilder) == len(titel), (len(bilder), len(titel))
nr_phys = [i + 1 for i, s in enumerate(prs.slides) if s._element.get("show") != "0"]
import json
folien_js = json.dumps([[b64(b, "image/jpeg"), t, n] for b, t, n in zip(bilder, titel, nr_phys)], ensure_ascii=False).replace("</", "<\\/")
optionen = "".join(f"<option value='{i}'>Folie {n} · {e(t)}</option>" for i, (t, n) in enumerate(zip(titel, nr_phys)))

STIL = open(os.path.join(HERE, "loesungsmappe_stil.css"), encoding="utf-8").read()
seite = f"""<title>Lösungsmappe Erkenntnisgewinn 1AB</title>
<meta name="claude:viewport-fit-sides" content="cover">
<link rel="stylesheet" href="../../schriften/schriften.css">
<style>{STIL}</style>
<div class="wrap">
<header class="lese">
  <div class="eyebrow">NAWI · 1AB (1AK, 1BK) · K1.4 Erkenntnisgewinn · nur für dich</div>
  <h1>Lösungen und Fachwissen: der Papierflieger-Test</h1>
  <p class="sub">Aufgaben 1–14 (11–14 Zusatz), Versuch {e(I.VS_ID)}, Herrichte-Liste zum Abhaken, die Unterrichts-PPP als Folien und das Fachwissen. Buch: Trauner NAWI HAK 1, S. 28–32. Stand {e(I.STAND)} · {e(I.V)}</p>
  <p class="warn">Privat in deinem Konto – nicht an Kadetten weitergeben. Die Kadetten sehen Musterlösungen erst nach der PDF bzw. nach drei Fehlversuchen.</p>
</header>
<nav class="tabs" role="tablist" aria-label="Bereiche">
  <button role="tab" id="t-ueberblick" aria-selected="true">Überblick</button>
  <button role="tab" id="t-herrichten" aria-selected="false">Herrichten</button>
  <button role="tab" id="t-aufgaben" aria-selected="false">Aufgaben ({len(I.T)})</button>
  <button role="tab" id="t-versuch" aria-selected="false">Versuch</button>
  <button role="tab" id="t-folien" aria-selected="false">Folien</button>
  <button role="tab" id="t-fachwissen" aria-selected="false">Fachwissen</button>
</nav>
<section class="panel lese" id="p-ueberblick">{uber}</section>
<section class="panel lese" id="p-herrichten" hidden>{herr}</section>
<section class="panel lese" id="p-aufgaben" hidden>{aufgaben}</section>
<section class="panel lese" id="p-versuch" hidden>{vs}</section>
<section class="panel" id="p-folien" hidden>
  <div class="steuer"><button class="btn ghost" id="f-zur" type="button">◀ zurück</button><button class="btn" id="f-vor" type="button">weiter ▶</button>
  <select id="f-wahl" aria-label="Folie wählen">{optionen}</select><button class="btn ghost" id="f-voll" type="button">Vollbild</button></div>
  <div class="buehne" id="buehne" tabindex="0"><img id="f-img" alt=""></div>
  <p class="hilfe">Die {len(titel)} sichtbaren Folien der TH-Unterricht (Lösungen sind hier schon aufgedeckt). Weiter: Tippen rechts, Pfeiltaste oder Leertaste · zurück: Tippen links · Vollbild: Knopf oder Taste F. Die Foliennummer entspricht der Nummer in PowerPoint.</p>
</section>
<section class="panel lese" id="p-fachwissen" hidden>{fw}</section>
<footer class="lese">Gebaut aus werkzeug/quellen/1ab-k1-4-erkenntnis (eine Datenquelle für Seiten, PPP, Word und diese Mappe) · Stand {e(I.STAND)}</footer>
</div>
<script>
const F={folien_js};
const tabs=[...document.querySelectorAll('nav.tabs button')];
function zeige(id){{tabs.forEach(b=>{{const on=b.id==='t-'+id;b.setAttribute('aria-selected',on);document.getElementById('p-'+b.id.slice(2)).hidden=!on;}});try{{localStorage.setItem('k14-tab',id)}}catch(e){{}}}}
tabs.forEach(b=>b.addEventListener('click',()=>zeige(b.id.slice(2))));
try{{const t=localStorage.getItem('k14-tab');if(t&&document.getElementById('p-'+t))zeige(t)}}catch(e){{}}
// Herrichten: Haken merken
document.querySelectorAll('input[data-k]').forEach(c=>{{try{{c.checked=localStorage.getItem('k14-'+c.dataset.k)==='1'}}catch(e){{}}
  c.addEventListener('change',()=>{{try{{localStorage.setItem('k14-'+c.dataset.k,c.checked?'1':'0')}}catch(e){{}}}})}});
// Aufgaben: Suche und Filter
const karten=[...document.querySelectorAll('#p-aufgaben details.card')];let fil=null;
function filtern(){{const q=document.getElementById('suche').value.toLowerCase().trim();let n=0;
  karten.forEach(k=>{{const ok=(!q||k.dataset.s.includes(q))&&(fil===null||k.dataset.h===fil);k.hidden=!ok;if(ok)n++;}});document.getElementById('leer').hidden=n>0;}}
document.getElementById('suche').addEventListener('input',filtern);
document.querySelectorAll('#filter button').forEach(b=>b.addEventListener('click',()=>{{fil=(fil===b.dataset.f)?null:b.dataset.f;
  document.querySelectorAll('#filter button').forEach(x=>x.setAttribute('aria-pressed',x.dataset.f===fil));filtern();}}));
// Folien
let i=0;const img=document.getElementById('f-img'),wahl=document.getElementById('f-wahl'),bu=document.getElementById('buehne');
function geh(n){{i=Math.max(0,Math.min(F.length-1,n));img.src=F[i][0];img.alt='Folie '+F[i][2]+': '+F[i][1];wahl.value=i;}}
document.getElementById('f-vor').onclick=()=>geh(i+1);document.getElementById('f-zur').onclick=()=>geh(i-1);wahl.onchange=()=>geh(+wahl.value);
bu.addEventListener('click',ev=>{{const r=bu.getBoundingClientRect();geh(ev.clientX-r.left<r.width/3?i-1:i+1)}});
document.getElementById('f-voll').onclick=()=>{{try{{bu.requestFullscreen&&bu.requestFullscreen().catch(()=>{{}})}}catch(e){{}}}};
document.addEventListener('keydown',ev=>{{if(document.getElementById('p-folien').hidden)return;
  if(['ArrowRight',' ','PageDown'].includes(ev.key)){{ev.preventDefault();geh(i+1)}}else if(['ArrowLeft','PageUp'].includes(ev.key)){{ev.preventDefault();geh(i-1)}}
  else if(ev.key==='f'||ev.key==='F')document.getElementById('f-voll').click();}});
geh(0);
</script>
"""
open(ZIEL, "w", encoding="utf-8").write(seite)
print("ok", ZIEL, len(seite) // 1024, "KB")
