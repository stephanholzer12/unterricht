"""Unterrichts-PPP K1.4 Erkenntnisgewinn (TH-Unterricht) – die kurze, geführte PPP für den Frontalunterricht im Saal.
Stephans Auftrag (29.10.2026): gekürzte TH-Lehrer, in der GENAU steht, was die Kadetten in welcher Reihenfolge tun; Theorie nur dort
wiederholt, wo es Schwierigkeiten oder falsche Vorstellungen gibt; die wichtigsten Aufgaben und der Versuch werden gemeinsam durchgegangen.
Aufbau nach Regelwerk »bhak-ppp-gestaltung« Abschnitt 10 und der Wellen-PPP v6 (ChatGPT): ausgeblendete Lehrer-Info »So führt die PPP …«,
ausgeblendete Regie je Klasse und Stunde (Minute · Was läuft jetzt? · PPP-Folie · Wozu?), je Teil »Heute« → »Lesen: <Abschnitt>« → »KURZ KLÄREN« →
Arbeitsauftrag → Besprechen; unten links ein Phasenhinweis (JETZT … · Minuten / Danach: …).
Physische Reihenfolge = 1AK (Teil 1→2→3→4); 1BK (Teil 1→3→2→4) über benutzerdefinierte Präsentationen »1BK gesamt« und »1BK <Datum> Stunde n«.
Alle Minuten kommen aus plan.py (eine Quelle mit Fahrplan, Stundenkarten und Lösungsmappe).
Aufruf im Repository-Ordner: python3 werkzeug/quellen/1ab-k1-4-erkenntnis/unterricht_ppp.py <BHAK_Folienmaster_Vorlage.pptx> <Zielordner>"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import bauen_ppp as B
import inhalt as I, plan as P, theorie_daten as TD, versuche_daten as VD
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

ZIEL = sys.argv[2]
OLIVE, INK, TEAL, RUST, MUTED, WHITE, GREEN, ORANGE, CARD, RED = B.OLIVE, B.INK, B.TEAL, B.RUST, B.MUTED, B.WHITE, B.GREEN, B.ORANGE, B.CARD, B.RED
X0, W = 0.5, 9.0
YC = 4.52                         # Inhalt endet hier; darunter der Phasenhinweis (4.6–5.0), dann die Master-Fußzeile
UNIT = B.UNIT
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

FARBE = {"LESEN": TEAL, "KURZ KLÄREN": OLIVE, "AUFGABE": GREEN, "VERSUCH": RED, "BESPRECHEN": ORANGE, "BAUEN": GREEN,
         "START": MUTED, "TEST": MUTED, "ABGABE": RUST, "WEG": MUTED, "EINSTIEG": ORANGE, "WIEDERHOLEN": TEAL, "ZU ZWEIT": GREEN}
TITEL = {it["id"]: it["title"] for p in TD.PARTS for it in p[3] if isinstance(it, dict)}
TITEL.update({it["id"]: it["title"] for it in TD.EINSTIEG})
T = {t["nr"]: t for t in I.T}
V = VD.VS[0]


# ================================================================ Bausteine
def phasen(teil, k="1AK"):
    st = [x for x in P.PLAN[k] if x["teile"] == [teil]][0]
    return st["phasen"]


def minuten(teil, idx):
    """Minuten der Phasen idx im Teil; unterscheiden sich die Klassen, beide nennen."""
    m = {k: sum(phasen(teil, k)[i][1] for i in idx) for k in I.KLASSEN}
    a, b = m["1AK"], m["1BK"]
    return f"{a} min" if a == b else f"{a} bzw. {b} min"


def hinweis(s, art, mins, danach):
    """Phasenhinweis unten links: was JETZT läuft und was danach kommt."""
    col = FARBE.get(art, OLIVE)
    lab = ("JETZT " + art if art not in ("START", "ABGABE", "TEST", "WEG") else art) + (" · " + mins if mins else "")
    w1 = max(2.2, min(4.6, B.FONT_B.getlength(lab) / 100 * 14 / 72 + 0.45))
    c = B.box(s, X0, 4.62, w1, 0.38, fill=col)
    B.text(c, [lab], size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
    if danach:
        B.tb(s, X0 + w1 + 0.1, 4.62, W - w1 - 0.1, 0.38, [[("Danach: ", {"bold": True, "color": col}), (danach, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE, after=0)


def karte(s, x, y, w, h, kopf, zeilen, akzent=GREEN, size=20):
    c = B.box(s, x, y, w, h, fill=CARD)
    band = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(0.08))
    band.fill.solid(); band.fill.fore_color.rgb = akzent; band.line.fill.background(); band.shadow.inherit = False
    c.text_frame.margin_top = Inches(0.14)
    B.text(c, [[(kopf, {"bold": True, "color": OLIVE})]] + list(zeilen), size=size, after=2)
    return c


def merke(s, y, txt, size=20, h=None):
    h = h or B.th(txt, size, W - 0.1, bold=True, pad=0.24)
    c = B.box(s, X0, y, W, h, fill=OLIVE)
    B.text(c, [txt], size=size, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return c, h


def frage_klick(s, y, frage, antwort, w=W, x=X0, size=20):
    """Frage fett, darunter Antwortkasten; die Antwort erscheint per Klick (Lehrer). Gibt neue y zurück."""
    hq = B.th(frage, size, w, True, pad=0.06)
    B.tb(s, x, y, w, hq, [frage], size=size, bold=True, after=0)
    ha = B.th(antwort, size, w - 0.2, pad=0.16)
    fr = B.box(s, x, y + hq, w, ha, fill=RGBColor.from_string("FFFFFF"), line=ORANGE, lw=1.5)
    a = B.tb(s, x + 0.08, y + hq + 0.02, w - 0.16, ha - 0.04, [antwort], size=size, color=RED, anchor=MSO_ANCHOR.MIDDLE, after=0)
    s._anim.append([a])
    return y + hq + ha + 0.1


def qr(s, key, x, y, size, label):
    s.shapes.add_picture(B.qr_png(B.URL[key]), Inches(x), Inches(y), Inches(size), Inches(size))
    B.tb(s, x - 0.3, y + size, size + 0.6, 0.3, [label], size=14, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)


def schritte(s, x, y, w, items, size=20, farbe=TEAL):
    for i, it in enumerate(items):
        h = B.th(it, size, w - 0.6, pad=0.06)
        c = B.box(s, x, y + 0.04, 0.42, 0.42, fill=farbe)
        B.text(c, [str(i + 1)], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, x + 0.55, y, w - 0.55, h, [it], size=size, after=0); y += max(0.5, h + 0.06)
    return y


PROBLEME = []


def titel_von(s):
    return next((ph.text_frame.text for ph in s.placeholders if ph.placeholder_format.type == 1), "?")


def check(s):
    """Kontrolle: nichts unter die Inhaltsgrenze (außer Phasenhinweis, Fußzeile, Pill)."""
    for sh in s.shapes:
        if sh.is_placeholder or sh.name == "Buchseite": continue
        unten = (sh.top + sh.height) / 914400
        if unten > YC + 0.02 and sh.top / 914400 < 4.6:
            PROBLEME.append(f"»{titel_von(s)}«: {sh.name} {sh.text_frame.text[:40] if sh.has_text_frame else ''!r} bis {unten:.2f} in")


# ================================================================ Folien je Teil
# Jede Folie: (Schlüssel, Phasen-Indizes in plan.py, Art für den Hinweis, Schritt für die »Heute«-Folie oder None, Wozu? für die Regie, Zeichenfunktion)
SPEC = {}
FOLIE = {}            # Schlüssel → Folie
NACH = {}             # Schlüssel → Text »Danach: …«


def heute_folie(teil):
    s = d.slide("Nur Titel", f"Teil {teil} · Heute")
    wann = " · ".join(f"{k}: {I.DATEN[k][I.TEIL_STUNDE[teil][k]]}" for k in I.KLASSEN)
    B.tb(s, X0, 0.98, W, 0.32, [[(I.TEILNAME[teil] + " · " + wann, {"bold": True, "color": TEAL})]], size=16, after=0)
    ziel = HEUTE_ZIEL[teil]
    c = B.box(s, X0, 1.36, 7.0, B.th("Ziel: " + ziel, 20, 6.9, pad=0.16), fill=OLIVE)
    B.text(c, [[("Ziel: ", {"bold": True}), (ziel, {})]], size=20, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, after=0)
    y = 1.36 + c.height / 914400 + 0.12
    st = [(sp[3], sp[1]) for sp in SPEC[teil] if sp[3]]
    items = [f"{lab} · {minuten(teil, idx)}" for lab, idx in st]
    if len(items) <= 4: schritte(s, X0, y, 6.6, items, size=18, farbe=GREEN)
    else:
        h1 = (len(items) + 1) // 2
        for col, part in enumerate((items[:h1], items[h1:])):
            yy = y
            for i, it in enumerate(part):
                n = i + 1 + col * h1
                c = B.box(s, X0 + col * 3.6, yy + 0.04, 0.4, 0.4, fill=GREEN)
                B.text(c, [str(n)], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
                hh = B.th(it, 18, 3.1, pad=0.04)
                B.tb(s, X0 + col * 3.6 + 0.48, yy, 3.1, hh, [it], size=18, after=0); yy += max(0.5, hh + 0.04)
    qr(s, "aufgaben", 7.75, 1.36, 1.05, "Arbeitsblatt")
    qr(s, "theorie" if teil != 2 else "versuche", 7.75, 2.85, 1.05, "Theorieteil" if teil != 2 else "Versuche")
    return s


def lesen_folie(teil, buch, abschnitte, leitfrage, idx):
    titel = "Lesen: " + (TITEL[abschnitte[0]] if len(abschnitte) == 1 else f"Theorieteil Teil {teil}")
    if len(titel) > 40: titel = f"Lesen: Theorieteil Teil {teil}"
    s = d.slide("Nur Titel", titel)
    y = 1.05
    if buch:
        bh = max(0.62, B.th(buch, 20, W - 0.95, pad=0.2))
        c = B.box(s, X0, y, W, bh, fill=CARD, line=OLIVE, lw=1.5)
        ic = B.box(s, X0 + 0.1, y + bh / 2 - 0.27, 0.62, 0.54, fill=OLIVE)
        B.text(ic, [[("BUCH", {"size": 12, "bold": True, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, X0 + 0.82, y, W - 0.9, bh, [buch], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0); y += bh + 0.12
    B.tb(s, X0, y, 6.9, 0.36, [[("Im Theorieteil genau diese Abschnitte:", {"bold": True, "color": TEAL})]], size=18, after=0); y += 0.4
    for a in abschnitte:
        c = B.box(s, X0, y, 6.9, 0.5, fill=RGBColor.from_string("E8F1F3"), line=TEAL, lw=1.25)
        B.text(c, [[("»" + TITEL[a] + "«", {"bold": True})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0); y += 0.58
    qr(s, "theorie", 7.75, 1.05 + (bh + 0.12 if buch else 0), 1.2, "Theorieteil")
    y += 0.08
    hq = B.th("Leitfrage: " + leitfrage, 20, W - 0.1, pad=0.2)
    c = B.box(s, X0, y, W, hq, fill=RGBColor.from_string("FFF4EA"), line=ORANGE, lw=1.5)
    B.text(c, [[("Leitfrage: ", {"bold": True, "color": ORANGE}), (leitfrage, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


def auftrag_folie(titel, nrs, tun, erwartet, ohne_handy=False, layout="Aufgabe"):
    """Arbeitsphase: welche Aufgaben, was genau tun, wie lange, was am Ende dastehen soll."""
    s = d.slide(layout, titel)
    rows = [["Nr.", "Aufgabe", "Zeit", "So arbeitest du"]]
    for n in nrs:
        rows.append([str(n), T[n]["title"], f"{T[n]['zeit']} min", tun.get(n, "allein am Handy")])
    cells, h = B.auto_grid(s, X0, 1.05, [0.6, 4.0, 0.9, 3.5], rows, size=18, minh=0.42)
    y = 1.05 + h + 0.12
    if ohne_handy:
        c = B.box(s, X0, y, W, 0.45, fill=RGBColor.from_string("F6E3DF"), line=RUST)
        B.text(c, [[("Handy bleibt zuerst weg: ", {"bold": True, "color": RUST}), (ohne_handy, {})]], size=18, anchor=MSO_ANCHOR.MIDDLE, after=0); y += 0.55
    if erwartet:
        he = B.th("Am Ende steht da: " + erwartet, 20, W - 0.2, pad=0.18)
        c = B.box(s, X0, y, W, he, fill=RGBColor.from_string("FFFFFF"), line=GREEN, lw=1.5)
        B.text(c, [[("Am Ende steht da: ", {"bold": True, "color": GREEN}), (erwartet, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


# ---------------------------------------------------------------- Inhalte (kurz, nur das Wichtigste)
HEUTE_ZIEL = {
    1: "Beobachtung, Hypothese und Vorhersage auseinanderhalten – und einen fairen Test planen, bevor geworfen wird.",
    2: "Den Papierflieger-Test nach festen Regeln durchziehen: drei Vergleichspaare, alle Werte aufschreiben.",
    3: "Verstehen, dass Forschung ein Kreislauf ist – und alte Daten kritisch lesen.",
    4: "Eigene Daten auswerten und vorsichtig schließen.",
}


def t1_einstieg():
    s = d.slide("Nur Titel", "Macht die Klammer den Flieger besser?")
    B.picture(s, UNIT + "img/flieger.png", X0, 1.05, 5.0, 3.35)
    karte(s, 5.75, 1.05, 3.75, 1.75, "Behauptung", ["»Eine Büroklammer an der Nase macht **jeden** Papierflieger besser.«"], ORANGE)
    c = karte(s, 5.75, 2.95, 3.75, 1.1, "Tipp ab – per Handzeichen", ["weiter · kürzer · gleich weit?"], GREEN)
    a = B.tb(s, 5.75, 4.1, 3.75, 0.36, ["Aufgelöst wird erst nach dem Test."], size=16, color=RED, bold=True, after=0)
    s._anim.append([a])
    return s


def t1_klaeren_hyp():
    s = d.slide("Nur Titel", "Kurz klären: Hypothese ≠ Vorhersage")
    tw = 2.9
    for i, (kopf, bsp, was, col) in enumerate([("Beobachtung", "»Wurf 1 landete bei 6,4 m.«", "was wir sehen oder messen", TEAL),
                                                ("Hypothese", "»Mit Klammer weiter, **weil** die Nase schwerer ist.«", "begründete, prüfbare Vermutung", GREEN),
                                                ("Vorhersage", "»Bei drei Paaren hat B den größeren Mittelwert.«", "was beim nächsten Test herauskommen sollte", ORANGE)]):
        karte(s, X0 + i * (tw + 0.15), 1.05, tw, 2.1, kopf, [bsp, [(was, {"color": MUTED, "size": 18})]], col)
    c, h = merke(s, 3.3, "Die Hypothese sagt WARUM, die Vorhersage sagt, WAS beim Test herauskommt – und beides steht fest, bevor geworfen wird.")
    s._anim.append([c])
    return s


def t1_klaeren_theorie():
    s = d.slide("Nur Titel", "Kurz klären: Theorie ist keine Vermutung")
    B.picture(s, UNIT + "img/theorie.png", X0, 1.05, 4.2, 3.0)
    y = 1.05
    y = frage_klick(s, y, "»Das ist doch nur eine Theorie!« – Stimmt das?", "Nein. Eine Theorie erklärt **viele geprüfte Befunde** (z. B. Zelltheorie).", w=4.6, x=4.9, size=20)
    y = frage_klick(s, y, "Wird eine oft bestätigte Hypothese zum Gesetz?", "Nein. Gesetz **beschreibt**, Theorie **erklärt** – keine Leiter.", w=4.6, x=4.9, size=20)
    B.tb(s, X0, 4.1, 4.2, 0.4, ["Buch S. 29 klingt nach Stufenleiter – das stimmt so nicht."], size=14, color=MUTED, after=0)
    return s


def t1_besprechen():
    s = d.slide("Aufgabe", "Besprechen: Aufgaben 1–3")
    y = 1.05
    y = frage_klick(s, y, "A1 e) »Vielleicht flog er weiter, weil Kadett Huber fester geworfen hat.«", "Hypothese – eine Vermutung über eine Ursache. Prüfbar: zweiter Werfer.")
    y = frage_klick(s, y, "A2 d) »Eine wissenschaftliche Theorie ist bloß eine geratene Vermutung.«", "stimmt nicht – sie verbindet viele geprüfte Befunde und erlaubt Vorhersagen.")
    y = frage_klick(s, y, "A3 d) »Soll das Bundesheer mehr Geld für Hubschrauber ausgeben?«", "keine NAWI-Frage – das entscheiden Menschen. Messungen liefern nur Fakten dazu.")
    return s


def t1_plan():
    s = d.slide("Nur Titel", "Unser Plan: nur EINE Sache ändern")
    rows = [["", "Unser Test"], ["Wir ändern", "nur die Klammer: A ohne, B mit"], ["Wir messen", "die Weite bis zum **ersten Bodenkontakt**, auf 0,1 m"],
            ["Bleibt gleich", "derselbe Flieger, Papier, Werfer, Abwurflinie und -höhe"], ["Wie oft", "drei Vergleichspaare, A und B abwechselnd"]]
    cells, h = B.auto_grid(s, X0, 1.05, [2.2, 6.8], [[B.plain(c) for c in r] for r in rows], size=20, minh=0.48)
    c, hm = merke(s, 1.05 + h + 0.15, "Achtung: Die Klammer ändert Masse UND Schwerpunkt. Wir sehen, ob diese EINE Änderung wirkt – nicht, welche Ursache dahinter steckt.", size=20)
    s._anim.append([c])
    return s


def t1_a4():
    s = auftrag_folie("Jetzt: Aufgabe 4 – vor dem ersten Wurf", [4], {4: "in der Gruppe, einer tippt"},
                      "a) Hypothese mit »weil …« · b) Vorhersage für A und B · c) was dagegen spräche. Danach wird NICHTS mehr geändert.")
    return s


def t1_bauen():
    s = d.slide("Nur Titel", "Bauen: Planarflieger")
    B.picture(s, UNIT + "img/faltanleitung.png", X0, 1.0, 6.0, 3.5)
    karte(s, 6.7, 1.05, 2.8, 3.35, "Wichtig", ["Alle bauen **genau diese** Form.", "Gruppennummer mit Bleistift auf den Flügel.", "Flach in die Schachtel – nicht knicken."], GREEN, size=20)
    return s


def t1_ende():
    s = d.slide("Nur Titel", "Flieger abgeben – wie geht's weiter?")
    karte(s, X0, 1.05, 4.4, 1.9, "1AK", ["Gleich in der **2. Stunde**: Papierflieger-Test im Eingangsbereich."], TEAL)
    karte(s, X0 + 4.6, 1.05, 4.4, 1.9, "1BK", ["Am **Fr 06.11.** Test + Teil 3.", "Geflogen wird am **Mo 09.11.**"], ORANGE)
    c, h = merke(s, 3.15, "Eure Hypothese aus Aufgabe 4 bleibt, wie sie ist – auch wenn ihr inzwischen eine andere Vermutung habt.", size=20)
    return s


def t2_heute_extra(s): pass


def t2_a5():
    s = d.slide("Aufgabe", "Jetzt: Aufgabe 5 – Zählt der Wurf?")
    B.tb(s, X0, 1.0, W, 0.4, [[("Allein am Handy, 3 min. ", {"bold": True, "color": GREEN}), ("Dann die heiklen Fälle gemeinsam:", {})]], size=20, after=0)
    y = 1.45
    y = frage_klick(s, y, "a) Der Flieger geht nach 2 m steil zu Boden.", "gültig – kurze Flüge zählen auch. Sonst behalten wir nur die guten.")
    y = frage_klick(s, y, "c) Er streift die Decke und landet bei 5,1 m.", "zählt – Wert bleibt, Vermerk »Decke«. Nie löschen, nie 0 eintragen.")
    y = frage_klick(s, y, "f) Gemessen wird …", "bis zum ersten Bodenkontakt – nicht dort, wo er nach dem Rutschen liegt.")
    return s


def t2_weg():
    s = d.slide("Nur Titel", "Ab in den Eingangsbereich")
    B.picture(s, UNIT + "img/messbahn.png", X0, 1.0, 5.6, 3.5)
    karte(s, 6.3, 1.05, 3.2, 3.35, "Mitnehmen", ["Flieger + Klammer", "Protokoll-Handy (Versuchsseite offen)", "Papierblatt der Gruppe", "Bleistift"], TEAL, size=20)
    return s


def t2_sicher():
    s = d.slide("Nur Titel", "Sicherheit und Rollen")
    y = 1.05
    for h_, t_ in [("Bahn frei", "nur werfen, wenn niemand in der Flugbahn ist"), ("Flach werfen", "aus Hüfthöhe, kein Kraftwurf – die Decke ist niedrig"),
                   ("Nicht rennen", "erst nach der Landung zum Flieger"), ("Stolperfallen", "Bandmaß am Rand, Taschen weg, Fluchtweg frei")]:
        c = B.box(s, X0, y + 0.05, 0.42, 0.42, fill=ORANGE)
        B.text(c, ["!"], size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        hh = B.th(h_ + ": " + t_, 18, 4.3, pad=0.06)
        B.tb(s, X0 + 0.52, y, 4.3, hh, [[(h_ + ": ", {"bold": True, "color": OLIVE}), (t_, {})]], size=20, after=0); y += max(0.55, hh + 0.04)
    karte(s, 5.55, 1.05, 3.95, 3.35, "Rollen in der Gruppe", ["**Werfer** – bleibt die ganze Reihe gleich", "**Messer** – am Landepunkt", "**Protokollführer** – tippt sofort",
                                                             "**Prüfer** (Vierergruppe) – Linie, Bahn frei"], GREEN, size=20)
    return s


def t2_ablauf():
    s = d.slide("Nur Titel", f"Versuch {I.VS_ID}: drei Vergleichspaare")
    rows = [["Gruppen", "Paar 1", "Paar 2", "Paar 3"], ["ungerade (1, 3, 5, 7)", "A, B", "B, A", "A, B"], ["gerade (2, 4, 6, 8)", "B, A", "A, B", "B, A"]]
    cells, h = B.auto_grid(s, X0, 1.05, [3.3, 1.9, 1.9, 1.9], rows, size=20, minh=0.46)
    y = 1.05 + h + 0.15
    y = schritte(s, X0, y, W, ["Probewurf mit A – zählt nicht. Danach wird nichts mehr verändert.",
                               "Flach aus Hüfthöhe, Fuß hinter der Linie. Messen bis zum ersten Bodenkontakt.",
                               "Sofort eintragen – auch kurze Würfe und Deckenkontakte (Vermerk)."], size=20, farbe=RED)
    return s


def t2_gegenprobe():
    s = d.slide("Nur Titel", "Gegenprobe: ein zweiter Werfer")
    karte(s, X0, 1.05, 4.4, 2.4, "Wer?", ["Die 1–2 Gruppen, die Stephan aufruft.", "Ein **anderer** Werfer wirft dieselben drei Paare."], TEAL)
    karte(s, X0 + 4.6, 1.05, 4.4, 2.4, "Wozu?", ["Zeigt sich das Ergebnis auch bei jemand anderem?", "Eingetragen wird **getrennt** – nicht zusammen mitteln."], ORANGE)
    merke(s, 3.65, "Ein Ergebnis wird glaubwürdiger, wenn es sich wiederholen lässt.", size=20)
    return s


def t2_ende():
    s = d.slide("Nur Titel", "Zurück und abgeben")
    schritte(s, X0, 1.1, W, ["Flieger mit Gruppennummer in die Schachtel.", "Papierblatt mit allen Werten bei Stephan abgeben.",
                             "Protokoll-Handy: Seite nicht löschen – ausgewertet wird in Stunde 4.", "Büroklammern zurück, Bandmaß aufrollen."], size=20, farbe=OLIVE)
    return s


def t3_heute_test(s): pass


def t3_klaeren():
    s = d.slide("Nur Titel", "Kurz klären: Forschung ist ein Kreislauf")
    B.picture(s, UNIT + "img/kreislauf.png", X0, 1.0, 4.4, 3.45)
    y = 1.05
    y = frage_klick(s, y, "Läuft Forschung von Schritt 1 bis 8 geradeaus?", "Nein – man geht zurück, ändert Bedingungen, findet neue Fragen.", w=4.4, x=5.1, size=20)
    y = frage_klick(s, y, "Garantiert eine Begutachtung die Wahrheit?", "Nein – sie senkt Fehler. Vertrauen entsteht durch Wiederholung.", w=4.4, x=5.1, size=20)
    return s


def t3_a6():
    s = d.slide("Aufgabe", "Jetzt: Aufgabe 6 – alte Daten lesen")
    B.picture(s, UNIT + "img/vorjahr.png", X0, 1.0, 5.3, 3.45)
    karte(s, 6.0, 1.05, 3.5, 3.35, "Zu zweit · 7 min", ["Erst **ohne Handy** an der Tabelle reden.", "Dann a)–d) eintragen.", [("Je Gruppe nur EIN Wurf je Flieger!", {"bold": True, "color": RUST})]], GREEN, size=20)
    return s


def t3_karten():
    s = d.slide("Aufgabe", "Jetzt: Aufgabe 7 – Partnercheck")
    B.tb(s, X0, 1.0, W, 0.4, [[("Ohne Handy: ", {"bold": True, "color": RUST}), ("A fragt, B antwortet – dann tauschen. Lösung per Klick, wenn alle fertig sind.", {})]], size=20, after=0)
    rows = [["Frage", "Lösung"]] + [[q, a] for q, a in I.PARTNERKARTEN[:4]]
    cells, h = B.auto_grid(s, X0, 1.45, [4.4, 4.6], rows, size=16, minh=0.36)
    for r in cells[1:]: s._anim.append([r[1]])
    return s


def t3_besprechen():
    s = d.slide("Aufgabe", "Besprechen: Aufgabe 6")
    y = 1.05
    y = frage_klick(s, y, "b) Alte 1BK: Welche Gruppe passt nicht ins Muster?", "Gruppe 3 – dort flog der kurze Flieger viel weiter.")
    y = frage_klick(s, y, "Hat Gruppe 3 also falsch gemessen?", "Nicht automatisch. Erst Bedingungen und Protokoll prüfen.")
    y = frage_klick(s, y, "d) Reichen die Daten für eine Ursache?", "Nein – nur ein Wurf je Flieger, und vieles war gleichzeitig verschieden.")
    return s


def t4_weisst():
    s = d.slide("Nur Titel", "Weißt du noch? Eure Hypothese")
    B.tb(s, X0, 1.0, W, 0.4, [[("Protokoll-Handy öffnen: ", {"bold": True, "color": TEAL}), ("Was habt ihr in Aufgabe 4 vorhergesagt?", {})]], size=20, after=0)
    y = 1.5
    for q, a in I.WEISST_DU_NOCH[4]: y = frage_klick(s, y, q, a)
    return s


def t4_klaeren():
    s = d.slide("Nur Titel", "Kurz klären: Mittelwert und Spannweite")
    B.picture(s, UNIT + "img/mittelwert.png", X0, 1.0, 4.6, 2.66)
    B.tb(s, X0, 3.7, 4.6, 0.8, ["Mittelwert = Summe : Anzahl", "Spannweite = größter − kleinster Wert"], size=20, bold=True, color=OLIVE, after=2)
    y = 1.05
    y = frage_klick(s, y, "A: 5,1 m · 6,0 m · 5,4 m. Mittelwert?", "(5,1 + 6,0 + 5,4) m : 3 = 16,5 m : 3 = 5,5 m", w=4.2, x=5.3, size=20)
    y = frage_klick(s, y, "Werte überlappen stark. Klare Sache?", "Nein – eher »noch keine Entscheidung«. Das ist erlaubt.", w=4.2, x=5.3, size=20)
    return s


def t4_auswertung():
    s = d.slide("Nur Titel", "Jetzt: Protokoll auswerten")
    y = schritte(s, X0, 1.05, 6.9, ["Mittelwert A und B, Spannweite – in die Tabelle.", "**Erkenntnis:** Wer flog im Mittel weiter, um wie viel? Kontakte?",
                                    "**Interpretation:** unterstützt · spricht dagegen · unklar – mit Daten begründen.",
                                    "**Schlussfolgerung:** eine Grenze + eine Verbesserung."], size=20, farbe=RED)
    qr(s, "versuche", 7.75, 1.05, 1.2, "Versuche")
    B.tb(s, 7.45, 2.65, 2.05, 1.5, ["ein Protokoll-Handy je Gruppe"], size=14, color=MUTED, align=PP_ALIGN.CENTER, after=0)
    return s


def t4_uebersicht():
    s = d.slide("Nur Titel", "Was sagt die ganze Klasse?")
    rows = [["Gruppe"] + [str(g) for g in range(1, VD.GRUPPEN + 1)], ["Mittelwert A in m"] + [""] * VD.GRUPPEN, ["Mittelwert B in m"] + [""] * VD.GRUPPEN, ["Kontakte?"] + [""] * VD.GRUPPEN]
    cells, h = B.auto_grid(s, X0, 1.05, [2.2] + [6.8 / VD.GRUPPEN] * VD.GRUPPEN, rows, size=20, minh=0.42)
    c = B.box(s, X0, 1.05 + h + 0.12, W, 0.95, fill=RGBColor.from_string("FFF4EA"), line=ORANGE, lw=1.5)
    B.text(c, [[("Frage an alle: ", {"bold": True, "color": ORANGE}), ("Flog B in den meisten Gruppen weiter? Streuen die Gruppen stark? Was heißt das für die Behauptung vom Anfang?", {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


def t4_auftrag():
    return auftrag_folie("Jetzt: Aufgaben 8–10", [8, 9, 10], {8: "allein am Handy", 9: "allein am Handy", 10: "zu zweit, erst reden, dann tippen"},
                         "8: drei Ergebnisse mit Einheit · 9: welche KI-Aussage stimmt · 10: Zuordnung + eure nächste Änderung.")


def t4_besprechen():
    s = d.slide("Aufgabe", "Besprechen: KI-Check (Aufgabe 9)")
    y = 1.05
    y = frage_klick(s, y, "1. »Wenn B weiter fliegt, ist eure Hypothese bewiesen.«", "stimmt nicht – sie wird unterstützt, nicht bewiesen.")
    y = frage_klick(s, y, "2. »Ein einziger kurzer B-Flug widerlegt sie.«", "stimmt nicht – es geht um den Mittelwert der ganzen Reihe.")
    y = frage_klick(s, y, "4. »Oft bestätigt → automatisch Gesetz.«", "stimmt nicht – Gesetz beschreibt, Theorie erklärt.")
    return s


def t4_abgabe():
    s = d.slide("Nur Titel", "Abgabe: beide PDFs in LMS")
    y = schritte(s, X0, 1.05, 6.9, ["Arbeitsblatt: »PDF erstellen« → in LMS hochladen.", "Versuch: Protokollführer macht für jeden eine PDF (»Ich gebe ab als …«) und schickt sie in die Gruppe.",
                                    "Jeder lädt seine Versuchs-PDF selbst in LMS.", "Geht LMS nicht: PDF per E-Mail an Stephan."], size=20, farbe=RUST)
    qr(s, "aufgaben", 7.75, 1.05, 1.15, "Arbeitsblatt"); qr(s, "versuche", 7.75, 2.75, 1.15, "Versuche")
    return s


def zusatz_folie():
    s = d.slide("Nur Titel", "OPTIONAL: Für Schnelle", hidden=True)
    rows = [["Nr.", "Zusatzaufgabe", "Zeit"]] + [[str(n), T[n]["title"], f"{T[n]['zeit']} min"] for n in (11, 12, 13, 14)]
    cells, h = B.auto_grid(s, X0, 1.05, [0.7, 6.9, 1.4], rows, size=20, minh=0.46)
    B.tb(s, X0, 1.05 + h + 0.15, W, 0.8, ["Nur, wenn die Pflicht fertig ist. Aufgabe 14 (Video, Podcast oder Plakat) ist freiwillig – Abgabe als Datei in LMS."], size=18, after=0)
    return s


# ---------------------------------------------------------------- Ablauf je Teil: (Schlüssel, Phasen, Art, Schritt auf »Heute«, Wozu?, Funktion)
SPEC = {
 1: [("t1_einst", [0, 1], "EINSTIEG", None, "Neugier wecken, Vorhersage per Handzeichen – noch nicht auflösen", t1_einstieg),
     ("t1_heute", [2], "START", None, "Ziel und Ablauf zeigen, Gruppen 1–8 einteilen", lambda: heute_folie(1)),
     ("t1_lesen", [3], "LESEN", "Lesen", "Begriffe selbst lesen – Stephan erklärt nicht vor",
      lambda: lesen_folie(1, "**Buch S. 29**, Abschnitt **4.1 »Von der Hypothese zur Theorie«**", ["hypothese", "theorie"],
                          "Was ist der Unterschied zwischen einer Hypothese und einer Vorhersage?", [3])),
     ("t1_hyp", [4], "KURZ KLÄREN", "Kurz klären", "Fehlvorstellung »Hypothese = Vorhersage« ausräumen", t1_klaeren_hyp),
     ("t1_theo", [4], "KURZ KLÄREN", None, "Fehlvorstellungen »Theorie = Vermutung« und »Stufenleiter« ausräumen", t1_klaeren_theorie),
     ("t1_a13", [5, 6, 7], "AUFGABE", "Aufgaben 1–3", "Begriffe anwenden (alles automatisch prüfbar)",
      lambda: auftrag_folie("Jetzt: Aufgaben 1–3", [1, 2, 3], {}, "15 Auswahlfelder, alle mit »Prüfen« kontrolliert. Wer fertig ist: Aufgabe 2 noch einmal begründen.")),
     ("t1_bespr", [8], "BESPRECHEN", None, "Die drei häufigsten Fehler aus A1–A3 klären", t1_besprechen),
     ("t1_plan", [9], "KURZ KLÄREN", "Plan", "Fairen Test gemeinsam festlegen (Variable, Messgröße, Konstanten)", t1_plan),
     ("t1_a4", [10], "AUFGABE", "Aufgabe 4", "Hypothese VOR dem Versuch festhalten", t1_a4),
     ("t1_bau", [11], "BAUEN", "Flieger bauen", "Alle Gruppen bauen dieselbe Form", t1_bauen),
     ("t1_ende", [12, 13], "ABGABE", None, "Flieger einsammeln, Ausblick je Klasse; Reserve", t1_ende)],
 2: [("t2_heute", [0], "START", None, "Flieger austeilen, Ziel zeigen", lambda: heute_folie(2)),
     ("t2_lesen", [1], "LESEN", "Lesen", "Regeln selbst lesen, bevor es ernst wird",
      lambda: lesen_folie(2, None, ["regeln", "ablauf"], "Bis wohin wird gemessen – und was passiert mit einem Deckenkontakt?", [1])),
     ("t2_a5", [2], "AUFGABE", "Aufgabe 5", "Regeln anwenden; heikle Fälle sofort klären (kurze Flüge, Decke, Messpunkt)", t2_a5),
     ("t2_weg", [3], "WEG", "Weg", "Material mitnehmen – im Eingangsbereich gibt es keinen Beamer", t2_weg),
     ("t2_sicher", [4], "KURZ KLÄREN", None, "Sicherheit und Rollen VOR dem Weg zeigen, flachen Abwurf vor Ort vorzeigen", t2_sicher),
     ("t2_ablauf", [5], "VERSUCH", "Versuch", "Reihenfolge der Würfe festlegen (im Saal zeigen, vor Ort gilt das Papierblatt)", t2_ablauf),
     ("t2_gegen", [6], "VERSUCH", "Gegenprobe", "Wiederholbarkeit zeigen", t2_gegenprobe),
     ("t2_ende", [7, 8], "ABGABE", None, "Flieger und Papierblätter einsammeln; Reserve", t2_ende)],
 3: [("t3_heute", [0, 1], "TEST", None, "Test (20 min); danach Ziel von Teil 3", lambda: heute_folie(3)),
     ("t3_lesen", [2], "LESEN", "Lesen", "Semmelweis und Kreislauf selbst lesen",
      lambda: lesen_folie(3, "**Buch S. 29**, Abschnitt **4.2 »Wissenschaftliche Forschung«**", ["s0c", "kreislauf"],
                          "Warum wurde Semmelweis nicht sofort geglaubt, obwohl seine Daten stimmten?", [2])),
     ("t3_klaeren", [3], "KURZ KLÄREN", "Kurz klären", "Fehlvorstellungen »Forschung geradeaus« und »Begutachtung = Wahrheit«", t3_klaeren),
     ("t3_a6", [4], "ZU ZWEIT", "Aufgabe 6", "Datenkritik ohne Handy", t3_a6),
     ("t3_a7", [5], "ZU ZWEIT", "Aufgabe 7", "Wichtigste Begriffe gegenseitig abfragen", t3_karten),
     ("t3_bespr", [6, 7], "BESPRECHEN", None, "»Abweichende Gruppe = falsch gemessen?« klären; Reserve", t3_besprechen)],
 4: [("t4_heute", [0], "START", None, "Ziel zeigen, Protokoll-Handys öffnen lassen", lambda: heute_folie(4)),
     ("t4_weiss", [1], "WIEDERHOLEN", "Hypothese von damals", "An die eigene Vorhersage erinnern (nicht mehr ändern)", t4_weisst),
     ("t4_lesen", [2], "LESEN", "Lesen", "Rechenweg und die drei möglichen Ergebnisse selbst lesen",
      lambda: lesen_folie(4, None, ["mittelwert", "schluss"], "Wann ist die ehrliche Antwort »noch keine Entscheidung möglich«?", [2])),
     ("t4_klaeren", [3], "KURZ KLÄREN", "Kurz klären", "Rechenweg einmal vorzeigen (andere Zahlen als A8), Überlappung = unklar", t4_klaeren),
     ("t4_ausw", [4], "VERSUCH", "Protokoll auswerten", "Eigene Daten auswerten – Kern der Einheit", t4_auswertung),
     ("t4_klasse", [5], "KURZ KLÄREN", "Klassenübersicht", "Alle Gruppen nebeneinander: Streuung zwischen Gruppen sichtbar machen", t4_uebersicht),
     ("t4_a810", [6, 7, 8], "AUFGABE", "Aufgaben 8–10", "Rechnen, KI-Aussagen prüfen, Konstruktionszyklus", t4_auftrag),
     ("t4_bespr", [9], "BESPRECHEN", None, "»Bewiesen / widerlegt / Stufenleiter« ein letztes Mal geraderücken", t4_besprechen),
     ("t4_abgabe", [10, 11], "ABGABE", "Abgabe", "Abgabe in LMS; Reserve", t4_abgabe)],
}

# Plan-Kontrolle: jede Phase jedes Teils ist genau einer Folie (oder mehreren Folien derselben Phase) zugeordnet
for teil, sp in SPEC.items():
    alle = sorted({i for x in sp for i in x[1]})
    assert alle == list(range(len(phasen(teil)))), (teil, alle)
    for k in I.KLASSEN: assert len(phasen(teil, k)) == len(phasen(teil)), (teil, k)

DANACH_TEIL = {1: "1AK: Teil 2 (Stunde 2) · 1BK: Teil 3 am Fr 06.11.", 2: "1AK: Teil 3 am Fr 06.11. · 1BK: Teil 4 am Mi 11.11.",
               3: "1AK: Teil 4 am Do 12.11. · 1BK: Teil 2 am Mo 09.11.", 4: "Ende der Einheit – nächstes Kapitel: Kraft (Buch S. 45)"}


# ================================================================ Deck bauen
def bauen():
    global d
    d = B.Deck(True, "NAWI · 1AB · K1.4 Erkenntnisgewinn · Unterricht"); B.d = d
    d.section("Start")
    s = d.slide("Titelfolie", "Erkenntnisgewinn")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2:
            B.text(ph, ["NAWI · 1AB · K1.4 · Unterrichts-PPP: Hypothese, Experiment, Theorie – der Papierflieger-Test · Buch S. 28–32",
                        [(f"Stand {I.STAND} · {I.V} · TH-Unterricht (Ablauf) – zum Nachlesen: TH-Lehrer", {"size": 12})]], size=18, color=RGBColor.from_string("C9D6B0"))
    titel = s
    bedien = d.slide("Nur Titel", "Lehrer-Info · So führt die PPP", hidden=True)
    regie = {}
    for k in I.KLASSEN:
        for st in P.PLAN[k]:
            regie[(k, st["std"])] = d.slide("Nur Titel", f"Regie {k} · Stunde {st['std']} · {st['datum'].split(' (')[0]}", hidden=True)
    quellen = d.slide("Nur Titel", "Lehrer-Info · Quellen", hidden=True)
    for teil in (1, 2, 3, 4):
        d.section(f"Teil {teil}")
        sec = d.slide("Abschnitt", f"Teil {teil} · " + I.TEILNAME[teil].split(" · ")[1])
        for ph in sec.placeholders:
            if ph.placeholder_format.type == 2: ph.text_frame.text = " · ".join(f"{k} {I.DATEN[k][I.TEIL_STUNDE[teil][k]].split(' (')[0]}" for k in I.KLASSEN)
        FOLIE[f"t{teil}_sec"] = sec
        sp = SPEC[teil]
        for j, (key, idx, art, schritt, wozu, fn) in enumerate(sp):
            s = fn(); FOLIE[key] = s
            nxt = sp[j + 1] if j + 1 < len(sp) else None
            if nxt:
                nn = nxt[3] or FOLIE_TITEL(nxt)
            else:
                nn = DANACH_TEIL[teil]
            hinweis(s, art, minuten(teil, idx), nn)
            B.pill(s, I.TEILBUCH[teil], w=2.2)
            check(s)
    d.section("Optional")
    zusatz = zusatz_folie()

    # ---------------- Foliennummern, Regie, Bedienung, Quellen
    ids = [el.get("id") for el in d.p.slides._sldIdLst]
    nr_von = {s.slide_id: i + 1 for i, s in enumerate(d.p.slides)}
    NR = lambda key: nr_von[FOLIE[key].slide_id]
    folgen = {}
    for k in I.KLASSEN:
        for st in P.PLAN[k]:
            teil = st["teile"][0]
            folgen[(k, st["std"])] = [FOLIE[f"t{teil}_sec"]] + [FOLIE[x[0]] for x in SPEC[teil]]
    for k in I.KLASSEN:
        for st in P.PLAN[k]:
            s = regie[(k, st["std"])]; teil = st["teile"][0]
            rows = [["Minute", "Was läuft jetzt?", "Folie", "Wozu?"]]; t0 = 0
            zu = {}
            for key, idx, art, schritt, wozu, fn in SPEC[teil]:
                for i in idx: zu.setdefault(i, []).append((key, wozu))
            for i, (name, mi, oh, a) in enumerate(st["phasen"]):
                keys = [x for x, _ in zu[i]]
                fol = ", ".join(dict.fromkeys(str(NR(x)) for x in keys))
                wz = zu[i][0][1] if a != "reserve" else "Puffer – nicht verplanen"
                rows.append([f"{t0}–{t0 + mi}", name + (" · ohne Handy" if oh and a not in ("org", "reserve", "test", "bespr", "erkl") else ""), fol, wz]); t0 += mi
            assert t0 == 50
            B.tb(s, X0, 0.98, W, 0.3, [[(f"{k} · {st['datum']} · Teil {teil} · {I.TEILNAME[teil].split(' · ')[1]} · Benutzerdefinierte Präsentation »{k} {st['datum'].split(' (')[0]} Stunde {st['std']}«", {"bold": True, "color": RUST})]], size=11, after=0)
            kn = P.KNAPP.get((k, st["std"]))
            hk = B.th("Wenn es knapp wird: " + kn, 10, W, pad=0.04) if kn else 0
            size = 11
            while size > 8 and B.table_h([0.8, 4.0, 0.7, 3.5], rows, size=size, minh=0.2) + hk > 3.95: size -= 1
            cells, h = B.auto_grid(s, X0, 1.28, [0.8, 4.0, 0.7, 3.5], rows, size=size, minh=0.2)
            if kn: B.tb(s, X0, 1.3 + h, W, hk, [[("Wenn es knapp wird: ", {"bold": True, "color": RUST}), (kn, {})]], size=10, after=0)
    B.tb(bedien, X0, 0.98, W, 4.1, [
        [("Eine Datei, beide Klassen. ", {"bold": True, "color": OLIVE}), ("F5 zeigt die sichtbaren Folien in der Reihenfolge der 1AK (Teil 1 → 2 → 3 → 4). Die 1BK macht Teil 1 → 3 → 2 → 4.", {})],
        [("Pro Stunde starten: ", {"bold": True, "color": OLIVE}), ("Bildschirmpräsentation → Benutzerdefinierte Bildschirmpräsentation → z. B. »1BK Fr 06.11. Stunde 2«. Ohne Menü: im Vortrag die Foliennummer tippen und Enter (Nummern auf den Regie-Folien).", {})],
        [("Jede Stunde läuft gleich: ", {"bold": True, "color": OLIVE}), ("»Heute« → »Lesen: …« (Abschnitte im Theorieteil und im Buch, mit Leitfrage) → »Kurz klären« (nur Fehlvorstellungen, Antwort per Klick) → »Jetzt: Aufgabe …« → »Besprechen« (Lösung per Klick, erst nach der Arbeitszeit).", {})],
        [("Unten links ", {"bold": True, "color": OLIVE}), ("steht immer, was JETZT läuft (mit Minuten aus dem Zeitplan) und was DANACH kommt – so sehen auch die Kadetten den Ablauf.", {})],
        [("Teil 2 im Eingangsbereich: ", {"bold": True, "color": OLIVE}), ("Dort gibt es keinen Beamer. Folien »Sicherheit und Rollen« und »drei Vergleichspaare« im Saal zeigen, vor Ort gilt das Papierblatt (VS-AB).", {})],
        [("Nachschlagen statt zeigen: ", {"bold": True, "color": OLIVE}), ("Die ausführliche TH-Lehrer-PPP und der Theorieteil bleiben vollständig – hier steht nur das Wichtigste. Zusatzaufgaben 11–14: ausgeblendete Folie »OPTIONAL«, eigene Präsentation »Optional: Für Schnelle«.", {})],
        [("Regiepläne: ", {"bold": True, "color": OLIVE}), (f"Folien {nr_von[regie[('1AK', 1)].slide_id]}–{nr_von[regie[('1BK', 4)].slide_id]} (ausgeblendet), zusammen als Präsentation »Lehrer-Regiepläne«.", {})]],
        size=13, after=5)
    B.tb(quellen, X0, 0.98, W, 4.1, ["· " + x for x in [
        "Inhalt und Zeiten: dieselbe Datenquelle wie Theorieteil, Arbeitsblatt, Versuchsseite und TH-Lehrer (werkzeug/quellen/1ab-k1-4-erkenntnis: inhalt.py, theorie_daten.py, plan.py).",
        "Lesestopps: Trauner NAWI HAK 1, S. 29 (4.1, 4.2); Abschnittstitel exakt wie im Theorieteil.",
        "Kreislauf-Idee: Hölzel NAWI HAK 1, S. 14 (eigene Zeichnung). Faltung: W. Gruber, »Ballons, Raketen und Papierflieger«, S. 34–35 (eigene Zeichnung).",
        "Vorjahresdaten (Aufgabe 6): bisherige Unterrichts-PPP von Stephan, Folien 13–31.",
        "Alle Bilder selbst gezeichnet (skizzen.py) – keine Buchabbildungen, darum auch als Datei weitergebbar.",
        "Grundlage: ChatGPT-Übergabe E3 v4 (Planung); gebaut und geprüft von Claude. Abspieltest in PowerPoint steht aus (nur LibreOffice gerendert)."]], size=13, after=5)

    # ---------------- benutzerdefinierte Präsentationen
    pres = d.p.part._element
    csl = etree.Element("{%s}custShowLst" % NS_P)
    anker = pres.find("{%s}defaultTextStyle" % NS_P)
    if anker is not None: anker.addprevious(csl)
    else: pres.append(csl)
    rid = {s.slide_id: el.get("{%s}id" % NS_R) for s, el in zip(d.p.slides, d.p.slides._sldIdLst)}
    nid = [0]

    def show(name, folien):
        cs = etree.SubElement(csl, "{%s}custShow" % NS_P, name=name, id=str(nid[0])); nid[0] += 1
        l = etree.SubElement(cs, "{%s}sldLst" % NS_P)
        for sl in folien: etree.SubElement(l, "{%s}sld" % NS_P).set("{%s}id" % NS_R, rid[sl.slide_id])

    for k in I.KLASSEN:
        show(f"{k} gesamt", [titel] + sum([folgen[(k, st["std"])] for st in P.PLAN[k]], []))
    for k in I.KLASSEN:
        for st in P.PLAN[k]:
            show(f"{k} {st['datum'].split(' (')[0]} Stunde {st['std']}", ([titel] if st["std"] == 1 else []) + folgen[(k, st["std"])])
    show("Lehrer-Regiepläne", [bedien] + list(regie.values()))
    show("Optional: Für Schnelle", [zusatz])

    os.makedirs(ZIEL, exist_ok=True)
    name = f"{I.PRAEFIX}-TH-Unterricht-{I.V}.pptx"
    d.save(os.path.join(ZIEL, name))
    sicht = [s for s in d.p.slides if s._element.get("show") != "0"]
    for x in PROBLEME: print("ÜBERLAUF", x)
    print("ok", name, len(d.p.slides._sldIdLst), "Folien,", len(sicht), "sichtbar ·", nid[0], "Präsentationen")
    for k in I.KLASSEN:
        for st in P.PLAN[k]:
            print(" ", k, st["std"], st["datum"], "Folien", [nr_von[x.slide_id] for x in folgen[(k, st["std"])]])


def FOLIE_TITEL(spec):
    """Kurzer »Danach«-Text für Folien ohne eigenen Schritt."""
    return {"t1_theo": "Kurz klären: Theorie", "t1_bespr": "Besprechen A1–A3", "t1_ende": "Flieger abgeben", "t2_sicher": "Sicherheit und Rollen",
            "t2_ende": "zurück, abgeben", "t3_bespr": "Besprechen A6", "t4_bespr": "Besprechen A9", "t1_heute": "Heute"}.get(spec[0], spec[0])


if __name__ == "__main__":
    bauen()
