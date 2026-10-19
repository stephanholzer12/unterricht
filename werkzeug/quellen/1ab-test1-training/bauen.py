"""Test-Training NAWI 1AB · 1. Test – Quiz aus werkzeug/vorlagen/quiz_vorlage.html.
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/1ab-test1-training/bauen.py
Schreibt 1ab-nawi/test1-training/quiz.html. Fragen in eigenen Worten; Seiten nach Trauner NAWI HAK 1.
Für eine andere Klasse: Ordner kopieren, ORDNER/EYEBROW/TITEL/QUIZ anpassen."""
import base64, io, json, os, re
import qrcode

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
WZ = os.path.join(ROOT, "werkzeug")
BASE = "https://stephanholzer12.github.io/unterricht/"
ORDNER = "1ab-nawi/test1-training"
EYEBROW = "NAWI · 1AB · Training für den 1. Test"
TITEL = "Test-Training"
KLASSE = "1AB"

import sys as _s; _s.path.insert(0, HERE)
from fragen import QUIZ

b64 = lambda p: base64.b64encode(open(p, "rb").read()).decode()
nb = lambda s: re.sub(r"(?<=\d) (?=\d{3}\b)", " ", s)


def qr_b64(url):
    q = qrcode.QRCode(border=2, box_size=12, error_correction=qrcode.constants.ERROR_CORRECT_M); q.add_data(url); q.make()
    b = io.BytesIO(); q.make_image(fill_color="#2C3424", back_color="white").save(b, format="PNG"); return base64.b64encode(b.getvalue()).decode()


def wrap(src):
    te = src.index("</title>") + 8; rest = src[te:]; bs = rest.index("<header")
    return ('<!doctype html>\n<html lang="de">\n<head>\n<meta charset="utf-8">\n<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
            '<meta name="robots" content="noindex">\n' + src[:te] + "\n" + rest[:bs] + "</head>\n<body>\n" + rest[bs:] + "\n</body>\n</html>\n")


assert all(len(a) == 4 for _, _, a, _, _ in QUIZ)
qs = ",\n".join("  { t: %s, q: %s, a: %s, e: %s, p: %s }" % tuple(json.dumps(x, ensure_ascii=False) for x in (t, nb(q), [nb(z) for z in a], nb(e), p)) for t, q, a, e, p in QUIZ)
qv = open(os.path.join(WZ, "vorlagen", "quiz_vorlage.html"), encoding="utf-8").read()
qv = qv.replace("<title>Bindungs-Check</title>", f"<title>{TITEL} – {KLASSE}</title>")
qv = qv.replace("NAWI · 2AB · K1.2 Die Bindungsarten", EYEBROW)
qv = qv.replace("<h1>Bindungs-Check</h1>", f"<h1>{TITEL}</h1>").replace("<h2>Bindungs-Check am Handy</h2>", f"<h2>{TITEL} am Handy</h2>")
qv = qv.replace('alt="QR-Code zum Bindungs-Check"', f'alt="QR-Code zum {TITEL}"')
qv = qv.replace("Frage 1 von 16", f"Frage 1 von {len(QUIZ)}")
a, b = qv.index("const Q = ["), qv.index("];", qv.index("const Q = [")) + 2
qv = qv[:a] + "const Q = [\n" + qs + "\n];" + qv[b:]
n = len(QUIZ)
qv = qv.replace("if (p >= 13) return", f"if (p >= {round(n * 0.85)}) return").replace("if (p >= 9) return", f"if (p >= {round(n * 0.6)}) return")
if "[hidden]" not in qv: qv = qv.replace("<style>", "<style>\n[hidden] { display: none !important; }", 1)
for k, v in {"__CAMO__": b64(WZ + "/camo.png"), "__LOGO__": b64(WZ + "/logo.png"), "__SMILEY__": b64(WZ + "/smiley.png"),
             "__QR__": qr_b64(BASE + ORDNER + "/quiz.html")}.items():
    qv = qv.replace(k, v)
assert not re.search(r"__[A-Z]+__", qv) and "Bindung" not in qv and "4AB" not in qv
os.makedirs(os.path.join(ROOT, ORDNER), exist_ok=True)
open(os.path.join(ROOT, ORDNER, "quiz.html"), "w", encoding="utf-8").write(wrap(qv))
print("quiz.html:", n, "Fragen")


# ---- Klassenspiel: dieselben Fragen als Kahoot-/Blooket-Tabelle (Regelwerk 10) ----
# Aufruf mit Zielordner:  python3 werkzeug/quellen/1ab-test1-training/bauen.py <ordner>
import random, sys, csv
if len(sys.argv) > 1:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    ziel = sys.argv[1]; os.makedirs(ziel, exist_ok=True)
    random.seed(1)
    rows = []
    for i, (t, q, a, e, p) in enumerate(QUIZ, 1):
        assert len(q) <= 120, ("Frage zu lang", i, len(q)); assert all(len(x) <= 75 for x in a), ("Antwort zu lang", i)
        pos = random.randrange(4); ans = a[1:]; ans.insert(pos, a[0])
        zeit = 30 if re.search(r"\d", q) and any(w in q for w in ("Wie groß", "Wie viel", "Mittelwert", "Masse hat")) else 20
        rows.append([i, q, *ans, zeit, pos + 1, t, p])
    wb = openpyxl.Workbook(); an = wb.active; an.title = "Anleitung"
    for z in ["Klassenspiel 1AB · 1. Test NAWI – dieselben Fragen wie im Test-Training",
              "Kahoot: kahoot.com → Erstellen → Fragen importieren (Tabelle) → Blatt »Kahoot« als eigene Datei speichern und hochladen.",
              "Blooket: blooket.com → Create → Spreadsheet Import → Blatt »Blooket« bzw. die CSV-Datei.",
              "Am Lehrer-PC am Beamer spielen; Kadetten nur mit Spitznamen, kein Konto. Handyfrei: Team-Modus (ein Gerät pro Gruppe) oder ABCD-Karten.",
              "Die richtige Antwort steht gemischt an Platz 1–4 (Spalte »Richtig«)."]:
        an.append([z])
    an.column_dimensions["A"].width = 120
    kopf = ["Nr.", "Frage", "Antwort 1", "Antwort 2", "Antwort 3", "Antwort 4", "Zeit (s)", "Richtig", "Thema", "Buch"]
    for name in ("Kahoot", "Blooket"):
        ws = wb.create_sheet(name); ws.append(kopf)
        for r in rows: ws.append(r)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="556B2F")
        for col, w in zip("ABCDEFGHIJ", (5, 60, 30, 30, 30, 30, 8, 8, 18, 22)): ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(os.path.join(ziel, "26-27-1AB-Test1-Kahoot-v1.xlsx"))
    with open(os.path.join(ziel, "26-27-1AB-Test1-Blooket-v1.csv"), "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh); w.writerow(["Question #", "Question Text", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Time Limit (sec)", "Correct Answer(s)"])
        for r in rows: w.writerow(r[:8])
    print("Kahoot-Tabelle und Blooket-CSV:", len(rows), "Fragen")
