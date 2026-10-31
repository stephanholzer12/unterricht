"""Präsentationen NAWI 3AB · K2.1 Formen der Landwirtschaft mit dem BHAK-Folienmaster: TH-Lehrer und TH-Kadetten (keine Versuche → keine VS-PPP).
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/3ab-k2-1-landbau/bauen_ppp.py <BHAK_Folienmaster_Vorlage.pptx> <Zielordner> [--kadetten]
Theorie-Folien kommen aus theorie_daten.py – dieselbe Quelle wie die Theorie-Seite (GitHub und PPP gleich).
Lehrer: Lösungen per Klick, ausgeblendete Lehrer-Infos (Lehrplanbezug, Stundenvorschlag, Stolpersteine). Kadetten: gleiche Folien, leere Antwortkästen."""
import copy, io, os, re, sys, uuid
import qrcode
from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, HERE)
import inhalt as I
import theorie_daten as TD
import fachwissen as F
import plan as PL

MASTER, OUT = sys.argv[1], sys.argv[2]
UNIT = os.path.join(ROOT, I.ORDNER) + "/"
VER = os.environ.get("PPP_V", I.V)
OLIVE, INK, CARD, GREEN, ORANGE, TEAL, RUST, MUTED, WHITE, PALE = (RGBColor.from_string(x) for x in
    ["3B4A2F", "1F2A1C", "EEF3E4", "7FA83A", "D9772B", "2E7D8C", "8C3B2E", "5E6652", "FFFFFF", "F7F9F2"])
RED = RGBColor.from_string("A33A2B")
GREY = RGBColor.from_string("DCDFD6")
FONT_R = ImageFont.truetype("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf", 100)
FONT_B = ImageFont.truetype("/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf", 100)
URL = {k: I.URL + k + ".html" for k in ("theorie", "aufgaben", "quiz")}
X0, X1, Y0, Y1 = 0.5, 9.5, 1.12, 5.08


# ================================================================ Messen
def plain(t):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", t)


def lines_needed(t, size, w_in, bold=False):
    """Zeilen, die ein Text bei Schriftgröße size (pt) in einer Box der Breite w_in (Zoll) braucht (Calibri ≈ Carlito)."""
    font = FONT_B if bold else FONT_R
    maxw = (w_in - 0.22) * 72 / size * 100 * 0.97
    n = 0
    for para in plain(t).split("\n"):
        words, cur = para.split(" "), ""
        n += 1
        for wd in words:
            test = (cur + " " + wd).strip()
            if font.getlength(test) > maxw and cur:
                n += 1; cur = wd
            else:
                cur = test
    return n


def th(t, size, w_in, bold=False, ls=1.17, pad=0.1):
    return lines_needed(t, size, w_in, bold) * size * ls / 72 + pad


# ================================================================ Deck
def qr_png(url):
    q = qrcode.QRCode(border=1, box_size=10); q.add_data(url); q.make(fit=True)
    b = io.BytesIO(); q.make_image(fill_color="#1F2A1C", back_color="white").save(b, "PNG"); b.seek(0); return b


class Deck:
    def __init__(self, lehrer, fuss):
        self.lehrer = lehrer
        self.p = Presentation(MASTER)
        lst = self.p.slides._sldIdLst
        self.pill_proto = None
        for sh in self.p.slides[2].shapes:
            if sh.name == "Buchseite": self.pill_proto = copy.deepcopy(sh._element)
        for sid in list(lst):
            self.p.part.drop_rel(sid.get(qn("r:id"))); lst.remove(sid)
        self.L = {l.name: l for l in self.p.slide_layouts}
        self.sections = []
        for part in self.p.part.package.iter_parts():
            if "slideLayout" in str(part.partname) or "slideMaster" in str(part.partname):
                for t in part._element.iter(qn("a:t")):
                    if t.text and "Fach · Jahrgang · Thema" in t.text: t.text = fuss

    def section(self, name):
        self.sections.append((name, []))

    def slide(self, layout, title, hidden=False):
        s = self.p.slides.add_slide(self.L[layout])
        for ph in s.placeholders:
            if ph.placeholder_format.type == 1:
                ph.text_frame.text = title
                if layout in ("Nur Titel", "Aufgabe") and len(title) > 36:   # lange Titel kleiner, damit sie den Inhalt nicht berühren
                    for r_ in ph.text_frame.paragraphs[0].runs: r_.font.size = Pt(24 if len(title) <= 44 else 21)
        if hidden: s._element.set("show", "0")
        self.sections[-1][1].append(self.p.slides._sldIdLst[-1].get("id"))
        s._anim = []
        return s

    def save(self, path):
        for s in self.p.slides:
            if getattr(s, "_anim", None): add_clicks(s, s._anim)
        prs = self.p.part._element
        ext_lst = prs.find(qn("p:extLst"))
        if ext_lst is None: ext_lst = etree.SubElement(prs, qn("p:extLst"))
        P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        for old in ext_lst.findall(qn("p:ext")):
            if old.get("uri") == "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}": ext_lst.remove(old)
        ext = etree.SubElement(ext_lst, qn("p:ext")); ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        sl = etree.SubElement(ext, "{%s}sectionLst" % P14, nsmap={"p14": P14})
        for name, ids in self.sections:
            se = etree.SubElement(sl, "{%s}section" % P14); se.set("name", name); se.set("id", "{%s}" % str(uuid.uuid4()).upper())
            il = etree.SubElement(se, "{%s}sldIdLst" % P14)
            for i in ids: etree.SubElement(il, "{%s}sldId" % P14).set("id", i)
        self.p.save(path)


def add_clicks(slide, groups):
    nid = [2]
    def cid():
        nid[0] += 1; return str(nid[0])
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    clicks = ""
    for grp in groups:
        inner = ""
        for k, sh in enumerate(grp):
            spid = sh.shape_id
            inner += (f'<p:par><p:cTn id="{cid()}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{"clickEffect" if k == 0 else "withEffect"}">'
                      f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                      f'<p:set><p:cBhvr><p:cTn id="{cid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                      f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                      f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                      f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid()}" dur="400"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                      f'</p:childTnLst></p:cTn></p:par>')
        clicks += (f'<p:par><p:cTn id="{cid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{cid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'
                   f'</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{"".join(f"<p:bldP spid=\"{sh.shape_id}\" grpId=\"0\" animBg=\"1\"/>" for grp in groups for sh in grp)}</p:bldLst></p:timing>')
    el = etree.fromstring(xml)
    sld = slide._element
    ext = sld.find(qn("p:extLst"))
    if ext is not None: ext.addprevious(el)
    else: sld.append(el)


def box(s, x, y, w, h, fill=None, line=None, radius=True, lw=1.25):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: shp.adjustments[0] = 0.08
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    return shp


def runs_md(t):
    out, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*", t):
        if m.start() > pos: out.append((t[pos:m.start()], {}))
        out.append((m.group(1), {"bold": True})); pos = m.end()
    if pos < len(t): out.append((t[pos:], {}))
    return out or [("", {})]


def text(target, parts, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4):
    tf = target.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    first = True
    for para in parts:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(after)
        runs = para if isinstance(para, list) else runs_md(para)
        for t, o in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.color.rgb = o.get("color", color)
            if o.get("italic"): f.italic = True
    return target


def tb(s, x, y, w, h, parts, **kw):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.margin_left = t.text_frame.margin_right = Inches(0.05)
    return text(t, parts, **kw)


def pill(s, label, x=3.55, y=5.18, w=2.2, kind="theory"):
    el = copy.deepcopy(d.pill_proto)
    ids = [int(v) for v in s.shapes._spTree.xpath("//p:cNvPr/@id")]
    el.xpath(".//p:cNvPr")[0].set("id", str(max(ids + [1]) + 1))
    s.shapes._spTree.append(el)
    shp = s.shapes[-1]
    shp.left, shp.top, shp.width = Inches(x), Inches(y), Inches(w)
    col = {"theory": RGBColor.from_string("6E7F4E"), "book": TEAL, "versuch": RED}[kind]
    shp.fill.solid(); shp.fill.fore_color.rgb = col
    r = shp.text_frame.paragraphs[0].runs
    r[0].text = label
    for extra in r[1:]: extra.text = ""
    return shp


def kind_of(src):
    return "versuch" if src.startswith("Versuch im Buch") else ("book" if src.startswith("Buch-Aufgabe") else "theory")


def tag(s, label, color=OLIVE, x=7.6, y=0.95, w=1.9):
    t = box(s, x, y, w, 0.32, fill=color)
    text(t, [label], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return t


def grid(s, x, y, colw, rows, rowh, head=True, size=16, fills=None):
    cells, yy = [], y
    for ri, row in enumerate(rows):
        xx, rc = x, []
        h = rowh[ri] if isinstance(rowh, (list, tuple)) else rowh
        for ci, val in enumerate(row):
            hd = head and ri == 0
            fill = OLIVE if hd else (fills[ri] if fills else (CARD if ri % 2 else WHITE))
            c = box(s, xx, yy, colw[ci], h, fill=fill, line=RGBColor.from_string("B9C4A6"), radius=False, lw=0.75)
            text(c, [val] if isinstance(val, str) else val, size=size, color=WHITE if hd else INK, bold=hd, anchor=MSO_ANCHOR.MIDDLE, after=0)
            rc.append(c); xx += colw[ci]
        cells.append(rc); yy += h
    return cells


def auto_grid(s, x, y, colw, rows, size=16, head=True, minh=0.36, fills=None):
    """Tabelle mit gemessenen Zeilenhöhen. Gibt (Zellen, Gesamthöhe) zurück."""
    hs = [max(minh, max(th(str(c), size, w, bold=(head and ri == 0), pad=0.1) for c, w in zip(r, colw))) for ri, r in enumerate(rows)]
    return grid(s, x, y, colw, rows, hs, head=head, size=size, fills=fills), sum(hs)


def table_h(colw, rows, size=16, head=True, minh=0.36):
    return sum(max(minh, max(th(str(c), size, w, bold=(head and ri == 0), pad=0.1) for c, w in zip(r, colw))) for ri, r in enumerate(rows))


def answer(s, x, y, w, h, title, body, lehrer, hint="Deine Antwort – am Arbeitsblatt eintragen", size=20):
    if lehrer:
        b = box(s, x, y, w, h, fill=WHITE, line=ORANGE, lw=2)
        text(b, ([[(title, {"bold": True, "color": ORANGE, "size": 16})]] if title else []) + (body if isinstance(body, list) else [body]), size=size, after=2)
        s._anim.append([b])
    else:
        b = box(s, x, y, w, h, fill=PALE, line=RGBColor.from_string("B9C4A6"), lw=1.5)
        text(b, [[(hint, {"italic": True, "color": MUTED, "size": 16})]], size=16)
    return b


def ppt_qr(s, url, x, y, size, step, label, color=OLIVE, linksize=8):
    s.shapes.add_picture(qr_png(url), Inches(x), Inches(y), Inches(size), Inches(size))
    c = box(s, x + size / 2 - 0.2, y - 0.25, 0.4, 0.4, fill=color)
    text(c, [step], size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    tb(s, x - 0.3, y + size + 0.02, size + 0.6, 0.32, [label], size=16, bold=True, color=color, align=PP_ALIGN.CENTER, after=0)
    tb(s, x - 0.45, y + size + 0.34, size + 0.9, 0.5, [url], size=linksize, color=MUTED, align=PP_ALIGN.CENTER, after=0)


def picture(s, path, x, y, w=None, h=None):
    from PIL import Image
    im = Image.open(path); iw, ih = im.size
    if w and h:
        k = min(w / iw, h / ih); w2, h2 = iw * k, ih * k
        return s.shapes.add_picture(path, Inches(x + (w - w2) / 2), Inches(y), Inches(w2), Inches(h2)), h2
    if w: return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(w * ih / iw)), w * ih / iw
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(h * iw / ih), Inches(h)), h



# ---------------------------------------------------------------- abwechslungsreiche Listen (Stil 1AB »Messen«)
AKZENT = [GREEN, TEAL, ORANGE]
_akz = [0]


def akzent():
    _akz[0] += 1
    return AKZENT[_akz[0] % len(AKZENT)]


def lead_split(item):
    m = re.match(r"\*\*(.+?)\*\*\s*[:.–-]?\s*(.+)$", item, re.S)
    if m and len(m.group(1)) <= 48 and len(m.group(2)) > 12:
        return m.group(1).rstrip(":"), m.group(2)
    return None, item


def ul_mode(items, w):
    if w >= 8 and 2 <= len(items) <= 3 and all(len(plain(x)) <= 210 for x in items): return "spalten"
    return "zeilen"


def spalten_h(items, w):
    n = len(items); cw = (w - 0.2 * (n - 1)) / n; hs = []
    for it in items:
        lead, body = lead_split(it)
        hs.append(0.16 + (th(lead, 20, cw - 0.25, True, pad=0.04) if lead else 0) + th(body, 20, cw - 0.25, pad=0.16))
    return max(hs)


def zeilen_h(items, w):
    return sum(max(0.52, th(it, 20, w - 0.8, pad=0.14)) + 0.08 for it in items)


def ul_h(items, w):
    return spalten_h(items, w) + 0.04 if ul_mode(items, w) == "spalten" else zeilen_h(items, w)


def draw_ul(s, items, x, y, w, nummern=False):
    col = akzent()
    if ul_mode(items, w) == "spalten":
        n = len(items); cw = (w - 0.2 * (n - 1)) / n; h = spalten_h(items, w)
        for i, it in enumerate(items):
            c_ = AKZENT[(_akz[0] + i) % len(AKZENT)]
            xx = x + i * (cw + 0.2)
            box(s, xx, y, cw, h, fill=CARD)
            band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(xx), Inches(y), Inches(cw), Inches(0.09))
            band.fill.solid(); band.fill.fore_color.rgb = c_; band.line.fill.background(); band.shadow.inherit = False
            lead, body = lead_split(it)
            parts = ([[(lead, {"bold": True, "color": OLIVE})]] if lead else []) + [body]
            tb(s, xx + 0.1, y + 0.14, cw - 0.2, h - 0.18, parts, size=20, after=3)
        return h + 0.04
    yy = y
    for i, it in enumerate(items):
        h = max(0.52, th(it, 20, w - 0.8, pad=0.14))
        box(s, x, yy, w, h, fill=PALE if i % 2 else CARD)
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.1), Inches(yy + h / 2 - 0.2), Inches(0.4), Inches(0.4))
        c.fill.solid(); c.fill.fore_color.rgb = col; c.line.fill.background(); c.shadow.inherit = False
        c.text_frame.margin_left = c.text_frame.margin_right = c.text_frame.margin_top = c.text_frame.margin_bottom = 0
        text(c, [str(i + 1) if nummern else "›"], size=16 if nummern else 22, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        tb(s, x + 0.62, yy + 0.02, w - 0.72, h - 0.04, [it], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
        yy += h + 0.08
    return yy - y

# ================================================================ Theorie-Blöcke → Folien
def blk_h(b, w):
    k = b[0]
    if k in ("ul", "ol"): return ul_h(b[1], w)
    if k == "merke": return th(b[1], 20, w - 0.1, bold=True, pad=0.22)
    if k == "hint": return th(b[1], 20, w - 0.2, pad=0.14)
    if k == "table":
        cw = col_widths(b[1], w); return table_h(cw, b[1])
    if k == "tiles":
        n = len(b[1]); per = 3 if n in (3, 6) else (n if n <= 3 else 2); tw = (w - 0.15 * (per - 1)) / per
        rows = [b[1][i:i + per] for i in range(0, n, per)]
        return sum(max(th(h_, 20, tw, True, pad=0.05) + sum(th(x, 20, tw, pad=0.04) for x in (bb if isinstance(bb, list) else [bb])) + 0.22 for h_, bb in r) + 0.12 for r in rows)
    if k == "video": return 0.95
    if k == "buch": return max(0.62, th(b[1], 20, w - 0.95, pad=0.2))
    if k == "versuch": return th(b[2], 20, w - 0.3, True, pad=0.05) + th(b[3], 20, w - 0.3, pad=0.05) + 0.55
    if k == "reveal": return sum(th(q, 20, w, True, pad=0.05) + max(0.75, th(a, 20, w - 0.2, pad=0.5)) + 0.1 for q, a in b[1])
    if k == "karten": return 0
    raise ValueError(k)


def col_widths(rows, w, size=16):
    """Spaltenbreiten nach Textmenge – aber jede Spalte mindestens so breit wie ihr längstes Wort (sonst »Begri-ff«)."""
    n = len(rows[0]); lens = [max(len(plain(str(r[i]))) for r in rows) for i in range(n)]
    tot = sum(max(l, 8) for l in lens); cw = [w * max(l, 8) / tot for l in lens]
    wort = [max(FONT_B.getlength(x) for r in rows for x in plain(str(r[i])).replace("\n", " ").split(" ") if x) / 100 * size / 72 + 0.25 for i in range(n)]
    for _ in range(5):
        zu = [i for i in range(n) if cw[i] < wort[i]]
        if not zu: break
        fehlt = sum(wort[i] - cw[i] for i in zu)
        for i in zu: cw[i] = wort[i]
        rest = [i for i in range(n) if i not in zu]
        rs = sum(cw[i] for i in rest)
        for i in rest: cw[i] -= fehlt * cw[i] / rs
    return cw


def draw_blk(s, b, x, y, w, L):
    k = b[0]
    if k in ("ul", "ol"):
        return draw_ul(s, b[1], x, y, w, nummern=(k == "ol"))
    if k == "merke":
        h = blk_h(b, w); c = box(s, x, y, w, h - 0.05, fill=OLIVE); text(c, [b[1]], size=20, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0); return h
    if k == "hint":
        h = blk_h(b, w)
        strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.04), Inches(0.07), Inches(h - 0.1))
        strip.fill.solid(); strip.fill.fore_color.rgb = TEAL; strip.line.fill.background(); strip.shadow.inherit = False
        tb(s, x + 0.17, y, w - 0.2, h, [[(t_, dict(o, color=MUTED)) for t_, o in runs_md(b[1])]], size=20); return h
    if k == "table":
        cw = col_widths(b[1], w); cells, h = auto_grid(s, x, y, cw, b[1]); return h + 0.08
    if k == "tiles":
        n = len(b[1]); per = 3 if n in (3, 6) else (n if n <= 3 else 2); tw = (w - 0.15 * (per - 1)) / per; yy = y
        for i in range(0, n, per):
            r = b[1][i:i + per]
            rh = max(th(h_, 20, tw, True, pad=0.05) + sum(th(x_, 20, tw, pad=0.04) for x_ in (bb if isinstance(bb, list) else [bb])) + 0.22 for h_, bb in r)
            for j, (h_, bb) in enumerate(r):
                c = box(s, x + j * (tw + 0.15), yy, tw, rh, fill=CARD)
                band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + j * (tw + 0.15)), Inches(yy), Inches(tw), Inches(0.08))
                band.fill.solid(); band.fill.fore_color.rgb = AKZENT[(i // per + j) % 3]; band.line.fill.background(); band.shadow.inherit = False
                c.text_frame.margin_top = Inches(0.13)
                text(c, [[(h_, {"bold": True, "color": OLIVE})]] + list(bb if isinstance(bb, list) else [bb]), size=20, after=2)
            yy += rh + 0.12
        return yy - y
    if k == "video":
        t_, q, du, u = I.VIDEOS[b[1]][:4]
        art = I.VIDEOS[b[1]][4] if len(I.VIDEOS[b[1]]) > 4 else "video"; sim = art == "sim"
        q2 = q + ("" if du == "kurz" else " · " + du) + (" · NICHT GEPRÜFT – vorher ansehen, erst dann zeigen" if art == "buch" and b[1] not in I.FREIGEGEBEN else "")
        c = box(s, x, y, w, 0.88, fill=RGBColor.from_string("E8F1F3" if sim else "FFF4EA"), line=TEAL if sim else ORANGE, lw=1.5)
        c.text_frame.margin_right = Inches(0.95)
        text(c, [[({"sim": "Simulation: ", "buch": "Aus dem Buch: ", "audio": "Podcast: "}.get(art, "Video: ") + t_, {"bold": True})], [(q2, {"size": 16 if art != "buch" else 13, "color": RED if art == "buch" and b[1] not in I.FREIGEGEBEN else MUTED})], [(u, {"size": 11, "color": MUTED})]], size=20, after=0)
        s.shapes.add_picture(qr_png(u), Inches(x + w - 0.84), Inches(y + 0.04), Inches(0.8), Inches(0.8))
        return 0.95
    if k == "buch":
        h = blk_h(b, w)
        c = box(s, x, y, w, h - 0.05, fill=CARD, line=OLIVE, lw=1.5)
        ic = box(s, x + 0.1, y + (h - 0.05) / 2 - 0.27, 0.62, 0.54, fill=OLIVE)
        text(ic, [[("BUCH", {"size": 12, "bold": True, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        tb(s, x + 0.82, y, w - 0.9, h - 0.05, [b[1]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
        return h
    if k == "versuch":
        h = blk_h(b, w)
        c = box(s, x, y, w, h - 0.05, fill=RGBColor.from_string("FBF1EE"), line=RED, lw=2)
        text(c, [[(b[1] + "  ", {"bold": True, "color": RED}), (b[2], {"bold": True})], b[3], [("Protokoll: Versuchsseite (QR auf »Dein Lernpfad«) · Papierblatt als Ersatz", {"size": 14, "color": MUTED})]], size=20, after=2)
        return h
    if k == "reveal":
        yy = y
        for q, a in b[1]:
            hq = th(q, 20, w, True, pad=0.05); tb(s, x, yy, w, hq, [q], size=20, bold=True); yy += hq
            ha = max(0.75, th(a, 20, w - 0.2, pad=0.5))
            answer(s, x, yy, w, ha, "Lösung", a, L, hint="Erst selbst überlegen – Lösung auf der Theorie-Seite: »Lösung anzeigen«"); yy += ha + 0.1
        return yy - y
    raise ValueError(k)


def zerlegen(blocks):
    out = []
    for b in blocks:
        if False: pass
        elif b[0] == "tiles":
            n = len(b[1]); per = 3 if n in (3, 6) else 2
            out += [("tiles", b[1][i:i + per]) if per == len(b[1][i:i + per]) or n <= per else ("tiles", b[1][i:i + per]) for i in range(0, n, per)]
        elif b[0] == "reveal" and len(b[1]) > 1: out += [("reveal", [x]) for x in b[1]]
        elif b[0] == "table" and len(b[1][0]) > 3 and len(b[1]) > 5:  # breite, lange Tabelle: in zwei Hälften mit Kopfzeile
            body = b[1][1:]; m = (len(body) + 1) // 2
            out += [("table", [b[1][0]] + body[:m]), ("table", [b[1][0]] + body[m:])]
        else: out.append(b)
    return out


def theorie_folien(item, L):
    """Theorie-Folie(n) aus einem Eintrag von theorie_daten (Bild links, Text rechts; zu viel Inhalt → Fortsetzungsfolie)."""
    blocks = [b for b in zerlegen(item["blocks"]) if not (b[0] == "video" and not I.zeigen(b[1], L))]
    img = next((b for b in blocks if b[0] == "img"), None)
    rest = [b for b in blocks if b[0] not in ("img", "karten")]
    karten = next((b for b in blocks if b[0] == "karten"), None)
    khint = None
    if karten and all(b[0] == "hint" for b in rest): khint, rest = rest[0][1], []
    def packen(blocks, w, avail, erste=None):
        """Blöcke auf Seiten verteilen; Listen werden zwischen den Punkten geteilt statt eine fast leere Folie zu erzeugen."""
        pages_, cur_, used_ = [], [], 0.0
        av = erste if erste is not None else avail
        q = list(blocks)
        while q:
            b = q.pop(0); h = blk_h(b, w if not (erste is not None and not pages_) else w_erste) + 0.1
            ww = w_erste if (erste is not None and not pages_) else w
            h = blk_h(b, ww) + 0.1
            if used_ + h <= av + 0.12:
                cur_.append(b); used_ += h; continue
            if b[0] in ("ul", "ol") and len(b[1]) > 1:
                for k_ in range(len(b[1]) - 1, 0, -1):
                    if used_ + blk_h((b[0], b[1][:k_]), ww) + 0.1 <= av + 0.12:
                        cur_.append((b[0], b[1][:k_])); q.insert(0, (b[0], b[1][k_:])); break
                else:
                    if not cur_: cur_.append(b); q_rest = None
                    else: q.insert(0, b)
                pages_.append(cur_); cur_, used_, av = [], 0.0, avail; continue
            if cur_:
                pages_.append(cur_); cur_, used_, av = [], 0.0, avail; q.insert(0, b); continue
            cur_.append(b); used_ += h
        if cur_: pages_.append(cur_)
        # fast leere Folien mit der folgenden zusammenlegen
        hh = lambda pg: sum(blk_h(b_, w) + 0.1 for b_ in pg)
        k_ = 1 if erste is not None else 0
        while k_ < len(pages_) - 1:
            if hh(pages_[k_]) + hh(pages_[k_ + 1]) <= avail + 0.12:
                pages_[k_] = pages_[k_] + pages_.pop(k_ + 1)
            else: k_ += 1
        return pages_

    pages = []
    w_erste = 4.35
    if img:
        if rest and not (rest[0][0] == "table" and len(rest[0][1][0]) > 2) and blk_h(rest[0], 4.35) + 0.1 <= Y1 - Y0 + 0.05:
            pk = packen(rest, X1 - X0, Y1 - Y0, erste=Y1 - Y0)
            pages.append(("bild", img, pk[0]))
            pages += [("text", None, p_) for p_ in pk[1:]]
        else:
            pages.append(("bild", img, []))
            pages += [("text", None, p_) for p_ in packen(rest, X1 - X0, Y1 - Y0)]
    elif rest:
        pages += [("text", None, p_) for p_ in packen(rest, X1 - X0, Y1 - Y0)]
    out = []
    for n, (art, im, bl) in enumerate(pages):
        title = item["title"] + (f" ({n + 1})" if len(pages) > 1 and n else "")
        s = d.slide("Nur Titel", title)
        if art == "bild":
            path = UNIT + im[1].replace(".gif", ".gif")
            if bl:
                pic, hh = picture(s, path, X0, Y0, 4.4, 3.15)
                tb(s, X0, Y0 + hh + 0.05, 4.4, min(0.7, 4.98 - Y0 - hh), [im[3]], size=14, color=MUTED, align=PP_ALIGN.CENTER, after=0)
                y = Y0
                for b in bl: y += draw_blk(s, b, 5.15, y, 4.35, L) + 0.1
            else:
                cap_h = th(im[3], 20, 9.0, pad=0.05)
                pic, hh = picture(s, path, X0 + 0.6, Y0, 7.8, min(3.5, Y1 - Y0 - cap_h - 0.05))
                tb(s, X0, Y0 + hh + 0.05, 9.0, cap_h, [im[3]], size=20, color=MUTED, align=PP_ALIGN.CENTER, after=0)
        elif len(bl) == 1 and bl[0][0] in ("merke", "hint") and n > 0:
            # einzelner Merksatz bzw. Hinweis als Folgefolie: groß als Kernaussage statt verloren oben am Rand
            c = box(s, 1.0, 1.9, 8.0, 1.9, fill=OLIVE if bl[0][0] == "merke" else CARD, line=None if bl[0][0] == "merke" else TEAL, lw=2)
            text(c, [bl[0][1]], size=24, color=WHITE if bl[0][0] == "merke" else INK, bold=bl[0][0] == "merke", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        else:
            y = Y0
            for b in bl: y += draw_blk(s, b, X0, y, X1 - X0, L) + 0.1
        if item.get("pill"): pill(s, item["pill"], kind="versuch" if item["pill"].startswith("Versuch im Buch") else "theory")
        out.append(s)
    if karten:
        for ki, part in enumerate([karten[1][:2]] + [karten[1][j:j + 3] for j in range(2, len(karten[1]), 3)]):
            s = d.slide("Nur Titel", item["title"] + (f" ({ki + 1})" if ki else ""))
            rows = [["Frage", "Lösung (für den Fragenden)"]] + [[q, a if L else ""] for q, a in part]
            y0 = Y0
            if khint and ki == 0:
                hh = th(khint, 20, 9); tb(s, X0, Y0, 9, hh, [khint], size=20, color=MUTED); y0 += hh
            cells, _ = auto_grid(s, X0, y0, [4.6, 4.4], rows, size=16)
            if L:
                for r in cells[1:]: s._anim.append([r[1]])
            pill(s, item["pill"])
    return out


# ================================================================ Aufgaben
def fields_rows(t):
    """Für Auswahl-Aufgaben: [(Feld-Label, richtige Antwort)]"""
    rows = []
    for f in t.get("fields", []):
        c = (t.get("chk") or {}).get(f[0], {})
        rows.append((f[1], c.get("sel")))
    return rows


def sol_lines(t):
    out = []
    for head, items in t["sol"]:
        for x in items: out.append(x if head in ("Lösung", "Beispiel", "Bewertung") else head + " " + x)
    return out


def teile_der_frage(q):
    """Frage in Einleitung und Teilfragen a), b) … zerlegen (für die Folien kurz und gegliedert)."""
    m = re.search(r"(?:(?<=\s)|^)a\)\s", q)
    if not m: return q.strip(), {}
    stem, rest = q[:m.start()].strip(), q[m.start():]
    parts = {}
    for mm in re.finditer(r"(?:(?<=\s)|^)([a-h])\)\s(.*?)(?=(?:\s(?:Zusatz:\s)?[a-h]\)\s)|$)", rest, re.S):
        parts[mm.group(1)] = re.sub(r"\s*(Zusatz|Zusatz \(freiwillig\)):?\s*$", "", mm.group(2).strip())
    zus = re.search(r"Zusatz:\s*([a-h])\)", rest)
    if zus and zus.group(1) in parts: parts[zus.group(1)] = "(freiwillig) " + parts[zus.group(1)]
    return stem, parts


def loesung_je_feld(t):
    """Lösung je Feld: Auswahl → richtige Option, sonst die Lösungszeile mit demselben Buchstaben. Rest = übrige Lösungszeilen."""
    sol = sol_lines(t); used = set(); out = {}
    chk = t.get("chk") or {}
    if t.get("fsol"):  # K1.4–1.5: kurze Lösung je Feld für die Folie (Auswahl bleibt die richtige Option)
        out = {f[0]: ((chk[f[0]]["sel"], "sel") if "sel" in chk.get(f[0], {}) else (t["fsol"].get(f[0]), "num" if "num" in chk.get(f[0], {}) else "text")) for f in t.get("fields", [])}
        return out, list(t.get("fsol_rest", []))
    for f in t.get("fields", []):
        k, lab = f[0], f[1]
        c = chk.get(k, {})
        if k not in chk and "tipp" in k:
            out[k] = ("Vorhersage – zählt nicht", "tipp"); continue
        mm = re.match(r"\s*\(?([a-h])\)", lab)
        line = None
        if mm:
            for n, x in enumerate(sol):
                if n not in used and re.match(r"\s*" + mm.group(1) + r"\)", x):
                    line = re.sub(r"^\s*[a-h]\)(?:\s*[–-]\s*[a-h]\))?\s*", "", x); used.add(n); break
        if "sel" in c and c["sel"] is not None: out[k] = (c["sel"], "sel")
        elif line: out[k] = (line, "text")
        elif "num" in c: out[k] = (str(c["num"]).replace(".", ",") + (" " + c["unit"] if c.get("unit") else ""), "num")
        else: out[k] = (None, "text")
    rest = [x for n, x in enumerate(sol) if n not in used]
    return out, rest


def quelle_badge(s, src, x, y):
    kind = kind_of(src)
    col = {"theory": RGBColor.from_string("6E7F4E"), "book": TEAL, "versuch": RED}[kind]
    a, _, b = src.rpartition(" S. ")
    c = box(s, x, y, 1.25, 0.72, fill=col)
    text(c, [[(a or src, {"size": 12, "bold": True, "color": WHITE})], [("S. " + b if a else "", {"size": 16, "bold": True, "color": WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)


def aufgabe(t, L):
    """Aufgaben-Folie im Stil der 1AB-Messen-PPP: Layout »Aufgabe« (hellgrün), Auftragskarte mit Nummer und Quelle,
    darunter jede Teilfrage als eigene Zeile mit Antwortkasten (Lehrer: Lösung per Klick, Kadetten: leer)."""
    tit = f"Aufgabe {t['nr']}: " + re.sub(r"\s*\((Demo|Versuch|Zusatz)[^)]*\)\s*$", "", t["title"])
    stem, parts = teile_der_frage(t["q"])
    sols, rest = loesung_je_feld(t)
    fields = [f for f in t.get("fields", [])]
    chk = t.get("chk") or {}
    seite = [0]

    def neue_folie():
        seite[0] += 1
        sl = d.slide("Aufgabe", tit + (f" ({seite[0]})" if seite[0] > 1 else ""))
        for ph in sl.placeholders:
            if ph.placeholder_format.type == 1:
                n_ = len(ph.text_frame.text)
                if n_ > 34:
                    for r_ in ph.text_frame.paragraphs[0].runs: r_.font.size = Pt(24 if n_ <= 42 else 21)
        lab = ("ZUSATZ" if t["tag"] == "Zusatz" or not t["h"] else t["tag"].upper()) + f" · ca. {t['zeit']} min"
        mz = re.search(r"\(((?:Demo|Versuch) [^)]*)\)", t["title"])
        if mz: lab += " · " + mz.group(1)
        if L and t.get("besprechen"): lab += " · BESPRECHEN"
        lab = lab.replace("ca. ", "")
        pill(sl, lab, x=2.75 if len(lab) > 30 else 3.35, w=4.6 if len(lab) > 30 else 3.1, kind="theory" if t["h"] and t["tag"] != "Zusatz" else "book")
        return sl

    s = neue_folie()
    # gemeinsame Auswahl (alle Felder mit denselben Optionen) in die Karte
    opts = [tuple(f[2]) for f in fields if len(f) > 2 and isinstance(f[2], (list, tuple))]
    gemeinsam = opts[0] if opts and len({frozenset(o) for o in opts}) == 1 and len(opts) == len(fields) else None
    zuletzt = [None]
    kt = stem if (stem and (parts or fields)) else t["q"]
    if t["type"] == "sketch": kt += " (mit der Hand, Foto ins Arbeitsblatt)" if "Foto" not in kt else ""
    lang = th(kt, 20, 6.45) > 2.0
    size = 20 if (lang or th(kt, 20, 6.45) <= 1.45) else 18
    zus = []
    if gemeinsam: zus = [[("Wähle: ", {"bold": True, "color": TEAL, "size": 20}), (" · ".join(gemeinsam), {"size": 20})]]
    ch = max(0.9, th(kt, size, 6.45, pad=0.16) + (th("Wähle: " + " · ".join(gemeinsam), 20, 6.45, pad=0.02) if gemeinsam else 0))
    if lang:
        size = 20
        while size > 16 and th(kt, size, 6.45, pad=0.16) > Y1 - Y0 - 0.05: size -= 1
        ch = Y1 - Y0 + 0.02
    has_img = bool(t.get("img"))
    card = box(s, X0, Y0 - 0.02, 9.0, ch, fill=WHITE, line=RGBColor.from_string("C9D3B6"), lw=1)
    ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(X0 + 0.14), Inches(Y0 - 0.02 + ch / 2 - 0.3), Inches(0.6), Inches(0.6))
    ov.fill.solid(); ov.fill.fore_color.rgb = GREEN; ov.line.fill.background(); ov.shadow.inherit = False
    ov.text_frame.margin_left = ov.text_frame.margin_right = ov.text_frame.margin_top = ov.text_frame.margin_bottom = 0; ov.text_frame.word_wrap = False
    text(ov, [str(t["nr"])], size=20 if t["nr"] < 10 else 17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    tb(s, X0 + 0.88, Y0 + 0.03, 6.45, ch - 0.1, [kt] + zus, size=size, anchor=MSO_ANCHOR.MIDDLE, after=2)
    quelle_badge(s, t["src"], X0 + 9.0 - 1.37, Y0 - 0.02 + ch / 2 - 0.36)
    y = Y0 + ch + 0.14
    xr, wr = X0, 9.0
    if lang: y = Y1 + 1  # langer Auftrag: Karte füllt die Folie, Antworten auf der nächsten
    if t.get("tab"):  # K1.4–1.5: Modelldaten als eigene Tabellenfolie (16 pt), Antworten danach
        s = neue_folie()
        cells_, hh_ = auto_grid(s, X0, Y0 - 0.02, [4.9, 2.05, 2.05], t["tab"], size=16, minh=0.34)
        tb(s, X0, Y0 + hh_ + 0.04, 9, Y1 - (Y0 + hh_ + 0.04), [t.get("tabnote_folie") or t.get("tabnote", "")], size=20, bold=True, color=RUST, after=0)
        y = Y1 + 1
    if has_img:
        pic, hh = picture(s, UNIT + t["img"], X0, y, 4.1, Y1 - y)
        xr, wr = X0 + 4.3, 4.7

    def zeile(label, ans, art, opt_txt):
        nonlocal s, y, xr, wr
        stacked = has_img or art == "text"
        lw_ = wr if stacked else wr * 0.58
        aw = wr if stacked else wr - lw_ - 0.12
        lh = th(label, 20, lw_ - 0.1, pad=0.06) + (th(opt_txt, 20, lw_ - 0.1, pad=0.02) if opt_txt else 0)
        atxt = ans if (L and ans) else ""
        ah = max(0.42, th(atxt or "x", 20, aw - 0.2, bold=True, pad=0.14)) if art != "text" else max(0.5, th(atxt or "x", 20, aw - 0.2, pad=0.16))
        rh = (lh + ah + 0.04) if stacked else max(lh, ah)
        if y + rh > Y1 + 0.02:
            s = neue_folie(); y = Y0; xr, wr = X0, 9.0
            if has_img: pass
        parts_ = [[(label, {"size": 20, "bold": True, "color": INK})]]
        if opt_txt: parts_.append([(opt_txt, {"size": 20, "color": MUTED})])
        tb(s, xr, y, lw_, lh, parts_, size=20, after=1)
        ax, ay = (xr, y + lh + 0.02) if stacked else (xr + lw_ + 0.12, y)
        frame = box(s, ax, ay, aw, ah, fill=WHITE, line=TEAL if not L else ORANGE, lw=1.5)
        frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH  # gestrichelt wie in der 1AB-PPP
        if L and ans:
            col = MUTED if art == "tipp" else RED
            a_ = tb(s, ax + 0.05, ay + 0.02, aw - 0.1, ah - 0.04, [ans], size=20, bold=art != "text", color=col, anchor=MSO_ANCHOR.MIDDLE, after=0)
            if art != "tipp": s._anim.append([a_])
        y += rh + 0.1

    for f in fields:
        k = f[0]
        mm = re.match(r"\s*\(?([a-h])\)", f[1])
        lab = f[1]
        if mm and mm.group(1) in parts and len(parts[mm.group(1)]) <= 170 and sum(1 for g in fields if re.match(r"\s*\(?" + mm.group(1) + r"\)", g[1])) == 1:
            lab = mm.group(1) + ") " + parts[mm.group(1)]
        ans, art = sols.get(k, (None, "text"))
        if art == "text" and "kw" in chk.get(k, {}) and not chk[k].get("man") and ans and len(ans) < 60: art = "num"
        o_txt = None
        if len(f) > 2 and isinstance(f[2], (list, tuple)) and not gemeinsam and frozenset(f[2]) != zuletzt[0]:
            o_txt = "Wähle: " + " · ".join(f[2]); zuletzt[0] = frozenset(f[2])
        zeile(lab, ans, art, o_txt)
    if not fields:
        zeile("Eure Antwort", " ".join(rest) if rest else "", "text", None); rest = []
    if L and rest:
        txt = rest
        h = sum(th(x, 20, 8.7, pad=0.04) for x in txt) + 0.45
        if y + h > Y1 + 0.02:
            if all(len(f) > 2 for f in fields) and h > 1.6: return
            s = neue_folie(); y = Y0
        b = box(s, X0, y, 9.0, min(h, Y1 - y + 0.02), fill=WHITE, line=ORANGE, lw=2)
        text(b, [[("Lösung" + (" · BESPRECHEN" if t.get("besprechen") else ""), {"bold": True, "color": ORANGE, "size": 20})]] + txt, size=20, after=2)
        s._anim.append([b])


def ueberblick(nrs, titel):
    s = d.slide("Nur Titel", "Aufgaben im Überblick · " + titel)
    rows = [["Nr.", "Aufgabe", "Was ist zu tun?", "Quelle", "Zeit"]]
    was = {"multi": "auswählen, kurz schreiben", "sketch": "mit der Hand zeichnen/schreiben, Foto", "text": "schreiben"}
    for t in I.T:
        if t["nr"] in nrs:
            rows.append([str(t["nr"]), t["title"], was[t["type"]], t["src"], f"{t['zeit']} min"])
    cells, h = auto_grid(s, X0, Y0, [0.55, 3.45, 2.7, 1.6, 0.7], rows, size=16, minh=0.34)
    tb(s, X0, Y0 + h + 0.05, 9, 0.35, [f"Zusammen ca. {sum(t['zeit'] for t in I.T if t['nr'] in nrs)} min · Arbeitsblatt: QR auf »Dein Lernpfad«"], size=16, color=MUTED)


def weisst_du_noch(st):
    s = d.slide("Nur Titel", "Weißt du noch?")
    y = Y0
    for q, a in I.WEISST_DU_NOCH[st]:
        hq = th(q, 20, 9, True, pad=0.06); tb(s, X0, y, 9, hq, [q], size=20, bold=True); y += hq
        ha = max(0.62, th(a, 20, 8.8, pad=0.42))
        answer(s, X0, y, 9, ha, "", a, d.lehrer, hint="Mini-Whiteboard: aufschreiben und hochhalten"); y += ha + 0.12
    tb(s, X0, y, 9, 0.35, ["ohne Handy: Antworten auf das Mini-Whiteboard schreiben und auf Zeichen hochhalten"], size=16, color=MUTED)


def stopp(st, L):
    s = d.slide("Nur Titel", "Stopp – kannst du das?")
    draw_blk(s, ("reveal", I.STOPPFRAGEN[st]), X0, Y0, 9, L)


def jetzt(txt):
    s = d.slide("Nur Titel", txt)
    c = box(s, 1.5, 2.0, 7.0, 1.2, fill=GREEN); text(c, [txt.replace("Jetzt: ", "Jetzt: ") + " →"], size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    s.shapes.add_picture(qr_png(URL["aufgaben"]), Inches(4.35), Inches(3.35), Inches(1.3), Inches(1.3))
    tb(s, 1.5, 4.68, 7, 0.3, [URL["aufgaben"]], size=10, color=MUTED, align=PP_ALIGN.CENTER)


def lernpfad():
    s = d.slide("Nur Titel", "Dein Lernpfad")
    for i, (k, step, lab, col) in enumerate([("theorie", "1", "Theorie", OLIVE), ("aufgaben", "2", "Arbeitsblatt", OLIVE), ("quiz", "3", "Quiz", OLIVE)]):
        ppt_qr(s, URL[k], 1.1 + i * 2.85, 1.45, 1.6, step, lab, col)
    tb(s, X0, 4.62, 9, 0.4, ["Theorieteil lesen → Aufgaben im Arbeitsblatt → Quiz (freiwillig, zum Wiederholen)."], size=16, color=MUTED, after=0)


def fahrplan_folien():
    """3AK und 3BK haben denselben Ablauf – eine Tabelle, eine Zeile je Stunde (Inhalt aus einheiten.json)."""
    import json
    u = [x for x in json.load(open(os.path.join(ROOT, "werkzeug/einheiten.json"))) if x["ordner"] == I.ORDNER][0]
    z = u["fahrplan"]["zeilen"]
    for k, teil in enumerate((z[:2], z[2:])):
        s = d.slide("Nur Titel", "Dein Fahrplan" + (f" ({k + 1}/2)"))
        tb(s, X0, 1.0, 9, 0.4, [[("4 Stunden (3AK und 3BK gleich): ", {"bold": True, "color": OLIVE}), ("Theorieteil lesen → Aufgaben 1–9 → am 16.11. PDF in LMS.", {})]], size=16, after=0)
        rows = [["Wann", "Das machen wir", "Aufgaben", "Ohne Handy", "Abgabe"]]
        for r in teil:
            wann = r["wann"] + ("\n" + r["datum"] if r.get("datum") else "")
            rows.append([wann, r["lesen"], r["aufgaben"], r.get("wie", "–").replace("ohne Handy: ", "").replace("zuerst Test; ", "Test; "), r["abgabe"]])
        auto_grid(s, X0, 1.42, [1.2, 2.9, 1.15, 2.3, 1.45], rows, size=16, minh=0.34)


def abgabe_folie():
    s = d.slide("Nur Titel", "Fertig? So gibst du ab")
    steps = [("1", "Arbeitsblatt öffnen (QR) → Name und Klasse eintragen"), ("2", "»PDF erstellen« tippen → PDF in LMS hochladen"),
             ("3", "Kein Download? »Teilen / Speichern« oder »PDF öffnen« tippen. Geht LMS nicht: PDF per E-Mail an Stephan.")]
    y = 1.0
    for n, t in steps:
        c = box(s, X0, y + 0.04, 0.5, 0.5, fill=OLIVE); text(c, [n], size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        h = th(t, 20, 5.8, pad=0.06); tb(s, X0 + 0.65, y, 5.8, h, [t], size=20); y += max(0.6, h + 0.04)
    s.shapes.add_picture(qr_png(URL["aufgaben"]), Inches(7.1), Inches(1.15), Inches(1.7), Inches(1.7))
    tb(s, 6.6, 2.9, 2.7, 0.3, ["Arbeitsblatt"], size=16, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)
    tb(s, 6.4, 3.25, 3.1, 0.5, [URL["aufgaben"]], size=9, color=MUTED, align=PP_ALIGN.CENTER, after=0)
    tb(s, X0, max(4.5, min(y + 0.02, 4.62)), 9, 0.4, ["Abgabe: eine PDF mit Aufgaben 1–9 am Mo 16.11. (3AK und 3BK)"], size=20, bold=True, color=RUST, after=0)


def lehrer_infos():
    s = d.slide("Nur Titel", "Lehrer-Info · Lehrplanbezug", hidden=True)
    tag(s, "AUSGEBLENDET", RUST, x=7.6, w=1.9)
    parts = [[(lab, {"bold": True, "color": RUST if lab.startswith("Abw") else OLIVE}), (x, {})] for lab, x in F.LEHRPLAN]
    size = 15
    while sum(th(a + b, size, 9, pad=0.04) for (a, _), (b, _) in parts) > Y1 - Y0 - 0.1 and size > 11: size -= 1
    tb(s, X0, Y0 + 0.1, 9, Y1 - Y0 - 0.1, parts, size=size, after=4)
    s = d.slide("Nur Titel", "Lehrer-Info · Stundenvorschlag", hidden=True)
    tb(s, X0, 0.98, 9, 0.3, ["AUSGEBLENDET · Kern in 4 Stunden (Jahresplan E4: 3 Std.) · Rest optional · 3AK und 3BK gleich"], size=14, bold=True, color=RUST, after=0)
    rows = [r[:1] + r[1:] for r in F.STUNDENVORSCHLAG[I.KLASSEN[0]]]
    rows += [["Zusatz", "Für Schnelle (Theorieteil)", "–", "10–13 (OPTIONAL)", "nur wer fertig ist", "Partnercheck"],
             ["Besprechung", "A1, A2, A4, A5, A6, A8, A9", "–", "BESPRECHEN", "in jeder Stunde eingeplant", "Tafel, Skala"]]
    cells, h = auto_grid(s, X0, 1.32, [1.15, 1.9, 0.9, 1.0, 2.4, 1.65], [["Std.", "Theorie", "Versuch", "Aufgaben", "Zeit (min)", "Methode"]] + rows, size=10, minh=0.24)
    tb(s, X0, 1.36 + h, 9, 0.3, ["KERN = Aufgaben 1–9 · OPTIONAL = Zusatz 10–13 · keine Hausübung. Minutenplan: Regiefolien in der TH-Unterricht."], size=12, color=MUTED, after=0)
    s = d.slide("Nur Titel", "Lehrer-Info · Die 5 größten Stolpersteine", hidden=True)
    rows = [["Stolperstein", "Richtig ist …"]] + [[a, plain(b)] for a, b in I.STOLPERSTEINE]
    auto_grid(s, X0, 1.1, [2.4, 6.6], rows, size=14)


def titelfolie(titel, unter, L):
    s = d.slide("Titelfolie", titel)
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2:
            text(ph, [unter, [(f"Stand {I.STAND} · {VER} · " + ("Lehrer" if L else "Kadetten"), {"size": 12})]], size=18, color=RGBColor.from_string("C9D6B0"))


# ================================================================ TH-Deck
TEILE = {1: ["acker", "formen", "bio", "dok"], 2: ["pflanzenschutz", "kennzeichnung"], 3: ["bezug", "sof", "gene", "genschere", "sorte"], 4: ["seralini", "studie", "gtwyst"]}
AUFG = {1: [1, 2], 2: [3, 4], 3: [5, 6, 7], 4: [8, 9]}
ZUSATZ = [10, 11, 12, 13]


def teil_abschnitt(n):
    pid = f"teil{n}"
    eyebrow, titel = [(p[1], p[2]) for p in TD.PARTS if p[0] == pid][0]
    s = d.slide("Abschnitt", titel)
    kurz = I.TEILBUCH[n] + " · " + I.wann(n)
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2: ph.text_frame.text = kurz


def build_th(L):
    global d
    d = Deck(L, "NAWI · 3AB · K2.1 Formen der Landwirtschaft")
    d.section("Start")
    titelfolie("Formen der Landwirtschaft", "NAWI · 3AB (3AK und 3BK) · K2.1 · Mi 04.11. bis Mo 16.11. · Buch S. 38–48", L)
    lernpfad(); fahrplan_folien(); abgabe_folie()
    if L: lehrer_infos()
    d.section("Einstieg")
    for it in TD.EINSTIEG: theorie_folien(it, L)
    s = d.slide("Nur Titel", "Das kannst du nach dieser Einheit")
    tb(s, X0, Y0, 9, Y1 - Y0, [f"{i + 1}. " + z for i, z in enumerate(I.LERNZIELE)], size=20, after=8)
    parts = {p[0]: p[3] for p in TD.PARTS}
    item = lambda pid, iid: [x for x in parts[pid] if isinstance(x, dict) and x["id"] == iid][0]
    for n in range(1, 5):
        d.section(f"Stunde {n}")
        teil_abschnitt(n)
        if n > 1 and n in I.WEISST_DU_NOCH: weisst_du_noch(n)
        for x in TEILE[n]:
            theorie_folien(item(f"teil{n}", x), L)
        stopp(n, L); ueberblick(AUFG[n], f"Stunde {n}")
        for t in I.T:
            if t["nr"] in AUFG[n]: aufgabe(t, L)
    s = d.slide("Nur Titel", "Quiz: Landbau-Check – zum Wiederholen")
    s.shapes.add_picture(qr_png(URL["quiz"]), Inches(3.6), Inches(1.25), Inches(2.8), Inches(2.8))
    tb(s, X0, 4.15, 9, 0.4, [URL["quiz"]], size=12, color=MUTED, align=PP_ALIGN.CENTER)
    tb(s, X0, 4.5, 9, 0.5, [f"{len(I.QUIZ)} Fragen · freiwilliger Selbsttest – zählt nicht für die Note"], size=18, align=PP_ALIGN.CENTER)

    d.section("Zum Schluss")
    rows = I.FACHBEGRIFFE
    # 3 Folien mit je 8 Zeilen liefen in die Fußzeile – darum gleichmäßig verteilt, höchstens 5 Zeilen je Folie
    body = rows[1:]; nf = -(-len(body) // 5); q, r = divmod(len(body), nf)
    gr = [q + (1 if i < r else 0) for i in range(nf)]
    teile = [body[sum(gr[:i]):sum(gr[:i + 1])] for i in range(nf)]
    for k, part in enumerate(teile):
        s = d.slide("Nur Titel", "Fachbegriffe" + (f" ({k + 1}/{len(teile)})" if k else ""))
        auto_grid(s, X0, Y0, [2.8, 6.2], [rows[0]] + part, size=16)
    s = d.slide("Nur Titel", "Das Wichtigste in 5 Sätzen")
    saetze = [f"{i + 1}. " + z for i, z in enumerate(I.WICHTIGSTE)]
    size = 20
    while size > 16 and sum(th(z, size, 9, pad=0.02) + 4 / 72 for z in saetze) > Y1 - Y0 - 0.1: size -= 1
    tb(s, X0, Y0, 9, Y1 - Y0, saetze, size=size, after=4)
    s = d.slide("Nur Titel", "Kann ich das?")
    rows = [["Ich kann …", "sicher", "fast", "noch nicht"]] + [[plain(z), "☐", "☐", "☐"] for z in I.LERNZIELE]
    auto_grid(s, X0, Y0, [6.3, 0.9, 0.9, 0.9], rows, size=16)

    d.section("Für Schnelle und Neugierige")
    for it in TD.SCHNELLE:
        if isinstance(it, dict):
            sl = theorie_folien(it, L)
            for x in sl: tag(x, "OPTIONAL", MUTED, x=6.2, y=5.2, w=1.5)
    ueberblick(ZUSATZ, "Für Schnelle (OPTIONAL)")
    for t in I.T:
        if t["nr"] in ZUSATZ: aufgabe(t, L)

    d.section("Abgabe")
    abgabe_folie()
    name = f"{I.PRAEFIX}-TH-{'Lehrer' if L else 'Kadetten'}-{VER}.pptx"
    d.save(os.path.join(OUT, name)); print("ok", name, len(d.p.slides._sldIdLst), "Folien")


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for L in ((True, False) if "--kadetten" in sys.argv else (True,)): build_th(L)
