"""Unterrichts-PPP K1.4–1.5 Ökodesign und Ökobilanz (TH-Unterricht) – die kurze, geführte PPP für den Frontalunterricht im Saal.
Stephans Auftrag (29.10.2026): gekürzte TH-Lehrer, in der GENAU steht, was die Kadetten in welcher Reihenfolge tun; Theorie nur dort
wiederholt, wo es Schwierigkeiten oder falsche Vorstellungen gibt; die wichtigsten Aufgaben werden gemeinsam durchgegangen.
Aufbau nach »bhak-ppp-gestaltung« Abschnitt 10–14 (Muster Wellen v6 bzw. 1AB K1.4): ausgeblendete Lehrer-Info »So führt die PPP …«,
ausgeblendete Regie je Klasse (Minute · Was läuft jetzt? · Folie · Wozu?), »Heute« → Einstieg → »Lesen: …« → »KURZ KLÄREN« → Arbeitsauftrag → Besprechen;
unten links ein Phasenhinweis (JETZT … · Minuten / Danach: …). Beide Klassen haben denselben Ablauf (5BK Di 03.11., 5AK Do 05.11.).
Alle Minuten kommen aus plan.py (eine Quelle mit Fahrplan, Stundenkarte und Lösungsmappe).
Aufruf im Repository-Ordner: python3 werkzeug/quellen/5ab-k1-4-5-oeko/unterricht_ppp.py <BHAK_Folienmaster_Vorlage.pptx> <Zielordner>"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import bauen_ppp as B
import inhalt as I, plan as P, theorie_daten as TD, fachwissen as F
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

ZIEL = sys.argv[2]
OLIVE, INK, TEAL, RUST, MUTED, WHITE, GREEN, ORANGE, CARD, RED = B.OLIVE, B.INK, B.TEAL, B.RUST, B.MUTED, B.WHITE, B.GREEN, B.ORANGE, B.CARD, B.RED
X0, W = 0.5, 9.0
YC = 4.52                         # Inhalt endet hier; darunter der Phasenhinweis (4.6–5.0), dann die Master-Fußzeile
UNIT = B.UNIT
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

FARBE = {"LESEN": TEAL, "KURZ KLÄREN": OLIVE, "AUFGABE": GREEN, "VERSUCH": RED, "BESPRECHEN": ORANGE, "BAUEN": GREEN,
         "START": MUTED, "TEST": MUTED, "ABGABE": RUST, "WEG": MUTED, "EINSTIEG": ORANGE, "WIEDERHOLEN": TEAL, "ZU ZWEIT": GREEN}
TITEL = {it["id"]: it["title"] for p in TD.PARTS for it in p[3] if isinstance(it, dict)}
TITEL.update({it["id"]: it["title"] for it in TD.EINSTIEG})
T = {t["nr"]: t for t in I.T}


# ================================================================ Bausteine
def phasen(k=None):
    return P.PLAN[k or I.KLASSEN[0]][0]["phasen"]


def minuten(idx):
    m = sum(phasen()[i][1] for i in idx)
    assert all(sum(phasen(k)[i][1] for i in idx) == m for k in I.KLASSEN)
    return f"{m} min"


def hinweis(s, art, mins, danach):
    """Phasenhinweis unten links: was JETZT läuft und was danach kommt."""
    col = FARBE.get(art, OLIVE)
    lab = ("JETZT " + art if art not in ("START", "ABGABE", "TEST", "WEG") else art) + (" · " + mins if mins else "")
    w1 = max(2.2, min(4.6, B.FONT_B.getlength(lab) / 100 * 14 / 72 + 0.45))
    c = B.box(s, X0, 4.62, w1, 0.38, fill=col)
    B.text(c, [lab], size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
    if danach:
        B.tb(s, X0 + w1 + 0.1, 4.62, W - w1 - 0.1, 0.38, [[("Danach: ", {"bold": True, "color": col}), (danach, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE, after=0)


def karte(s, x, y, w, h, kopf, zeilen, akzent=GREEN, size=20):
    c = B.box(s, x, y, w, h, fill=CARD)
    band = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(0.08))
    band.fill.solid(); band.fill.fore_color.rgb = akzent; band.line.fill.background(); band.shadow.inherit = False
    c.text_frame.margin_top = Inches(0.14)
    B.text(c, [[(kopf, {"bold": True, "color": OLIVE})]] + list(zeilen), size=size, after=2)
    return c


def merke(s, y, txt, size=20, h=None):
    h = h or B.th(txt, size, W - 0.1, bold=True, pad=0.24)
    c = B.box(s, X0, y, W, h, fill=OLIVE)
    B.text(c, [txt], size=size, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return c, h


def frage_klick(s, y, frage, antwort, w=W, x=X0, size=20):
    """Frage fett, darunter Antwortkasten; die Antwort erscheint per Klick (Lehrer). Gibt neue y zurück."""
    hq = B.th(frage, size, w, True, pad=0.06)
    B.tb(s, x, y, w, hq, [frage], size=size, bold=True, after=0)
    ha = B.th(antwort, size, w - 0.2, pad=0.16)
    fr = B.box(s, x, y + hq, w, ha, fill=RGBColor.from_string("FFFFFF"), line=ORANGE, lw=1.5)
    a = B.tb(s, x + 0.08, y + hq + 0.02, w - 0.16, ha - 0.04, [antwort], size=size, color=RED, anchor=MSO_ANCHOR.MIDDLE, after=0)
    s._anim.append([a])
    return y + hq + ha + 0.1


def qr(s, key, x, y, size, label):
    s.shapes.add_picture(B.qr_png(B.URL[key]), Inches(x), Inches(y), Inches(size), Inches(size))
    B.tb(s, x - 0.3, y + size, size + 0.6, 0.3, [label], size=14, bold=True, color=OLIVE, align=PP_ALIGN.CENTER, after=0)


def schritte(s, x, y, w, items, size=20, farbe=TEAL):
    for i, it in enumerate(items):
        h = B.th(it, size, w - 0.6, pad=0.06)
        c = B.box(s, x, y + 0.04, 0.42, 0.42, fill=farbe)
        B.text(c, [str(i + 1)], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, x + 0.55, y, w - 0.55, h, [it], size=size, after=0); y += max(0.5, h + 0.06)
    return y


PROBLEME = []


def titel_von(s):
    return next((ph.text_frame.text for ph in s.placeholders if ph.placeholder_format.type == 1), "?")


def check(s):
    """Kontrolle: nichts unter die Inhaltsgrenze (außer Phasenhinweis, Fußzeile, Pill)."""
    for sh in s.shapes:
        if sh.is_placeholder or sh.name == "Buchseite": continue
        unten = (sh.top + sh.height) / 914400
        if unten > YC + 0.02 and sh.top / 914400 < 4.6:
            PROBLEME.append(f"»{titel_von(s)}«: {sh.name} {sh.text_frame.text[:40] if sh.has_text_frame else ''!r} bis {unten:.2f} in")


# ================================================================ Folien je Teil
# Jede Folie: (Schlüssel, Phasen-Indizes in plan.py, Art für den Hinweis, Schritt für die »Heute«-Folie oder None, Wozu? für die Regie, Zeichenfunktion)
SPEC = {}
FOLIE = {}            # Schlüssel → Folie
NACH = {}             # Schlüssel → Text »Danach: …«


def lesen_folie(teil, buch, abschnitte, leitfrage, idx):
    titel = f"Lesen: Theorieteil Teil {teil} · " + I.TEILNAME[teil].split(" · ")[1]
    s = d.slide("Nur Titel", titel)
    y = 1.05
    if buch:
        bh = max(0.62, B.th(buch, 20, W - 0.95, pad=0.2))
        c = B.box(s, X0, y, W, bh, fill=CARD, line=OLIVE, lw=1.5)
        ic = B.box(s, X0 + 0.1, y + bh / 2 - 0.27, 0.62, 0.54, fill=OLIVE)
        B.text(ic, [[("BUCH", {"size": 12, "bold": True, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
        B.tb(s, X0 + 0.82, y, W - 0.9, bh, [buch], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0); y += bh + 0.12
    B.tb(s, X0, y, 6.9, 0.36, [[("Im Theorieteil genau diese Abschnitte:", {"bold": True, "color": TEAL})]], size=18, after=0); y += 0.4
    for a in abschnitte:
        c = B.box(s, X0, y, 6.9, 0.44, fill=RGBColor.from_string("E8F1F3"), line=TEAL, lw=1.25)
        hh = max(0.44, B.th("»" + TITEL[a] + "«", 20, 6.4, True, pad=0.16))
        c.height = Inches(hh)
        B.text(c, [[("»" + TITEL[a] + "«", {"bold": True})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0); y += hh + 0.06
    qr(s, "theorie", 7.75, 1.05 + (bh + 0.12 if buch else 0), 1.2, "Theorieteil")
    y += 0.08
    hq = B.th("Leitfrage: " + leitfrage, 20, W - 0.1, pad=0.2)
    c = B.box(s, X0, y, W, hq, fill=RGBColor.from_string("FFF4EA"), line=ORANGE, lw=1.5)
    B.text(c, [[("Leitfrage: ", {"bold": True, "color": ORANGE}), (leitfrage, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


def auftrag_folie(titel, nrs, tun, erwartet, ohne_handy=False, layout="Aufgabe"):
    """Arbeitsphase: welche Aufgaben, was genau tun, wie lange, was am Ende dastehen soll."""
    s = d.slide(layout, titel)
    rows = [["Nr.", "Aufgabe", "Zeit", "So arbeitest du"]]
    for n in nrs:
        rows.append([str(n), T[n]["title"], f"{T[n]['zeit']} min", tun.get(n, "allein am Handy")])
    cells, h = B.auto_grid(s, X0, 1.05, [0.6, 4.0, 0.9, 3.5], rows, size=18, minh=0.42)
    y = 1.05 + h + 0.12
    if ohne_handy:
        c = B.box(s, X0, y, W, 0.45, fill=RGBColor.from_string("F6E3DF"), line=RUST)
        B.text(c, [[("Handy bleibt zuerst weg: ", {"bold": True, "color": RUST}), (ohne_handy, {})]], size=18, anchor=MSO_ANCHOR.MIDDLE, after=0); y += 0.55
    if erwartet:
        he = B.th("Am Ende steht da: " + erwartet, 20, W - 0.2, pad=0.18)
        c = B.box(s, X0, y, W, he, fill=RGBColor.from_string("FFFFFF"), line=GREEN, lw=1.5)
        B.text(c, [[("Am Ende steht da: ", {"bold": True, "color": GREEN}), (erwartet, {})]], size=20, anchor=MSO_ANCHOR.MIDDLE, after=0)
    return s


# ---------------------------------------------------------------- Inhalte (kurz, nur das Wichtigste)
GR = RGBColor.from_string("FFFFFF")


def heute_folie():
    s = d.slide("Nur Titel", "Heute: Ökodesign und Ökobilanz")
    B.tb(s, X0, 0.98, W, 0.32, [[(I.wann() + " · 1 Stunde · " + I.BUCH, {"bold": True, "color": TEAL})]], size=16, after=0)
    c = B.box(s, X0, 1.36, 7.0, B.th("Ziel: " + I.LEITFRAGE, 20, 6.9, pad=0.16), fill=OLIVE)
    B.text(c, [[("Ziel: ", {"bold": True}), (I.LEITFRAGE, {})]], size=20, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, after=0)
    y = 1.36 + c.height / 914400 + 0.12
    items = ["Lesen Teil 1 · 4 min", "Aufgabe 1 zu zweit · 5 min", "Lesen Teil 2 · 4 min",
             "Aufgabe 2 zu zweit · 11 min", "Aufgabe 3 allein · 4 min", "PDF in LMS · 2 min"]
    h1 = 3
    for col, part in enumerate((items[:h1], items[h1:])):
        yy = y
        for i, it in enumerate(part):
            n = i + 1 + col * h1
            c = B.box(s, X0 + col * 3.6, yy + 0.04, 0.4, 0.4, fill=GREEN)
            B.text(c, [str(n)], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
            hh = B.th(it, 18, 3.1, pad=0.04)
            B.tb(s, X0 + col * 3.6 + 0.48, yy, 3.1, hh, [it], size=18, after=0); yy += max(0.5, hh + 0.04)
    qr(s, "theorie", 7.75, 1.36, 1.05, "Theorieteil")
    qr(s, "aufgaben", 7.75, 2.85, 1.05, "Arbeitsblatt")
    return s


def einstieg():
    s = d.slide("Nur Titel", "Welche Flasche würdest du beschaffen?")
    B.picture(s, UNIT + "img/flaschen.png", X0, 1.0, 5.1, 3.2)
    karte(s, 5.8, 1.0, 3.7, 1.25, "Handzeichen", ["A oder B – für Sport und Ausbildung?"], ORANGE)
    c = karte(s, 5.8, 2.4, 3.7, 1.95, "Welche Information fehlt dir?", [""], GREEN)
    a = B.tb(s, 5.9, 3.0, 3.5, 1.3, ["Wie lange hält sie? Was kostet sie über die ganze Nutzung? Wie belastet sie die Umwelt?"], size=18, color=RED, after=0)
    s._anim.append([a])
    B.tb(s, X0, 4.2, 5.1, 0.3, ["Erfundene Flaschen – keine echten Produkte."], size=14, color=MUTED, after=0)
    return s


def klaeren_1():
    s = d.slide("Nur Titel", "Kurz klären: Ökodesign denkt den ganzen Lebenszyklus")
    B.picture(s, UNIT + "img/lebenszyklus.png", X0, 1.0, 4.6, 3.3)
    y = 1.0
    y = frage_klick(s, y, "Grünes Blatt auf die Flasche – ist das Ökodesign?", "Nein. Ökodesign ändert Material, Haltbarkeit, Reparatur, Trennbarkeit. Nur Aufdruck = eher Greenwashing.", w=4.3, x=5.2, size=18)
    y = frage_klick(s, y, "Was gilt seit 2025 für neue Smartphones in der EU?", "Label mit Reparierbarkeit A–E; Ersatzteile bis zu 7 Jahre, Updates mind. 5 Jahre.", w=4.3, x=5.2, size=18)
    return s


def a1_auftrag():
    return auftrag_folie("Jetzt: Aufgabe 1 – zu zweit am Whiteboard", [1], {1: "zu zweit, Mini-Whiteboard, dann Foto ins Arbeitsblatt"},
                         "3 Zeilen »Verbesserung · Abschnitt · Wirkung« und 1 Zielkonflikt", ohne_handy="Tabelle zuerst auf das Mini-Whiteboard, erst danach fotografieren.")


def a1_loesung():
    s = d.slide("Aufgabe", "Whiteboards hoch: Aufgabe 1")
    rows = [["Verbesserung", "Abschnitt", "Wirkung"],
            ["Deckel und Dichtung nachkaufbar", "Nutzung und Pflege", "kein Neukauf bei kleinem Defekt"],
            ["weite Öffnung, keine engen Ecken", "Nutzung und Pflege", "leichter hygienisch zu reinigen"],
            ["trennbare, gekennzeichnete Materialien", "Herstellung, Nutzungsende", "sortenrein verwertbar"]]
    cells, h = B.auto_grid(s, X0, 1.0, [3.6, 2.4, 3.0], rows, size=18, minh=0.36)
    for r in cells[1:]:
        s._anim.append([c for c in r])
    c, hm = merke(s, 1.0 + h + 0.1, "Zielkonflikt: dünnere Wand spart Material – hält aber Stöße schlechter aus.", size=20)
    s._anim.append([c])
    return s


def klaeren_2a():
    s = d.slide("Nur Titel", "Kurz klären: vier Schritte einer Ökobilanz")
    B.picture(s, UNIT + "img/schritte.png", X0, 1.0, 9.0, 3.1)
    a = B.tb(s, X0, 4.1, W, 0.4, ["Lebenszyklus = WAS wir untersuchen · vier Schritte = WIE wir untersuchen"], size=18, bold=True, color=RUST, align=PP_ALIGN.CENTER, after=0)
    s._anim.append([a])
    return s


def klaeren_2b():
    s = d.slide("Nur Titel", "Kurz klären: fair vergleichen")
    karte(s, X0, 1.0, 4.4, 1.55, "Funktionale Einheit", ["dieselbe Leistung: **1.000 Nutzungen** einer Ein-Liter-Flasche – nicht 1 gegen 1 Flasche"], TEAL, size=18)
    karte(s, X0 + 4.6, 1.0, 4.4, 1.55, "Systemgrenze", ["was wir einrechnen (Herstellung, Reinigung, Nutzungsende) – und was nicht"], OLIVE, size=18)
    y = 2.7
    y = frage_klick(s, y, "Was zeigen kg CO₂-Äquivalente (Klimawirkung, Einheit kg)?", "Nur die Klimawirkung – nicht Wasser, Rohstoffe oder Schadstoffe.", size=20)
    return s


def a2_auftrag():
    return auftrag_folie("Jetzt: Aufgabe 2 – zu zweit rechnen", [2], {2: "zu zweit am Handy, Rechenweg im Heft"},
                         "a) Vergleichsleistung + Grenzen · b) 6 Zahlen (Flaschen, kg CO₂-Äq., €) · c) Empfehlung mit 2 Argumenten und 1 Grenze")


def a2_daten():
    s = d.slide("Aufgabe", "Aufgabe 2: Modelldaten und Rechenweg")
    rows = [["", "A", "B"], ["Nutzungen pro Flasche", "200", "1.000"], ["Kaufpreis", "8 €", "30 €"], ["Herstellung + Lieferung (je Flasche)", "3 kg", "6 kg"],
            ["Endbehandlung (je Flasche)", "0,2 kg", "0,5 kg"], ["Reinigung (je Nutzung)", "0,01 kg · 0,02 €", "0,01 kg · 0,02 €"]]
    cells, h = B.auto_grid(s, X0, 1.0, [3.1, 1.85, 1.85], rows, size=16, minh=0.36)
    B.tb(s, X0, 1.02 + h, 6.8, 0.6, ["kg = kg CO₂-Äquivalente · erfundene Modelldaten · Trinkwasser nicht eingerechnet, keine Recycling-Gutschrift"], size=14, color=RUST, bold=True, after=0)
    karte(s, 7.0, 1.0, 2.5, 3.45, "Rechenweg", [[("Flaschen", {"bold": True})], "1.000 : Nutzungen", [("Klima", {"bold": True})], "Flaschen × (Herst. + Ende) + 1.000 × Reinig.",
                                               [("Kosten", {"bold": True})], "Flaschen × Preis + 1.000 × 0,02 €"], GREEN, size=16)
    return s


def a2_bespr():
    s = d.slide("Aufgabe", "Besprechen: Aufgabe 2")
    rows = [["für 1.000 Nutzungen", "Flasche A", "Flasche B"],
            ["Flaschen", "1.000 : 200 = 5", "1.000 : 1.000 = 1"],
            ["Klima (kg CO₂-Äq.)", "5 × 3,2 + 10 = 26", "1 × 6,5 + 10 = 16,5"],
            ["Kosten", "5 × 8 € + 20 € = 60 €", "30 € + 20 € = 50 €"]]
    cells, h = B.auto_grid(s, X0, 1.05, [3.0, 3.0, 3.0], rows, size=20, minh=0.5)
    for r in cells[1:]:
        s._anim.append([r[1], r[2]])
    c, hm = merke(s, 1.05 + h + 0.15, "Empfehlung im Modell: B – 9,5 kg CO₂-Äq. und 10 € weniger. Für ein ganzes Umwelturteil fehlen z. B. Wasser, Rohstoffe, echte Lebensdauer.", size=20)
    s._anim.append([c])
    return s


def a2_nicht_erlaubt():
    s = d.slide("Aufgabe", "Welche Behauptung wäre nicht erlaubt?")
    y = 1.05
    y = frage_klick(s, y, "»B ist in jeder Hinsicht umweltfreundlicher.«", "nicht erlaubt – gerechnet wurde nur die Klimawirkung.")
    y = frage_klick(s, y, "»Metall ist immer besser als Kunststoff.«", "nicht erlaubt – das Material ist im Modell gar nicht festgelegt.")
    y = frage_klick(s, y, "»Für 1.000 Nutzungen hat B im Modell die geringere Klimawirkung.«", "erlaubt – mit Leistung und Grenze formuliert.")
    return s


def a3_auftrag():
    return auftrag_folie("Jetzt: Aufgabe 3 – allein, ohne nachzuschauen", [3], {3: "allein am Handy"}, "2 Lückensätze, 1 Satz zu »unfair«, 1 Umweltfrage, die offen bleibt")


def abgabe():
    s = d.slide("Nur Titel", "PDF abgeben – und bis zum Test wiederholen")
    y = schritte(s, X0, 1.05, 6.6, ["Arbeitsblatt: »PDF erstellen« → in LMS hochladen.", "Kein Download? »Teilen / Speichern« oder »PDF öffnen«. Geht LMS nicht: PDF per E-Mail an Stephan.",
                                    "Test am " + I.TEST + ": Quiz und Partnercheck-Karten (Theorieteil, Für Schnelle) zum Wiederholen."], size=20, farbe=RUST)
    qr(s, "aufgaben", 7.6, 1.0, 1.2, "Arbeitsblatt"); qr(s, "quiz", 7.6, 2.75, 1.2, "Quiz")
    return s


def zusatz_folie():
    s = d.slide("Nur Titel", "OPTIONAL: Für Schnelle", hidden=True)
    rows = [["Nr.", "Zusatzaufgabe", "Zeit"]] + [[str(n), T[n]["title"], f"{T[n]['zeit']} min"] for n in (4, 5, 6, 7, 8, 9)]
    cells, h = B.auto_grid(s, X0, 1.05, [0.7, 6.9, 1.4], rows, size=18, minh=0.4)
    B.tb(s, X0, 1.05 + h + 0.1, W, 0.6, ["Nur, wenn die Pflicht fertig ist. Keine Hausübung. Aufgabe 9 (Plakat, Video, Podcast) ist freiwillig."], size=16, after=0)
    return s


# ---------------------------------------------------------------- Ablauf: (Schlüssel, Phasen-Indizes in plan.STUNDE, Art, Abschnitt, Wozu?, Funktion)
SPEC = [
 ("heute", [0], "START", "Einstieg", "Klassenbuch; Ziel und Ablauf zeigen (QR für Theorieteil und Arbeitsblatt)", heute_folie),
 ("einst", [1], "EINSTIEG", "Einstieg", "Vorwissen holen: per Handzeichen A oder B, dann »Welche Information fehlt?« – noch nicht auflösen", einstieg),
 ("lesen1", [2], "LESEN", "Teil 1", "Begriffe selbst lesen – Stephan erklärt nicht vor",
  lambda: lesen_folie(1, None, ["oekodesign", "kriterien"], "Welche zwei Eigenschaften einer Flasche werden erst während der Nutzung wichtig?", [2])),
 ("klaeren1", [3], "KURZ KLÄREN", "Teil 1", "Lebenszyklus zeigen; Fehlvorstellung »grüner Aufdruck = Ökodesign«; Smartphone als aktuelles Beispiel", klaeren_1),
 ("a1", [4], "ZU ZWEIT", "Teil 1", "Ökodesign anwenden, Zielkonflikt erkennen – ohne Handy", a1_auftrag),
 ("a1l", [4], "BESPRECHEN", "Teil 1", "Whiteboards hochhalten lassen, zwei zeigen, Beispiel per Klick (letzte Minute von Aufgabe 1)", a1_loesung),
 ("lesen2", [5], "LESEN", "Teil 2", "Methode selbst lesen",
  lambda: lesen_folie(2, None, ["oekobilanz", "schritte"], "Welche gemeinsame Leistung vergleichen wir – und was unterscheidet Sachbilanz und Wirkungsabschätzung?", [5])),
 ("klaeren2a", [6], "KURZ KLÄREN", "Teil 2", "Vier Schritte von den Lebenszyklus-Abschnitten unterscheiden", klaeren_2a),
 ("klaeren2b", [6], "KURZ KLÄREN", "Teil 2", "Funktionale Einheit, Systemgrenze, CO₂-Äquivalent = nur Klima", klaeren_2b),
 ("a2", [7], "AUFGABE", "Teil 2", "Kern der Stunde: gleiche Leistung, Rechnen, begrenztes Urteil (K02)", a2_auftrag),
 ("a2d", [7], "AUFGABE", "Teil 2", "Diese Folie bleibt während der Arbeitszeit stehen", a2_daten),
 ("a2b", [8], "BESPRECHEN", "Teil 2", "Rechenweg zeilenweise per Klick", a2_bespr),
 ("a2n", [8], "BESPRECHEN", "Teil 2", "Grenzen der Aussage – nie streichen", a2_nicht_erlaubt),
 ("a3", [9], "AUFGABE", "Abschluss", "Diagnose: Wer kann Ökodesign/Ökobilanz trennen und »unfair« begründen?", a3_auftrag),
 ("abgabe", [10, 11], "ABGABE", "Abschluss", "PDF in LMS; Testhinweis, Quiz zum Wiederholen; Reserve", abgabe),
]
assert sorted({i for x in SPEC for i in x[1]}) == list(range(len(phasen())))
DANACH = {"heute": "Einstieg", "einst": "Lesen Teil 1", "lesen1": "Kurz klären", "klaeren1": "Aufgabe 1", "a1": "Whiteboards hoch", "a1l": "Lesen Teil 2",
          "lesen2": "Kurz klären", "klaeren2a": "fair vergleichen", "klaeren2b": "Aufgabe 2", "a2": "Modelldaten", "a2d": "Besprechen",
          "a2b": "nicht erlaubte Behauptungen", "a2n": "Aufgabe 3", "a3": "PDF abgeben", "abgabe": "Ende – nächstes Thema laut Plan: Umweltmanagement (Buch S. 26)"}


# ================================================================ Deck bauen
def bauen():
    global d
    d = B.Deck(True, "TÖW · 5AB · K1.4–1.5 Ökodesign und Ökobilanz · Unterricht"); B.d = d
    d.section("Start")
    s = d.slide("Titelfolie", "Ökodesign und Ökobilanz")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 2:
            B.text(ph, ["TÖW · 5AB · K1.4–1.5 · Unterrichts-PPP · " + I.wann() + " · Buch S. 19–25",
                        [(f"Stand {I.STAND} · {I.V} · TH-Unterricht (Ablauf) – zum Nachlesen: TH-Lehrer", {"size": 12})]], size=18, color=RGBColor.from_string("C9D6B0"))
    titel = s
    bedien = d.slide("Nur Titel", "Lehrer-Info · So führt die PPP", hidden=True)
    regie = {k: d.slide("Nur Titel", f"Regie {k} · {I.DATEN[k][1]} · Stunde 1", hidden=True) for k in I.KLASSEN}
    ziele = d.slide("Nur Titel", "Lehrer-Info · Lernziele und Abdeckung", hidden=True)
    abschnitt = None
    for j, (key, idx, art, absch, wozu, fn) in enumerate(SPEC):
        if absch != abschnitt:
            d.section(absch); abschnitt = absch
            if absch.startswith("Teil"):
                n = int(absch[-1])
                sec = d.slide("Abschnitt", I.TEILNAME[n])
                for ph in sec.placeholders:
                    if ph.placeholder_format.type == 2: ph.text_frame.text = I.TEILBUCH[n] + " · " + I.wann()
                FOLIE["sec" + str(n)] = sec
        s = fn(); FOLIE[key] = s
        hinweis(s, art, minuten(idx), DANACH[key])
        B.pill(s, I.TEILBUCH[2] if absch in ("Teil 2", "Abschluss") else I.TEILBUCH[1], w=2.2)
        check(s)
    d.section("Optional")
    zusatz = zusatz_folie()

    nr_von = {s.slide_id: i + 1 for i, s in enumerate(d.p.slides)}
    NR = lambda key: nr_von[FOLIE[key].slide_id]
    ablauf = [FOLIE[x[0]] for x in SPEC[:2]] + [FOLIE["sec1"]] + [FOLIE[x[0]] for x in SPEC[2:6]] + [FOLIE["sec2"]] + [FOLIE[x[0]] for x in SPEC[6:]]
    for k in I.KLASSEN:
        s = regie[k]
        rows = [["Minute", "Was läuft jetzt?", "Folie", "Wozu?"]]; t0 = 0
        zu = {}
        for key, idx, art, absch, wozu, fn in SPEC:
            for i in idx: zu.setdefault(i, []).append((key, wozu))
        for i, (name, mi, oh, a) in enumerate(phasen(k)):
            fol = ", ".join(dict.fromkeys(str(NR(x)) for x, _ in zu[i]))
            wz = " / ".join(dict.fromkeys(w for _, w in zu[i])) if a != "reserve" else "Puffer – nichts Neues beginnen"
            rows.append([f"{t0}–{t0 + mi}", name, fol, wz]); t0 += mi
        assert t0 == 50
        B.tb(s, X0, 0.98, W, 0.3, [[(f"{k} · {I.DATEN[k][1]} · Benutzerdefinierte Präsentation »{k} {I.DATEN[k][1]} Stunde 1« · Foliennummern gelten für diese Datei", {"bold": True, "color": RUST})]], size=11, after=0)
        kn = P.KNAPP[(k, 1)]
        hk = B.th("Wenn es knapp wird: " + kn, 10, W, pad=0.04)
        size = 11
        while size > 8 and B.table_h([0.8, 4.0, 0.7, 3.5], rows, size=size, minh=0.2) + hk > 3.95: size -= 1
        cells, h = B.auto_grid(s, X0, 1.28, [0.8, 4.0, 0.7, 3.5], rows, size=size, minh=0.2)
        B.tb(s, X0, 1.3 + h, W, hk, [[("Wenn es knapp wird: ", {"bold": True, "color": RUST}), (kn, {})]], size=10, after=0)
    B.tb(bedien, X0, 0.98, W, 4.1, [
        [("Eine Datei, beide Klassen, gleicher Ablauf. ", {"bold": True, "color": OLIVE}), ("F5 zeigt alle sichtbaren Folien in Unterrichtsreihenfolge. Benutzerdefinierte Präsentationen »5BK Di 03.11. Stunde 1« und »5AK Do 05.11. Stunde 1« (gleich, nur Datum) – Menü Bildschirmpräsentation.", {})],
        [("Ablauf: ", {"bold": True, "color": OLIVE}), ("»Heute« (Klassenbuch) → Einstieg (Handzeichen) → »Lesen: Teil 1« → »Kurz klären« → Aufgabe 1 am Whiteboard → »Lesen: Teil 2« → »Kurz klären« (2 Folien) → Aufgabe 2 (Datenfolie bleibt stehen) → Besprechen (per Klick) → Aufgabe 3 → PDF.", {})],
        [("Unten links ", {"bold": True, "color": OLIVE}), ("steht, was JETZT läuft (Minuten aus plan.py) und was DANACH kommt.", {})],
        [("Lösungen ", {"bold": True, "color": OLIVE}), ("erscheinen erst per Klick – nach der Arbeitszeit. Vollständige Lösungen: TH-Loesung und private Lösungsmappe.", {})],
        [("Zeit: ", {"bold": True, "color": OLIVE}), ("50 min (Annahme) ohne Luft – Klassenbuch 2 min, Reserve 2 min. Zeitstopp und Kürzungsweg auf der Regiefolie.", {})],
        [("Nachschlagen statt zeigen: ", {"bold": True, "color": OLIVE}), ("TH-Lehrer und Theorieteil bleiben vollständig (EU-Recht, Bilanzarten, Wirkungskategorien). Zusatz 4–9: ausgeblendete Folie »OPTIONAL«.", {})],
        [("Regiepläne: ", {"bold": True, "color": OLIVE}), (f"Folien {nr_von[regie[I.KLASSEN[0]].slide_id]}–{nr_von[regie[I.KLASSEN[-1]].slide_id]} (ausgeblendet), Präsentation »Lehrer-Regiepläne«.", {})]],
        size=13, after=5)
    rows = [["Lernziel (Kadetten)", "Wo?"], [B.plain(I.LERNZIELE[0]), "Lesen 1/2, Klären, A3 a"], [B.plain(I.LERNZIELE[1]), "Klären 2, A2 a"],
            [B.plain(I.LERNZIELE[2]), "A2 b–c, Besprechung"], [B.plain(I.LERNZIELE[3]), "Klären 2, A2 c, A3 c"]]
    cells, h = B.auto_grid(ziele, X0, 1.0, [6.4, 2.6], rows, size=13, minh=0.3)
    B.tb(ziele, X0, 1.05 + h, W, 1.5, ["K01 → Teil 1/2, A1, A2 a, A3 · K02 → A2 b–c · N10 (Warenprobe) in dieser Stunde nicht praktisch abgedeckt · Jahresplan E3: 2 Std., hier 1 Std. (Stephans Vorgabe)."], size=13, after=0)

    pres = d.p.part._element
    csl = etree.Element("{%s}custShowLst" % NS_P)
    anker = pres.find("{%s}defaultTextStyle" % NS_P)
    if anker is not None: anker.addprevious(csl)
    else: pres.append(csl)
    rid = {s.slide_id: el.get("{%s}id" % NS_R) for s, el in zip(d.p.slides, d.p.slides._sldIdLst)}
    nid = [0]

    def show(name, folien):
        cs = etree.SubElement(csl, "{%s}custShow" % NS_P, name=name, id=str(nid[0])); nid[0] += 1
        l = etree.SubElement(cs, "{%s}sldLst" % NS_P)
        for sl in folien: etree.SubElement(l, "{%s}sld" % NS_P).set("{%s}id" % NS_R, rid[sl.slide_id])

    for k in I.KLASSEN:
        show(f"{k} {I.DATEN[k][1]} Stunde 1", [titel] + ablauf)
    show("Lehrer-Regiepläne", [bedien] + list(regie.values()) + [ziele])
    show("Optional: Für Schnelle", [zusatz])

    os.makedirs(ZIEL, exist_ok=True)
    name = f"{I.PRAEFIX}-TH-Unterricht-{I.V}.pptx"
    d.save(os.path.join(ZIEL, name))
    sicht = [s for s in d.p.slides if s._element.get("show") != "0"]
    for x in PROBLEME: print("ÜBERLAUF", x)
    print("ok", name, len(d.p.slides._sldIdLst), "Folien,", len(sicht), "sichtbar ·", nid[0], "Präsentationen")


if __name__ == "__main__":
    bauen()
