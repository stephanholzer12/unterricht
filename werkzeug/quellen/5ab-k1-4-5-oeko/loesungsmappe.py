"""Private Lehrer-Seite (Lösungsmappe) K1.4–1.5 Ökodesign und Ökobilanz (5AB) – als Artifact, nur für Stephan. NICHT ins öffentliche Repository legen.
Aus derselben Datenquelle wie alle anderen Ausgaben (inhalt, plan, fachwissen). Keine Versuche in dieser Einheit.
Aufruf: python3 loesungsmappe.py <TH-Unterricht.pptx> <Ordner mit Folienbildern s-NN.jpg> <Ziel.html>
Die Folienbilder sind die sichtbaren Folien der TH-Unterricht (LibreOffice-Render), damit die Seite am Beamer-PC als Ersatz taugt."""
import base64, glob, html, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
import inhalt as I, plan as P, fachwissen as F
from pptx import Presentation

PPTX, FOLDIR, ZIEL = sys.argv[1:4]
UNIT = os.path.join(ROOT, I.ORDNER) + "/"
e = lambda t: html.escape(str(t), quote=False)


def md(t):
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", e(t))


def b64(path, mime):
    return f"data:{mime};base64," + base64.b64encode(open(path, "rb").read()).decode()


def ul(items):
    return "<ul>" + "".join(f"<li>{md(x)}</li>" for x in items) + "</ul>"


def tab(head, rows, cls=""):
    h = "".join(f"<th>{e(x)}</th>" for x in head)
    b = "".join("<tr>" + "".join(f"<td>{md(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f"<div class='scroll'><table class='{cls}'><tr>{h}</tr>{b}</table></div>"


# ---------------------------------------------------------------- Überblick
offen = [
    "**Einheit einblenden** vor der ersten Stunde (5BK Di 03.11.): GitHub → Actions → »Einheit ein-/ausblenden« → 5ab-toew/k1-4-5-oeko → einblenden. Die Direktlinks und QR-Codes funktionieren auch versteckt.",
    "**Zeit:** 50 min ohne Luft (Annahme aus der Übergabe). Zeitstopp: Um Minute 37 beginnt die Besprechung von Aufgabe 2 – auch wenn noch nicht alle fertig sind.",
    "**Mini-Whiteboards** (1 je Paar) für Aufgabe 1 bereitlegen.",
    "**Test am Do 12.11.:** Dauer und Teststoff sind offen. Ob K1.4–1.5 dazugehört, entscheidest du – es gibt keinen Teststoff im Paket.",
    "**Jahresplan:** E3 hat 2 Stunden, du hast 1 Stunde vorgegeben. Die Lehrer-Seite zeigt deshalb »bitte abgleichen« – das ist gewollt.",
    "**LMS-Abgabeordner** für 5AK und 5BK anlegen (PDF am Stundenende).",
]
uber = "<h2>Was du vor Ort noch prüfen musst</h2>" + ul(offen)
uber += "<h2>Ablauf (beide Klassen gleich)</h2><p>" + md(P.FAHRPLAN_TEXT) + "</p>"
uber += tab(["Std.", "Theorie", "Versuch", "Aufgaben", "Zeit", "Methode"], [[" · ".join(f"{k} {I.DATEN[k][1]}" for k in I.KLASSEN)] + F.STUNDENVORSCHLAG[I.KLASSEN[0]][0][1:]])
uber += "<h2>Stundenkarte (Minuten je Phase, Summe 50)</h2>"
for k in I.KLASSEN:
    for st in P.PLAN[k]:
        t0, li = 0, []
        for name, mi, oh, a in st["phasen"]:
            li.append(f"<tr><td class='num'>{t0}–{t0 + mi}</td><td>{e(name)}{' <span class=chip>ohne Handy</span>' if oh and a not in ('org', 'reserve', 'test', 'bespr', 'erkl') else ''}</td></tr>"); t0 += mi
        kn = P.KNAPP.get((k, st["std"]))
        uber += (f"<details class='card'><summary><span class='nr'>{k}</span><span class='ttl'>Stunde {st['std']} · {e(st['datum'])} · {e(I.TEILNAME[st['teile'][0]])}</span></summary>"
                 f"<div class='body'><div class='scroll'><table>{''.join(li)}</table></div>" + (f"<p class='warnline'><b>Wenn es knapp wird:</b> {e(kn)}</p>" if kn else "") + "</div></details>")
uber += "<h2>Die 5 größten Stolpersteine</h2>" + "".join(f"<div class='fehler'><b class='k'>{e(a)}</b>{md(b)}</div>" for a, b in I.STOLPERSTEINE)
uber += "<h2>So bewertet die Seite</h2>" + ul([
    "**Automatisch geprüft:** A2 a (Vergleichsleistung, zwei Vorgänge), A2 b (sechs Zahlen, ±2 %, »26 kg CO₂-Äq.« wird erkannt), A2 c Auswahl, A3 a (Auswahl) und c (Umweltfrage), Zusatz 5 a–c und 6 a–d.",
    "**Offene Texte bewertest du:** A1 (Foto der Tabelle + Zielkonflikt), A2 Ausschluss und die drei Argumente in c, A3 b, Zusatz 4, 5 d, 7, 8, 9. Die Seite zeigt nur eine Schreibhilfe und gibt dafür keine Punkte.",
    "**PDF:** Kasten »Für LMS: x von y Punkten · Notenvorschlag« (88/75/62/50 %); Aufgaben mit angesehener Lösung zählen 0 Punkte. Zusatzaufgaben zählen nur, wenn bearbeitet.",
    "**Keine Versuche** in dieser Einheit (V5.3 bleibt für später).",
])
uber += "<h2>Links der Kadetten-Seiten</h2><ul>" + "".join(f"<li><span class='mono'>{I.URL}{k}.html</span></li>" for k in ("theorie", "aufgaben", "quiz")) + "</ul>"

herr = "<h2>Vor der Stunde</h2><ul class='check'>"
for n, x in enumerate(F.CHECKLISTE, 1):
    herr += f"<li><label><input type='checkbox' id='h{n}' data-k='h{n}'> {md(x)}</label></li>"
herr += "</ul>"
uber = herr + uber

# ---------------------------------------------------------------- Aufgaben
auf = []
for t in I.T:
    man = [k for k, v in (t.get("chk") or {}).items() if v.get("man")]
    tags = [f"<span class='tag{' zus' if t['h'] == 0 else ''}'>{e(t['tag'])} · {t['zeit']} min</span>"]
    if t.get("besprechen"): tags.append("<span class='tag besp'>BESPRECHEN</span>")
    if man: tags.append("<span class='tag bew'>du bewertest</span>")
    felder = "".join(f"<li>{e(f[1])}" + (f" <span class='klein'>({' · '.join(e(o) for o in f[2])})</span>" if len(f) > 2 and isinstance(f[2], (list, tuple)) else "") + "</li>" for f in t.get("fields", []))
    sol = "".join(f"<p>{'' if h in ('Lösung', 'Beispiel', 'Bewertung') else '<b>' + e(h) + '</b> '}{md(x)}</p>" for h, xs in t["sol"] for x in xs)
    bild = f"<div class='bild'><img src='{b64(UNIT + t['img'], 'image/png')}' alt=''></div>" if t.get("img") else ""
    if t.get("tab"): bild += tab(t["tab"][0], t["tab"][1:]) + f"<p class='klein'>{e(t.get('tabnote', ''))}</p>"
    auf.append(f"<details class='card' data-h='{t['h']}' data-s='{html.escape((t["title"] + " " + t["q"]).lower(), quote=True)}'><summary><span class='nr'>{t['nr']}</span><span class='ttl'>{e(t['title'])}</span>{''.join(tags)}</summary>"
               f"<div class='body'><p class='q'>{md(t['q'])}</p>{bild}<p class='klein'>{e(t['src'])}</p><ul class='felder'>{felder}</ul>"
               f"<div class='box sol'><b class='k'>Lösung</b>{sol}</div>"
               + (f"<div class='box bew'><b class='k'>Bewertungshilfe</b>Offene Felder ({', '.join(man)}): Die Seite prüft nur Stichwörter als Schreibhilfe – fachlich entscheidest du. Musterantwort oben.</div>" if man else "")
               + (f"<p class='klein'>Hinweis für Kadetten: {md(t['hint'])}</p>" if t.get("hint") else "") + "</div></details>")
filt = "".join(f"<button type='button' data-f='{h}' aria-pressed='false'>{e(I.GRP[h].split(' · ')[0] + (' · ' + I.GRP[h].split(' · ')[1] if h else ''))}</button>" for h in (1, 2, 0))
aufgaben = (f"<input class='search' type='search' id='suche' placeholder='Aufgaben durchsuchen' aria-label='Aufgaben durchsuchen'><div class='filter' id='filter'>{filt}</div>"
            + "".join(auf) + "<p class='leer' id='leer' hidden>Nichts gefunden.</p>"
            + "<h2>Partnercheck-Karten (Aufgabe 4)</h2>" + tab(["Frage", "Lösung"], I.PARTNERKARTEN))

# ---------------------------------------------------------------- Fachwissen
fw = "<h2>Lehrplanbezug</h2><ul>" + "".join(f"<li><b>{e(a)}</b> {md(b)}</li>" for a, b in F.LEHRPLAN) + "</ul>"
for kap in F.KAPITEL:
    fw += f"<h2>{e(kap['titel'])}</h2>" + "".join(f"<p>{md(x)}</p>" for x in kap["text"])
    if kap.get("bild"): fw += f"<div class='bild'><img src='{b64(UNIT + kap['bild'], 'image/png')}' alt=''></div>"
    if kap.get("nur"): fw += "<div class='box fach'><b class='k'>Nur für dich</b>" + "".join(f"<p>{md(y)}</p>" for y in kap["nur"]) + "</div>"
    if kap.get("stolper"): fw += f"<div class='fehler'><b class='k'>Stolperstein</b>{md(kap['stolper'])}</div>"
fw += "<h2>Rechenbeispiele</h2>" + "".join(f"<p>{md(x)}</p>" for x in F.RECHNEN)
fw += "<h2>Häufige Fehlvorstellungen</h2>" + tab(["Fehlvorstellung", "Richtig ist", "So zeigst du es"], F.FEHLVORSTELLUNGEN)
fw += "<h2>Fragen, die Kadetten stellen</h2>" + "".join(f"<p><b>{e(q)}</b><br>{md(a)}</p>" for q, a in F.FRAGEN)
fw += "<h2>Videos und Podcasts</h2>" + "".join(f"<p>{md(x)}</p>" for x in F.VIDEO_NOTIZ.values())
fw += "<h2>Ausblick</h2><p>" + md(F.AUSBLICK) + "</p><h2>Quellen</h2>" + ul(F.QUELLEN)

# ---------------------------------------------------------------- Folien (TH-Unterricht als Bilder)
prs = Presentation(PPTX)
titel = [next((ph.text_frame.text for ph in s.placeholders if ph.placeholder_format.type == 1), "") for s in prs.slides if s._element.get("show") != "0"]
bilder = sorted(glob.glob(os.path.join(FOLDIR, "s-*.jpg")))
assert len(bilder) == len(titel), (len(bilder), len(titel))
nr_phys = [i + 1 for i, s in enumerate(prs.slides) if s._element.get("show") != "0"]
import json
folien_js = json.dumps([[b64(b, "image/jpeg"), t, n] for b, t, n in zip(bilder, titel, nr_phys)], ensure_ascii=False).replace("</", "<\\/")
optionen = "".join(f"<option value='{i}'>Folie {n} · {e(t)}</option>" for i, (t, n) in enumerate(zip(titel, nr_phys)))

STIL = open(os.path.join(HERE, "loesungsmappe_stil.css"), encoding="utf-8").read()
seite = f"""<title>Lösungsmappe Ökodesign 5AB</title>
<meta name="claude:viewport-fit-sides" content="cover">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;700&family=Source+Serif+4:wght@700&display=swap">
<style>{STIL}</style>
<div class="wrap">
<header class="lese">
  <div class="eyebrow">TÖW · 5AB (5BK Di 03.11. · 5AK Do 05.11.) · K1.4–1.5 Ökodesign und Ökobilanz · nur für dich</div>
  <h1>Lösungen und Fachwissen: Ökodesign und Ökobilanz</h1>
  <p class="sub">Aufgaben 1–9 (4–9 Zusatz), die Unterrichts-PPP als Folien und das Fachwissen. Keine Versuche. Buch: Trauner TÖW, S. 19–25. Stand {e(I.STAND)} · {e(I.V)}</p>
  <p class="warn">Privat in deinem Konto – nicht an Kadetten weitergeben. Die Kadetten sehen Musterlösungen erst nach der PDF bzw. nach drei Fehlversuchen.</p>
</header>
<nav class="tabs" role="tablist" aria-label="Bereiche">
  <button role="tab" id="t-ueberblick" aria-selected="true">Überblick</button>
  <button role="tab" id="t-aufgaben" aria-selected="false">Aufgaben ({len(I.T)})</button>
  <button role="tab" id="t-versuch" aria-selected="false">Versuche</button>
  <button role="tab" id="t-folien" aria-selected="false">Folien</button>
  <button role="tab" id="t-fachwissen" aria-selected="false">Fachwissen</button>
</nav>
<section class="panel lese" id="p-ueberblick">{uber}</section>
<section class="panel lese" id="p-aufgaben" hidden>{aufgaben}</section>
<section class="panel lese" id="p-versuch" hidden><h2>Keine Versuche in dieser Einheit</h2><p>Die Stunde arbeitet mit vorgegebenen, erfundenen Modelldaten (Aufgabe 2) – das ist kein Versuch und bekommt keine Versuchs-ID.</p><p>Für später: <b>Jahresplan V5.3 »Vereinfachte Lebenszyklusanalyse eines Produkts«</b> (45 min, Schülerarbeit; Waage, Datenblätter, Taschenrechner; zwei vergleichbare Produkte). Funktionale Einheit und Systemgrenze festlegen, Massen, Transport und Energiedaten mit Quelle dokumentieren, Ergebnis als vereinfachte Analyse ausweisen. Vor einer Durchführung Inventar prüfen.</p></section>
<section class="panel" id="p-folien" hidden>
  <div class="steuer"><button class="btn ghost" id="f-zur" type="button">◀ zurück</button><button class="btn" id="f-vor" type="button">weiter ▶</button>
  <select id="f-wahl" aria-label="Folie wählen">{optionen}</select><button class="btn ghost" id="f-voll" type="button">Vollbild</button></div>
  <div class="buehne" id="buehne" tabindex="0"><img id="f-img" alt=""></div>
  <p class="hilfe">Die {len(titel)} sichtbaren Folien der TH-Unterricht (Lösungen sind hier schon aufgedeckt). Weiter: Tippen rechts, Pfeiltaste oder Leertaste · zurück: Tippen links · Vollbild: Knopf oder Taste F. Die Foliennummer entspricht der Nummer in PowerPoint.</p>
</section>
<section class="panel lese" id="p-fachwissen" hidden>{fw}</section>
<footer class="lese">Gebaut aus werkzeug/quellen/5ab-k1-4-5-oeko (eine Datenquelle für Seiten, PPP, Word und diese Mappe) · Stand {e(I.STAND)}</footer>
</div>
<script>
const F={folien_js};
const tabs=[...document.querySelectorAll('nav.tabs button')];
function zeige(id){{tabs.forEach(b=>{{const on=b.id==='t-'+id;b.setAttribute('aria-selected',on);document.getElementById('p-'+b.id.slice(2)).hidden=!on;}});try{{localStorage.setItem('k145oeko-tab',id)}}catch(e){{}}}}
tabs.forEach(b=>b.addEventListener('click',()=>zeige(b.id.slice(2))));
try{{const t=localStorage.getItem('k145oeko-tab');if(t&&document.getElementById('p-'+t))zeige(t)}}catch(e){{}}
// Herrichten: Haken merken
document.querySelectorAll('input[data-k]').forEach(c=>{{try{{c.checked=localStorage.getItem('k145oeko-'+c.dataset.k)==='1'}}catch(e){{}}
  c.addEventListener('change',()=>{{try{{localStorage.setItem('k145oeko-'+c.dataset.k,c.checked?'1':'0')}}catch(e){{}}}})}});
// Aufgaben: Suche und Filter
const karten=[...document.querySelectorAll('#p-aufgaben details.card')];let fil=null;
function filtern(){{const q=document.getElementById('suche').value.toLowerCase().trim();let n=0;
  karten.forEach(k=>{{const ok=(!q||k.dataset.s.includes(q))&&(fil===null||k.dataset.h===fil);k.hidden=!ok;if(ok)n++;}});document.getElementById('leer').hidden=n>0;}}
document.getElementById('suche').addEventListener('input',filtern);
document.querySelectorAll('#filter button').forEach(b=>b.addEventListener('click',()=>{{fil=(fil===b.dataset.f)?null:b.dataset.f;
  document.querySelectorAll('#filter button').forEach(x=>x.setAttribute('aria-pressed',x.dataset.f===fil));filtern();}}));
// Folien
let i=0;const img=document.getElementById('f-img'),wahl=document.getElementById('f-wahl'),bu=document.getElementById('buehne');
function geh(n){{i=Math.max(0,Math.min(F.length-1,n));img.src=F[i][0];img.alt='Folie '+F[i][2]+': '+F[i][1];wahl.value=i;}}
document.getElementById('f-vor').onclick=()=>geh(i+1);document.getElementById('f-zur').onclick=()=>geh(i-1);wahl.onchange=()=>geh(+wahl.value);
bu.addEventListener('click',ev=>{{const r=bu.getBoundingClientRect();geh(ev.clientX-r.left<r.width/3?i-1:i+1)}});
document.getElementById('f-voll').onclick=()=>{{try{{bu.requestFullscreen&&bu.requestFullscreen().catch(()=>{{}})}}catch(e){{}}}};
document.addEventListener('keydown',ev=>{{if(document.getElementById('p-folien').hidden)return;
  if(['ArrowRight',' ','PageDown'].includes(ev.key)){{ev.preventDefault();geh(i+1)}}else if(['ArrowLeft','PageUp'].includes(ev.key)){{ev.preventDefault();geh(i-1)}}
  else if(ev.key==='f'||ev.key==='F')document.getElementById('f-voll').click();}});
geh(0);
</script>
"""
open(ZIEL, "w", encoding="utf-8").write(seite)
print("ok", ZIEL, len(seite) // 1024, "KB")
