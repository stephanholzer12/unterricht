"""Test-Training NAWI 2AB · 1. Test (05.11.2026) – Quiz aus werkzeug/vorlagen/quiz_vorlage.html + Kahoot-/Blooket-Tabelle.
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/2ab-test1-training/bauen.py [<Zielordner für Kahoot-xlsx>]
Schreibt 2ab-nawi/test1-training/quiz.html (und 26-27-2AB-Test1-Kahoot-v1.xlsx in den Zielordner). Fragen in eigenen Worten; Seiten nach Trauner Chemie 2.
Aufgebaut nach quellen/4ab-test1-training/bauen.py."""
import base64, io, json, os, re, sys
import qrcode

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
WZ = os.path.join(ROOT, "werkzeug")
BASE = "https://stephanholzer12.github.io/unterricht/"
ORDNER = "2ab-nawi/test1-training"
EYEBROW = "NAWI · 2AB · Training für den 1. Test (05.11.)"
TITEL = "Test-Training"

# (Thema, Frage, [richtig, falsch, falsch, falsch], Erklärung, Seite) – richtige Antwort zuerst, die Anzeige mischt
QUIZ = [
 ("Atombau", "Welches Teilchen bestimmt, um welches Element es sich handelt?", ["die Zahl der Protonen", "die Zahl der Neutronen", "die Zahl der Elektronen in der Außenschale", "die Masse des Atoms"],
  "Die Ordnungszahl im PSE ist die Protonenzahl. Ändert sich die Neutronenzahl, bleibt es dasselbe Element (Isotop).", "Buch S. 10"),
 ("Atombau", "Isotope eines Elements unterscheiden sich in der Zahl der …", ["Neutronen", "Protonen", "Elektronen", "Außenelektronen"],
  "Isotope haben gleich viele Protonen, aber unterschiedlich viele Neutronen – z. B. C-12 und C-14.", "Buch S. 10"),
 ("Atommodelle", "Dalton stellte sich Atome vor als …", ["kleine, unteilbare Kugeln", "Kugel mit Elektronen wie Rosinen", "winzigen Kern mit fast leerer Hülle", "Kern mit Elektronen auf festen Bahnen"],
  "Bei Dalton sind Atome massive, unteilbare Kugeln; die Atome eines Elements sind gleich.", "Buch S. 12"),
 ("Atommodelle", "Welches Modell heißt auch »Rosinenkuchenmodell«?", ["das Modell von Thomson", "das Modell von Dalton", "das Modell von Rutherford", "das Modell von Bohr"],
  "Thomson: eine positiv geladene Masse, in der die negativen Elektronen stecken wie Rosinen im Kuchen.", "Buch S. 12"),
 ("Atommodelle", "Beim Goldfolienversuch fliegen fast alle Teilchen durch die Folie. Was folgt daraus?", ["Das Atom ist fast leer, die Masse sitzt im winzigen Kern.", "Gold ist ein besonders weiches Metall.", "Atome sind massive Kugeln.", "Elektronen sind größer als der Kern."],
  "Rutherford schloss: Ein winziger, positiver Kern trägt fast die ganze Masse, die Hülle ist praktisch leer.", "Buch S. 13"),
 ("Atommodelle", "Was ist die große Schwäche des Modells von Rutherford?", ["Es erklärt nicht, warum die Elektronen nicht in den Kern stürzen.", "Es kennt keinen Atomkern.", "Es kennt keine Elektronen.", "Es sagt, Atome seien unteilbar."],
  "Kreisende Ladungen müssten Energie abgeben und in den Kern stürzen. Bohr löste das mit festen Energiestufen.", "Buch S. 13–14"),
 ("Atommodelle", "Wie viele Elektronen passen nach Bohr höchstens auf die 2. Schale (L-Schale)?", ["8", "2", "18", "32"],
  "Faustregel 2·n²: 1. Schale 2, 2. Schale 8, 3. Schale 18 Elektronen.", "Buch S. 14"),
 ("Orbitalmodell", "Was ist ein Orbital?", ["Raum, in dem sich ein Elektron sehr wahrscheinlich aufhält", "feste Kreisbahn eines Elektrons", "anderer Name für den Atomkern", "Teilchen im Atomkern"],
  "Im Orbitalmodell gibt es keine Bahnen mehr – nur Aufenthaltsräume mit hoher Wahrscheinlichkeit.", "Buch S. 14–15"),
 ("Orbitalmodell", "Wie viele Elektronen passen in alle p-Orbitale einer Schale zusammen?", ["6", "2", "10", "14"],
  "s: 2 · p: 6 · d: 10 · f: 14 Elektronen. Jedes einzelne Orbital fasst 2 Elektronen.", "Buch S. 15–16"),
 ("Orbitalmodell", "Welche Kurzschreibweise passt zu Natrium (11 Elektronen)?", ["[Ne] 3s¹", "[Ne] 3p¹", "[Ar] 4s¹", "1s² 2s² 2p⁷"],
  "Na: 1s² 2s² 2p⁶ 3s¹ – die ersten 10 Elektronen wie Neon, dann ein Elektron im 3s-Orbital.", "Buch S. 16–17"),
 ("Orbitalmodell", "Welches Orbital wird zuerst besetzt?", ["4s vor 3d", "3d vor 4s", "4p vor 4s", "3d vor 3p"],
  "Das 4s-Orbital ist energieärmer als 3d und wird deshalb vorher gefüllt (z. B. Kalium [Ar] 4s¹).", "Buch S. 16"),
 ("Bindungsarten", "Warum gehen Atome Bindungen ein?", ["Sie streben den stabilen Edelgaszustand (8 Außenelektronen) an.", "Sie wollen schwerer werden.", "Sie verlieren sonst ihre Protonen.", "Gleich geladene Atome ziehen sich an."],
  "Oktettregel: Eine volle Außenschale wie bei den Edelgasen ist energiearm und stabil.", "Buch S. 20"),
 ("Bindungsarten", "Metall + Nichtmetall: Welche Bindung entsteht?", ["Ionenbindung", "Atombindung", "Metallbindung", "Wasserstoffbrücke"],
  "Das Metall gibt Elektronen ab, das Nichtmetall nimmt sie auf – es entstehen Ionen, z. B. NaCl.", "Buch S. 20–21"),
 ("Bindungsarten", "Welche Bindung liegt in Kohlenstoffdioxid CO₂ vor?", ["Atombindung", "Ionenbindung", "Metallbindung", "keine Bindung"],
  "C und O sind beide Nichtmetalle – sie teilen sich Elektronenpaare.", "Buch S. 20, 22"),
 ("Bindungsarten", "Welche Bindung hält eine Patronenhülse aus Messing (Kupfer + Zink) zusammen?", ["Metallbindung", "Ionenbindung", "Atombindung", "Van-der-Waals-Kräfte"],
  "Metall + Metall = Metallbindung. Messing ist eine Legierung.", "Buch S. 25"),
 ("Ionenbindung", "Ein Natrium-Atom gibt ein Elektron ab. Was entsteht?", ["ein positives Ion (Kation) Na⁺", "ein negatives Ion (Anion) Na⁻", "ein Natrium-Molekül Na₂", "ein Edelgas-Atom"],
  "Wer Elektronen abgibt, hat mehr Protonen als Elektronen und wird positiv: Kation.", "Buch S. 21"),
 ("Ionenbindung", "Warum leitet festes Kochsalz den Strom nicht, Salzwasser aber schon?", ["Erst gelöst sind die Ionen frei beweglich.", "Im Wasser entstehen Elektronen.", "Festes Salz hat keine Ionen.", "Wasser allein leitet sehr gut."],
  "Im Ionengitter sitzen die Ionen fest. Gelöst können sie wandern und den Strom tragen (Elektrolyt).", "Buch S. 21"),
 ("Ionenbindung", "Warum zerspringt ein Salzkristall unter dem Hammer?", ["Die Ionenschichten verschieben sich, gleiche Ladungen stoßen sich ab.", "Salz ist innen hohl.", "Die Elektronen fliegen davon.", "Salz ist ein Metall."],
  "Nach dem Verschieben liegen Plus über Plus und Minus über Minus – der Kristall bricht.", "Buch S. 21"),
 ("Atombindung", "Was teilen sich zwei Atome bei der Atombindung?", ["ein oder mehrere Elektronenpaare", "ihre Protonen", "ihre Atomkerne", "frei bewegliche Ionen"],
  "Bei der Atombindung (Elektronenpaarbindung) gehören gemeinsame Elektronenpaare beiden Atomen.", "Buch S. 22"),
 ("Atombindung", "Wann ist eine Atombindung polar?", ["wenn ein Atom die Elektronen stärker anzieht (EN-Unterschied)", "wenn beide Atome gleich sind", "wenn Metalle beteiligt sind", "wenn das Molekül sehr groß ist"],
  "Das Atom mit der größeren Elektronegativität bekommt δ–, das andere δ+. Bei Cl–Cl ist die Bindung unpolar.", "Buch S. 22–23"),
 ("Atombindung", "Warum ist Wasser ein Dipol?", ["polare O–H-Bindungen und gewinkelte Form", "es enthält Ionen", "es ist ein Metall", "es ist linear gebaut"],
  "Die Teilladungen heben sich nur bei symmetrischen Molekülen auf. Wasser ist gewinkelt – eine Seite δ–, die andere δ+.", "Buch S. 23"),
 ("Metallbindung", "Warum leitet ein Kupferkabel den Strom?", ["Die Außenelektronen sind frei beweglich (Elektronengas).", "Kupfer besteht aus Molekülen.", "Zwischen Kupferatomen gibt es Wasserstoffbrücken.", "Kupfer enthält gelöste Ionen."],
  "Im Metallgitter bewegen sich die Elektronen frei zwischen den Kationen und tragen Strom und Wärme.", "Buch S. 25"),
 ("Metallbindung", "Warum lässt sich Stahl schmieden, ohne zu brechen?", ["Das Elektronengas hält die verschobenen Kationen weiter zusammen.", "Stahl enthält Ionen wie Salz.", "Stahl ist eigentlich flüssig.", "Die Atome sind unteilbare Kugeln."],
  "Beim Verformen gleiten die Schichten ab, das Elektronengas »klebt« sie sofort wieder zusammen.", "Buch S. 25"),
 ("Wasserstoffbrücken", "Wasserstoffbrücken entstehen, wenn H gebunden ist an …", ["Sauerstoff, Stickstoff oder Fluor", "Kohlenstoff", "ein Metall", "jedes beliebige Atom"],
  "Nur bei O, N und F ist der EN-Unterschied groß genug. Das δ+ H zieht das δ– Atom des Nachbarn an.", "Buch S. 26"),
 ("Wasserstoffbrücken", "Warum schwimmt Eis auf Wasser?", ["Wasserstoffbrücken halten die Moleküle im Eis auf Abstand.", "Im Eis sind Luftblasen eingeschlossen.", "Kalte Stoffe sind immer leichter.", "Eis besteht aus anderen Molekülen."],
  "Eis bildet ein weitmaschiges Gitter (0,92 g/cm³) und ist leichter als Wasser. Teiche frieren von oben zu.", "Buch S. 26"),
 ("Wasserstoffbrücken", "Warum kann eine Büroklammer auf Wasser schwimmen?", ["Oberflächenspannung durch Wasserstoffbrücken", "Metall ist leichter als Wasser", "Die Klammer ist magnetisch", "Wasser stößt Metalle ab"],
  "Die Moleküle an der Oberfläche halten über Wasserstoffbrücken zusammen wie eine Haut. Waschmittel zerstört sie.", "Buch S. 26"),
 ("Van-der-Waals-Kräfte", "Butan siedet bei –1 °C, Octan bei 126 °C. Warum?", ["Längere Moleküle haben stärkere Van-der-Waals-Kräfte.", "Octan enthält Ionenbindungen.", "Octan bildet Wasserstoffbrücken.", "Butan ist ein Metall."],
  "Je länger die Kette, desto stärker die Van-der-Waals-Kräfte – und desto höher die Siedetemperatur.", "Buch S. 27"),
 ("Löslichkeit", "Warum mischen sich Öl und Wasser nicht?", ["Öl ist unpolar, Wasser ist polar.", "Öl ist schwerer als Wasser.", "Wasser ist ein Metall.", "Öl besteht aus Ionen."],
  "Ähnliches löst sich in Ähnlichem: polar mit polar, unpolar mit unpolar. Öl schwimmt oben, weil es leichter ist.", "Buch S. 27"),
 ("Löslichkeit", "Warum löst sich Kochsalz in Wasser, aber nicht in Öl?", ["Die Ionen werden von den Wasser-Dipolen angezogen.", "Öl ist zu kalt.", "Salz ist unpolar.", "Öl enthält schon zu viel Salz."],
  "Die geladenen Ionen werden von den polaren Wassermolekülen umhüllt – unpolares Öl kann das nicht.", "Buch S. 21, 27"),
]
b64 = lambda p: base64.b64encode(open(p, "rb").read()).decode()
nb = lambda s: re.sub(r"(?<=\d) (?=\d{3}\b)", " ", s)


def qr_b64(url):
    q = qrcode.QRCode(border=2, box_size=12, error_correction=qrcode.constants.ERROR_CORRECT_M); q.add_data(url); q.make()
    b = io.BytesIO(); q.make_image(fill_color="#2C3424", back_color="white").save(b, format="PNG"); return base64.b64encode(b.getvalue()).decode()


def wrap(src):
    te = src.index("</title>") + 8; rest = src[te:]; bs = rest.index("<header")
    return ('<!doctype html>\n<html lang="de">\n<head>\n<meta charset="utf-8">\n<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
            '<meta name="robots" content="noindex">\n' + src[:te] + "\n" + rest[:bs] + "</head>\n<body>\n" + rest[bs:] + "\n</body>\n</html>\n")


assert all(len(a) == 4 for _, _, a, _, _ in QUIZ)
qs = ",\n".join("  { t: %s, q: %s, a: %s, e: %s, p: %s }" % tuple(json.dumps(x, ensure_ascii=False) for x in (t, nb(q), [nb(z) for z in a], nb(e), p)) for t, q, a, e, p in QUIZ)
qv = open(os.path.join(WZ, "vorlagen", "quiz_vorlage.html"), encoding="utf-8").read()
qv = qv.replace("<title>Bindungs-Check</title>", f"<title>{TITEL} – 2AB</title>")
qv = qv.replace("NAWI · 2AB · K1.2 Die Bindungsarten", EYEBROW)
qv = qv.replace("<h1>Bindungs-Check</h1>", f"<h1>{TITEL}</h1>").replace("<h2>Bindungs-Check am Handy</h2>", f"<h2>{TITEL} am Handy</h2>")
qv = qv.replace('alt="QR-Code zum Bindungs-Check"', f'alt="QR-Code zum {TITEL}"')
qv = qv.replace("Frage 1 von 16", f"Frage 1 von {len(QUIZ)}")
a, b = qv.index("const Q = ["), qv.index("];", qv.index("const Q = [")) + 2
qv = qv[:a] + "const Q = [\n" + qs + "\n];" + qv[b:]
n = len(QUIZ)
qv = qv.replace("if (p >= 13) return", f"if (p >= {round(n * 0.85)}) return").replace("if (p >= 9) return", f"if (p >= {round(n * 0.6)}) return")
if "[hidden]" not in qv: qv = qv.replace("<style>", "<style>\n[hidden] { display: none !important; }", 1)
for k, v in {"__CAMO__": b64(WZ + "/camo.png"), "__LOGO__": b64(WZ + "/logo.png"), "__SMILEY__": b64(WZ + "/smiley.png"),
             "__QR__": qr_b64(BASE + ORDNER + "/quiz.html")}.items():
    qv = qv.replace(k, v)
assert not re.search(r"__[A-Z]+__", qv) and "Bindungs-Check" not in qv
os.makedirs(os.path.join(ROOT, ORDNER), exist_ok=True)
open(os.path.join(ROOT, ORDNER, "quiz.html"), "w", encoding="utf-8").write(wrap(qv))
print("quiz.html:", n, "Fragen")

# ---------------------------------------------------------------- Kahoot-/Blooket-Tabelle (Regelwerk 10)
if len(sys.argv) > 1:
    import random
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    random.seed(5112026)
    wb = Workbook(); an = wb.active; an.title = "Anleitung"
    for r in [["Klassenspiel 2AB · 1. Test (05.11.2026) – dieselben Fragen wie das Test-Training"], [],
              ["Kahoot: kahoot.com → Erstellen → Quiz → »Tabelle importieren« → Blatt »Kahoot« als eigene .xlsx speichern und hochladen."],
              ["Blooket: Set erstellen → »Spreadsheet Import« → Blatt »Blooket« verwenden."],
              ["Gespielt wird am Lehrer-PC am Beamer; Kadetten nur mit Spitznamen, kein Konto. Handyfrei: Team-Modus (ein Gerät pro Gruppe) oder ABCD-Karten."],
              ["Die richtige Antwort steht gemischt an Platz 1–4 (Spalte »Richtig«). Fragen ≤ 120 Zeichen, Antworten ≤ 75 Zeichen."]]:
        an.append(r)
    an["A1"].font = Font(bold=True, size=14, color="3B4A2F"); an.column_dimensions["A"].width = 120
    ka = wb.create_sheet("Kahoot"); bl = wb.create_sheet("Blooket")
    hk = ["Nr.", "Frage", "Antwort 1", "Antwort 2", "Antwort 3", "Antwort 4", "Zeit (s)", "Richtig", "Thema", "Buch"]
    ka.append(hk); bl.append(["Question #", "Question Text", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Time Limit (sec)", "Correct Answer(s)"])
    for i, (t, q, a, e, p) in enumerate(QUIZ, 1):
        assert len(q) <= 120, (i, len(q), q); assert all(len(x) <= 75 for x in a), (i, [len(x) for x in a])
        order = list(range(4)); random.shuffle(order); ans = [a[k] for k in order]; ok = order.index(0) + 1
        ka.append([i, q] + ans + [20, ok, t, p]); bl.append([i, q] + ans + [20, ok])
    for ws in (ka, bl):
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="3B4A2F")
        for col, w in zip("ABCDEFGHIJ", [5, 60, 32, 32, 32, 32, 9, 9, 18, 12]): ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    out = os.path.join(sys.argv[1], "26-27-2AB-Test1-Kahoot-v1.xlsx"); wb.save(out); print("Kahoot:", out)
