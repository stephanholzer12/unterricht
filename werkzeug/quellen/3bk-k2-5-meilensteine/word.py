"""Word-Dateien PBGW 3BK · K2.1–5.4 Meilensteine: TH-AB, TH-Loesung, TH-Fachwissen (keine Versuche, kein Kahoot).
Aufruf im Repository-Ordner:  python3 werkzeug/quellen/3bk-k2-5-meilensteine/word.py <Zielordner>
Layout wie Muster-TH-AB (Kopf mit Wappen, Titelzeile mit olivgrüner Linie). Daten aus inhalt.py, plan.py, fachwissen.py."""
import json, os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, HERE)
import inhalt as I
import fachwissen as F
import plan as PL

OUT = sys.argv[1]
VER = I.V
PRE = I.PRAEFIX
OLIVE, TEAL, RED, MUTED, RUST, INK = RGBColor(0x3B, 0x4A, 0x2F), RGBColor(0x2E, 0x7D, 0x8C), RGBColor(0xC0, 0x00, 0x00), RGBColor(0x5E, 0x66, 0x52), RGBColor(0x8C, 0x3B, 0x2E), RGBColor(0x1F, 0x2A, 0x1C)
THEORY, BOOKRED = RGBColor(0x6E, 0x7F, 0x4E), RGBColor(0xA3, 0x3A, 0x2B)
LOGO = os.path.join(ROOT, "werkzeug/logo.png")
UNIT = os.path.join(ROOT, I.ORDNER) + "/"
EINH = [u for u in json.load(open(os.path.join(ROOT, "werkzeug/einheiten.json"))) if u["ordner"] == I.ORDNER][0]


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcol); tcPr.append(s)


def borders(table, color="B9C4A6", size=6):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:color"), color); b.append(el)
    tblPr.append(b)


def noborders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblPr.append(b)


def pbottom(p, color="3B4A2F", sz=12):
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "1"); b.set(qn("w:color"), color); bd.append(b); pPr.append(bd)


def run(p, text, size=10.5, bold=False, color=INK, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color; r.font.name = "Calibri"
    return r


def md_runs(p, text, size=10.5, color=INK, bold=False):
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > pos: run(p, text[pos:m.start()], size, bold, color)
        run(p, m.group(1), size, True, color); pos = m.end()
    if pos < len(text): run(p, text[pos:], size, bold, color)


def para(doc, text="", size=10.5, bold=False, color=INK, italic=False, after=2, before=0, align=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.05
    if keep: pf.keep_with_next = True
    if align: p.alignment = align
    if text: (md_runs(p, text, size, color, bold) if not italic else run(p, text, size, bold, color, italic))
    return p


def label(p, lab, text, size=10.5, color=INK, labcol=OLIVE):
    run(p, lab, size, True, labcol); md_runs(p, text, size, color)


def widths(table, cms):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None: lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), cms): gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for c, w in zip(row.cells, cms): c.width = Cm(w)


def cellp(cell, text="", size=10, bold=False, color=INK, align=None):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    if align: p.alignment = align
    if text: md_runs(p, text, size, color, bold)
    return p


def lines(doc, n):
    for _ in range(n):
        p = para(doc, "", after=0); p.paragraph_format.line_spacing = 1.0; pbottom(p, "B9C4A6", 4)
        p.paragraph_format.space_before = Pt(6)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_doc(kopf, titel, fuss, quer=False):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.4); sec.footer_distance = Cm(0.6)
    t = doc.add_table(rows=1, cols=2); noborders(t); widths(t, [1.6, 15.8])
    t.cell(0, 0).paragraphs[0].add_run().add_picture(LOGO, width=Cm(1.2))
    c = t.cell(0, 1); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cellp(c, kopf, 8.5, True, TEAL); p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); run(p2, titel, 16, True, OLIVE)
    p = para(doc, "", after=4); pbottom(p, "3B4A2F", 14)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(fp, fuss + " · Seite ", 8, False, MUTED)
    r = fp.add_run(); r.font.size = Pt(8); r.font.color.rgb = MUTED
    for tg, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tg:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tg); r._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
    return doc


def box(doc, title, rows_fn, fill="EEF3E4", line="B9C4A6"):
    t = doc.add_table(rows=1, cols=1); borders(t, line, 8); widths(t, [17.4]); c = t.cell(0, 0); shade(c, fill)
    cellp(c, title, 10, True, OLIVE); rows_fn(c)
    para(doc, "", after=2)
    return t


def table(doc, head, rows, cms, size=9.5, headfill="3B4A2F", zebra=True):
    t = doc.add_table(rows=1, cols=len(head)); borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; shade(c, headfill); cellp(c, h, size, True, RGBColor(0xFF, 0xFF, 0xFF))
    for k, r in enumerate(rows):
        cells = t.add_row().cells
        for i, x in enumerate(r):
            if isinstance(x, tuple): cellp(cells[i], x[0], size, x[1] if len(x) > 1 else False, x[2] if len(x) > 2 else INK)
            else: cellp(cells[i], str(x), size)
            if zebra and k % 2 == 1: shade(cells[i], "F6F9F0")
    widths(t, cms)
    return t


# ================================================================ gemeinsame Kästen
def fahrplan_box(doc):
    z = EINH["fahrplan"]["zeilen"]
    je = {}
    for r in z:
        if r.get("klasse") and r["wann"].startswith("Stunde"): je[r["klasse"]] = je.get(r["klasse"], 0) + 1
    def fill(c):
        p = c.add_paragraph(); run(p, EINH["fahrplan"]["text"], 9.5)
        for k in list(je) + [None]:
            zz = [r for r in z if r.get("klasse") == k]
            if not zz: continue
            p = c.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1); run(p, k or "3BK", 9.5, True, OLIVE)
            it = c.add_table(rows=1, cols=5); borders(it, "D3DCC2", 4)
            for i, h in enumerate(["Wann", "Das machen wir", "Aufgaben · Zeit", "Wie", "Abgabe"]): cellp(it.rows[0].cells[i], h, 8, True, OLIVE)
            for r in zz:
                cs = it.add_row().cells
                cellp(cs[0], r["wann"] + (" · " + r["datum"] if r.get("datum") else ""), 8, True); cellp(cs[1], r["lesen"], 8)
                cellp(cs[2], r.get("aufgaben", "–") + (" · " + r["zeit"] if r.get("zeit") else ""), 8); cellp(cs[3], r.get("wie") or "–", 8); cellp(cs[4], r.get("abgabe", "–"), 8)
            widths(it, [2.2, 4.6, 4.2, 4.2, 1.8])
    box(doc, "Dein Fahrplan", fill)


def lehrplan_box(doc):
    def fill(c):
        for lab, x in F.LEHRPLAN:
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1); label(p, lab, x, 9)
    box(doc, "Lehrplanbezug", fill)


def stundenvorschlag(doc):
    para(doc, "AUSGEBLENDET · Kern in 3 Stunden (Jahresplan E3–E5: 6 Stunden brutto) · Rest optional · 3BK", 11, True, OLIVE, after=2, before=4, keep=True)
    rows = F.STUNDENVORSCHLAG[I.KLASSEN[0]] + [["Zusatz", "Für Schnelle (Theorieteil)", "–", "9–12 (OPTIONAL)", "nur wer fertig ist", "Partnercheck"],
                                               ["Besprechung", "in jeder Stunde eingeplant", "–", ", ".join(map(str, I.BESPRECHEN)) + " (BESPRECHEN)", "siehe Minutenplan", "Tafel, Skala"]]
    table(doc, ["Std.", "Theorie", "Kern-Versuch", "Aufgaben", "Zeit (min)", "Methode (ohne Handy)"], rows, [2.0, 3.4, 2.2, 2.2, 3.8, 3.8], 8.5)
    para(doc, "KERN = Pflicht für alle (Aufgaben 1–8) · OPTIONAL = Zusatz 9–12 · keine Hausübung, kein Versuch (PBGW). "
              "50 min je Stunde (Annahme), davon 2 min Klassenbuch; Summen geprüft (plan.py). Am 23.11. zuerst Test (20 min) und Organisation (4 min).", 8.5, False, MUTED, after=3)
    for st in PL.PLAN[I.KLASSEN[0]]:
        para(doc, f"Minutenplan Stunde {st['std']} · {st['datum']} (3BK)", 10, True, TEAL, after=1, keep=True)
        t0, rows = 0, []
        for name, m, oh, art in st["phasen"]:
            rows.append([f"{t0}–{t0 + m}", name, f"{m}", "ohne Handy" if oh else "am Handy"]); t0 += m
        table(doc, ["Minute", "Phase", "min", "Wie"], rows, [1.8, 11.4, 1.4, 2.8], 8.5)


def stolpersteine(doc):
    para(doc, "Die 5 größten Stolpersteine", 11, True, OLIVE, after=2, before=4, keep=True)
    table(doc, ["Stolperstein", "Richtig ist …"], [[a, b] for a, b in I.STOLPERSTEINE], [4.4, 13.0], 9)


def src_color(src):
    return BOOKRED if src.startswith("Versuch im Buch") else (TEAL if src.startswith("Buch-Aufgabe") else THEORY)


# ================================================================ TH-AB und TH-Loesung
def optionen(t, f):
    """Antwortmöglichkeiten wie am Handy gemischt (feste Reihenfolge nur, wenn das 4. Element False ist) – richtige Antwort nie immer zuerst."""
    import random
    if len(f) > 3 and f[3] is False: return list(f[2])
    o = list(f[2]); random.Random(t["id"] + f[0]).shuffle(o)
    if o[0] == f[2][0] and len(o) > 1: o = o[1:] + o[:1]
    return o


def aufgabe(doc, t, loesung):
    tb = doc.add_table(rows=1, cols=2); noborders(tb); widths(tb, [12.6, 4.8])
    p = cellp(tb.cell(0, 0), "", 11); run(p, f"Aufgabe {t['nr']}  ", 11, True, OLIVE); run(p, t["title"], 11, True, INK)
    p = cellp(tb.cell(0, 1), "", 8.5, align=WD_ALIGN_PARAGRAPH.RIGHT); run(p, t["src"], 8.5, True, src_color(t["src"]))
    p2 = tb.cell(0, 1).add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p2.paragraph_format.space_after = Pt(0)
    run(p2, f"{t['tag']} · ca. {t['zeit']} min" + (" (im Versuch)" if t.get("invs") else "") + (" · BESPRECHEN" if loesung and t.get("besprechen") else ""), 8.5, False, MUTED)
    para(doc, t["q"], 10, after=2)
    if t.get("hint"): para(doc, t["hint"], 8.5, False, MUTED, italic=True, after=2)
    if t.get("img"):
        p = para(doc, "", after=2, align=WD_ALIGN_PARAGRAPH.CENTER); p.add_run().add_picture(UNIT + t["img"], width=Cm(11))
    if t.get("tab"):
        table(doc, t["tab"][0], t["tab"][1:], {2: [1.6, 15.8], 3: [9.4, 4.0, 4.0], 4: [3.6, 5.2, 4.3, 4.3]}[len(t["tab"][0])], 9)
        if t.get("tabnote"): para(doc, t["tabnote"], 8.5, True, RUST, after=3)
    if loesung:
        for head, items in t["sol"]:
            for x in items:
                p = para(doc, "", after=1); run(p, ("" if head in ("Lösung", "Beispiel", "Bewertung") else head + " "), 9.5, True, RED); run(p, x, 9.5, False, RED)
        para(doc, "", after=3)
        return
    if t.get("raster"):  # K2.1 Aufgabe 1: 4 Merkmale × 3 Formen als Raster zum Ausfüllen
        table(doc, [""] + t["raster"][1], [[z, "", "", ""] for z in t["raster"][0]], [3.4, 4.7, 4.7, 4.6], 9, zebra=False)
        for r in doc.tables[-1].rows[1:]: r.height = Cm(1.6)
        para(doc, "", after=3)
        return
    sel = [f for f in t.get("fields", []) if len(f) > 2]
    if t["type"] == "multi" and sel and len(sel) == len(t["fields"]) and len({tuple(f[2]) for f in sel}) == 1:
        opts = sel[0][2]
        rows = [[f[1], "   ".join("☐ " + o for o in opts)] for f in sel]
        table(doc, ["", "Kreuze an"], rows, [9.4, 8.0], 9, "6E7F4E", zebra=False)
    else:
        for f in t.get("fields", []):
            if len(f) > 2:
                p = para(doc, "", after=1); run(p, f[1] + ":  ", 9.5, True); run(p, "   ".join("☐ " + o for o in optionen(t, f)), 9)
            else:
                p = para(doc, "", after=0); run(p, f[1], 9.5, True)
                lines(doc, 2 if len(f[1]) < 60 else 3)
        if t["type"] == "sketch":
            if t.get("ptab"):
                table(doc, t["ptab"], [["", "", ""] for _ in range(3)], [5.8, 5.8, 5.8], 9, zebra=False)
                for r in doc.tables[-1].rows[1:]: r.height = Cm(1.3)
            para(doc, "Am Handy: Foto (Mini-Whiteboard oder Papier) im Arbeitsblatt einfügen.", 9, False, MUTED, italic=True)
    para(doc, "", after=3)


def th_docs():
    fuss = f"K2.1–5.4 Meilensteine · Stand {I.STAND} · {VER}"
    for loesung in (False, True):
        doc = new_doc(("LÖSUNGEN · " if loesung else "") + "ARBEITSBLATT · PBGW · 3. JAHRGANG (3BK) · K2.1–5.4 · BUCH S. 23–44, ERGÄNZUNGSBUCH S. 58–73",
                      "Meilensteine: Vom Vorrat zum Staat" + (" · Lösungen" if loesung else ""), fuss + (" · Lösungen" if loesung else ""))
        if loesung:
            lehrplan_box(doc)
        if not loesung:
            p = para(doc, "", after=2); run(p, "Name: ", 10, True); run(p, "__________________________   ", 10); run(p, "Klasse: ", 10, True); run(p, "______", 10)
            p = para(doc, "", after=4); run(p, "Am Handy: ", 9, True, OLIVE); run(p, I.URL + "aufgaben.html", 9, False, TEAL)
        fahrplan_box(doc)
        if loesung:
            stundenvorschlag(doc); stolpersteine(doc)
        last = None
        for t in I.T:
            if t["h"] != last:
                para(doc, I.GRP[t["h"]], 11.5, True, OLIVE, after=3, before=6, keep=True).paragraph_format.keep_with_next = True
                last = t["h"]
            aufgabe(doc, t, loesung)
        if loesung:
            para(doc, "Partnercheck-Karten (Aufgabe 12)", 11, True, OLIVE, after=2, before=4, keep=True)
            table(doc, ["Frage", "Lösung"], [[q, a] for q, a in I.PARTNERKARTEN], [8.7, 8.7], 9)
        name = f"{PRE}-TH-{'Loesung' if loesung else 'AB'}-{VER}.docx"
        doc.save(os.path.join(OUT, name)); print(name)


# ================================================================ TH-Fachwissen
def th_fachwissen():
    fuss = f"K2.1–5.4 Meilensteine · Stand {I.STAND} · {VER} · Fachwissen"
    doc = new_doc("NUR FÜR DICH · PBGW · 3. JAHRGANG · K2.1–5.4 · BUCH S. 23–44, ERGÄNZUNGSBUCH S. 58–73", "Meilensteine – Fachwissen für die Lehrkraft", fuss)
    para(doc, "Stundenkarten (3BK)", 13, True, OLIVE, after=2)
    for tit, steps, mat, ohne in F.STUNDENKARTEN:
        def fill(c, steps=steps, mat=mat, ohne=ohne):
            for i, s in enumerate(steps, 1):
                p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0.5); run(p, f"{i}  ", 9, True, OLIVE); run(p, s, 9)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); label(p, "Material: ", mat, 8.5)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); label(p, "Ohne Handy: ", ohne, 8.5)
        box(doc, tit, fill)
    p = para(doc, "", after=1); label(p, "Links: ", " · ".join(I.URL + k + ".html" for k in ("theorie", "aufgaben", "quiz")), 8, TEAL)
    para(doc, "Vor der Stunde", 11, True, OLIVE, after=1, before=4)
    for x in F.CHECKLISTE:
        p = para(doc, "", after=1); run(p, "☐  ", 10, True, OLIVE); md_runs(p, x, 9.5)
    page_break(doc)
    para(doc, "Die Einheit auf einen Blick", 13, True, OLIVE, after=2)
    table(doc, ["Teil · Stunde", "Kernaussage", "Material"], F.AUF_EINEN_BLICK, [3.2, 8.8, 5.4], 9)
    lehrplan_box(doc); stundenvorschlag(doc)
    for kap in F.KAPITEL:
        para(doc, kap["titel"], 12, True, OLIVE, after=2, before=6, keep=True)
        for x in kap["text"]: para(doc, x, 10, after=2)
        if kap.get("bild"):
            pp = para(doc, "", after=2, align=WD_ALIGN_PARAGRAPH.CENTER); pp.add_run().add_picture(UNIT + kap["bild"], width=Cm(12))
        if kap.get("nur"):
            def fill(c, x=kap["nur"]):
                for y in x:
                    p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1); md_runs(p, y, 9.5)
            box(doc, "Nur für dich", fill, "ECECEC", "BDBDBD")
        if kap.get("stolper"):
            def fill(c, x=kap["stolper"]):
                p = c.add_paragraph(); md_runs(p, x, 9.5)
            box(doc, "Stolperstein", fill, "FBF1E6", "E0B48A")
    para(doc, "Zeitrechnung", 12, True, OLIVE, after=2, before=6, keep=True)
    for x in F.RECHNEN: para(doc, x, 10, after=2)
    stolpersteine(doc)
    para(doc, "Häufige Fehlvorstellungen", 12, True, OLIVE, after=2, before=6, keep=True)
    table(doc, ["Fehlvorstellung", "Richtig ist", "So zeigst du es"], F.FEHLVORSTELLUNGEN, [4.8, 6.6, 6.0], 9)
    para(doc, "Fragen, die Kadetten stellen – mit Antworten", 12, True, OLIVE, after=2, before=6, keep=True)
    for q, a in F.FRAGEN:
        p = para(doc, "", after=0, keep=True); run(p, q, 10, True)
        para(doc, a, 9.5, after=3)
    para(doc, "Videos und Podcasts – Stand und Hinweise", 12, True, OLIVE, after=2, before=6, keep=True)
    for k, x in F.VIDEO_NOTIZ.items():
        para(doc, x, 9.5, after=2)
    para(doc, "Ausblick", 12, True, OLIVE, after=2, before=6, keep=True)
    para(doc, F.AUSBLICK, 10, after=2)
    para(doc, "Quellen – geprüfte Bücher und was übernommen wurde", 12, True, OLIVE, after=2, before=6, keep=True)
    for x in F.QUELLEN:
        p = para(doc, "", after=1); run(p, "·  ", 9.5, True, OLIVE); md_runs(p, x, 9)
    f = f"{PRE}-TH-Fachwissen-{VER}.docx"; doc.save(os.path.join(OUT, f)); print(f)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    th_docs(); th_fachwissen()
